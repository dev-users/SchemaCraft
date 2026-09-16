import base64
import copy
import io
import unittest
from unittest import mock
from openpyxl import Workbook
import SchemaCraft as APP
from schemacraft_workspace import use_context
from schemacraft_io import inspect_import_sheets, parse_import_sheets
import test_release3_abc as fixtures
from test_backend import CAT_IDENTITY,CAT_CHILDREN,FLD_NAME,FLD_STATUS,FLD_CHILD_NAME


class ImportReviewTests(unittest.TestCase):
    setUp=fixtures.Release3ABCFeatureTests.setUp
    tearDown=fixtures.Release3ABCFeatureTests.tearDown

    def payload(self, conflict=False):
        workbook=Workbook();book=workbook.active;book.title='People'
        book.append(['ID','Name']);book.append(['R3000001','Changed name']);book.append(['R3000002','New person'])
        if conflict:book.append(['R3000001','Conflicting name'])
        sheet=workbook.create_sheet('Status');sheet.append(['ID','Status']);sheet.append(['R3000001','متوقف']);sheet.append(['R3000002','نشط'])
        sheet=workbook.create_sheet('Children');sheet.append(['ID','minor_id','Name']);sheet.append(['R3000001',1,'Changed child']);sheet.append(['R3000001',None,'Another child']);sheet.append(['R3000002',None,'New child'])
        output=io.BytesIO();workbook.save(output);workbook.close()
        return {'file_data':base64.b64encode(output.getvalue()).decode(),'filename':'all-sheets.xlsx',
            'schema_revision':APP.read_schema_file()['revision'],'generate_missing_ids':False,
            'duplicate_policy':'update','clear_blank_values':False,'sheet_mappings':[
                {'sheet_name':'People','category_id':CAT_IDENTITY,'mapping':{'1':'__record_code__','2':FLD_NAME}},
                {'sheet_name':'Status','category_id':CAT_IDENTITY,'mapping':{'1':'__record_code__','2':FLD_STATUS}},
                {'sheet_name':'Children','category_id':CAT_CHILDREN,'repeated_mode':'merge','mapping':{'1':'__record_code__','2':'__minor_id__','3':FLD_CHILD_NAME}}]}

    def seed_child(self):
        record=APP.load_record('R3000001')
        APP.save_record({'mode':'update','record_code':'R3000001','expected_updated_at':record['updated_at'],'main':record['main'],'related':{CAT_CHILDREN:[{'values':{FLD_CHILD_NAME:'Original child'}}]}})

    def test_inspection_reads_all_sheets_and_review_does_not_write(self):
        with use_context(self.primary):
            self.seed_child();payload=self.payload()
            inspection=APP.inspect_import(payload)
            self.assertEqual([s['sheet_name'] for s in inspection['sheet_inspections']],['People','Status','Children'])
            before=self.primary.workbook_path.read_bytes()
            with mock.patch.object(APP,'create_backup',side_effect=AssertionError('preview must not write')):
                review=APP.commit_import(payload,_preview=True)
            self.assertEqual(before,self.primary.workbook_path.read_bytes())
            self.assertEqual(review['rejected'],0)
            self.assertEqual([e['action'] for e in review['entries']],['updated','added'])
            update=review['entries'][0]
            name=next(c for c in update['changes'] if c.get('field_id')==FLD_NAME)
            self.assertEqual(name['before'],'طالب أول');self.assertEqual(name['after'],'Changed name')
            selections=[{'record_code':update['record_code'],'change_ids':[c['id'] for c in update['changes'] if c.get('field_id') not in {FLD_NAME,FLD_CHILD_NAME}]}]
            result=APP.apply_import_review({'review_token':review['review_token'],'selections':selections})
            self.assertEqual((result['imported'],result['updated']),(0,1))
            record=APP.load_record('R3000001');self.assertEqual(record['main'][FLD_NAME],'طالب أول')
            self.assertEqual(record['main'][FLD_STATUS],'متوقف')
            self.assertEqual([r['values'][FLD_CHILD_NAME] for r in record['related'][CAT_CHILDREN]],['Original child','Another child'])
            self.assertFalse(APP.require_workspace().identity_search('R3000002'))
            self.assertTrue(result['history']['details'][0]['changes'])
            with self.assertRaises(APP.ApplicationError):APP.apply_import_review({'review_token':review['review_token'],'selections':selections})

    def test_approve_all_and_reject_stale_preview(self):
        with use_context(self.primary):
            self.seed_child();payload=self.payload();review=APP.commit_import(payload,_preview=True)
            selections=[{'record_code':e['record_code'],'change_ids':[c['id'] for c in e['changes']]} for e in review['entries']]
            result=APP.apply_import_review({'review_token':review['review_token'],'selections':selections})
            self.assertEqual((result['imported'],result['updated']),(1,1))
            self.assertEqual(APP.load_record('R3000002')['main'][FLD_NAME],'New person')
            review=APP.commit_import(payload,_preview=True)
            record=APP.load_record('R3000001')
            APP.save_record({'mode':'update','record_code':'R3000001','expected_updated_at':record['updated_at'],'main':{**record['main'],FLD_NAME:'Newer edit'},'related':record['related']})
            with self.assertRaisesRegex(APP.ApplicationError,'تغيّرت'):APP.apply_import_review({'review_token':review['review_token'],'selections':selections})

    def test_conflicts_unknown_cards_and_direct_commit_are_rejected(self):
        with use_context(self.primary):
            payload=self.payload(conflict=True);review=APP.commit_import(payload,_preview=True)
            self.assertTrue(any(e['record_code']=='R3000001' for e in review['errors']))
            payload=self.payload();review=APP.commit_import(payload,_preview=True)
            self.assertTrue(any('رقم البطاقة' in e['message'] for e in review['errors']))
            with self.assertRaises(APP.ApplicationError):APP.commit_import(payload)
            payload['sheet_mappings'][0]['mapping']['2']=FLD_CHILD_NAME
            with self.assertRaises(APP.ApplicationError):APP.commit_import(payload,_preview=True)

    def test_dates_survive_review_and_serialize_for_the_browser(self):
        import json
        from test_backend import FLD_GREGORIAN
        with use_context(self.primary):
            self.seed_child()
            record=APP.load_record('R3000001')
            APP.save_record({'mode':'update','record_code':'R3000001','expected_updated_at':record['updated_at'],'main':{**record['main'],FLD_GREGORIAN:'2020-01-02'},'related':record['related']})
            payload=self.payload();review=APP.commit_import(payload,_preview=True)
            json.dumps(review)
            self.assertFalse(any(c.get('field_id')==FLD_GREGORIAN for c in review['entries'][0]['changes']))

    def test_nested_repeated_categories_link_parent_cards_across_sheets(self):
        with use_context(self.primary):
            self.seed_child()
            schema=APP.read_schema_file()
            category=copy.deepcopy(next(c for c in schema['categories'] if c['id']==CAT_CHILDREN))
            category.update(id='cat_999999999991',label='Nested',parent_category_id=CAT_CHILDREN)
            category['fields']=[copy.deepcopy(category['fields'][0])]
            category['fields'][0]['id']='fld_999999999991'
            schema['categories'].append(category);APP.save_schema(schema)
            payload=self.payload()
            from openpyxl import load_workbook
            w=load_workbook(io.BytesIO(base64.b64decode(payload['file_data'])))
            sheet=w.create_sheet('Nested');sheet.append(['ID','parent_minor_id','Name']);sheet.append(['R3000001',1,'Nested value'])
            output=io.BytesIO();w.save(output);w.close();payload['file_data']=base64.b64encode(output.getvalue()).decode()
            payload['sheet_mappings'].append({'sheet_name':'Nested','category_id':category['id'],'mapping':{'1':'__record_code__','2':'__parent_minor_id__','3':category['fields'][0]['id']}})
            review=APP.commit_import(payload,_preview=True)
            self.assertEqual(review['rejected'],0,str(review['errors']))
            selections=[{'record_code':entry['record_code'],'change_ids':[change['id'] for change in entry['changes']]} for entry in review['entries']]
            APP.apply_import_review({'review_token':review['review_token'],'selections':selections})
            result=APP.load_record('R3000001')
            self.assertEqual(result['related'][category['id']][0]['parent_child_id'],result['related'][CAT_CHILDREN][0]['_child_id'])

    def test_replace_cards_requires_review_and_can_be_declined(self):
        with use_context(self.primary):
            self.seed_child();payload=self.payload();payload['sheet_mappings'][2]['repeated_mode']='replace'
            review=APP.commit_import(payload,_preview=True);entry=review['entries'][0]
            self.assertTrue(any(c['kind']=='card' and c['after'] is None for c in entry['changes']))
            result=APP.apply_import_review({'review_token':review['review_token'],'selections':[{'record_code':entry['record_code'],'change_ids':[c['id'] for c in entry['changes'] if 'category_id' not in c]}]})
            self.assertEqual(len(APP.load_record('R3000001')['related'][CAT_CHILDREN]),1)
            self.assertEqual(result['updated'],1)

if __name__=='__main__':unittest.main()
