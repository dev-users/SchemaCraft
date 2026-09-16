from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from test_backend import CAT_DOCUMENTS, FLD_DOC_TYPE, FLD_NAME, configured_schema

import SchemaCraft as APP
from schemacraft_advanced import AuditUserStore
from schemacraft_io import EXPORT_MAIN_SHEET, export_workbook_bytes


class Alpha8FeatureTests(unittest.TestCase):
    def test_audit_user_is_session_scoped_but_names_are_suggested(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            first = AuditUserStore(Path(root))
            self.assertEqual(first.select("ليلى")["current_user"], "ليلى")
            restarted = AuditUserStore(Path(root))
            payload = restarted.read()
            self.assertEqual(payload["current_user"], "")
            self.assertIn("ليلى", payload["users"])

    def test_repeatable_cards_export_to_their_own_sheet(self) -> None:
        schema = configured_schema()
        documents = next(category for category in schema["categories"] if category["id"] == CAT_DOCUMENTS)
        records = [{
            "record_code": "A8000001",
            "archived": False,
            "created_at": "2026-08-11T12:00:00Z",
            "updated_at": "2026-08-11T12:00:00Z",
            "values": {FLD_NAME: "ليلى"},
            "related": {CAT_DOCUMENTS: [
                {"values": {FLD_DOC_TYPE: "قديم"}},
                {"values": {FLD_DOC_TYPE: "حالي"}},
            ]},
        }]
        plain = load_workbook(io.BytesIO(export_workbook_bytes(
            schema, records, [FLD_NAME, FLD_DOC_TYPE], include_related=False,
        )))
        self.assertNotIn(FLD_DOC_TYPE, [cell.value for cell in plain[EXPORT_MAIN_SHEET][1]])

        related = load_workbook(io.BytesIO(export_workbook_bytes(
            schema,
            records,
            [FLD_NAME, FLD_DOC_TYPE],
            include_related=True,
        )))
        sheet = related[documents["label"]]
        headers = [cell.value for cell in sheet[1]]
        self.assertIn(FLD_DOC_TYPE, headers)
        column = headers.index(FLD_DOC_TYPE) + 1
        self.assertEqual(sheet.cell(3, column).value, "قديم")
        self.assertEqual(sheet.cell(4, column).value, "حالي")

    def test_nested_global_tree_updates_configuration_and_preserves_ids(self) -> None:
        schema = {
            "categories": [
                {
                    "id": "cat_aaaaaaaaaaaa",
                    "label": "جذر قديم",
                    "kind": "main",
                    "fields": [{
                        "id": "fld_aaaaaaaaaaaa",
                        "label": "حقل قديم",
                        "type": "text",
                        "global_tree_ref": "gcat_aaaaaaaaaaaa",
                        "global_tree_key": "field-root",
                    }],
                    "global_tree_ref": "gcat_aaaaaaaaaaaa",
                    "global_tree_key": "root",
                    "parent_category_id": None,
                    "anchor_field_id": None,
                },
                {
                    "id": "cat_bbbbbbbbbbbb",
                    "label": "فرع قديم",
                    "kind": "repeatable",
                    "fields": [],
                    "global_tree_ref": "gcat_aaaaaaaaaaaa",
                    "global_tree_key": "child",
                    "parent_category_id": "cat_aaaaaaaaaaaa",
                    "anchor_field_id": "fld_aaaaaaaaaaaa",
                },
            ]
        }
        definition = {"category_tree": [
            {"key": "root", "definition": {"label": "جذر محدث", "kind": "main"}, "fields": [
                {"key": "field-root", "definition": {"label": "حقل محدث", "type": "textarea"}}
            ]},
            {"key": "child", "definition": {"label": "فرع محدث", "kind": "repeatable"}, "fields": []},
        ]}
        self.assertTrue(APP._merge_global_category_tree(schema, definition, "gcat_aaaaaaaaaaaa"))
        root, child = schema["categories"]
        self.assertEqual((root["id"], root["label"]), ("cat_aaaaaaaaaaaa", "جذر محدث"))
        self.assertEqual((root["fields"][0]["id"], root["fields"][0]["label"]), ("fld_aaaaaaaaaaaa", "حقل محدث"))
        self.assertEqual(child["parent_category_id"], "cat_aaaaaaaaaaaa")
        self.assertEqual(child["anchor_field_id"], "fld_aaaaaaaaaaaa")


if __name__ == "__main__":
    unittest.main()
