from __future__ import annotations

import base64
import copy
import io
import tempfile
import unittest
import json
from unittest import mock
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from test_backend import FLD_NAME, FLD_STATUS, configured_schema

import SchemaCraft as APP
from schemacraft_advanced import AdvancedFeatureError, inspect_portable_package
from schemacraft_io import EXPORT_MAIN_SHEET
from schemacraft_workspace import use_context

GLOBAL_NAME = "gfld_300000000001"
GLOBAL_CATEGORY = "gcat_300000000001"


class Release3ABCFeatureTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.data = self.root / "data"
        names = (
            "DATA_DIR",
            "SCHEMA_PATH",
            "WORKBOOK_PATH",
            "BACKUP_DIR",
            "BUILDER_AUTH_PATH",
            "STARTUP_ERROR_LOG",
            "WORKSPACE_SETTINGS_PATH",
            "WORKSPACE_MANAGER",
            "GLOBAL_DEFINITIONS",
            "EXPORT_HISTORY",
            "IMPORT_HISTORY",
            "DEVELOPER_MODE",
            "_BUILDER_UNLOCKED",
            "_DATASET_SNAPSHOT",
            "_DATASET_SNAPSHOTS",
        )
        self.originals = {name: getattr(APP, name) for name in names}
        APP.DATA_DIR = self.data
        APP.SCHEMA_PATH = self.data / "schema.json"
        APP.WORKBOOK_PATH = self.data / "database.xlsx"
        APP.BACKUP_DIR = self.root / "backups"
        APP.BUILDER_AUTH_PATH = self.root / "builder-auth.json"
        APP.STARTUP_ERROR_LOG = self.root / "startup-error.log"
        APP.WORKSPACE_SETTINGS_PATH = self.data / "workspace-settings.json"
        APP.WORKSPACE_MANAGER = None
        APP.GLOBAL_DEFINITIONS = None
        APP.EXPORT_HISTORY = None
        APP.IMPORT_HISTORY = None
        APP.DEVELOPER_MODE = True
        APP._BUILDER_UNLOCKED = True
        APP._DATASET_SNAPSHOT = None
        APP._DATASET_SNAPSHOTS = {}
        APP.ensure_storage()
        APP.save_schema(configured_schema())
        APP.save_record(
            {
                "mode": "create",
                "record_code": "R3000001",
                "main": {FLD_NAME: "طالب أول", FLD_STATUS: "نشط"},
                "related": {},
            }
        )
        APP.initialize_workspace()
        self.manager = APP.require_workspace()
        self.primary = self.manager.context()

    def tearDown(self) -> None:
        for name, value in self.originals.items():
            setattr(APP, name, value)
        self.temporary.cleanup()

    def test_general_delete_uses_metadata_recovery_and_retains_record_indexes(self) -> None:
        APP.GLOBAL_DEFINITIONS.save_definition("field", GLOBAL_NAME, {"label": "اسم عام", "type": "text"}, expected_revision=0)
        with use_context(self.primary):
            previous = APP.read_schema_file()
            self._field(previous, FLD_NAME)["global_ref"] = GLOBAL_NAME
            APP.atomic_write_json(APP._schema_path(), previous)
            cached = APP._dataset_snapshot_unlocked(previous)
            workbook_bytes = self.primary.workbook_path.read_bytes()
        with mock.patch.object(APP, "create_backup", side_effect=AssertionError("must not copy the workspace")), mock.patch.object(APP, "save_schema", side_effect=AssertionError("must not migrate unchanged values")), mock.patch.object(APP, "schedule_workbook_schema_sync") as sync:
            removed = APP.manage_global_definition({"action":"delete", "kind":"field", "global_ref":GLOBAL_NAME, "expected_revision":1})
        self.assertTrue(removed["backup"]["metadata_only"])
        self.assertEqual(removed["detached_schema_ids"], [self.primary.schema_id])
        self.assertEqual(self.primary.workbook_path.read_bytes(), workbook_bytes)
        sync.assert_called_once()
        with use_context(self.primary):
            current = APP.read_schema_file()
            self.assertIsNone(self._field(current, FLD_NAME)["global_ref"])
            self.assertEqual(current["revision"], previous["revision"] + 1)
            self.assertIs(APP._dataset_snapshot_unlocked(current).records, cached.records)
            self.assertEqual(APP.load_record("R3000001")["main"][FLD_NAME], "طالب أول")
        with zipfile.ZipFile(APP.BACKUP_DIR / removed["backup"]["filename"]) as recovery:
            self.assertIn(GLOBAL_NAME, json.loads(recovery.read("global-definitions.json"))["fields"])
            saved_schema = json.loads(recovery.read(self.primary.schema_path.relative_to(APP.DATA_DIR).as_posix()))
            self.assertEqual(self._field(saved_schema, FLD_NAME)["global_ref"], GLOBAL_NAME)
            self.assertFalse(any(name.endswith(".xlsx") or "attachments/" in name for name in recovery.namelist()))

    def test_general_order_commit_checks_revision_and_membership(self) -> None:
        refs = ["gfld_300000000001", "gfld_300000000002"]
        for index, ref in enumerate(refs):
            APP.GLOBAL_DEFINITIONS.save_definition("field", ref, {"label":str(index),"type":"text"}, expected_revision=index)
        result = APP.manage_global_definition({"action":"reorder_all","kind":"field","order":list(reversed(refs)),"expected_revision":2})
        self.assertEqual(list(result["global_definitions"]["fields"]), list(reversed(refs)))
        with self.assertRaises(APP.ApplicationError):
            APP.manage_global_definition({"action":"reorder_all","kind":"field","order":refs,"expected_revision":2})
        with self.assertRaises(APP.ApplicationError):
            APP.manage_global_definition({"action":"reorder_all","kind":"field","order":[refs[0],refs[0]],"expected_revision":3})
        self.assertEqual(list(APP.GLOBAL_DEFINITIONS.read()["fields"]), list(reversed(refs)))

    @staticmethod
    def _field(schema: dict, field_id: str) -> dict:
        return next(
            field
            for category in schema["categories"]
            for field in category["fields"]
            if field["id"] == field_id
        )

    def _link_global_name(self, context) -> None:
        with use_context(context):
            schema = APP.read_schema_file()
            self._field(schema, FLD_NAME)["global_ref"] = GLOBAL_NAME
            APP.save_schema(schema)

    def test_global_definition_search_propagation_and_safe_detach(self) -> None:
        self._link_global_name(self.primary)
        with use_context(self.primary):
            duplicate = APP.manage_workspace_schema(
                {
                    "action": "duplicate",
                    "name": "العاملون",
                    "schema_id": self.primary.schema_id,
                    "template_schema_id": self.primary.schema_id,
                }
            )
        worker = self.manager.context(duplicate["schema_id"])
        APP.create_linked_profile(
            {
                "source_schema_id": self.primary.schema_id,
                "target_schema_id": worker.schema_id,
                "record_code": "R3000001",
                "mappings": [
                    {"source_field_id": FLD_NAME, "target_field_id": FLD_NAME}
                ],
            }
        )

        with use_context(self.primary):
            definition = copy.deepcopy(self._field(APP.read_schema_file(), FLD_NAME))
        saved = APP.manage_global_definition(
            {
                "action": "save",
                "kind": "field",
                "global_ref": GLOBAL_NAME,
                "expected_revision": 0,
                "definition": {
                    **definition,
                    "label": "الاسم العام",
                },
            }
        )
        self.assertEqual(set(saved["updated_schema_ids"]), {self.primary.schema_id, worker.schema_id})
        for context in (self.primary, worker):
            with use_context(context):
                field = self._field(APP.read_schema_file(), FLD_NAME)
                self.assertEqual(field["label"], "الاسم العام")
                self.assertEqual(field["global_ref"], GLOBAL_NAME)

        with use_context(worker):
            worker_record = APP.load_record("R3000001")
            worker_main = copy.deepcopy(worker_record["main"])
            worker_main[FLD_NAME] = "قيمة مختلفة"
            APP.save_record(
                {
                    "mode": "update",
                    "record_code": "R3000001",
                    "expected_updated_at": worker_record["updated_at"],
                    "main": worker_main,
                    "related": worker_record["related"],
                }
            )
        with use_context(self.primary):
            target_result = APP.manage_workspace_schema(
                {
                    "action": "duplicate",
                    "name": "التصميم الهدف",
                    "schema_id": self.primary.schema_id,
                    "template_schema_id": self.primary.schema_id,
                }
            )
        conflicts = APP.inspect_identity_profiles(
            {
                "record_code": "R3000001",
                "target_schema_id": target_result["schema_id"],
            }
        )["global_conflicts"]
        self.assertEqual(len(conflicts), 1)
        self.assertEqual(conflicts[0]["global_ref"], GLOBAL_NAME)
        self.assertEqual(len(conflicts[0]["sources"]), 2)

        found = APP.global_field_search(
            {"criteria": {GLOBAL_NAME: "طالب"}, "column_refs": [GLOBAL_NAME]}
        )
        self.assertEqual(sorted(item["total"] for item in found["results"]), [0, 0, 1])

        with use_context(self.primary):
            record = APP.load_record("R3000001")
            main = copy.deepcopy(record["main"])
            main[FLD_NAME] = "اسم موحّد"
            response = APP.save_record(
                {
                    "mode": "update",
                    "record_code": "R3000001",
                    "expected_updated_at": record["updated_at"],
                    "main": main,
                    "related": record["related"],
                    "propagate_global_values": True,
                    "propagate_global_refs": [GLOBAL_NAME],
                }
            )
        self.assertEqual(response["propagation"]["updated_profiles"], 1)
        with use_context(worker):
            self.assertEqual(APP.load_record("R3000001")["main"][FLD_NAME], "اسم موحّد")

        removed = APP.manage_global_definition(
            {
                "action": "delete",
                "kind": "field",
                "global_ref": GLOBAL_NAME,
                "expected_revision": 1,
            }
        )
        self.assertEqual(
            set(removed["detached_schema_ids"]),
            {self.primary.schema_id, worker.schema_id, target_result["schema_id"]},
        )
        for context in (self.primary, worker):
            with use_context(context):
                self.assertIsNone(self._field(APP.read_schema_file(), FLD_NAME)["global_ref"])
                self.assertEqual(APP.load_record("R3000001")["main"][FLD_NAME], "اسم موحّد")

    def test_global_category_configuration_propagates_and_local_edit_detaches(self) -> None:
        with use_context(self.primary):
            schema = APP.read_schema_file()
            category = schema["categories"][0]
            category["global_ref"] = GLOBAL_CATEGORY
            APP.save_schema(schema)
            definition = {
                key: copy.deepcopy(value)
                for key, value in category.items()
                if key not in {"id", "global_ref", "fields", "parent_category_id", "anchor_field_id"}
            }
        saved = APP.manage_global_definition(
            {
                "action": "save",
                "kind": "category",
                "global_ref": GLOBAL_CATEGORY,
                "expected_revision": 0,
                "definition": {**definition, "label": "الفئة العامة"},
            }
        )
        self.assertIn(self.primary.schema_id, saved["updated_schema_ids"])
        with use_context(self.primary):
            schema = APP.read_schema_file()
            self.assertEqual(schema["categories"][0]["label"], "الفئة العامة")
            schema["categories"][0]["label"] = "تعديل محلي"
            schema["categories"][0]["global_ref"] = None
            APP.save_schema(schema)
        APP.manage_global_definition(
            {
                "action": "save",
                "kind": "category",
                "global_ref": GLOBAL_CATEGORY,
                "expected_revision": 1,
                "definition": {**definition, "label": "تعديل عام لاحق"},
            }
        )
        with use_context(self.primary):
            self.assertEqual(APP.read_schema_file()["categories"][0]["label"], "تعديل محلي")

    def test_excel_update_preserves_blanks_and_requires_missing_id_confirmation(self) -> None:
        with use_context(self.primary):
            _name, content, _count = APP.create_filtered_export(
                {
                    "criteria": {"_allow_empty": True},
                    "field_ids": [FLD_NAME, FLD_STATUS],
                    "include_related": False,
                }
            )
            workbook = load_workbook(io.BytesIO(content))
            sheet = workbook[EXPORT_MAIN_SHEET]
            headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
            sheet.cell(3, headers.index(FLD_NAME) + 1).value = ""
            sheet.cell(3, headers.index(FLD_STATUS) + 1).value = "متوقف"
            output = io.BytesIO()
            workbook.save(output)
            workbook.close()
            encoded = base64.b64encode(output.getvalue()).decode("ascii")
            inspection = APP.inspect_import({"file_data": encoded})
            mapping = {
                str(column["column"]): column["suggested_target"]
                for column in inspection["columns"]
                if column["suggested_target"]
            }
            result = APP.commit_import(
                {
                    "file_data": encoded,
                    "sheet_name": inspection["sheet_name"],
                    "mapping": mapping,
                    "schema_revision": APP.read_schema_file()["revision"],
                    "duplicate_policy": "update",
                    "clear_blank_values": False,
                    "generate_missing_ids": False,
                }
            )
            self.assertEqual(result["updated"], 1)
            record = APP.load_record("R3000001")
            self.assertEqual(record["main"][FLD_NAME], "طالب أول")
            self.assertEqual(record["main"][FLD_STATUS], "متوقف")

            workbook = load_workbook(io.BytesIO(output.getvalue()))
            sheet = workbook[EXPORT_MAIN_SHEET]
            sheet.cell(3, headers.index("record_code") + 1).value = ""
            missing_id = io.BytesIO()
            workbook.save(missing_id)
            workbook.close()
            missing_encoded = base64.b64encode(missing_id.getvalue()).decode("ascii")
            rejected = APP.commit_import(
                {
                    "file_data": missing_encoded,
                    "sheet_name": inspection["sheet_name"],
                    "mapping": mapping,
                    "schema_revision": APP.read_schema_file()["revision"],
                    "duplicate_policy": "update",
                    "clear_blank_values": False,
                    "generate_missing_ids": False,
                }
            )
            self.assertEqual(rejected["rejected"], 1)
            self.assertIn("أكّد توليد", rejected["errors"][0]["message"])

    def test_batch_pdf_exports_independent_documents_and_audits_the_batch(self):
        from pypdf import PdfReader
        with use_context(self.primary):
            APP.save_record({"mode":"create", "record_code":"R3000002", "main":{FLD_NAME:"Second person", FLD_STATUS:"نشط"}, "related":{}})
            payload = {"type":"profile_pdf_batch", "record_codes":["r3000001", "R3000002", "R3000001"],
                "schema_ids_by_record":{code:[self.primary.schema_id] for code in ["R3000001","R3000002"]},
                "field_ids_by_schema":{self.primary.schema_id:[FLD_NAME]},
                "show_profile_image":True, "show_attachments":False,
                "destination":str(self.root/'batch.zip'), "notes":"Batch test"}
            with mock.patch.object(APP, "profile_pdf_bytes", wraps=APP.profile_pdf_bytes) as render:
                result = APP.save_export(payload)
            self.assertEqual(render.call_count, 2)
            profiles = [call.args[1] for call in render.call_args_list]
            self.assertIn("طالب أول", str(profiles[0]))
            self.assertNotIn("Second person", str(profiles[0]))
            self.assertIn("Second person", str(profiles[1]))
            self.assertNotIn("طالب أول", str(profiles[1]))
            self.assertEqual(result['row_count'],2)
            with zipfile.ZipFile(self.root/'batch.zip') as archive:
                self.assertEqual(archive.namelist(),['SchemaCraft-profile-R3000001.pdf','SchemaCraft-profile-R3000002.pdf'])
                for code in ['R3000001','R3000002']:
                    reader=PdfReader(io.BytesIO(archive.read(f'SchemaCraft-profile-{code}.pdf')))
                    # PDF metadata preserves the identity independently of extractor bidi behavior.
                    text=reader.metadata.title
                    self.assertIn(code,text)
                    other='R3000002' if code=='R3000001' else 'R3000001'
                    self.assertNotIn(other,text)
            self.assertEqual(result['history']['configuration']['record_codes'],payload['record_codes'])
            self.assertEqual(result['history']['configuration']['schema_ids_by_record'],payload['schema_ids_by_record'])
            self.assertTrue(result['history']['configuration']['show_profile_image'])
            with mock.patch.object(APP,'choose_export_destination',return_value=self.root/'chosen.zip') as choose:
                APP.choose_export_destination_response(payload)
                self.assertTrue(choose.call_args.args[0].endswith('.zip'))
                self.assertEqual(choose.call_args.args[1][0][1],'*.zip')

    def test_batch_pdf_failure_leaves_destination_untouched(self):
        with use_context(self.primary):
            target=self.root/'batch.zip'
            target.write_bytes(b'previous export')
            payload={"type":"profile_pdf_batch", "record_codes":["R3000001","R3999999"],
                "schema_ids":[self.primary.schema_id], "destination":str(target)}
            with self.assertRaisesRegex(APP.ApplicationError,'R3999999'):
                APP.save_export(payload)
            self.assertEqual(target.read_bytes(),b'previous export')
            with self.assertRaises(APP.ApplicationError):
                APP.create_profile_pdf_batch_export({**payload,"record_codes":[]})
            with self.assertRaises(APP.ApplicationError):
                APP.create_profile_pdf_batch_export({**payload,"record_codes":["../bad"]})
            with self.assertRaisesRegex(APP.ApplicationError,'R3000001'):
                APP.create_profile_pdf_batch_export({**payload,"schema_ids_by_record":{"R3000001":[]}})

    def test_portable_package_pdf_history_and_shortcuts_round_trip(self) -> None:
        attachment = self.primary.attachments_path / "documents" / "proof.txt"
        attachment.parent.mkdir(parents=True, exist_ok=True)
        attachment.write_text("portable attachment", encoding="utf-8")
        with use_context(self.primary):
            filename, package_bytes, count, schemas = APP.create_portable_export(
                {"schema_id": self.primary.schema_id}
            )
            self.assertTrue(filename.endswith(".zip"))
            self.assertEqual(count, 1)
            self.assertEqual(schemas, [self.primary.schema_id])
            inspected = inspect_portable_package(package_bytes)
            self.assertEqual(inspected["manifest"]["schema_id"], self.primary.schema_id)
            self.assertEqual(inspected["schema"]["app"]["title"], "نظام تجريبي")
            self.assertIn("attachments/documents/proof.txt", inspected["attachments"])
            damaged = io.BytesIO()
            with zipfile.ZipFile(io.BytesIO(package_bytes)) as source, zipfile.ZipFile(
                damaged, "w", compression=zipfile.ZIP_DEFLATED
            ) as destination:
                for name in source.namelist():
                    content = source.read(name)
                    destination.writestr(
                        name,
                        content + b"damage" if name == "database.xlsx" else content,
                    )
            with self.assertRaisesRegex(AdvancedFeatureError, "فشل التحقق"):
                inspect_portable_package(damaged.getvalue())

            pdf_name, pdf, pdf_count, pdf_schemas = APP.create_profile_pdf_export(
                {
                    "record_code": "R3000001",
                    "schema_ids": [self.primary.schema_id],
                    "field_ids_by_schema": {self.primary.schema_id: [FLD_NAME, FLD_STATUS]},
                }
            )
            self.assertTrue(pdf_name.endswith(".pdf"))
            self.assertTrue(pdf.startswith(b"%PDF-"))
            self.assertGreater(len(pdf), 10_000)
            self.assertEqual((pdf_count, pdf_schemas), (1, [self.primary.schema_id]))

            record = APP.load_record("R3000001")
            changed_main = copy.deepcopy(record["main"])
            changed_main[FLD_NAME] = "تعديل محلي"
            APP.save_record(
                {
                    "mode": "update",
                    "record_code": "R3000001",
                    "expected_updated_at": record["updated_at"],
                    "main": changed_main,
                    "related": record["related"],
                }
            )
            encoded_package = base64.b64encode(package_bytes).decode("ascii")
            attachment.unlink()
            preview = APP.inspect_portable_import({"file_data": encoded_package})
            self.assertEqual(preview["matching_count"], 1)
            newer = APP.commit_portable_import(
                {"file_data": encoded_package, "conflict_policy": "newer"}
            )
            self.assertEqual(newer["skipped"], 1)
            self.assertEqual(APP.load_record("R3000001")["main"][FLD_NAME], "تعديل محلي")
            self.assertEqual(attachment.read_text(encoding="utf-8"), "portable attachment")
            overwritten = APP.commit_portable_import(
                {"file_data": encoded_package, "conflict_policy": "overwrite"}
            )
            self.assertEqual(overwritten["updated"], 1)
            self.assertEqual(APP.load_record("R3000001")["main"][FLD_NAME], "طالب أول")

        entry = APP.EXPORT_HISTORY.append(
            {
                "type": "portable_zip",
                "schemas": schemas,
                "row_count": count,
                "filename": filename,
                "destination": str(self.root / filename),
                "checksum": "abc",
            }
        )
        history = APP.export_history_response({"limit": 10})
        self.assertEqual(history["total"], 1)
        self.assertEqual(history["entries"][0]["id"], entry["id"])

        settings = APP.save_workspace_settings(
            {"shortcuts": {"Ctrl+Alt+N": "new_record", "F6": "next_schema"}}
        )
        self.assertEqual(settings["workspace_settings"]["shortcuts"]["F6"], "next_schema")
        self.assertTrue((self.data / "workspace-settings.json").is_file())
        self.assertEqual(APP.read_workspace_settings()["shortcuts"]["Ctrl+Alt+N"], "new_record")
        preferences = APP.save_workspace_settings(
            {
                "shortcuts": settings["workspace_settings"]["shortcuts"],
                "home_entry_history_limit": 5,
                "home_builder_history_limit": 4,
                "home_search_history_limit": 3,
                "home_import_history_limit": 2,
                "home_export_history_limit": 1,
                "entry_history_limit": 4,
                "search_history_limit": 12,
                "import_history_limit": 2,
                "export_history_limit": 6,
                "show_explanations": True,
            }
        )["workspace_settings"]
        self.assertEqual(preferences["home_entry_history_limit"], 5)
        self.assertEqual(preferences["home_export_history_limit"], 1)
        self.assertEqual(preferences["entry_history_limit"], 4)
        self.assertTrue(preferences["show_explanations"])
        imported_history = APP.import_history_response({"limit": 10})
        self.assertGreaterEqual(imported_history["total"], 2)
        history_id = imported_history["entries"][0]["id"]
        changed = APP.update_operation_history_notes(
            {"kind": "import", "id": history_id, "notes": "مراجعة ناجحة"}
        )
        self.assertEqual(changed["entry"]["notes"], "مراجعة ناجحة")


if __name__ == "__main__":
    unittest.main()
