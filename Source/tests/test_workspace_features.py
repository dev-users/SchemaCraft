from __future__ import annotations

import base64
import copy
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from test_backend import (
    CAT_IDENTITY,
    FLD_CUSTOM_FILE,
    FLD_NAME,
    FLD_NUMBER,
    configured_schema,
    field,
)

import SchemaCraft as APP


class WorkspaceFeatureTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.originals = {
            "DATA_DIR": APP.DATA_DIR,
            "SCHEMA_PATH": APP.SCHEMA_PATH,
            "WORKBOOK_PATH": APP.WORKBOOK_PATH,
            "BACKUP_DIR": APP.BACKUP_DIR,
            "STARTUP_ERROR_LOG": APP.STARTUP_ERROR_LOG,
            "BUILDER_AUTH_PATH": APP.BUILDER_AUTH_PATH,
            "DEVELOPER_MODE": APP.DEVELOPER_MODE,
            "_BUILDER_UNLOCKED": APP._BUILDER_UNLOCKED,
            "_DATASET_SNAPSHOT": APP._DATASET_SNAPSHOT,
        }
        APP.DATA_DIR = self.root
        APP.SCHEMA_PATH = self.root / "schema.json"
        APP.WORKBOOK_PATH = self.root / "database.xlsx"
        APP.BACKUP_DIR = self.root / "backups"
        APP.STARTUP_ERROR_LOG = self.root / "startup-error.log"
        APP.BUILDER_AUTH_PATH = self.root / "builder-auth.json"
        APP.DEVELOPER_MODE = True
        APP._BUILDER_UNLOCKED = False
        APP.invalidate_dataset_cache()
        APP.ensure_storage()

    def tearDown(self) -> None:
        for name, value in self.originals.items():
            setattr(APP, name, value)
        self.temporary.cleanup()

    def save_simple_record(
        self,
        code: str,
        *,
        name: str = "سجل اختبار",
        extra: dict | None = None,
        related: dict | None = None,
    ) -> dict:
        main = {FLD_NAME: name}
        main.update(extra or {})
        return APP.save_record(
            {
                "mode": "create",
                "record_code": code,
                "main": main,
                "related": related or {},
            }
        )

    def test_id_search_paging_and_stale_editor_protection(self) -> None:
        APP.save_schema(configured_schema())
        first = self.save_simple_record("W0000001", name="الأول")
        self.save_simple_record("W0000002", name="الثاني")

        result = APP.search_records(
            {
                "_record_code": "w000000",
                "_limit": 1,
                "_offset": 0,
            }
        )
        self.assertEqual(result["total"], 2)
        self.assertEqual(result["matches"][0]["record_code"], "W0000001")
        self.assertTrue(result["truncated"])

        page = APP.search_records(
            {"_record_code": "W000000", "_limit": 1, "_offset": 1}
        )
        self.assertEqual(page["total"], 2)
        self.assertEqual(len(page["matches"]), 1)
        self.assertEqual(page["offset"], 1)

        archived = APP.archive_record(
            "W0000001", True, expected_updated_at=first["updated_at"]
        )
        self.assertTrue(archived["archived"])
        restored = APP.archive_record(
            "W0000001", False, expected_updated_at=archived["updated_at"]
        )
        self.assertFalse(restored["archived"])

        with self.assertRaisesRegex(APP.ApplicationError, "نافذة أخرى"):
            APP.save_record(
                {
                    "mode": "update",
                    "record_code": "W0000001",
                    "expected_updated_at": first["updated_at"],
                    "main": {FLD_NAME: "تعديل قديم"},
                    "related": {},
                }
            )

    def test_full_search_can_choose_result_columns_independently(self) -> None:
        APP.save_schema(configured_schema())
        self.save_simple_record(
            "C0000001",
            name="عنوان السجل",
            extra={FLD_NUMBER: "001234"},
        )

        selected = APP.search_records(
            {
                "_allow_empty": True,
                "_result_field_ids": [FLD_NUMBER],
            }
        )
        self.assertEqual(selected["matches"][0]["title"], "عنوان السجل")
        self.assertEqual(
            [item["field_id"] for item in selected["matches"][0]["details"]],
            [FLD_NUMBER],
        )

        no_optional_columns = APP.search_records(
            {"_allow_empty": True, "_result_field_ids": []}
        )
        self.assertEqual(no_optional_columns["matches"][0]["details"], [])
        self.assertEqual(
            no_optional_columns["matches"][0]["title"],
            "عنوان السجل",
        )

        with self.assertRaisesRegex(APP.ApplicationError, "غير متاح"):
            APP.search_records(
                {
                    "_allow_empty": True,
                    "_result_field_ids": [FLD_CUSTOM_FILE],
                }
            )

    def test_number_grouping_exact_zeroes_and_allowed_symbols_round_trip(self) -> None:
        schema = configured_schema()
        number_field = next(
            item
            for item in schema["categories"][0]["fields"]
            if item["id"] == FLD_NUMBER
        )
        number_field["number_behavior"] = {
            "storage_mode": "numeric",
            "format_thousands": True,
            "preserve_leading_zeros": False,
            "allowed_special_characters": "",
        }
        exact_field_id = "fld_200000000001"
        exact_field = field(
            exact_field_id,
            "رقم مرن",
            "number",
            searchable=True,
        )
        exact_field["number_behavior"] = {
            "storage_mode": "numeric",
            "format_thousands": False,
            "preserve_leading_zeros": True,
            "allowed_special_characters": "/-_$%*",
        }
        schema["categories"][0]["fields"].append(exact_field)
        saved_schema = APP.save_schema(schema)
        saved_exact = next(
            item
            for category in saved_schema["categories"]
            for item in category["fields"]
            if item["id"] == exact_field_id
        )
        self.assertEqual(saved_exact["number_behavior"]["storage_mode"], "text")

        exact_value = "00012-_$%*/"
        self.save_simple_record(
            "N0000001",
            extra={
                FLD_NUMBER: "5,623,423,423",
                exact_field_id: exact_value,
            },
        )
        loaded = APP.load_record("N0000001")
        self.assertEqual(loaded["main"][FLD_NUMBER], 5623423423)
        self.assertEqual(loaded["main"][exact_field_id], exact_value)

        workbook = load_workbook(APP.WORKBOOK_PATH, data_only=False)
        try:
            worksheet = workbook[APP.MAIN_SHEET]
            headers = APP.technical_headers(worksheet)
            grouped_cell = worksheet.cell(3, headers.index(FLD_NUMBER) + 1)
            exact_cell = worksheet.cell(3, headers.index(exact_field_id) + 1)
            self.assertEqual(grouped_cell.value, 5623423423)
            self.assertEqual(grouped_cell.number_format, "#,##0.################")
            self.assertEqual(exact_cell.value, exact_value)
            self.assertEqual(exact_cell.number_format, "@")
        finally:
            workbook.close()

        found = APP.search_records(
            {
                exact_field_id: exact_value,
                "_search_field_ids": [exact_field_id],
            }
        )
        self.assertEqual(found["total"], 1)
        with self.assertRaisesRegex(APP.ApplicationError, "غير مسموحة"):
            self.save_simple_record(
                "N0000002", extra={exact_field_id: "00012A"}
            )

    def test_unique_repeatable_checkbox_is_evaluated_on_the_same_card(self) -> None:
        category_id = "cat_200000000001"
        relationship_id = "fld_200000000002"
        current_id = "fld_200000000003"
        schema = {
            "schema_version": 2,
            "revision": 0,
            "app": {
                "title": "سجل العلاقات",
                "entity_singular": "شخص",
                "entity_plural": "الأشخاص",
                "direction": "rtl",
                "primary_color": "#315F8A",
            },
            "categories": [
                {
                    "id": CAT_IDENTITY,
                    "label": "الهوية",
                    "description": "",
                    "kind": "main",
                    "add_label": "",
                    "auto_start": False,
                    "anchor_field_id": None,
                    "fields": [field(FLD_NAME, "الاسم", "text", searchable=True)],
                },
                {
                    "id": category_id,
                    "label": "العلاقات",
                    "description": "",
                    "kind": "repeatable",
                    "add_label": "إضافة علاقة",
                    "auto_start": False,
                    "anchor_field_id": None,
                    "fields": [
                        field(
                            relationship_id,
                            "صلة القرابة",
                            "text",
                            searchable=True,
                        ),
                        {
                            **field(current_id, "الحالية", "checkbox", searchable=True),
                            "unique_checked_across_cards": True,
                        },
                    ],
                },
            ],
            "conditions": [],
        }
        APP.save_schema(schema)
        self.save_simple_record(
            "M0000001",
            name="مطابقة",
            related={
                category_id: [
                    {
                        "values": {relationship_id: "زوجة", current_id: True},
                    },
                    {
                        "values": {relationship_id: "أخت", current_id: False},
                    },
                ]
            },
        )
        self.save_simple_record(
            "M0000002",
            name="غير مطابقة",
            related={
                category_id: [
                    {
                        "values": {relationship_id: "أخت", current_id: True},
                    },
                    {
                        "values": {relationship_id: "زوجة", current_id: False},
                    },
                ]
            },
        )

        current_wife = APP.search_records(
            {
                relationship_id: "زوجة",
                current_id: True,
                "_search_field_ids": [relationship_id, current_id],
            }
        )
        self.assertEqual(
            [match["record_code"] for match in current_wife["matches"]],
            ["M0000001"],
        )

        _filename, exported, count = APP.create_filtered_export(
            {
                "criteria": {"_record_code": "M0000001"},
                "field_ids": [FLD_NAME, relationship_id, current_id],
                "include_related": True,
            }
        )
        self.assertEqual(count, 1)
        round_trip = load_workbook(io.BytesIO(exported))
        try:
            round_trip["السجلات المصدرة"]["A3"] = "M0000003"
            related_sheet = round_trip["العلاقات"]
            for row_index in range(3, related_sheet.max_row + 1):
                related_sheet.cell(row_index, 1).value = "M0000003"
            output = io.BytesIO()
            round_trip.save(output)
            encoded = base64.b64encode(output.getvalue()).decode("ascii")
        finally:
            round_trip.close()
        inspection = APP.inspect_import({"file_data": encoded})
        mapping = {
            str(column["column"]): column["suggested_target"]
            for column in inspection["columns"]
            if column["suggested_target"]
        }
        imported = APP.commit_import(
            {
                "file_data": encoded,
                "sheet_name": inspection["sheet_name"],
                "mapping": mapping,
                "duplicate_policy": "skip",
                "schema_revision": inspection["schema_revision"],
            }
        )
        self.assertEqual(imported["imported"], 1)
        imported_record = APP.load_record("M0000003")
        self.assertTrue(
            imported_record["related"][category_id][0]["values"][current_id]
        )
        self.assertEqual(
            imported_record["related"][category_id][0]["values"][relationship_id],
            "زوجة",
        )

    def test_filtered_export_and_header_mapped_import_round_trip(self) -> None:
        APP.save_schema(configured_schema())
        created = self.save_simple_record(
            "I0000001", name="للتصدير", extra={FLD_NUMBER: "1234"}
        )
        APP.archive_record("I0000001", True, expected_updated_at=created["updated_at"])
        filename, exported, count = APP.create_filtered_export(
            {
                "criteria": {"_record_code": "I0000001", "_include_archived": True},
                "field_ids": [FLD_NAME, FLD_NUMBER],
                "include_related": False,
            }
        )
        self.assertTrue(filename.endswith(".xlsx"))
        self.assertEqual(count, 1)

        editable = load_workbook(io.BytesIO(exported))
        try:
            sheet = editable["السجلات المصدرة"]
            sheet["A3"] = "I0000002"
            technical_headers = [cell.value for cell in sheet[1]]
            sheet.cell(3, technical_headers.index(FLD_NAME) + 1).value = "مستورد من Excel"
            output = io.BytesIO()
            editable.save(output)
            encoded = base64.b64encode(output.getvalue()).decode("ascii")
        finally:
            editable.close()

        inspection = APP.inspect_import({"file_data": encoded})
        mapping = {
            str(column["column"]): column["suggested_target"]
            for column in inspection["columns"]
            if column["suggested_target"]
        }
        imported = APP.commit_import(
            {
                "file_data": encoded,
                "sheet_name": inspection["sheet_name"],
                "mapping": mapping,
                "duplicate_policy": "skip",
                "schema_revision": inspection["schema_revision"],
            }
        )
        self.assertEqual(imported["imported"], 1)
        loaded = APP.load_record("I0000002")
        self.assertEqual(loaded["main"][FLD_NAME], "مستورد من Excel")

    def test_settings_save_is_revision_checked_and_does_not_rewrite_excel(self) -> None:
        saved = APP.save_schema(configured_schema())
        workbook_before = APP.WORKBOOK_PATH.read_bytes()
        app_settings = copy.deepcopy(saved["app"])
        app_settings.update(
            {
                "title": "مساحة العمل المطورة",
                "language": "ar",
                "startup_page": "search",
                "search_page_size": 100,
                "draft_autosave": False,
                "shortcuts": ["new_record", "search"],
            }
        )
        updated = APP.save_app_settings(
            {"revision": saved["revision"], "app": app_settings}
        )
        self.assertEqual(updated["revision"], saved["revision"] + 1)
        self.assertEqual(updated["app"]["startup_page"], "search")
        self.assertEqual(APP.WORKBOOK_PATH.read_bytes(), workbook_before)
        with self.assertRaisesRegex(APP.ApplicationError, "نافذة أخرى"):
            APP.save_app_settings(
                {"revision": saved["revision"], "app": app_settings}
            )


if __name__ == "__main__":
    unittest.main()
