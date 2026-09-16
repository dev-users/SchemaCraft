"""Portable, metadata-driven Excel data-entry application.

The distributed application starts with an empty schema.  Users define
categories, fields, search behaviour, visibility rules, and attachment naming
from Builder mode.  The schema uses stable IDs; Excel labels and sheet names are
presentation details and may change without breaking saved data.
"""

from __future__ import annotations

import base64
import binascii
import copy
import ctypes
import hashlib
import io
import json
import logging
import mimetypes
import os
import re
import secrets
import shutil
import string
import subprocess
import sys
import tempfile
import threading
import time
import traceback
import unicodedata
import uuid
import zipfile
from collections.abc import Iterable
from contextlib import nullcontext
from datetime import date, datetime
from http.client import HTTPConnection
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from logging.handlers import RotatingFileHandler
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import parse_qs, quote, unquote, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ``SchemaCraft.py`` is also loaded through importlib by the portable test
# launcher.  Keep sibling modules discoverable in that execution mode as well
# as when the application is started normally or bundled.
_SOURCE_MODULE_DIR = str(Path(__file__).resolve().parent)
if _SOURCE_MODULE_DIR not in sys.path:
    sys.path.insert(0, _SOURCE_MODULE_DIR)
_VENDOR_MODULE_DIR = str(Path(_SOURCE_MODULE_DIR) / "vendor")
if Path(_VENDOR_MODULE_DIR).is_dir() and _VENDOR_MODULE_DIR not in sys.path:
    sys.path.insert(0, _VENDOR_MODULE_DIR)

from schemacraft_advanced import (
    GLOBAL_REFERENCE,
    AdvancedFeatureError,
    AuditUserStore,
    ExportHistoryStore,
    GlobalDefinitionStore,
    ImportHistoryStore,
    SearchHistoryStore,
    choose_directory,
    choose_export_destination,
    inspect_portable_package,
    normalize_export_destination,
    portable_package_bytes,
    profile_pdf_bytes,
    save_export_bytes,
)
from schemacraft_io import (
    WorkbookExchangeError,
    export_workbook_bytes,
    inspect_import_workbook,
    parse_import_rows,
    inspect_import_sheets,
    parse_import_sheets,
)
from schemacraft_security import (
    AuthenticationThrottle,
    BrowserSessionManager,
    bcrypt_hash_password,
    bcrypt_verify_password,
)
from schemacraft_workspace import (
    SchemaContext,
    WorkspaceError,
    WorkspaceManager,
    active_context,
    use_context,
)

HOST = "127.0.0.1"
APP_PORT_MIN = 51000
APP_PORT_SPAN = 10000
SCHEMA_VERSION = 2
SUPPORTED_SCHEMA_VERSIONS = {1, 2}
MAIN_SHEET = "السجلات"
META_SHEET = "_meta"
TECHNICAL_HEADER_ROW = 1
VISIBLE_HEADER_ROW = 2
FIRST_DATA_ROW = 3
BROWSER_ACTIVE_SECONDS = 90.0

MAIN_REQUIRED_INTERNAL_HEADERS = (
    "_record_id",
    "record_code",
    "created_at",
    "updated_at",
)
MAIN_INTERNAL_HEADERS = (
    *MAIN_REQUIRED_INTERNAL_HEADERS,
    "_archived",
    "_archived_at",
)
RELATED_INTERNAL_HEADERS = (
    "_child_id",
    "_record_id",
    "record_code",
    "minor_id",
    "created_at",
    "updated_at",
)
RELATED_LINK_HEADER = "_linked_record_code"
RELATED_PARENT_HEADER = "_parent_child_id"
RELATED_PERSON_MODE_SOURCE_PREFIX = "related_person_mode:"

FIELD_TYPES = {
    "spacer",
    "text",
    "textarea",
    "number",
    "select",
    "checkbox",
    "checkbox_group",
    "yes_no",
    "date_gregorian",
    "date_hijri",
    "date_persian",
    "file",
    "system_record_code",
    "system_created_at",
    "system_updated_at",
}
AUDIT_FIELD_TYPES = {"user_name"}
SYSTEM_FIELD_TYPES = {
    "system_record_code",
    "system_created_at",
    "system_updated_at",
}
LEGACY_COMPLETION_FIELD_TYPES = {"system_completed_at", "system_completion"}
CATEGORY_KINDS = {"main", "repeatable"}
FIELD_WIDTHS = {"1", "2", "3", "4", "5", "full"}
LEGACY_FIELD_WIDTHS = {"normal": "1", "wide": "4", "long": "4"}
SEARCH_MATCHES = {"contains", "exact"}
CONDITION_OPERATORS = {
    "equals",
    "not_equals",
    "contains",
    "not_contains",
    "greater_than",
    "greater_or_equal",
    "less_than",
    "less_or_equal",
    "before",
    "after",
    "on_or_before",
    "on_or_after",
    "empty",
    "not_empty",
}

CONDITION_OPERATORS_BY_FIELD_TYPE = {
    "text": {
        "equals",
        "not_equals",
        "contains",
        "not_contains",
        "empty",
        "not_empty",
    },
    "textarea": {
        "equals",
        "not_equals",
        "contains",
        "not_contains",
        "empty",
        "not_empty",
    },
    "user_name": {
        "equals",
        "not_equals",
        "contains",
        "not_contains",
        "empty",
        "not_empty",
    },
    "number": {
        "equals",
        "not_equals",
        "greater_than",
        "greater_or_equal",
        "less_than",
        "less_or_equal",
        "empty",
        "not_empty",
    },
    "select": {
        "equals",
        "not_equals",
        "empty",
        "not_empty",
    },
    "yes_no": {
        "equals",
        "not_equals",
        "empty",
        "not_empty",
    },
    "checkbox": {
        "equals",
        "not_equals",
    },
    "checkbox_group": {
        "contains",
        "not_contains",
        "empty",
        "not_empty",
    },
    "date_gregorian": {
        "equals",
        "not_equals",
        "before",
        "after",
        "on_or_before",
        "on_or_after",
        "empty",
        "not_empty",
    },
    "date_hijri": {
        "equals",
        "not_equals",
        "before",
        "after",
        "on_or_before",
        "on_or_after",
        "empty",
        "not_empty",
    },
    "date_persian": {
        "equals",
        "not_equals",
        "before",
        "after",
        "on_or_before",
        "on_or_after",
        "empty",
        "not_empty",
    },
    "file": {
        "empty",
        "not_empty",
    },
}
DEFINITION_ID_PATTERN = re.compile(
    r"^(cat|fld|cond|grp|opt|mark)_[a-f0-9]{12}$"
)
PERSON_CODE_PATTERN = re.compile(r"^[A-Z][A-Z0-9]{7}$")
PERSON_CODE_FIRST_CHARACTERS = string.ascii_uppercase
PERSON_CODE_OTHER_CHARACTERS = string.ascii_uppercase + string.digits
INTERNAL_ID_PATTERN = re.compile(r"^[a-f0-9]{32}$")

MAX_ATTACHMENT_BYTES = 100 * 1024 * 1024
MAX_REQUEST_BYTES = 140 * 1024 * 1024
MAX_BACKGROUND_IMAGE_BYTES = 15 * 1024 * 1024
MAX_SEARCH_RESULTS = 100
MAX_SEARCH_PAGE_SIZE = 250
EXCEL_DATE_FORMAT = "yyyy-mm-dd"
WINDOWS_ALREADY_EXISTS = 183

APP_PAGES = {"home", "entry", "search", "import", "export", "settings", "builder"}
APP_SHORTCUTS = {
    "context_new",
    "context_save",
    "context_edit",
    "context_delete",
    "context_focus_search",
    "context_archive",
    "next_workspace_tab",
    "new_record",
    "save_record",
    "open_by_id",
    "focus_entry_search",
    "open_home",
    "open_search",
    "open_entry",
    "open_import",
    "open_export",
    "open_builder",
    "save_builder",
    "archive_record",
    "previous_schema",
    "next_schema",
    "open_settings",
    "close_dialog",
    "close_application",
    "enter_admin",
    "select_user",
    "exit_admin",
}
DEFAULT_SHORTCUT_BINDINGS = {
    "Ctrl+N": "new_record",
    "Ctrl+S": "save_record",
    "Ctrl+O": "open_by_id",
    "Ctrl+F": "focus_entry_search",
    "Ctrl+Shift+F": "open_search",
}
DEFAULT_WORKSPACE_SHORTCUT_BINDINGS = {
    "Ctrl+N": "context_new",
    "Ctrl+S": "context_save",
    "Ctrl+E": "context_edit",
    "Ctrl+Delete": "context_delete",
    "Ctrl+F": "context_focus_search",
    "Ctrl+Shift+F": "open_search",
    "Ctrl+Q": "close_application",
    "Ctrl+Alt+A": "enter_admin",
    "Ctrl+Alt+U": "select_user",
}
DEFAULT_APP_SHORTCUTS = list(DEFAULT_SHORTCUT_BINDINGS.values())
NUMBER_STORAGE_MODES = {"numeric", "text"}
NUMBER_SPECIAL_CHARACTERS = "/-_$%*"

FIXED_BACKGROUND_COLOR = "#F4F7FB"
FIXED_SURFACE_COLOR = "#FFFFFF"
INLINE_ATTACHMENT_EXTENSIONS = {
    ".avif",
    ".bmp",
    ".gif",
    ".jpeg",
    ".jpg",
    ".pdf",
    ".png",
    ".webp",
}

ARABIC_DIACRITICS_PATTERN = re.compile(
    r"[\u0610-\u061A\u064B-\u065F\u0670\u06D6-\u06ED]"
)
ARABIC_SEARCH_TRANSLATION = {
    ord("أ"): "ا",
    ord("إ"): "ا",
    ord("آ"): "ا",
    ord("ٱ"): "ا",
    ord("ى"): "ي",
    ord("ـ"): "",
}
INVALID_WINDOWS_FILENAME_PATTERN = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
INVALID_EXCEL_SHEET_PATTERN = re.compile(r"[\[\]:*?/\\]")
HEX_COLOR_PATTERN = re.compile(r"^#[0-9A-Fa-f]{6}$")

BASE_DIR = (
    Path(sys.executable).resolve().parent
    if getattr(sys, "frozen", False)
    else Path(__file__).resolve().parent
)
def application_port() -> int:
    """Use the same local port whenever this application folder starts."""
    folder_hash = hashlib.sha256(
        str(BASE_DIR).casefold().encode("utf-8")
    ).digest()

    return APP_PORT_MIN + (
        int.from_bytes(folder_hash[:2], "big") % APP_PORT_SPAN
    )
APP_DIR = BASE_DIR / "app"
UI_TEXT_PATH = APP_DIR / "ui_text.json"
DATA_DIR = BASE_DIR / "data"
LOG_DIR = DATA_DIR / "logs"
LOG_FILE = LOG_DIR / "application.log"
SCHEMA_PATH = DATA_DIR / "schema.json"
WORKBOOK_PATH = DATA_DIR / "database.xlsx"
WORKBOOK_LOCK = threading.RLock()
WORKSPACE_MANAGER: WorkspaceManager | None = None
GLOBAL_DEFINITIONS: GlobalDefinitionStore | None = None
EXPORT_HISTORY: ExportHistoryStore | None = None
IMPORT_HISTORY: ImportHistoryStore | None = None
SEARCH_HISTORY: SearchHistoryStore | None = None
AUDIT_USERS: AuditUserStore | None = None
WORKSPACE_SETTINGS_PATH = DATA_DIR / "workspace-settings.json"
IMPORT_ARCHIVE_DIR = DATA_DIR / "imported-source-files"
PDF_FONT_PATH = BASE_DIR / "assets" / "fonts" / "DejaVuSans.ttf"

_UI_TEXT_CACHE_LOCK = threading.RLock()
_UI_TEXT_ASSET_CACHE: dict[tuple[str, int, int, int], bytes] = {}


def read_ui_text_catalog() -> dict[str, str]:
    """Read operator-editable application wording from ``app/ui_text.json``."""

    try:
        payload = json.loads(UI_TEXT_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    if not isinstance(payload, dict):
        return {}
    return {
        source: replacement
        for source, replacement in payload.items()
        if isinstance(source, str)
        and source
        and not source.startswith("_")
        and isinstance(replacement, str)
    }


def localized_ui_asset(path: Path) -> bytes:
    """Return a text asset with the current UI catalog applied at runtime."""

    asset_stat = path.stat()
    try:
        catalog_version = UI_TEXT_PATH.stat().st_mtime_ns
    except OSError:
        catalog_version = 0
    cache_key = (
        str(path.resolve()),
        asset_stat.st_mtime_ns,
        asset_stat.st_size,
        catalog_version,
    )
    with _UI_TEXT_CACHE_LOCK:
        cached = _UI_TEXT_ASSET_CACHE.get(cache_key)
        if cached is not None:
            return cached
    content = path.read_text(encoding="utf-8")
    catalog = read_ui_text_catalog()
    placeholders: list[tuple[str, str]] = []
    for index, source in enumerate(sorted(catalog, key=len, reverse=True)):
        token = f"__SC_RUNTIME_UI_TEXT_{index:04d}__"
        if token in content:
            continue
        content = content.replace(source, token)
        placeholders.append((token, catalog[source]))
    for token, replacement in placeholders:
        content = content.replace(token, replacement)
    result = content.encode("utf-8")
    with _UI_TEXT_CACHE_LOCK:
        if len(_UI_TEXT_ASSET_CACHE) > 24:
            _UI_TEXT_ASSET_CACHE.clear()
        _UI_TEXT_ASSET_CACHE[cache_key] = result
    return result


def _schema_path() -> Path:
    context = active_context()
    return context.schema_path if context is not None else SCHEMA_PATH


def _workbook_path() -> Path:
    context = active_context()
    return context.workbook_path if context is not None else WORKBOOK_PATH


def current_schema_id() -> str:
    context = active_context()
    return context.schema_id if context is not None else "legacy"


class DatasetSnapshot:
    """One coherent in-memory view of the Excel workbook and its indexes."""

    __slots__ = (
        "record_positions",
        "records",
        "records_by_code",
        "records_by_id",
        "schema_signature",
        "search_main",
        "search_related",
        "unique_indexes",
        "workbook_path",
        "workbook_signature",
    )

    def __init__(
        self,
        *,
        workbook_path: str,
        workbook_signature: tuple[int, int] | None,
        schema_signature: str,
        records: tuple[dict[str, Any], ...],
        records_by_code: dict[str, dict[str, Any]],
        records_by_id: dict[str, dict[str, Any]],
        record_positions: dict[str, int],
        unique_indexes: dict[str, dict[str, frozenset[str]]],
        search_main: dict[str, dict[str, Any]],
        search_related: dict[
            str, dict[str, tuple[dict[str, Any], ...]]
        ],
    ) -> None:
        self.workbook_path = workbook_path
        self.workbook_signature = workbook_signature
        self.schema_signature = schema_signature
        self.records = records
        self.records_by_code = records_by_code
        self.records_by_id = records_by_id
        self.record_positions = record_positions
        self.unique_indexes = unique_indexes
        self.search_main = search_main
        self.search_related = search_related


_DATASET_SNAPSHOT: DatasetSnapshot | None = None
_DATASET_SNAPSHOTS: dict[str, DatasetSnapshot] = {}
_WORKBOOK_SYNC_PENDING: dict[str, dict[str, Any]] = {}
_WORKBOOK_SYNC_CONDITION = threading.Condition()
_WORKBOOK_SYNC_WORKER: threading.Thread | None = None
_WORKBOOK_SYNC_DEBOUNCE_SECONDS = 0.75
DEVELOPER_ACCESS_PATH = BASE_DIR / "developer-access.key"
BACKUP_DIR = BASE_DIR / "backups"
BUILDER_AUTH_PATH = BASE_DIR / "builder-auth.json"
STARTUP_ERROR_LOG = BASE_DIR / "startup-error.log"
DEVELOPER_MODE = (
    "--builder" in sys.argv and DEVELOPER_ACCESS_PATH.is_file()
)
BUILDER_PASSWORD_ITERATIONS = 310_000
_BUILDER_SESSION_LOCK = threading.RLock()
_BUILDER_UNLOCKED = DEVELOPER_MODE
_BUILDER_AUTH_THROTTLE = AuthenticationThrottle()
_WINDOWS_MUTEX_HANDLE: int | None = None
_APPLICATION_BROWSER_PROCESS: subprocess.Popen[Any] | None = None
LOGGER = logging.getLogger("SchemaCraft")

WINDOWS_HIDDEN_ATTRIBUTE = 0x02
WINDOWS_INVALID_FILE_ATTRIBUTES = 0xFFFFFFFF


def hide_windows_path(path: Path) -> bool:
    """Add the Windows hidden attribute while preserving every other flag."""

    if os.name != "nt" or not path.exists():
        return False
    kernel32 = ctypes.windll.kernel32
    kernel32.GetFileAttributesW.argtypes = [ctypes.c_wchar_p]
    kernel32.GetFileAttributesW.restype = ctypes.c_uint32
    kernel32.SetFileAttributesW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32]
    kernel32.SetFileAttributesW.restype = ctypes.c_int
    attributes = kernel32.GetFileAttributesW(str(path))
    if attributes == WINDOWS_INVALID_FILE_ATTRIBUTES:
        return False
    return bool(
        kernel32.SetFileAttributesW(
            str(path),
            attributes | WINDOWS_HIDDEN_ATTRIBUTE,
        )
    )


def hide_packaged_support_paths() -> None:
    """Keep a frozen user package visually reduced to SchemaCraft.exe."""

    if not getattr(sys, "frozen", False):
        return
    for path in (
        APP_DIR,
        BASE_DIR / "assets",
        DATA_DIR,
        BACKUP_DIR,
        BUILDER_AUTH_PATH,
        DEVELOPER_ACCESS_PATH,
        STARTUP_ERROR_LOG,
    ):
        hide_windows_path(path)

class ApplicationError(Exception):
    """An error safe to display to the user."""
def configure_logging() -> None:
    LOG_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )
    hide_packaged_support_paths()

    LOGGER.setLevel(logging.INFO)
    LOGGER.propagate = False

    if LOGGER.handlers:
        return

    handler = RotatingFileHandler(
        LOG_FILE,
        maxBytes=5 * 1024 * 1024,
        backupCount=5,
        encoding="utf-8",
    )

    handler.setFormatter(
        logging.Formatter(
            "%(asctime)s | "
            "%(levelname)s | "
            "%(threadName)s | "
            "%(message)s"
        )
    )

    LOGGER.addHandler(handler)


def install_exception_logging() -> None:
    original_sys_hook = sys.excepthook
    original_thread_hook = threading.excepthook

    def application_exception_hook(
        exception_type,
        exception,
        traceback_object,
    ) -> None:
        if issubclass(exception_type, KeyboardInterrupt):
            original_sys_hook(
                exception_type,
                exception,
                traceback_object,
            )
            return

        LOGGER.critical(
            "Unhandled application exception",
            exc_info=(
                exception_type,
                exception,
                traceback_object,
            ),
        )

        original_sys_hook(
            exception_type,
            exception,
            traceback_object,
        )

    def thread_exception_hook(args) -> None:
        thread_name = (
            args.thread.name
            if args.thread is not None
            else "unknown"
        )

        LOGGER.critical(
            "Unhandled exception in thread %s",
            thread_name,
            exc_info=(
                args.exc_type,
                args.exc_value,
                args.exc_traceback,
            ),
        )

        original_thread_hook(args)

    sys.excepthook = application_exception_hook
    threading.excepthook = thread_exception_hook

def _read_builder_auth() -> dict[str, Any] | None:
    if not BUILDER_AUTH_PATH.is_file():
        return None
    try:
        payload = json.loads(BUILDER_AUTH_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ApplicationError(f"تعذّر قراءة إعداد كلمة مرور المصمّم: {exc}") from exc
    if not isinstance(payload, dict):
        raise ApplicationError("إعداد كلمة مرور المصمّم غير صالح.")
    algorithm = clean_text(payload.get("algorithm")).casefold()
    if algorithm == "bcrypt":
        encoded_hash = str(payload.get("hash") or "")
        if not encoded_hash.startswith(("$2a$", "$2b$", "$2y$")) or len(encoded_hash) > 100:
            raise ApplicationError("إعداد كلمة مرور المصمّم تالف.")
        return {"algorithm": "bcrypt", "hash": encoded_hash}

    # Read legacy PBKDF2 files so existing installations can upgrade without
    # resetting the administrator password. A successful verification rewrites
    # the file using bcrypt.
    try:
        iterations = int(payload.get("iterations", 0))
        salt = base64.b64decode(str(payload.get("salt", "")), validate=True)
        digest = base64.b64decode(str(payload.get("hash", "")), validate=True)
    except (ValueError, binascii.Error) as exc:
        raise ApplicationError("إعداد كلمة مرور المصمّم تالف.") from exc
    if (
        algorithm not in {"", "pbkdf2-sha256"}
        or iterations < 100_000
        or len(salt) < 16
        or len(digest) != 32
    ):
        raise ApplicationError("إعداد كلمة مرور المصمّم تالف.")
    return {
        "algorithm": "pbkdf2-sha256",
        "iterations": iterations,
        "salt": salt,
        "hash": digest,
    }


def builder_password_configured() -> bool:
    return BUILDER_AUTH_PATH.is_file()


def _password_digest(password: str, salt: bytes, iterations: int) -> bytes:
    return hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt,
        iterations,
        dklen=32,
    )


def verify_builder_password(password: Any) -> bool:
    text = str(password or "")
    auth = _read_builder_auth()
    if auth is None:
        return False
    _BUILDER_AUTH_THROTTLE.wait_before_attempt()
    if auth["algorithm"] == "bcrypt":
        verified = bcrypt_verify_password(text, auth["hash"])
    else:
        candidate = _password_digest(text, auth["salt"], auth["iterations"])
        verified = secrets.compare_digest(candidate, auth["hash"])
    _BUILDER_AUTH_THROTTLE.record(verified)
    if verified and auth["algorithm"] != "bcrypt":
        set_builder_password(text)
        LOGGER.info("Administrator password hash upgraded to bcrypt.")
    return verified


def validate_new_builder_password(password: Any) -> str:
    text = str(password or "")
    if len(text) < 8:
        raise ApplicationError("يجب أن تتكون كلمة مرور المصمّم من 8 أحرف على الأقل.")
    if len(text.encode("utf-8")) > 72:
        raise ApplicationError("كلمة مرور المصمّم طويلة جدًا.")
    return text


def set_builder_password(password: Any) -> None:
    text = validate_new_builder_password(password)
    payload = {
        "version": 2,
        "algorithm": "bcrypt",
        "hash": bcrypt_hash_password(text),
        "updated_at": now_iso(),
    }
    atomic_write_json(BUILDER_AUTH_PATH, payload)
    hide_packaged_support_paths()


def builder_is_unlocked() -> bool:
    if DEVELOPER_MODE:
        return True

    with _BUILDER_SESSION_LOCK:
        return _BUILDER_UNLOCKED


def unlock_builder(
    password: Any = None,
    *,
    initialize: bool = False,
) -> dict[str, Any]:
    global _BUILDER_UNLOCKED

    started_at = time.perf_counter()
    configured = builder_password_configured()

    try:
        if initialize:
            if configured:
                raise ApplicationError(
                    "تم إعداد كلمة مرور المصمّم مسبقًا."
                )
            set_builder_password(password)
        elif not configured:
            raise ApplicationError(
                "أنشئ كلمة مرور للمصمّم أولًا."
            )
        elif not verify_builder_password(password):
            raise ApplicationError(
                "كلمة مرور المصمّم غير صحيحة."
            )

        with _BUILDER_SESSION_LOCK:
            _BUILDER_UNLOCKED = True

        return builder_access_response()
    finally:
        LOGGER.info(
            "Builder authentication completed in %.3f seconds",
            time.perf_counter() - started_at,
        )


def lock_builder() -> dict[str, Any]:
    global _BUILDER_UNLOCKED

    with _BUILDER_SESSION_LOCK:
        _BUILDER_UNLOCKED = False
    return builder_access_response()


def change_builder_password(current_password: Any, new_password: Any) -> dict[str, Any]:
    if (
        builder_password_configured()
        and not builder_is_unlocked()
        and not verify_builder_password(current_password)
    ):
        raise ApplicationError("كلمة المرور الحالية غير صحيحة.")
    set_builder_password(new_password)
    return unlock_builder(new_password)


def require_builder_access() -> None:
    if not builder_is_unlocked():
        raise ApplicationError("المصمّم مقفل. أدخل كلمة المرور أولًا.")


def builder_access_response() -> dict[str, Any]:
    return {
        "configured": builder_password_configured(),
        "unlocked": builder_is_unlocked(),
        "session_minutes": 0,
        "session_timeout_enabled": False,
        "developer_key_available": DEVELOPER_ACCESS_PATH.is_file(),
    }


def default_schema() -> dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION,
        "revision": 0,
        "app": {
            "title": "نظام إدخال البيانات",
            "entity_singular": "سجل",
            "entity_plural": "السجلات",
            "direction": "rtl",
            "language": "ar",
            "primary_color": "#1F5F95",
            "background_color": FIXED_BACKGROUND_COLOR,
            "surface_color": FIXED_SURFACE_COLOR,
            "startup_page": "home",
            "search_page_size": 50,
            "show_entry_search": True,
            "draft_autosave": True,
            "shortcuts": copy.deepcopy(DEFAULT_SHORTCUT_BINDINGS),
        },
        "categories": [],
        "conditions": [],
    }


def new_definition_id(prefix: str) -> str:
    return f"{prefix}_{secrets.token_hex(6)}"


def new_internal_id() -> str:
    return uuid.uuid4().hex


def now_iso() -> str:
    # Sub-second precision is part of the optimistic-concurrency token used by
    # record editors.  Second-only timestamps can miss two saves performed in
    # quick succession and allow a stale window to overwrite newer data.
    return datetime.now().astimezone().isoformat(timespec="microseconds")


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_search_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(json_value(value)))
    text = ARABIC_DIACRITICS_PATTERN.sub("", text)
    text = text.translate(ARABIC_SEARCH_TRANSLATION)
    return " ".join(text.casefold().split())


def json_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time().isoformat() == "00:00:00":
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def validate_person_code(value: Any) -> str:
    code = clean_text(value)
    if not PERSON_CODE_PATTERN.fullmatch(code):
        raise ApplicationError("معرّف السجل غير صالح.")
    return code


def generate_person_code() -> str:
    return secrets.choice(PERSON_CODE_FIRST_CHARACTERS) + "".join(
        secrets.choice(PERSON_CODE_OTHER_CHARACTERS) for _ in range(7)
    )


def validate_internal_id(value: Any, label: str) -> str:
    identifier = clean_text(value)
    if not INTERNAL_ID_PATTERN.fullmatch(identifier):
        raise ApplicationError(f"{label} الداخلي غير صالح.")
    return identifier


def is_persian_leap_year(year: int) -> bool:
    return year % 33 in {1, 5, 9, 13, 17, 22, 26, 30}


def validate_persian_date(value: str, label: str) -> str:
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        raise ApplicationError(f'قيمة الحقل "{label}" ليست تاريخًا شمسيًا صالحًا.')
    year, month, day = (int(part) for part in value.split("-"))
    if not 1 <= month <= 12:
        raise ApplicationError(f'قيمة الحقل "{label}" ليست تاريخًا شمسيًا صالحًا.')
    maximum_day = (
        31
        if month <= 6
        else 30
        if month <= 11
        else 30
        if is_persian_leap_year(year)
        else 29
    )
    if not 1 <= day <= maximum_day:
        raise ApplicationError(f'قيمة الحقل "{label}" ليست تاريخًا شمسيًا صالحًا.')
    return value


def is_hijri_leap_year(year: int) -> bool:
    """Return the arithmetic Hijri leap-year result for validation.

    Actual observed month starts can differ by one day.  The application stores
    the user's chosen Hijri date as entered; this validation only prevents
    impossible selector combinations.
    """

    return year % 30 in {2, 5, 7, 10, 13, 16, 18, 21, 24, 26, 29}


def validate_hijri_date(value: str, label: str) -> str:
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        raise ApplicationError(f'قيمة الحقل "{label}" ليست تاريخًا هجريًا صالحًا.')
    year, month, day = (int(part) for part in value.split("-"))
    if not 1 <= month <= 12:
        raise ApplicationError(f'قيمة الحقل "{label}" ليست تاريخًا هجريًا صالحًا.')
    maximum_day = (
        30
        if month % 2 == 1
        else 30
        if month == 12 and is_hijri_leap_year(year)
        else 29
    )
    if not 1 <= day <= maximum_day:
        raise ApplicationError(f'قيمة الحقل "{label}" ليست تاريخًا هجريًا صالحًا.')
    return value


def normalize_field_value(value: Any, field: dict[str, Any]) -> Any:
    field_type = field["type"]

    if field_type == "file":
        return value
    if field_type == "checkbox":
        if isinstance(value, bool):
            return value
        text = clean_text(value).casefold()
        if text in {"", "false", "0", "لا", "no", "off"}:
            return False
        if text in {"true", "1", "نعم", "yes", "on"}:
            return True
        raise ApplicationError(f'قيمة الحقل "{field["label"]}" ليست اختيارًا صحيحًا.')
    if field_type == "checkbox_group":
        result: list[str] = []
        for item in parse_checkbox_group_value(value):
            label = normalize_choice_value(item, field)
            if label not in result:
                result.append(label)
        return result

    if value is None:
        return ""
    if isinstance(value, str):
        value = value.strip()
    if value == "":
        return ""

    if field_type == "date_gregorian":
        if isinstance(value, datetime):
            normalized: Any = value.date()
        elif isinstance(value, date):
            normalized = value
        else:
            try:
                normalized = date.fromisoformat(str(value))
            except ValueError as exc:
                raise ApplicationError(
                    f'قيمة الحقل "{field["label"]}" ليست تاريخًا ميلاديًا صالحًا.'
                ) from exc
        _validate_simple_constraints(normalized, field)
        return normalized

    if field_type == "date_persian":
        normalized = validate_persian_date(str(value), field["label"])
        _validate_simple_constraints(normalized, field)
        return normalized

    if field_type == "date_hijri":
        normalized = validate_hijri_date(str(value), field["label"])
        _validate_simple_constraints(normalized, field)
        return normalized

    if field_type == "number":
        if isinstance(value, bool):
            raise ApplicationError(f'قيمة الحقل "{field["label"]}" ليست رقمًا صالحًا.')
        if number_is_text(field):
            normalized_text = normalize_number_text(value, field)
            _validate_simple_constraints(normalized_text, field)
            return normalized_text
        try:
            number = float(
                clean_text(value).replace(",", "").replace("٬", "")
            )
        except (TypeError, ValueError) as exc:
            raise ApplicationError(
                f'قيمة الحقل "{field["label"]}" ليست رقمًا صالحًا.'
            ) from exc
        normalized = int(number) if number.is_integer() else number
        _validate_simple_constraints(normalized, field)
        return normalized

    if field_type == "yes_no":
        normalized = normalize_choice_value(value, field)
        _validate_simple_constraints(normalized, field)
        return normalized

    if field_type == "select":
        normalized = normalize_choice_value(value, field)
        _validate_simple_constraints(normalized, field)
        return normalized

    normalized = str(value)
    _validate_simple_constraints(normalized, field)
    return normalized


def _validate_simple_constraints(value: Any, field: dict[str, Any]) -> None:
    validation = field.get("validation", {})
    if not validation or value in {"", None}:
        return
    label = field["label"]
    field_type = field["type"]
    if field_type in {"text", "textarea", "user_name"}:
        length = len(str(value))
        minimum = validation.get("min_length")
        maximum = validation.get("max_length")
        if minimum is not None and length < minimum:
            raise ApplicationError(f'الحقل "{label}" يجب ألا يقل عن {minimum} أحرف.')
        if maximum is not None and length > maximum:
            raise ApplicationError(f'الحقل "{label}" يجب ألا يزيد على {maximum} أحرف.')
        pattern = validation.get("pattern")
        if pattern and re.fullmatch(pattern, str(value)) is None:
            raise ApplicationError(f'قيمة الحقل "{label}" لا تطابق النمط المطلوب.')
    elif field_type == "number" and not number_is_text(field):
        if validation.get("integer_only") and not float(value).is_integer():
            raise ApplicationError(f'الحقل "{label}" يقبل عددًا صحيحًا فقط.')
        minimum = validation.get("min")
        maximum = validation.get("max")
        if minimum is not None and float(value) < minimum:
            raise ApplicationError(f'قيمة الحقل "{label}" يجب ألا تقل عن {minimum}.')
        if maximum is not None and float(value) > maximum:
            raise ApplicationError(f'قيمة الحقل "{label}" يجب ألا تزيد على {maximum}.')
    elif field_type.startswith("date_"):
        current = value.isoformat() if isinstance(value, date) else str(value)
        minimum = validation.get("min_date")
        maximum = validation.get("max_date")
        if minimum and current < minimum:
            raise ApplicationError(f'تاريخ الحقل "{label}" أقدم من الحد المسموح.')
        if maximum and current > maximum:
            raise ApplicationError(f'تاريخ الحقل "{label}" أحدث من الحد المسموح.')


def _comparison_value(value: Any) -> str:
    if isinstance(value, date):
        return value.isoformat()
    return clean_text(value)


def validate_cross_field_constraints(
    fields: Iterable[dict[str, Any]], values: dict[str, Any]
) -> None:
    lookup = {field["id"]: field for field in fields}
    for field in lookup.values():
        validation = field.get("validation", {})
        other_id = validation.get("compare_field_id")
        operator = validation.get("compare_operator")
        if not other_id or not operator or other_id not in lookup:
            continue
        left = _comparison_value(values.get(field["id"], ""))
        right = _comparison_value(values.get(other_id, ""))
        if not left or not right:
            continue
        valid = {
            "before": left < right,
            "after": left > right,
            "on_or_before": left <= right,
            "on_or_after": left >= right,
        }.get(operator, True)
        if not valid:
            other = lookup[other_id]
            phrases = {
                "before": "قبل",
                "after": "بعد",
                "on_or_before": "في أو قبل",
                "on_or_after": "في أو بعد",
            }
            raise ApplicationError(
                f'يجب أن يكون الحقل "{field["label"]}" {phrases[operator]} '
                f'الحقل "{other["label"]}".'
            )


def excel_value(value: Any, field: dict[str, Any]) -> Any:
    normalized = normalize_field_value(value, field)
    if field["type"] == "checkbox_group":
        return " | ".join(normalized)
    if field["type"] == "checkbox":
        return "نعم" if normalized else "لا"
    return normalized



def _require_definition_id(value: Any, prefix: str, label: str) -> str:
    identifier = clean_text(value) or new_definition_id(prefix)
    if (
        not DEFINITION_ID_PATTERN.fullmatch(identifier)
        or not identifier.startswith(f"{prefix}_")
    ):
        raise ApplicationError(f"{label} الداخلي غير صالح.")
    return identifier


def _stable_option_id(field_id: str, label: str) -> str:
    digest = hashlib.sha256(f"{field_id}\0{label}".encode()).hexdigest()[:12]
    return f"opt_{digest}"


def _normalize_options(values: Any, field_id: str) -> list[dict[str, Any]]:
    if not isinstance(values, list):
        return []
    result: list[dict[str, Any]] = []
    ids: set[str] = set()
    labels: set[str] = set()
    for value in values:
        if isinstance(value, dict):
            label = clean_text(value.get("label"))
            option_id = clean_text(value.get("id")) or _stable_option_id(field_id, label)
            active = bool(value.get("active", True))
        else:
            label = clean_text(value)
            option_id = _stable_option_id(field_id, label)
            active = True
        if not label:
            continue
        normalized_label = normalize_search_text(label)
        if normalized_label in labels:
            continue
        option_id = _require_definition_id(option_id, "opt", "معرّف الخيار")
        if option_id in ids:
            raise ApplicationError("يوجد معرّف خيار مكرر.")
        ids.add(option_id)
        labels.add(normalized_label)
        result.append({"id": option_id, "label": label, "active": active})
    return result


def option_labels(field: dict[str, Any], *, active_only: bool = False) -> list[str]:
    return [
        option["label"]
        for option in field.get("options", [])
        if not active_only or option.get("active", True)
    ]


def option_by_id(field: dict[str, Any], option_id: Any) -> dict[str, Any] | None:
    text = clean_text(option_id)
    return next(
        (option for option in field.get("options", []) if option["id"] == text),
        None,
    )


def option_by_value(field: dict[str, Any], value: Any) -> dict[str, Any] | None:
    text = clean_text(value)
    return next(
        (
            option
            for option in field.get("options", [])
            if option["id"] == text or option["label"] == text
        ),
        None,
    )


def option_token(field: dict[str, Any], value: Any) -> str:
    if field.get("type") == "checkbox":
        if isinstance(value, bool):
            return "true" if value else "false"
        return "true" if clean_text(value).casefold() in {"true", "1", "نعم", "yes"} else "false"
    option = option_by_value(field, value)
    return option["id"] if option else clean_text(value)


def normalize_choice_value(value: Any, field: dict[str, Any]) -> str:
    option = option_by_value(field, value)
    if option is None or not option.get("active", True):
        raise ApplicationError(
            f'قيمة الحقل "{field["label"]}" ليست ضمن الخيارات المسموحة.'
        )
    return option["label"]


def parse_checkbox_group_value(value: Any) -> list[Any]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, tuple):
        return list(value)
    text = clean_text(value)
    if not text:
        return []
    if text.startswith("["):
        try:
            parsed = json.loads(text)
            if isinstance(parsed, list):
                return parsed
        except json.JSONDecodeError:
            pass
    return [part.strip() for part in text.split(" | ") if part.strip()]




def _optional_int(value: Any, label: str, *, minimum: int = 0) -> int | None:
    if value in {None, ""}:
        return None
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ApplicationError(f"{label} يجب أن يكون عددًا صحيحًا.") from exc
    if number < minimum:
        raise ApplicationError(f"{label} يجب ألا يقل عن {minimum}.")
    return number


def _optional_float(value: Any, label: str) -> float | None:
    if value in {None, ""}:
        return None
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ApplicationError(f"{label} يجب أن يكون رقمًا.") from exc


def normalize_number_behavior(raw: Any) -> dict[str, Any]:
    """Normalize optional number presentation/storage behavior.

    Numeric mode remains the default and stores real numbers in Excel.  Text
    mode is selected automatically when exact leading zeroes or identifier
    characters are requested, preventing Excel/Python from changing the value.
    """

    raw = raw if isinstance(raw, dict) else {}
    allowed = "".join(
        character
        for character in NUMBER_SPECIAL_CHARACTERS
        if character in clean_text(raw.get("allowed_special_characters"))
    )
    preserve = bool(raw.get("preserve_leading_zeros"))
    storage_mode = clean_text(raw.get("storage_mode"))
    if storage_mode not in NUMBER_STORAGE_MODES:
        storage_mode = "text" if preserve or allowed else "numeric"
    if preserve or allowed:
        storage_mode = "text"
    return {
        "storage_mode": storage_mode,
        "format_thousands": bool(raw.get("format_thousands")),
        "preserve_leading_zeros": preserve,
        "allowed_special_characters": allowed,
    }


def number_is_text(field: dict[str, Any]) -> bool:
    return (
        field.get("type") == "number"
        and field.get("number_behavior", {}).get("storage_mode") == "text"
    )


def normalize_number_text(value: Any, field: dict[str, Any]) -> str:
    text = clean_text(value)
    behavior = field.get("number_behavior", {})
    if behavior.get("format_thousands"):
        text = text.replace(",", "").replace("٬", "")
    text = text.translate(
        str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
    )
    allowed = set(clean_text(behavior.get("allowed_special_characters")))
    invalid = [character for character in text if not character.isdigit() and character not in allowed]
    if invalid:
        shown = " ".join(sorted(set(invalid)))
        raise ApplicationError(
            f'قيمة الحقل "{field["label"]}" تحتوي رموزًا غير مسموحة: {shown}.'
        )
    if not any(character.isdigit() for character in text):
        raise ApplicationError(f'قيمة الحقل "{field["label"]}" لا تحتوي رقمًا صالحًا.')
    return text


def _normalize_validation(
    raw: Any,
    field_type: str,
    field_label: str,
    number_behavior: dict[str, Any] | None = None,
) -> dict[str, Any]:
    raw = raw if isinstance(raw, dict) else {}
    result: dict[str, Any] = {}
    if field_type in {"text", "textarea", "user_name"}:
        minimum = _optional_int(raw.get("min_length"), f'الحد الأدنى للحقل "{field_label}"')
        maximum = _optional_int(raw.get("max_length"), f'الحد الأقصى للحقل "{field_label}"')
        if minimum is not None and maximum is not None and minimum > maximum:
            raise ApplicationError(f'حدود الطول للحقل "{field_label}" غير صحيحة.')
        pattern = clean_text(raw.get("pattern"))
        if pattern:
            try:
                re.compile(pattern)
            except re.error as exc:
                raise ApplicationError(f'نمط التحقق للحقل "{field_label}" غير صالح: {exc}') from exc
        result.update({"min_length": minimum, "max_length": maximum, "pattern": pattern})
    elif field_type == "number" and (
        not number_behavior
        or number_behavior.get("storage_mode") != "text"
    ):
        minimum = _optional_float(raw.get("min"), f'الحد الأدنى للحقل "{field_label}"')
        maximum = _optional_float(raw.get("max"), f'الحد الأقصى للحقل "{field_label}"')
        if minimum is not None and maximum is not None and minimum > maximum:
            raise ApplicationError(f'حدود الرقم للحقل "{field_label}" غير صحيحة.')
        result.update({"min": minimum, "max": maximum, "integer_only": bool(raw.get("integer_only"))})
    elif field_type.startswith("date_"):
        result.update(
            {
                "min_date": clean_text(raw.get("min_date")),
                "max_date": clean_text(raw.get("max_date")),
                "compare_field_id": clean_text(raw.get("compare_field_id")) or None,
                "compare_operator": raw.get("compare_operator")
                if raw.get("compare_operator") in {"before", "after", "on_or_before", "on_or_after"}
                else None,
            }
        )
    return result


def condition_operators_for_field(
    field_type: str,
) -> set[str]:
    return CONDITION_OPERATORS_BY_FIELD_TYPE.get(
        field_type,
        {"equals", "not_equals", "empty", "not_empty"},
    )


def stable_condition_group_id(
    target_type: str,
    target_id: str,
) -> str:
    digest = hashlib.sha256(
        f"{target_type}\0{target_id}".encode()
    ).hexdigest()[:12]

    return f"grp_{digest}"


def related_person_mode_source_id(category_id: str) -> str:
    return f"{RELATED_PERSON_MODE_SOURCE_PREFIX}{category_id}"


def related_person_mode_source_field(
    category: dict[str, Any],
) -> dict[str, Any] | None:
    if (
        category.get("kind") != "repeatable"
        or not category.get("related_person_enabled")
    ):
        return None
    return {
        "id": related_person_mode_source_id(category["id"]),
        "label": "هل لديه سجل؟",
        "type": "yes_no",
        "options": [
            {"id": "existing", "label": "لديه سجل", "active": True},
            {"id": "manual", "label": "ليس لديه سجل", "active": True},
        ],
    }


def normalize_condition_value(
    source_field: dict[str, Any],
    operator: str,
    raw_value: Any,
) -> str:
    if operator in {"empty", "not_empty"}:
        return ""

    field_type = source_field["type"]
    value = clean_text(raw_value)

    if not value:
        raise ApplicationError(
            f'أدخل قيمة شرط الحقل "{source_field["label"]}".'
        )

    if field_type == "checkbox":
        normalized = value.casefold()

        if normalized in {
            "true",
            "1",
            "yes",
            "نعم",
            "checked",
            "محدد",
        }:
            return "true"

        if normalized in {
            "false",
            "0",
            "no",
            "لا",
            "unchecked",
            "غير محدد",
        }:
            return "false"

        raise ApplicationError(
            f'قيمة شرط مربع الاختيار '
            f'"{source_field["label"]}" غير صالحة.'
        )

    if field_type in {
        "select",
        "yes_no",
        "checkbox_group",
    }:
        option = option_by_value(
            source_field,
            value,
        )

        if option is None:
            raise ApplicationError(
                f'قيمة شرط الحقل '
                f'"{source_field["label"]}" '
                "ليست ضمن خياراته."
            )

        # Store the stable option ID, not its editable label.
        return option["id"]

    if field_type == "number":
        if number_is_text(source_field):
            return normalize_number_text(value, source_field)
        try:
            number = float(value)
        except ValueError as exc:
            raise ApplicationError(
                f'قيمة شرط الحقل '
                f'"{source_field["label"]}" ليست رقمًا.'
            ) from exc

        return (
            str(int(number))
            if number.is_integer()
            else format(number, ".15g")
        )

    if field_type == "date_gregorian":
        try:
            return date.fromisoformat(value).isoformat()
        except ValueError as exc:
            raise ApplicationError(
                f'قيمة شرط التاريخ '
                f'"{source_field["label"]}" غير صالحة.'
            ) from exc

    if field_type == "date_hijri":
        return validate_hijri_date(
            value,
            source_field["label"],
        )

    if field_type == "date_persian":
        return validate_persian_date(
            value,
            source_field["label"],
        )

    return value

def validate_schema(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise ApplicationError("صيغة إعداد التطبيق غير صحيحة.")

    app_payload = payload.get("app")
    app_payload = app_payload if isinstance(app_payload, dict) else {}
    raw_shortcuts = app_payload.get("shortcuts", DEFAULT_SHORTCUT_BINDINGS)
    shortcuts: dict[str, str] = {}
    if isinstance(raw_shortcuts, dict):
        for raw_key, raw_action in raw_shortcuts.items():
            key = clean_text(raw_key)
            action = clean_text(raw_action)
            if key and action in APP_SHORTCUTS and action not in shortcuts.values():
                shortcuts[key] = action
    if isinstance(raw_shortcuts, list):
        legacy_actions = []
        for value in raw_shortcuts:
            action = clean_text(value)
            if action in APP_SHORTCUTS and action not in legacy_actions:
                legacy_actions.append(action)
        for key, action in DEFAULT_SHORTCUT_BINDINGS.items():
            if action in legacy_actions:
                shortcuts[key] = action
    startup_page = clean_text(app_payload.get("startup_page")) or "home"
    if startup_page not in APP_PAGES or startup_page in {"builder", "import", "export"}:
        startup_page = "home"
    try:
        search_page_size = int(app_payload.get("search_page_size", 50))
    except (TypeError, ValueError):
        search_page_size = 50
    search_page_size = min(100, max(10, search_page_size))
    app = {
        "title": clean_text(app_payload.get("title")) or "نظام إدخال البيانات",
        "entity_singular": clean_text(app_payload.get("entity_singular")) or "سجل",
        "entity_plural": clean_text(app_payload.get("entity_plural")) or "السجلات",
        "direction": "rtl",
        "language": "ar",
        "primary_color": (
            app_payload.get("primary_color")
            if HEX_COLOR_PATTERN.fullmatch(
                clean_text(app_payload.get("primary_color"))
            )
            else "#1F5F95"
        ),
        # These keys remain in Schema v1 for compatibility.  Their values are
        # fixed so a custom theme cannot hide controls or their text.
        "background_color": FIXED_BACKGROUND_COLOR,
        "surface_color": FIXED_SURFACE_COLOR,
        "startup_page": startup_page,
        "search_page_size": search_page_size,
        "show_entry_search": bool(app_payload.get("show_entry_search", True)),
        "draft_autosave": bool(app_payload.get("draft_autosave", True)),
        "shortcuts": shortcuts or copy.deepcopy(DEFAULT_SHORTCUT_BINDINGS),
    }

    raw_categories = payload.get("categories", [])
    if not isinstance(raw_categories, list):
        raise ApplicationError("قائمة الفئات غير صحيحة.")

    categories: list[dict[str, Any]] = []
    category_ids: set[str] = set()
    category_labels: set[str] = set()
    field_ids: set[str] = set()
    removed_legacy_field_ids: set[str] = set()
    field_lookup: dict[str, dict[str, Any]] = {}
    field_category: dict[str, str] = {}

    for raw_category in raw_categories:
        if not isinstance(raw_category, dict):
            raise ApplicationError("أحد تعريفات الفئات غير صحيح.")
        category_id = _require_definition_id(
            raw_category.get("id"), "cat", "معرّف الفئة"
        )
        if category_id in category_ids:
            raise ApplicationError("يوجد معرّف فئة مكرر.")
        category_ids.add(category_id)

        label = clean_text(raw_category.get("label"))
        if not label:
            raise ApplicationError("اسم الفئة مطلوب.")
        normalized_label = normalize_search_text(label)
        if normalized_label in category_labels:
            raise ApplicationError(f'اسم الفئة "{label}" مكرر.')
        category_labels.add(normalized_label)

        kind = raw_category.get("kind", "main")
        if kind not in CATEGORY_KINDS:
            raise ApplicationError(f'نوع الفئة "{label}" غير صالح.')

        related_person_enabled = kind == "repeatable" and bool(
            raw_category.get("related_person_enabled")
        )

        raw_fields = raw_category.get("fields", [])
        if not isinstance(raw_fields, list):
            raise ApplicationError(f'حقول الفئة "{label}" غير صحيحة.')

        fields: list[dict[str, Any]] = []
        local_labels: set[str] = set()
        for raw_field in raw_fields:
            if not isinstance(raw_field, dict):
                raise ApplicationError(f'أحد حقول الفئة "{label}" غير صحيح.')
            # Alpha 25 removes the former record-completion feature. Older
            # schemas remain readable, but their two display-only completion
            # fields are deliberately omitted from the active definition.
            if raw_field.get("type") in LEGACY_COMPLETION_FIELD_TYPES:
                legacy_id = clean_text(raw_field.get("id"))
                if legacy_id:
                    removed_legacy_field_ids.add(legacy_id)
                continue
            field_id = _require_definition_id(
                raw_field.get("id"), "fld", "معرّف الحقل"
            )
            if field_id in field_ids:
                raise ApplicationError("يوجد معرّف حقل مكرر.")
            field_ids.add(field_id)

            spacer = raw_field.get("type") == "spacer"
            if spacer:
                # Layout definitions never acquire values, validation or data rules.
                raw_field = {key: value for key, value in raw_field.items() if key in {
                    "id", "type", "width", "start_new_line", "global_ref", "global_tree_ref", "global_tree_key"
                }}
            field_label = clean_text(raw_field.get("label"))
            if not field_label and not spacer:
                raise ApplicationError(f'يوجد حقل بلا اسم في الفئة "{label}".')
            normalized_field_label = normalize_search_text(field_label)
            if normalized_field_label in local_labels and not spacer:
                raise ApplicationError(
                    f'اسم الحقل "{field_label}" مكرر داخل الفئة "{label}".'
                )
            if not spacer:
                local_labels.add(normalized_field_label)

            field_type = raw_field.get("type", "text")
            if field_type not in FIELD_TYPES | AUDIT_FIELD_TYPES:
                raise ApplicationError(f'نوع الحقل "{field_label}" غير صالح.')
            if field_type in SYSTEM_FIELD_TYPES and kind != "main":
                raise ApplicationError(
                    f'الحقل "{field_label}" يعرض بيانات السجل الرئيسي فقط.'
                )
            if field_type == "user_name" and kind != "main":
                raise ApplicationError(
                    f'حقل المستخدم "{field_label}" متاح داخل الفئات الرئيسية فقط.'
                )
            options = _normalize_options(raw_field.get("options", []), field_id)
            if field_type == "yes_no":
                existing = {option["label"]: option for option in options}
                options = [
                    existing.get(label)
                    or {
                        "id": _stable_option_id(field_id, label),
                        "label": label,
                        "active": True,
                    }
                    for label in ("نعم", "لا")
                ]
            if field_type in {"select", "checkbox_group"} and not any(
                option.get("active", True) for option in options
            ):
                raise ApplicationError(
                    f'أضف خيارًا واحدًا على الأقل للحقل "{field_label}".'
                )

            number_behavior = (
                normalize_number_behavior(raw_field.get("number_behavior"))
                if field_type == "number"
                else None
            )

            system_field = field_type in SYSTEM_FIELD_TYPES or spacer
            searchable = bool(raw_field.get("searchable")) and field_type != "file"
            show_in_results = bool(raw_field.get("show_in_results")) and field_type != "file"
            result_title = (
                bool(raw_field.get("result_title"))
                and show_in_results
                and kind == "main"
            )
            default_match = (
                "contains"
                if field_type in {"text", "textarea", "user_name"}
                else "exact"
            )
            search_match = raw_field.get("search_match", default_match)
            if search_match not in SEARCH_MATCHES:
                search_match = default_match

            user_value_mode = clean_text(raw_field.get("user_value_mode"))
            if user_value_mode not in {
                "created_by",
                "current_on_save",
                "current_on_checkbox",
            }:
                user_value_mode = "created_by"
            current_on_save = (
                field_type == "user_name" and user_value_mode == "current_on_save"
            )
            date_value_mode = clean_text(raw_field.get("date_value_mode"))
            if field_type != "date_gregorian" or date_value_mode not in {
                "manual",
                "on_checkbox",
            }:
                date_value_mode = "manual"
            field = {
                "id": field_id,
                "global_ref": (
                    clean_text(raw_field.get("global_ref"))
                    if GLOBAL_REFERENCE.fullmatch(clean_text(raw_field.get("global_ref")))
                    and clean_text(raw_field.get("global_ref")).startswith("gfld_")
                    else None
                ),
                "global_tree_ref": (
                    clean_text(raw_field.get("global_tree_ref"))
                    if GLOBAL_REFERENCE.fullmatch(clean_text(raw_field.get("global_tree_ref")))
                    and clean_text(raw_field.get("global_tree_ref")).startswith("gcat_")
                    else None
                ),
                "global_tree_key": clean_text(raw_field.get("global_tree_key"))[:120] or None,
                "label": field_label,
                "type": field_type,
                "start_new_line": bool(raw_field.get("start_new_line", False)),
                "checkbox_true_label": clean_text(raw_field.get("checkbox_true_label"))[:80] if field_type == "checkbox" else "",
                "checkbox_false_label": clean_text(raw_field.get("checkbox_false_label"))[:80] if field_type == "checkbox" else "",
                "required": (
                    bool(raw_field.get("required"))
                    and not system_field
                    and not current_on_save
                ),
                "placeholder": (
                    "" if system_field else clean_text(raw_field.get("placeholder"))
                ),
                "width": (
                    raw_field.get("width")
                    if raw_field.get("width") in FIELD_WIDTHS
                    else LEGACY_FIELD_WIDTHS.get(raw_field.get("width"), "1")
                ),
                "options": options,
                "searchable": searchable,
                "search_match": search_match,
                "show_in_results": show_in_results,
                "result_title": result_title,
                "unique": (
                    bool(raw_field.get("unique"))
                    and field_type not in {"file", "checkbox"}
                    and not system_field
                ),
                "user_editable": (
                    bool(raw_field.get("user_editable", True))
                    if field_type == "user_name"
                    else True
                ),
                "user_value_mode": (
                    user_value_mode if field_type == "user_name" else "created_by"
                ),
                "user_trigger_field_id": (
                    clean_text(raw_field.get("user_trigger_field_id")) or None
                    if field_type == "user_name"
                    and user_value_mode == "current_on_checkbox"
                    else None
                ),
                "date_value_mode": date_value_mode,
                "date_trigger_field_id": (
                    clean_text(raw_field.get("date_trigger_field_id")) or None
                    if date_value_mode == "on_checkbox"
                    else None
                ),
                "unique_checked_across_cards": (
                    bool(raw_field.get("unique_checked_across_cards"))
                    if field_type == "checkbox" and kind == "repeatable"
                    else False
                ),
                "validation": (
                    {}
                    if system_field
                    else _normalize_validation(
                        raw_field.get("validation"),
                        field_type,
                        field_label,
                        number_behavior,
                    )
                ),
                "option_filter": None,
                "_option_filter_raw": raw_field.get("option_filter"),
                "related_person_source_field_id": (
                    clean_text(raw_field.get("related_person_source_field_id"))
                    or None
                    if related_person_enabled
                    else None
                ),
                "related_person_source_checkbox_id": (
                    clean_text(
                        raw_field.get("related_person_source_checkbox_id")
                        or raw_field.get("related_person_source_marker_id")
                    )
                    or None
                    if related_person_enabled
                    else None
                ),
                "_auto_update_raw": raw_field.get("auto_update"),
            }

            if number_behavior is not None:
                field["number_behavior"] = number_behavior

            if field_type == "file":
                field["image_display"] = (
                    "profile"
                    if kind == "main"
                    and raw_field.get("image_display") == "profile"
                    else "card"
                    if kind == "repeatable"
                    and raw_field.get("image_display") in {"profile", "card"}
                    else None
                )
                raw_naming = raw_field.get("file_naming")
                raw_naming = raw_naming if isinstance(raw_naming, dict) else {}
                mode = raw_naming.get("mode", "original")
                if mode not in {"original", "template"}:
                    mode = "original"
                raw_parts = raw_naming.get("parts", [])
                parts = []
                if isinstance(raw_parts, list):
                    for raw_part in raw_parts:
                        if not isinstance(raw_part, dict):
                            continue
                        source_field_id = clean_text(raw_part.get("field_id"))
                        if source_field_id:
                            parts.append(
                                {
                                    "field_id": source_field_id,
                                    "prefix": str(
                                        raw_part.get("prefix", "")
                                    )[:30],
                                    "suffix": str(
                                        raw_part.get("suffix", "")
                                    )[:30],
                                }
                            )
                field["file_naming"] = {"mode": mode, "parts": parts}

            fields.append(field)
            if not spacer:
                field_lookup[field_id] = field
            field_category[field_id] = category_id

        # Category placement is now defined only by root order or an exact
        # parent field. The retired independent anchor selector is ignored.
        anchor_field_id = None
        parent_category_id = (
            clean_text(raw_category.get("parent_category_id")) or None
        )
        parent_field_id = clean_text(raw_category.get("parent_field_id")) or None
        categories.append(
            {
                "id": category_id,
                "global_ref": (
                    clean_text(raw_category.get("global_ref"))
                    if GLOBAL_REFERENCE.fullmatch(clean_text(raw_category.get("global_ref")))
                    and clean_text(raw_category.get("global_ref")).startswith("gcat_")
                    else None
                ),
                "global_tree_ref": (
                    clean_text(raw_category.get("global_tree_ref"))
                    if GLOBAL_REFERENCE.fullmatch(clean_text(raw_category.get("global_tree_ref")))
                    and clean_text(raw_category.get("global_tree_ref")).startswith("gcat_")
                    else None
                ),
                "global_tree_key": clean_text(raw_category.get("global_tree_key"))[:120] or None,
                "label": label,
                "description": clean_text(raw_category.get("description")),
                "kind": kind,
                "add_label": (
                    clean_text(raw_category.get("add_label"))
                    or f"إضافة {label}"
                ),
                "auto_start": bool(raw_category.get("auto_start")),
                "anchor_field_id": anchor_field_id,
                "parent_category_id": parent_category_id,
                "parent_field_id": parent_field_id,
                "related_person_enabled": related_person_enabled,
                "card_title_field_id": (
                    clean_text(raw_category.get("card_title_field_id")) or None
                    if kind == "repeatable"
                    else None
                ),
                "card_name_prefix": (
                    clean_text(raw_category.get("card_name_prefix"))[:80]
                    if kind == "repeatable"
                    else ""
                ),
                "card_sort": (
                    raw_category.get("card_sort")
                    if kind == "repeatable"
                    and isinstance(raw_category.get("card_sort"), dict)
                    else {"mode": "manual", "direction": "asc"}
                ),
                "fields": fields,
            }
        )

    category_lookup = {category["id"]: category for category in categories}

    for category in categories:
        parent_id = category["parent_category_id"]
        if parent_id and parent_id not in category_lookup:
            raise ApplicationError(
                f'الفئة الأم للفئة "{category["label"]}" محذوفة.'
            )
        if parent_id == category["id"]:
            raise ApplicationError("لا يمكن أن تكون الفئة أمًّا لنفسها.")
        parent_field_id = category.get("parent_field_id")
        if not parent_id:
            category["parent_field_id"] = None
        else:
            parent_field_ids = {
                field["id"] for field in category_lookup[parent_id]["fields"]
            }
            # A missing field means that the child category is placed at the
            # end of its parent. Otherwise the anchor must belong to the
            # selected parent category.
            if parent_field_id and parent_field_id not in parent_field_ids:
                raise ApplicationError(
                    f'اختر حقلًا صالحًا من الفئة الأم للفئة "{category["label"]}".'
                )

    for category in categories:
        visited = {category["id"]}
        current_id = category["parent_category_id"]
        while current_id:
            if current_id in visited:
                raise ApplicationError("يوجد تسلسل دائري بين الفئات الأم والفرعية.")
            visited.add(current_id)
            current_id = category_lookup[current_id]["parent_category_id"]

    for category in categories:
        if category["kind"] != "repeatable":
            category["card_title_field_id"] = None
            category["card_sort"] = {"mode": "manual", "direction": "asc"}
            continue
        local_field_ids = {field["id"] for field in data_fields(category)}
        title_id = category.get("card_title_field_id")
        if title_id not in local_field_ids or field_lookup.get(title_id, {}).get("type") == "file":
            category["card_title_field_id"] = None
        raw_sort = category.get("card_sort", {})
        mode = raw_sort.get("mode", "manual")
        if mode not in {"manual", "title", "field"}:
            mode = "manual"
        field_id = clean_text(raw_sort.get("field_id")) or None
        if field_id not in local_field_ids or field_lookup.get(field_id, {}).get("type") == "file":
            field_id = None
        if mode == "field" and not field_id:
            mode = "manual"
        category["card_sort"] = {
            "mode": mode,
            "field_id": field_id,
            "direction": "desc" if raw_sort.get("direction") == "desc" else "asc",
        }

    for category in categories:
        if not category.get("related_person_enabled"):
            continue
        for field in category["fields"]:
            source_id = field.get("related_person_source_field_id")
            if not source_id:
                continue
            source = field_lookup.get(source_id)
            source_category = category_lookup.get(field_category.get(source_id, ""))
            if not source or not source_category:
                raise ApplicationError(
                    f'حقل النسخ المرتبط بالحقل "{field["label"]}" غير متاح.'
                )
            if source["type"] == "file" or source["type"] != field["type"]:
                raise ApplicationError(
                    f'نوع حقل المصدر لا يطابق الحقل "{field["label"]}".'
                )
            if source_category["kind"] == "repeatable":
                selector_id = field.get("related_person_source_checkbox_id")
                selector = field_lookup.get(selector_id)
                if (
                    not selector
                    or field_category.get(selector_id) != source_category["id"]
                    or selector.get("type") != "checkbox"
                    or not selector.get("unique_checked_across_cards")
                ):
                    raise ApplicationError(
                        f'اختر مربع اختيار فريدًا لتحديد بطاقة المصدر للحقل "{field["label"]}".'
                    )
            else:
                field["related_person_source_checkbox_id"] = None

    for category in categories:
        for field in category["fields"]:
            trigger_specs = []
            if field.get("date_value_mode") == "on_checkbox":
                trigger_specs.append((field.get("date_trigger_field_id"), "التاريخ التلقائي"))
            if field.get("user_value_mode") == "current_on_checkbox":
                trigger_specs.append((field.get("user_trigger_field_id"), "اسم المستخدم التلقائي"))
            for trigger_id, trigger_label in trigger_specs:
                trigger = field_lookup.get(trigger_id)
                trigger_category = category_lookup.get(field_category.get(trigger_id, ""))
                if (
                    not trigger
                    or trigger.get("type") != "checkbox"
                    or not trigger_category
                    or trigger_category.get("kind") != "main"
                ):
                    raise ApplicationError(
                        f'اختر مربع اختيار من فئة رئيسية لتشغيل {trigger_label} في "{field["label"]}".'
                    )

    for category in categories:
        for field in category["fields"]:
            raw_rule = field.pop("_auto_update_raw", None)
            field["auto_update"] = None
            if not raw_rule:
                continue
            if field["type"] in SYSTEM_FIELD_TYPES:
                raise ApplicationError("لا يمكن أن يكون حقل بيانات السجل هدفًا لتحديث تلقائي.")
            if not isinstance(raw_rule, dict):
                raise ApplicationError(f'قاعدة التحديث التلقائي للحقل "{field["label"]}" غير صحيحة.')
            source_id = clean_text(raw_rule.get("source_field_id"))
            if source_id in removed_legacy_field_ids:
                continue
            if source_id not in field_lookup or source_id == field["id"]:
                raise ApplicationError(f'مصدر التحديث التلقائي للحقل "{field["label"]}" غير صالح.')
            operator = raw_rule.get("operator", "equals")
            if operator not in condition_operators_for_field(field_lookup[source_id]["type"]):
                raise ApplicationError(f'عملية التحديث التلقائي لا تناسب الحقل "{field_lookup[source_id]["label"]}".')
            action = raw_rule.get("action", "fixed")
            if action not in {"fixed", "current_user", "copy_source", "clear"}:
                action = "fixed"
            field["auto_update"] = {
                "source_field_id": source_id,
                "operator": operator,
                "value": normalize_condition_value(
                    field_lookup[source_id], operator, raw_rule.get("value", "")
                ),
                "action": action,
                "result_value": clean_text(raw_rule.get("result_value")),
            }

    profile_fields = [
        field
        for category in categories
        for field in category["fields"]
        if field.get("image_display") == "profile"
    ]
    if len(profile_fields) > 1:
        raise ApplicationError("يمكن تحديد حقل صورة شخصية واحد فقط.")

    system_type_counts = {
        field_type: sum(
            field["type"] == field_type
            for category in categories
            for field in category["fields"]
        )
        for field_type in SYSTEM_FIELD_TYPES
    }
    if any(count > 1 for count in system_type_counts.values()):
        raise ApplicationError(
            "يمكن إضافة حقل واحد فقط من كل نوع من حقول بيانات السجل."
        )

    main_field_ids = {
        field["id"]
        for category in categories
        if category["kind"] == "main"
        for field in category["fields"]
        if field["type"] not in SYSTEM_FIELD_TYPES
    }

    for category in categories:
        anchor_field_id = category["anchor_field_id"]
        if anchor_field_id and anchor_field_id not in main_field_ids:
            raise ApplicationError(
                f'موضع الفئة "{category["label"]}" يجب أن يرتبط بحقل رئيسي.'
            )
        for field in category["fields"]:
            if field["type"] != "file":
                continue
            naming = field["file_naming"]
            if naming["mode"] == "template" and not naming["parts"]:
                raise ApplicationError(
                    f'أضف أجزاء تسمية الملف للحقل "{field["label"]}".'
                )
            for part in naming["parts"]:
                source_id = part["field_id"]
                if source_id not in field_lookup:
                    raise ApplicationError(
                        f'صيغة اسم الملف في "{field["label"]}" تشير إلى حقل محذوف.'
                    )
                source_category = category_lookup[field_category[source_id]]
                if field_lookup[source_id]["type"] == "file":
                    raise ApplicationError(
                        f'صيغة اسم الملف في "{field["label"]}" لا يمكن أن تستخدم ملفًا آخر.'
                    )
                if field_lookup[source_id]["type"] in SYSTEM_FIELD_TYPES:
                    raise ApplicationError(
                        f'صيغة اسم الملف في "{field["label"]}" '
                        "لا تستخدم حقول بيانات السجل التقنية."
                    )
                if source_category["kind"] == "repeatable":
                    if source_category["id"] != category["id"]:
                        raise ApplicationError(
                            f'اسم الملف في "{field["label"]}" لا يمكنه استخدام حقل من فئة متكررة أخرى.'
                        )

    for category in categories:
        for field in category["fields"]:
            validation = field.get("validation", {})
            compare_id = validation.get("compare_field_id")
            if compare_id:
                if compare_id not in field_lookup:
                    raise ApplicationError(
                        f'مقارنة الحقل "{field["label"]}" تشير إلى حقل محذوف.'
                    )
                compare_category = category_lookup[field_category[compare_id]]
                if field_lookup[compare_id]["type"] != field["type"]:
                    raise ApplicationError(
                        f'مقارنة الحقل "{field["label"]}" يجب أن تكون مع تاريخ من النوع نفسه.'
                    )
                if compare_category["kind"] != "main" and compare_category["id"] != category["id"]:
                    raise ApplicationError(
                        "لا يمكن مقارنة حقل بصف من فئة متكررة أخرى."
                    )

            raw_filter = field.pop("_option_filter_raw", None)
            if not raw_filter:
                continue
            if field["type"] not in {"select", "checkbox_group"}:
                raise ApplicationError(
                    f'تصفية الخيارات متاحة فقط للقائمة أو مجموعة الاختيارات في "{field["label"]}".'
                )
            if not isinstance(raw_filter, dict):
                raise ApplicationError(f'تصفية خيارات الحقل "{field["label"]}" غير صحيحة.')
            source_id = clean_text(raw_filter.get("source_field_id"))
            if source_id in removed_legacy_field_ids:
                continue
            if source_id not in field_lookup:
                raise ApplicationError(f'مصدر تصفية الحقل "{field["label"]}" محذوف.')
            if source_id == field["id"]:
                raise ApplicationError("لا يمكن للحقل أن يصفّي خياراته بنفسه.")
            source = field_lookup[source_id]
            source_category = category_lookup[field_category[source_id]]
            if source["type"] not in {"select", "yes_no", "checkbox"}:
                raise ApplicationError(
                    f'الحقل المتحكم في "{field["label"]}" يجب أن يكون قائمة أو نعم/لا أو مربع اختيار.'
                )
            if source_category["kind"] != "main" and source_category["id"] != category["id"]:
                raise ApplicationError(
                    "تصفية فئة متكررة لا يمكن أن تعتمد على صف من فئة متكررة أخرى."
                )
            valid_source_tokens = (
                {"true", "false"}
                if source["type"] == "checkbox"
                else {option["id"] for option in source["options"]}
            )
            valid_target_ids = {option["id"] for option in field["options"]}
            raw_mappings = raw_filter.get("mappings", {})
            if not isinstance(raw_mappings, dict):
                raise ApplicationError(f'خريطة خيارات الحقل "{field["label"]}" غير صحيحة.')
            mappings: dict[str, list[str]] = {}
            for raw_token, raw_allowed in raw_mappings.items():
                token = clean_text(raw_token)
                if token not in valid_source_tokens:
                    continue
                if not isinstance(raw_allowed, list):
                    continue
                allowed: list[str] = []
                for option_id in raw_allowed:
                    option_id = clean_text(option_id)
                    if option_id in valid_target_ids and option_id not in allowed:
                        allowed.append(option_id)
                mappings[token] = allowed
            field["option_filter"] = {
                "source_field_id": source_id,
                "mappings": mappings,
                "unmatched": "none" if raw_filter.get("unmatched") == "none" else "all",
            }

    for category in categories:
        mode_source = related_person_mode_source_field(category)
        if mode_source:
            field_lookup[mode_source["id"]] = mode_source
            field_category[mode_source["id"]] = category["id"]

    raw_conditions = payload.get("conditions", [])

    if not isinstance(raw_conditions, list):
        raise ApplicationError(
            "قائمة الشروط غير صحيحة."
        )

    conditions: list[dict[str, Any]] = []
    condition_ids: set[str] = set()

    for raw_condition in raw_conditions:
        if not isinstance(raw_condition, dict):
            raise ApplicationError(
                "أحد شروط الظهور غير صحيح."
            )

        condition_id = _require_definition_id(
            raw_condition.get("id"),
            "cond",
            "معرّف الشرط",
        )

        if condition_id in condition_ids:
            raise ApplicationError(
                "يوجد معرّف شرط مكرر."
            )

        condition_ids.add(condition_id)

        target_type = raw_condition.get("target_type")
        target_id = clean_text(
            raw_condition.get("target_id")
        )
        source_field_id = clean_text(
            raw_condition.get("source_field_id")
        )

        if target_id in removed_legacy_field_ids or source_field_id in removed_legacy_field_ids:
            continue

        if target_type not in {"category", "field"}:
            raise ApplicationError(
                "نوع هدف الشرط غير صالح."
            )

        if (
            target_type == "category"
            and target_id not in category_lookup
        ):
            raise ApplicationError(
                "هدف أحد الشروط هو فئة محذوفة."
            )

        if (
            target_type == "field"
            and target_id not in field_ids
        ):
            raise ApplicationError(
                "هدف أحد الشروط هو حقل محذوف."
            )

        if source_field_id not in field_lookup:
            raise ApplicationError(
                "مصدر أحد الشروط هو حقل محذوف."
            )

        if (
            source_field_id.startswith(RELATED_PERSON_MODE_SOURCE_PREFIX)
            and target_type != "field"
        ):
            raise ApplicationError(
                "حالة سجل الشخص المرتبط تتحكم في حقول بطاقته فقط."
            )

        if (
            target_type == "field"
            and target_id == source_field_id
        ):
            raise ApplicationError(
                "لا يمكن للحقل أن يتحكم في ظهوره بنفسه."
            )

        source_field = field_lookup[source_field_id]
        operator = raw_condition.get(
            "operator",
            "equals",
        )

        allowed_operators = condition_operators_for_field(
            source_field["type"]
        )

        if operator not in allowed_operators:
            raise ApplicationError(
                f'العملية المختارة لا تناسب نوع الحقل '
                f'"{source_field["label"]}".'
            )

        target_category_id = (
            target_id
            if target_type == "category"
            else field_category[target_id]
        )

        source_category = category_lookup[
            field_category[source_field_id]
        ]
        target_category = category_lookup[
            target_category_id
        ]

        if (
            source_category["kind"] != "main"
            and source_category["id"]
            != target_category["id"]
        ):
            raise ApplicationError(
                "شرط فئة متكررة لا يمكن أن يعتمد "
                "على صف من فئة متكررة أخرى."
            )

        default_group_id = stable_condition_group_id(
            target_type,
            target_id,
        )

        group_id = _require_definition_id(
            raw_condition.get("group_id")
            or default_group_id,
            "grp",
            "معرّف مجموعة الشرط",
        )

        value = normalize_condition_value(
            source_field,
            operator,
            raw_condition.get("value", ""),
        )

        normalized_condition = {
                "id": condition_id,
                "group_id": group_id,
                "negate": bool(
                    raw_condition.get("negate", False)
                ),
                "target_type": target_type,
                "target_id": target_id,
                "source_field_id": source_field_id,
                "operator": operator,
                "value": value,
            }
        global_tree_ref = clean_text(raw_condition.get("global_tree_ref"))
        global_condition_key = clean_text(raw_condition.get("global_condition_key"))
        if GLOBAL_REFERENCE.fullmatch(global_tree_ref) and global_condition_key:
            normalized_condition["global_tree_ref"] = global_tree_ref
            normalized_condition["global_condition_key"] = global_condition_key[:160]
        conditions.append(normalized_condition)
    revision = payload.get("revision", 0)
    try:
        revision = int(revision)
    except (TypeError, ValueError):
        revision = 0

    return {
        "schema_version": SCHEMA_VERSION,
        "revision": max(0, revision),
        "app": app,
        "categories": categories,
        "conditions": conditions,
    }


def schema_indexes(schema: dict[str, Any]) -> dict[str, Any]:
    categories = {category["id"]: category for category in schema["categories"]}
    fields: dict[str, dict[str, Any]] = {}
    field_categories: dict[str, str] = {}
    for category in schema["categories"]:
        for field in data_fields(category):
            fields[field["id"]] = field
            field_categories[field["id"]] = category["id"]
        mode_source = related_person_mode_source_field(category)
        if mode_source:
            fields[mode_source["id"]] = mode_source
            field_categories[mode_source["id"]] = category["id"]
    conditions_by_target: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for condition in schema["conditions"]:
        key = (condition["target_type"], condition["target_id"])
        conditions_by_target.setdefault(key, []).append(condition)
    return {
        "categories": categories,
        "fields": fields,
        "field_categories": field_categories,
        "conditions_by_target": conditions_by_target,
    }


def read_schema_file() -> dict[str, Any]:
    schema_path = _schema_path()
    if not schema_path.is_file():
        return default_schema()
    try:
        payload = json.loads(schema_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ApplicationError(f"تعذّر قراءة ملف الإعدادات: {exc}") from exc
    try:
        version = int(payload.get("schema_version", 1))
    except (TypeError, ValueError):
        version = 0
    if version not in SUPPORTED_SCHEMA_VERSIONS:
        raise ApplicationError("إصدار ملف الإعدادات غير مدعوم.")
    schema = validate_schema(payload)
    if version != SCHEMA_VERSION or payload != schema:
        atomic_write_json(schema_path, schema)
    return schema


def _write_json_temp(path: Path, payload: dict[str, Any]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        prefix=f".{path.stem}-",
        suffix=".json",
        dir=path.parent,
        delete=False,
    ) as temporary:
        json.dump(payload, temporary, ensure_ascii=False, indent=2)
        temporary.write("\n")
        return Path(temporary.name)


def atomic_write_json(path: Path, payload: dict[str, Any]) -> None:
    temporary = _write_json_temp(path, payload)
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def atomic_write_bytes(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        mode="wb",
        prefix=f".{path.stem}-",
        dir=path.parent,
        delete=False,
    ) as temporary:
        temporary.write(content)
        temporary_path = Path(temporary.name)
    try:
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def sanitize_sheet_name(value: str) -> str:
    name = INVALID_EXCEL_SHEET_PATTERN.sub(" ", value)
    name = " ".join(name.split()).strip("' ")
    return name or "فئة"


def related_sheet_names(schema: dict[str, Any]) -> dict[str, str]:
    used = {MAIN_SHEET.casefold(), META_SHEET.casefold()}
    result: dict[str, str] = {}
    for category in schema["categories"]:
        if category["kind"] != "repeatable":
            continue
        base = sanitize_sheet_name(category["label"])
        candidate = base[:31]
        if candidate.casefold() in used:
            suffix = f"-{category['id'][-6:]}"
            candidate = f"{base[:31 - len(suffix)]}{suffix}"
        while candidate.casefold() in used:
            candidate = f"{base[:22]}-{secrets.token_hex(4)}"[:31]
        used.add(candidate.casefold())
        result[category["id"]] = candidate
    return result


def data_fields(category: dict[str, Any]) -> list[dict[str, Any]]:
    """Fields with record values; spacer definitions belong only to the layout."""
    return [field for field in category["fields"] if field["type"] != "spacer"]


def main_fields(schema: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        field
        for category in schema["categories"]
        if category["kind"] == "main"
        for field in data_fields(category)
        if field["type"] not in SYSTEM_FIELD_TYPES
    ]


def file_fields(schema: dict[str, Any]) -> list[tuple[dict[str, Any], dict[str, Any]]]:
    return [
        (category, field)
        for category in schema["categories"]
        for field in category["fields"]
        if field["type"] == "file"
    ]


def technical_headers(worksheet) -> list[str]:
    row = next(
        worksheet.iter_rows(
            min_row=TECHNICAL_HEADER_ROW,
            max_row=TECHNICAL_HEADER_ROW,
            values_only=True,
        ),
        (),
    )
    headers = [clean_text(value) for value in row]
    while headers and not headers[-1]:
        headers.pop()
    return headers


def workbook_schema_revision(path: Path | None = None) -> int | None:
    """Read only the workbook metadata revision, without loading record sheets."""

    workbook_path = path or _workbook_path()
    if not workbook_path.is_file():
        return None
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        if META_SHEET not in workbook.sheetnames:
            return None
        for key, value in workbook[META_SHEET].iter_rows(min_row=1, max_col=2, values_only=True):
            if clean_text(key) == "schema_revision":
                try:
                    return int(value)
                except (TypeError, ValueError):
                    return None
    finally:
        workbook.close()
    return None


def _sync_excel_field_labels_unlocked(
    schema: dict[str, Any],
) -> dict[str, Any]:
    """Import visible Excel label edits while preserving stable field IDs."""
    workbook_path = _workbook_path()
    if not workbook_path.is_file() or not schema["categories"]:
        return schema
    if workbook_schema_revision(workbook_path) not in {None, schema["revision"]}:
        # A schema edit may already be committed while its Excel projection is
        # being rebuilt in the background. Never re-import labels from that
        # older workbook revision.
        return schema
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    except Exception as exc:
        raise ApplicationError(f"تعذّر قراءة ملف Excel: {exc}") from exc

    updated = copy.deepcopy(schema)
    changed = False
    sheet_labels: dict[str, dict[str, str]] = {}
    try:
        category_sheets = related_sheet_names(updated)
        for category in updated["categories"]:
            if category["kind"] == "main":
                if MAIN_SHEET not in workbook.sheetnames:
                    continue
                worksheet = workbook[MAIN_SHEET]
            else:
                sheet_name = category_sheets[category["id"]]
                if sheet_name not in workbook.sheetnames:
                    continue
                worksheet = workbook[sheet_name]

            if worksheet.title not in sheet_labels:
                # A cell lookup on a read-only worksheet reparses its XML.
                # Read both header rows once, including shared main categories.
                rows = list(worksheet.iter_rows(
                    min_row=TECHNICAL_HEADER_ROW,
                    max_row=VISIBLE_HEADER_ROW,
                    values_only=True,
                ))
                headers = rows[0] if rows else ()
                labels = rows[1] if len(rows) > 1 else ()
                sheet_labels[worksheet.title] = {
                    clean_text(header): clean_text(label)
                    for header, label in zip(headers, labels)
                    if clean_text(header)
                }
            labels_by_id = sheet_labels[worksheet.title]
            for field in data_fields(category):
                excel_label = labels_by_id.get(field["id"], "")
                if excel_label and excel_label != field["label"]:
                    field["label"] = excel_label
                    changed = True
    finally:
        workbook.close()

    if not changed:
        return schema

    updated["revision"] = schema["revision"] + 1
    updated = validate_schema(updated)
    atomic_write_json(_schema_path(), updated)
    return updated


def read_schema_with_excel_labels() -> dict[str, Any]:
    with WORKBOOK_LOCK:
        return _sync_excel_field_labels_unlocked(read_schema_file())


def _style_headers(worksheet, column_count: int) -> None:
    worksheet.row_dimensions[TECHNICAL_HEADER_ROW].hidden = True
    worksheet.row_dimensions[VISIBLE_HEADER_ROW].height = 28
    worksheet.freeze_panes = f"A{FIRST_DATA_ROW}"
    worksheet.sheet_view.showGridLines = False
    for column in range(1, column_count + 1):
        cell = worksheet.cell(VISIBLE_HEADER_ROW, column)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.column_dimensions[get_column_letter(column)].width = max(
            13, min(34, len(clean_text(cell.value)) + 4)
        )


def _add_field_validations(
    worksheet,
    field_columns: dict[str, int],
    fields: Iterable[dict[str, Any]],
) -> None:
    for field in fields:
        if field["id"] not in field_columns:
            continue
        options = option_labels(field, active_only=True)
        if field["type"] not in {"select", "yes_no"} or not options:
            continue
        escaped = ",".join(option.replace('"', '""') for option in options)
        if len(escaped) > 250:
            continue
        validation = DataValidation(
            type="list",
            formula1=f'"{escaped}"',
            allow_blank=True,
        )
        worksheet.add_data_validation(validation)
        column_letter = get_column_letter(field_columns[field["id"]])
        validation.add(f"{column_letter}{FIRST_DATA_ROW}:{column_letter}1048576")


def _new_workbook_structure(
    schema: dict[str, Any],
) -> tuple[Workbook, dict[str, Any]]:
    workbook = Workbook()
    main_sheet = workbook.active
    main_sheet.title = MAIN_SHEET

    fields = main_fields(schema)
    main_headers = [*MAIN_INTERNAL_HEADERS, *(field["id"] for field in fields)]
    main_labels = [
        "المعرّف الداخلي",
        "ID",
        "تاريخ الإنشاء",
        "تاريخ التعديل",
        "مؤرشف",
        "تاريخ الأرشفة",
        *(field["label"] for field in fields),
    ]
    main_sheet.append(main_headers)
    main_sheet.append(main_labels)
    _style_headers(main_sheet, len(main_headers))
    main_columns = {
        header: index for index, header in enumerate(main_headers, start=1)
    }
    _add_field_validations(main_sheet, main_columns, fields)

    sheets: dict[str, Any] = {"main": main_sheet, "related": {}}
    names = related_sheet_names(schema)
    for category in schema["categories"]:
        if category["kind"] != "repeatable":
            continue
        worksheet = workbook.create_sheet(names[category["id"]])
        headers = [
            *RELATED_INTERNAL_HEADERS,
            RELATED_LINK_HEADER,
            RELATED_PARENT_HEADER,
            *(field["id"] for field in data_fields(category)),
        ]
        labels = [
            "معرّف الصف الداخلي",
            "معرّف السجل الداخلي",
            "ID",
            "minor_id",
            "تاريخ الإنشاء",
            "تاريخ التعديل",
            "ID الشخص المرتبط",
            "معرّف بطاقة الفئة الأم",
            *(field["label"] for field in data_fields(category)),
        ]
        worksheet.append(headers)
        worksheet.append(labels)
        _style_headers(worksheet, len(headers))
        columns = {header: index for index, header in enumerate(headers, start=1)}
        _add_field_validations(worksheet, columns, data_fields(category))
        sheets["related"][category["id"]] = worksheet

    meta = workbook.create_sheet(META_SHEET)
    meta.sheet_state = "hidden"
    meta.append(["schema_version", SCHEMA_VERSION])
    meta.append(["schema_revision", schema["revision"]])
    meta.append(["category_id", "sheet_name"])
    for category_id, sheet_name in names.items():
        meta.append([category_id, sheet_name])
    return workbook, sheets


def _apply_date_formats(
    worksheet,
    row_index: int,
    field_columns: dict[str, int],
    fields: Iterable[dict[str, Any]],
) -> None:
    for field in fields:
        if field["type"] == "date_gregorian" and field["id"] in field_columns:
            worksheet.cell(
                row=row_index, column=field_columns[field["id"]]
            ).number_format = EXCEL_DATE_FORMAT
        elif field["type"] == "number" and field["id"] in field_columns:
            cell = worksheet.cell(row=row_index, column=field_columns[field["id"]])
            if number_is_text(field):
                cell.number_format = "@"
            elif field.get("number_behavior", {}).get("format_thousands"):
                cell.number_format = "#,##0.################"


def write_dataset_workbook(
    schema: dict[str, Any],
    records: list[dict[str, Any]],
    destination: Path,
) -> None:
    workbook, sheets = _new_workbook_structure(schema)
    try:
        main_sheet = sheets["main"]
        main_headers = technical_headers(main_sheet)
        main_columns = {
            header: index for index, header in enumerate(main_headers, start=1)
        }
        fields = main_fields(schema)

        for record in records:
            row = [
                record["_record_id"],
                record["record_code"],
                record["created_at"],
                record["updated_at"],
                "نعم" if record.get("archived") else "لا",
                record.get("archived_at", ""),
                *[
                    excel_value(record.get("values", {}).get(field["id"], ""), field)
                    for field in fields
                ],
            ]
            main_sheet.append(row)
            _apply_date_formats(
                main_sheet, main_sheet.max_row, main_columns, fields
            )

        for category in schema["categories"]:
            if category["kind"] != "repeatable":
                continue
            worksheet = sheets["related"][category["id"]]
            headers = technical_headers(worksheet)
            columns = {
                header: index for index, header in enumerate(headers, start=1)
            }
            for record in records:
                rows = record.get("related", {}).get(category["id"], [])
                for sequence, child in enumerate(rows, start=1):
                    row = [
                        child["_child_id"],
                        record["_record_id"],
                        record["record_code"],
                        sequence,
                        child["created_at"],
                        child["updated_at"],
                        child.get("linked_record_code", ""),
                        child.get("parent_child_id", ""),
                        *[
                            excel_value(
                                child.get("values", {}).get(field["id"], ""),
                                field,
                            )
                            for field in data_fields(category)
                        ],
                    ]
                    worksheet.append(row)
                    _apply_date_formats(
                        worksheet,
                        worksheet.max_row,
                        columns,
                        data_fields(category),
                    )

        for worksheet in workbook.worksheets:
            if worksheet.title == META_SHEET:
                continue
            last_row = max(VISIBLE_HEADER_ROW, worksheet.max_row)
            last_column = max(1, worksheet.max_column)
            worksheet.auto_filter.ref = (
                f"A{VISIBLE_HEADER_ROW}:"
                f"{get_column_letter(last_column)}{last_row}"
            )
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
    finally:
        workbook.close()


def atomic_write_workbook(
    schema: dict[str, Any], records: list[dict[str, Any]]
) -> None:
    workbook_path = _workbook_path()
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=".database-",
        suffix=".xlsx",
        dir=workbook_path.parent,
        delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        started = time.perf_counter()
        write_dataset_workbook(schema, records, temporary_path)
        os.replace(temporary_path, workbook_path)
        invalidate_dataset_cache()
        LOGGER.info(
            "Workbook written: %.3f seconds; records=%d",
            time.perf_counter() - started,
            len(records),
        )
        try:
            synchronize_profile_dependents(current_schema_id(), records)
        except Exception:
            # The source commit is already durable. Never roll back its attachments
            # because another workbook is locked; retry from the persisted baseline.
            LOGGER.exception("Profile synchronization pending after source workbook commit")
            schedule_profile_sync_retry()
    finally:
        temporary_path.unlink(missing_ok=True)


_PROFILE_SYNC_RETRY_THREAD: threading.Thread | None = None
_PROFILE_SYNC_RETRY_LOCK = threading.Lock()


def synchronize_profile_dependents(source_id: str, source_records: list[dict[str, Any]]) -> None:
    """Apply source membership deltas, preserving independent target values/deletions."""
    manager = WORKSPACE_MANAGER
    if manager is None:
        return
    codes = {record["record_code"] for record in source_records}
    with WORKBOOK_LOCK:
        for target, seen in manager.profile_dependents(source_id):
            added, removed = codes - seen, seen - codes
            if not added and not removed:
                continue
            with use_context(target):
                schema = read_schema_file()
                if not target.workbook_path.exists():
                    ensure_storage()
                records = list(_dataset_snapshot_unlocked(schema).records)
                existing = {record["record_code"] for record in records}
                for code in sorted(removed & existing):
                    delete_record(code, synchronized=True)
                records = list(_dataset_snapshot_unlocked(schema).records)
                existing = {record["record_code"] for record in records}
                timestamp = now_iso()
                for code in sorted(added - existing):
                    records.append({
                        "_record_id": new_internal_id(), "record_code": code,
                        "created_at": timestamp, "updated_at": timestamp,
                        "archived": False, "archived_at": "", "values": {}, "related": {},
                    })
                if added - existing:
                    atomic_write_workbook(schema, records)
                    _safe_publish_dataset_snapshot(schema, records)
                # Register all present additions, including a replay after a crash.
                manager.register_profiles(sorted(added & {record["record_code"] for record in records}), target.schema_id)
                manager.mark_profile_source_seen(target.schema_id, codes)


def synchronize_workspace_profiles() -> None:
    manager = WORKSPACE_MANAGER
    if manager is None:
        return
    with WORKBOOK_LOCK:
        for context in manager.contexts():
            if not manager.profile_dependents(context.schema_id):
                continue
            with use_context(context):
                schema = read_schema_file()
                synchronize_profile_dependents(context.schema_id, list(_dataset_snapshot_unlocked(schema).records))


def schedule_profile_sync_retry() -> None:
    """Retry locked downstream workbooks; restart recovery uses the same baseline."""
    global _PROFILE_SYNC_RETRY_THREAD
    with _PROFILE_SYNC_RETRY_LOCK:
        if _PROFILE_SYNC_RETRY_THREAD and _PROFILE_SYNC_RETRY_THREAD.is_alive():
            return
        def retry() -> None:
            while True:
                time.sleep(3)
                try:
                    synchronize_workspace_profiles()
                    return
                except Exception:
                    LOGGER.exception("Profile synchronization still pending")
        _PROFILE_SYNC_RETRY_THREAD = threading.Thread(target=retry, daemon=True, name="profile-sync-retry")
        _PROFILE_SYNC_RETRY_THREAD.start()


def _temporary_copy(path: Path, prefix: str) -> Path:
    with tempfile.NamedTemporaryFile(
        prefix=prefix,
        suffix=path.suffix,
        dir=path.parent,
        delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        shutil.copy2(path, temporary_path)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise
    return temporary_path


def _replace_storage_pair(workbook_temp: Path, schema_temp: Path) -> None:
    """Commit workbook and schema together, rolling the workbook back on failure."""
    workbook_path = _workbook_path()
    schema_path = _schema_path()
    workbook_backup = _temporary_copy(
        workbook_path,
        ".database-rollback-",
    )
    workbook_replaced = False
    try:
        os.replace(workbook_temp, workbook_path)
        workbook_replaced = True
        os.replace(schema_temp, schema_path)
        invalidate_dataset_cache()
    except Exception:
        if workbook_replaced:
            try:
                os.replace(workbook_backup, workbook_path)
                workbook_backup = None
            except Exception as recovery_exc:
                raise ApplicationError(
                    "تعذّر إكمال تحديث التصميم وتعذّرت الاستعادة التلقائية. "
                    "لا تفتح التطبيق مجددًا قبل استعادة آخر نسخة احتياطية."
                ) from recovery_exc
        raise
    finally:
        workbook_temp.unlink(missing_ok=True)
        schema_temp.unlink(missing_ok=True)
        if workbook_backup is not None:
            workbook_backup.unlink(missing_ok=True)


def ensure_storage() -> None:
    schema_path = _schema_path()
    workbook_path = _workbook_path()
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not schema_path.is_file():
        atomic_write_json(schema_path, default_schema())
    schema = read_schema_file()
    if not workbook_path.is_file():
        atomic_write_workbook(schema, [])
    elif workbook_schema_revision(workbook_path) not in {None, schema["revision"]}:
        # Complete an interrupted background projection during cold launch.
        records = read_dataset_unlocked(schema)
        atomic_write_workbook(schema, records)
    attachments_directory().mkdir(parents=True, exist_ok=True)


def _write_schema_recovery_backup(task: dict[str, Any]) -> dict[str, str]:
    """Write a pre-edit schema recovery package without blocking the request."""

    context: SchemaContext = task["context"]
    old_schema = task.get("recovery_schema")
    old_records = task.get("recovery_records")
    if not old_schema or old_records is None:
        return {}
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    hide_packaged_support_paths()
    stamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S-%f")
    filename = f"GenericSchemaCraft-auto-backup-{stamp}.zip"
    destination = BACKUP_DIR / filename
    with tempfile.NamedTemporaryFile(
        prefix=".recovery-workbook-", suffix=".xlsx", dir=BACKUP_DIR, delete=False
    ) as temporary:
        workbook_temp = Path(temporary.name)
    with tempfile.NamedTemporaryFile(
        prefix=".recovery-package-", suffix=".zip", dir=BACKUP_DIR, delete=False
    ) as temporary:
        package_temp = Path(temporary.name)
    try:
        with use_context(context):
            write_dataset_workbook(old_schema, list(old_records), workbook_temp)
            with zipfile.ZipFile(package_temp, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr(
                    "recovery/manifest.json",
                    json.dumps(
                        {
                            "schema_id": context.schema_id,
                            "schema_name": context.name,
                            "schema_revision": old_schema["revision"],
                            "created_at": iso_now(),
                        },
                        ensure_ascii=False,
                        indent=2,
                    ) + "\n",
                )
                archive.writestr(
                    "recovery/schema.json",
                    json.dumps(old_schema, ensure_ascii=False, indent=2) + "\n",
                )
                archive.write(workbook_temp, "recovery/database.xlsx")
                for stored_path in sorted(task.get("recovery_paths") or set()):
                    source = attachment_absolute_path(stored_path)
                    if source is None or not source.is_file():
                        continue
                    relative = source.relative_to(context.attachments_path)
                    archive.write(source, (Path("recovery/attachments") / relative).as_posix())
        os.replace(package_temp, destination)
        automatic_backups = sorted(
            BACKUP_DIR.glob("GenericSchemaCraft-auto-backup-*.zip"),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        for old in automatic_backups[10:]:
            old.unlink(missing_ok=True)
        return {"filename": filename, "download_url": f"/api/backups/{filename}"}
    finally:
        workbook_temp.unlink(missing_ok=True)
        package_temp.unlink(missing_ok=True)


def _project_schema_workbook(task: dict[str, Any]) -> None:
    """Build Excel outside the shared lock, then atomically publish if current."""

    context: SchemaContext = task["context"]
    expected_revision = int(task["expected_revision"])
    workbook_path = context.workbook_path
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=".database-projection-",
        suffix=".xlsx",
        dir=workbook_path.parent,
        delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with use_context(context):
            with WORKBOOK_LOCK:
                current = read_schema_file()
                if current["revision"] != expected_revision:
                    return
                records = list(_dataset_snapshot_unlocked(current).records)
            started = time.perf_counter()
            write_dataset_workbook(current, records, temporary_path)
            with WORKBOOK_LOCK:
                latest = read_schema_file()
                if latest["revision"] != expected_revision:
                    return
                os.replace(temporary_path, workbook_path)
                # Keep the already-indexed records hot. The old implementation
                # invalidated this cache after every projection, so the next
                # Builder edit re-read and re-indexed the complete workbook.
                _safe_publish_dataset_snapshot(latest, records)
            LOGGER.info(
                "Background workbook projection completed in %.3f seconds; schema=%s records=%d",
                time.perf_counter() - started,
                context.schema_id,
                len(records),
            )
    finally:
        temporary_path.unlink(missing_ok=True)


def _workbook_sync_loop() -> None:
    global _WORKBOOK_SYNC_WORKER
    while True:
        time.sleep(_WORKBOOK_SYNC_DEBOUNCE_SECONDS)
        with _WORKBOOK_SYNC_CONDITION:
            if not _WORKBOOK_SYNC_PENDING:
                _WORKBOOK_SYNC_WORKER = None
                return
            tasks = list(_WORKBOOK_SYNC_PENDING.values())
            _WORKBOOK_SYNC_PENDING.clear()
        for task in tasks:
            recovery_ready = not bool(task.get("recovery_schema"))
            try:
                recovery_result = _write_schema_recovery_backup(task)
                recovery_ready = recovery_ready or bool(recovery_result)
            except Exception:
                LOGGER.exception(
                    "Background schema recovery backup failed: schema=%s",
                    task["context"].schema_id,
                )
            try:
                _project_schema_workbook(task)
            except Exception:
                LOGGER.exception(
                    "Background workbook schema sync failed: schema=%s revision=%s",
                    task["context"].schema_id,
                    task["expected_revision"],
                )
            if task.get("cleanup_paths") and recovery_ready:
                with use_context(task["context"]):
                    remove_attachment_files(set(task["cleanup_paths"]))


def schedule_workbook_schema_sync(
    context: SchemaContext,
    expected_revision: int,
    *,
    recovery_schema: dict[str, Any] | None = None,
    recovery_records: Iterable[dict[str, Any]] | None = None,
    recovery_paths: set[str] | None = None,
    cleanup_paths: set[str] | None = None,
) -> None:
    """Coalesce recovery/projection work and keep it outside HTTP requests."""

    global _WORKBOOK_SYNC_WORKER
    workbook_key = str(context.workbook_path.resolve())
    with _WORKBOOK_SYNC_CONDITION:
        existing = _WORKBOOK_SYNC_PENDING.get(workbook_key, {})
        task = {
            "context": context,
            "expected_revision": expected_revision,
            "recovery_schema": existing.get("recovery_schema") or recovery_schema,
            "recovery_records": existing.get("recovery_records") or recovery_records,
            "recovery_paths": existing.get("recovery_paths") or recovery_paths or set(),
            "cleanup_paths": set(existing.get("cleanup_paths") or set()) | set(cleanup_paths or set()),
        }
        _WORKBOOK_SYNC_PENDING[workbook_key] = task
        if _WORKBOOK_SYNC_WORKER is None or not _WORKBOOK_SYNC_WORKER.is_alive():
            _WORKBOOK_SYNC_WORKER = threading.Thread(
                target=_workbook_sync_loop,
                name="schema-workbook-maintenance",
                daemon=True,
            )
            _WORKBOOK_SYNC_WORKER.start()


def initialize_workspace() -> None:
    """Enable Release 3 storage while leaving direct legacy tests compatible."""
    global WORKSPACE_MANAGER, GLOBAL_DEFINITIONS, EXPORT_HISTORY, IMPORT_HISTORY
    global SEARCH_HISTORY, AUDIT_USERS, WORKSPACE_SETTINGS_PATH, IMPORT_ARCHIVE_DIR
    manager = WorkspaceManager(DATA_DIR)
    try:
        context = manager.initialize(default_schema())
    except WorkspaceError as exc:
        raise ApplicationError(str(exc)) from exc
    WORKSPACE_MANAGER = manager
    GLOBAL_DEFINITIONS = GlobalDefinitionStore(DATA_DIR)
    EXPORT_HISTORY = ExportHistoryStore(DATA_DIR)
    IMPORT_HISTORY = ImportHistoryStore(DATA_DIR)
    SEARCH_HISTORY = SearchHistoryStore(DATA_DIR)
    AUDIT_USERS = AuditUserStore(DATA_DIR)
    WORKSPACE_SETTINGS_PATH = DATA_DIR / "workspace-settings.json"
    IMPORT_ARCHIVE_DIR = DATA_DIR / "imported-source-files"
    GLOBAL_DEFINITIONS.read()
    with use_context(context):
        ensure_storage()
    try:
        synchronize_workspace_profiles()
    except Exception:
        LOGGER.exception("Profile synchronization pending on startup")
        schedule_profile_sync_retry()
    hide_packaged_support_paths()


def require_workspace() -> WorkspaceManager:
    if WORKSPACE_MANAGER is None:
        raise ApplicationError("مساحة العمل متعددة التصاميم غير مفعلة.")
    return WORKSPACE_MANAGER


def workspace_response() -> dict[str, Any]:
    manager = require_workspace()
    response = manager.response()
    definitions: dict[str, dict[str, Any]] = {}
    for context in manager.contexts(include_archived=True):
        with use_context(context):
            schema = read_schema_with_excel_labels()
            definition = copy.deepcopy(schema)
            definition["builder_access"] = builder_access_response()
            definition["developer_mode"] = DEVELOPER_MODE
            definition["schema_id"] = context.schema_id
            definition["schema_name"] = context.name
            definition["schema_archived"] = context.archived
            definitions[context.schema_id] = definition
    response["definitions"] = definitions
    response["builder_access"] = builder_access_response()
    response["global_definitions"] = (
        GLOBAL_DEFINITIONS.response() if GLOBAL_DEFINITIONS is not None else {}
    )
    response["workspace_settings"] = read_workspace_settings()
    response["audit_users"] = AUDIT_USERS.read() if AUDIT_USERS is not None else {
        "version": 1,
        "current_user": "",
        "users": [],
    }
    return response


def read_workspace_settings() -> dict[str, Any]:
    defaults = {
        "shortcuts": copy.deepcopy(DEFAULT_WORKSPACE_SHORTCUT_BINDINGS),
        "entry_history_limit": 8,
        "builder_history_limit": 20,
        "search_history_limit": 20,
        "import_history_limit": 20,
        "export_history_limit": 20,
        "home_entry_history_limit": 3,
        "home_builder_history_limit": 3,
        "home_search_history_limit": 3,
        "home_import_history_limit": 3,
        "home_export_history_limit": 3,
        "show_explanations": False,
        "home_custom_stats": [],
        "background_all_pages": False,
        "background_image_name": "",
        "background_image_custom": False,
        "background_image_version": "",
        "background_image_mime": "",
        "background_image_id": "default",
        "background_images": [],
    }
    payload: dict[str, Any] = {}
    if WORKSPACE_SETTINGS_PATH.is_file():
        try:
            loaded_payload = json.loads(
                WORKSPACE_SETTINGS_PATH.read_text(encoding="utf-8")
            )
            if isinstance(loaded_payload, dict):
                payload = loaded_payload
        except (OSError, json.JSONDecodeError):
            pass
    raw = payload.get("shortcuts")
    if not isinstance(raw, dict):
        raw = defaults["shortcuts"]
    shortcuts: dict[str, str | list[str]] = {}
    for raw_key, raw_action in raw.items():
        key = clean_text(raw_key)
        raw_actions = raw_action if isinstance(raw_action, list) else [raw_action]
        actions: list[str] = []
        for raw_item in raw_actions:
            action = clean_text(raw_item)
            if key == "Ctrl+Alt+U" and action == "exit_admin":
                action = "select_user"
            if action in APP_SHORTCUTS and action not in actions:
                actions.append(action)
        if key and actions:
            shortcuts[key] = actions[0] if len(actions) == 1 else actions
    bound_actions = {
        action
        for value in shortcuts.values()
        for action in (value if isinstance(value, list) else [value])
    }
    for key, action in DEFAULT_WORKSPACE_SHORTCUT_BINDINGS.items():
        if action in {"close_application", "enter_admin", "select_user"} and action not in bound_actions and key not in shortcuts:
            shortcuts[key] = action
    result = copy.deepcopy(defaults)
    primary_color = clean_text(payload.get("primary_color"))
    result["primary_color"] = primary_color if re.fullmatch(r"#[0-9a-fA-F]{6}", primary_color) else "#1F5F95"
    result["shortcuts"] = shortcuts or defaults["shortcuts"]
    legacy_operation_limit = payload.get("operation_history_limit", 20)
    legacy_home_limits = [
        payload.get("home_search_limit"), payload.get("home_entry_recent_limit"),
        payload.get("home_import_limit"), payload.get("home_export_limit"),
    ]
    migrated_home_limit = payload.get("home_history_limit")
    if migrated_home_limit is None:
        migrated_home_limit = next((value for value in legacy_home_limits if value is not None), 3)
    migrated_defaults = {
        "entry_history_limit": payload.get("home_entry_recent_limit", 8),
        "builder_history_limit": legacy_operation_limit,
        "search_history_limit": legacy_operation_limit,
        "import_history_limit": legacy_operation_limit,
        "export_history_limit": legacy_operation_limit,
        "home_entry_history_limit": migrated_home_limit,
        "home_builder_history_limit": migrated_home_limit,
        "home_search_history_limit": migrated_home_limit,
        "home_import_history_limit": migrated_home_limit,
        "home_export_history_limit": migrated_home_limit,
    }
    for key in migrated_defaults:
        try:
            raw_limit = payload.get(key, migrated_defaults[key])
            if clean_text(raw_limit).casefold() == "all":
                raw_limit = 100
            result[key] = max(1, min(100, int(raw_limit)))
        except (TypeError, ValueError):
            result[key] = defaults[key]
    result["show_explanations"] = bool(payload.get("show_explanations", False))
    result["home_custom_stats"] = normalize_home_custom_stats(
        payload.get("home_custom_stats", [])
    )
    result["background_all_pages"] = bool(
        payload.get("background_all_pages", False)
    )
    raw_images = payload.get("background_images")
    if not isinstance(raw_images, list):
        raw_images = []
    images: list[dict[str, str]] = []
    for raw_image in raw_images[:3]:
        if not isinstance(raw_image, dict):
            continue
        image_id = clean_text(raw_image.get("id"))
        image_path = workspace_background_image_file(image_id)
        if not re.fullmatch(r"[a-f0-9]{16}", image_id) or not image_path.is_file():
            continue
        try:
            with image_path.open("rb") as background_file:
                detected_mime = detect_background_image_mime(background_file.read(16))
            if detected_mime:
                images.append({
                    "id": image_id,
                    "name": clean_text(raw_image.get("name"))[:200] or "صورة مخصصة",
                    "mime": detected_mime,
                    "version": str(image_path.stat().st_mtime_ns),
                })
        except OSError:
            pass
    legacy_path = workspace_background_image_path()
    if not images and legacy_path.is_file():
        try:
            with legacy_path.open("rb") as background_file:
                detected_mime = detect_background_image_mime(background_file.read(16))
            if detected_mime:
                images.append({"id": "legacy", "name": clean_text(payload.get("background_image_name"))[:200] or "صورة مخصصة", "mime": detected_mime, "version": str(legacy_path.stat().st_mtime_ns)})
        except OSError:
            pass
    selected_id = clean_text(payload.get("background_image_id")) or (images[0]["id"] if images else "default")
    selected = next((item for item in images if item["id"] == selected_id), None)
    result["background_images"] = images
    result["background_image_id"] = selected["id"] if selected else "default"
    result["background_image_custom"] = selected is not None
    result["background_image_name"] = selected["name"] if selected else ""
    result["background_image_mime"] = selected["mime"] if selected else ""
    result["background_image_version"] = selected["version"] if selected else ""
    return result


def normalize_home_custom_stats(value: Any) -> list[dict[str, Any]]:
    """Validate the three persisted cross-schema Home statistic definitions."""

    if not isinstance(value, list):
        return []
    result: list[dict[str, Any]] = []
    for raw_item in value[:3]:
        if not isinstance(raw_item, dict):
            continue
        identifier = clean_text(raw_item.get("id"))[:160]
        label = clean_text(raw_item.get("label"))[:160]
        calculation = clean_text(raw_item.get("calculation"))
        if not identifier or not label:
            continue
        if calculation not in {"filled", "checked", "distinct"}:
            calculation = "filled"
        selections: list[dict[str, Any]] = []
        raw_selections = raw_item.get("selections")
        if isinstance(raw_selections, list):
            for raw_selection in raw_selections[:100]:
                if not isinstance(raw_selection, dict):
                    continue
                schema_id = clean_text(raw_selection.get("schemaId"))[:160]
                field_id = clean_text(raw_selection.get("fieldId"))[:160]
                if not schema_id or not field_id:
                    continue
                raw_values = raw_selection.get("values")
                values = []
                if isinstance(raw_values, list):
                    for raw_value in raw_values[:100]:
                        selected_value = clean_text(raw_value)[:500]
                        if selected_value and selected_value not in values:
                            values.append(selected_value)
                selections.append({
                    "schemaId": schema_id,
                    "fieldId": field_id,
                    "values": values,
                })
        if selections:
            result.append({
                "id": identifier,
                "label": label,
                "calculation": calculation,
                "selections": selections,
            })
    return result


def save_home_custom_stats(payload: Any) -> dict[str, Any]:
    """Persist administrator-defined Home statistics for every app session."""

    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("إعدادات إحصاءات الرئيسية غير صالحة.")
    settings = read_workspace_settings()
    settings["home_custom_stats"] = normalize_home_custom_stats(
        payload.get("home_custom_stats", [])
    )
    atomic_write_json(WORKSPACE_SETTINGS_PATH, settings)
    return {
        "ok": True,
        "home_custom_stats": copy.deepcopy(settings["home_custom_stats"]),
    }


def workspace_background_image_path() -> Path:
    return WORKSPACE_SETTINGS_PATH.parent / "workspace-background-image"


def workspace_background_images_dir() -> Path:
    return WORKSPACE_SETTINGS_PATH.parent / "workspace-backgrounds"


def workspace_background_image_file(image_id: str) -> Path:
    return workspace_background_images_dir() / image_id


def detect_background_image_mime(content: bytes) -> str:
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if content.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if content.startswith((b"GIF87a", b"GIF89a")):
        return "image/gif"
    if len(content) >= 12 and content.startswith(b"RIFF") and content[8:12] == b"WEBP":
        return "image/webp"
    return ""


def save_workspace_settings(payload: Any) -> dict[str, Any]:
    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("إعدادات مساحة العمل غير صالحة.")
    existing = read_workspace_settings()
    raw_shortcuts = payload.get("shortcuts", existing["shortcuts"])
    if not isinstance(raw_shortcuts, dict):
        raise ApplicationError("إعدادات اختصارات لوحة المفاتيح غير صالحة.")
    shortcuts: dict[str, str | list[str]] = {}
    for raw_key, raw_action in raw_shortcuts.items():
        key = clean_text(raw_key)
        raw_actions = raw_action if isinstance(raw_action, list) else [raw_action]
        actions: list[str] = []
        for raw_item in raw_actions:
            action = clean_text(raw_item)
            if action in APP_SHORTCUTS and action not in actions:
                actions.append(action)
        if not key or not actions:
            raise ApplicationError("أحد اختصارات لوحة المفاتيح غير صالح.")
        shortcuts[key] = actions[0] if len(actions) == 1 else actions
    settings = copy.deepcopy(existing)
    settings["shortcuts"] = shortcuts
    deleted_background_path: Path | None = None
    delete_background_id = clean_text(payload.get("delete_background_image_id"))
    if delete_background_id:
        available_images = {
            clean_text(item.get("id")): item
            for item in existing.get("background_images", [])
            if isinstance(item, dict)
        }
        if delete_background_id not in available_images:
            raise ApplicationError("صورة الخلفية المطلوب حذفها لم تعد متاحة.")
        settings["background_images"] = [
            item
            for item in existing.get("background_images", [])
            if isinstance(item, dict) and clean_text(item.get("id")) != delete_background_id
        ]
        if clean_text(settings.get("background_image_id")) == delete_background_id:
            settings["background_image_id"] = "default"
        deleted_background_path = (
            workspace_background_image_path()
            if delete_background_id == "legacy"
            else workspace_background_image_file(delete_background_id)
        )
    for key in (
        "entry_history_limit",
        "builder_history_limit",
        "search_history_limit",
        "import_history_limit",
        "export_history_limit",
        "home_entry_history_limit",
        "home_builder_history_limit",
        "home_search_history_limit",
        "home_import_history_limit",
        "home_export_history_limit",
    ):
        if key in payload:
            try:
                settings[key] = max(1, min(100, int(payload[key])))
            except (TypeError, ValueError) as exc:
                raise ApplicationError("عدد السجلات الظاهرة يجب أن يكون بين 1 و100.") from exc
    if "show_explanations" in payload:
        settings["show_explanations"] = bool(payload["show_explanations"])
    if "background_all_pages" in payload:
        settings["background_all_pages"] = bool(payload["background_all_pages"])
    if "background_image_data" in payload:
        encoded = payload.get("background_image_data")
        if not isinstance(encoded, str) or not encoded:
            raise ApplicationError("بيانات صورة الخلفية غير مكتملة.")
        try:
            content = base64.b64decode(encoded, validate=True)
        except (ValueError, binascii.Error) as exc:
            raise ApplicationError("تعذّر قراءة صورة الخلفية.") from exc
        if not content or len(content) > MAX_BACKGROUND_IMAGE_BYTES:
            raise ApplicationError("يجب ألا يتجاوز حجم صورة الخلفية 15 ميغابايت.")
        detected_mime = detect_background_image_mime(content[:16])
        if not detected_mime:
            raise ApplicationError("صيغة صورة الخلفية غير مدعومة.")
        image_id = hashlib.sha256(content).hexdigest()[:16]
        image_path = workspace_background_image_file(image_id)
        image_path.parent.mkdir(parents=True, exist_ok=True)
        atomic_write_bytes(image_path, content)
        # Keep the alpha.11 current-image path as a compatibility mirror for
        # existing integrations while the gallery retains the three originals.
        atomic_write_bytes(workspace_background_image_path(), content)
        settings["background_image_mime"] = detected_mime
        image_name = sanitize_filename_text(payload.get("background_image_name"))[:200] or "صورة مخصصة"
        previous_images = [item for item in settings.get("background_images", []) if isinstance(item, dict) and clean_text(item.get("id")) not in {image_id, "legacy"}]
        settings["background_images"] = [{"id": image_id, "name": image_name, "mime": detected_mime}, *previous_images][:3]
        settings["background_image_id"] = image_id
        keep_ids = {clean_text(item.get("id")) for item in settings["background_images"]}
        if workspace_background_images_dir().is_dir():
            for candidate in workspace_background_images_dir().iterdir():
                if candidate.is_file() and candidate.name not in keep_ids:
                    candidate.unlink(missing_ok=True)
    elif "background_image_id" in payload:
        selected_id = clean_text(payload.get("background_image_id")) or "default"
        available = {clean_text(item.get("id")) for item in settings.get("background_images", []) if isinstance(item, dict)}
        if selected_id != "default" and selected_id not in available:
            raise ApplicationError("صورة الخلفية المختارة لم تعد متاحة.")
        settings["background_image_id"] = selected_id
    settings.pop("background_image_custom", None)
    settings.pop("background_image_version", None)
    atomic_write_json(WORKSPACE_SETTINGS_PATH, settings)
    if deleted_background_path is not None:
        deleted_background_path.unlink(missing_ok=True)
    return {"ok": True, "workspace_settings": read_workspace_settings()}


DEFAULT_APP_HISTORY_FILES = {
    "audit-users.json",
    "export-history.json",
    "import-history.json",
    "search-history.json",
}
DEFAULT_APP_HISTORY_DIRECTORIES = {
    "exported-files",
    "imported-source-files",
}
DEFAULT_APP_COPY_EXCLUSIONS = {
    ".git",
    ".pytest_cache",
    ".ruff_cache",
    "__pycache__",
    "backups",
    "node_modules",
    "release",
    "startup-error.log",
    "developer-access.key",
}


def choose_default_app_parent() -> dict[str, Any]:
    """Choose the parent folder for a sanitized application copy."""

    require_builder_access()
    try:
        selected = choose_directory("اختر مجلد إنشاء التطبيق الافتراضي")
    except AdvancedFeatureError as exc:
        raise ApplicationError(str(exc)) from exc
    return {"ok": True, "cancelled": selected is None, "destination": str(selected or "")}


def _default_app_copy_ignore(_directory: str, names: list[str]) -> set[str]:
    return {
        name
        for name in names
        if name in DEFAULT_APP_COPY_EXCLUSIONS
        or name.startswith(".build-")
        or name.endswith(".pyc")
    }


def _remove_cloned_path(path: Path) -> None:
    if path.is_dir():
        shutil.rmtree(path)
    else:
        path.unlink(missing_ok=True)


def _clear_cloned_history(data_dir: Path) -> None:
    for name in DEFAULT_APP_HISTORY_FILES:
        (data_dir / name).unlink(missing_ok=True)
    for name in DEFAULT_APP_HISTORY_DIRECTORIES:
        _remove_cloned_path(data_dir / name)


def _clear_cloned_records(data_dir: Path) -> None:
    """Remove every record repository while preserving schema definitions."""

    _clear_cloned_history(data_dir)
    for name in ("database.xlsx", "identity-registry.xlsx"):
        (data_dir / name).unlink(missing_ok=True)
    _remove_cloned_path(data_dir / "attachments")
    _remove_cloned_path(data_dir / "release-2-original")
    schemas_dir = data_dir / "schemas"
    if schemas_dir.is_dir():
        for schema_dir in schemas_dir.iterdir():
            if not schema_dir.is_dir():
                continue
            for workbook in schema_dir.glob("*.xlsx"):
                workbook.unlink(missing_ok=True)
            attachments = schema_dir / "attachments"
            _remove_cloned_path(attachments)
            attachments.mkdir(parents=True, exist_ok=True)


def create_default_app_copy(payload: Any) -> dict[str, Any]:
    """Create a guarded copy of the running app with optional data removal."""

    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("خيارات إنشاء التطبيق الافتراضي غير صالحة.")
    raw_destination = clean_text(payload.get("destination"))
    if not raw_destination:
        raise ApplicationError("اختر مجلد إنشاء التطبيق الجديد أولًا.")
    destination = Path(raw_destination).expanduser().resolve()
    if not destination.is_dir():
        raise ApplicationError("مجلد إنشاء التطبيق الجديد غير موجود.")
    source = BASE_DIR.resolve()
    if destination == source or source in destination.parents:
        raise ApplicationError("اختر مجلدًا خارج مجلد التطبيق الحالي.")

    clear_schema = bool(payload.get("clear_schema"))
    clear_records = clear_schema or bool(payload.get("clear_records"))
    clear_history = clear_records or bool(payload.get("clear_history"))
    timestamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S")
    target = destination / f"SchemaCraft-Default-{timestamp}"
    counter = 2
    while target.exists():
        target = destination / f"SchemaCraft-Default-{timestamp}-{counter}"
        counter += 1

    try:
        shutil.copytree(source, target, ignore=_default_app_copy_ignore)
        cloned_data = target / "data"
        _remove_cloned_path(cloned_data / "logs")
        if clear_schema:
            _remove_cloned_path(cloned_data)
            # A schema-free copy is a genuinely new application. It must ask
            # its new administrator to create a password on first use instead
            # of inheriting the current application's credential hash.
            (target / BUILDER_AUTH_PATH.name).unlink(missing_ok=True)
        elif clear_records:
            _clear_cloned_records(cloned_data)
        elif clear_history:
            _clear_cloned_history(cloned_data)
    except (OSError, shutil.Error) as exc:
        _remove_cloned_path(target)
        raise ApplicationError(f"تعذّر إنشاء التطبيق الافتراضي: {exc}") from exc

    return {
        "ok": True,
        "destination": str(target),
        "clear_history": clear_history,
        "clear_records": clear_records,
        "clear_schema": clear_schema,
    }


def current_audit_user() -> str:
    if AUDIT_USERS is None:
        return ""
    try:
        return clean_text(AUDIT_USERS.read().get("current_user"))[:160]
    except AdvancedFeatureError:
        return ""


def select_audit_user(payload: Any) -> dict[str, Any]:
    if AUDIT_USERS is None or not isinstance(payload, dict):
        raise ApplicationError("طلب اسم المستخدم غير صالح.")
    try:
        users = AUDIT_USERS.select(payload.get("name"))
    except AdvancedFeatureError as exc:
        raise ApplicationError(str(exc)) from exc
    return {"ok": True, "audit_users": users}


def select_workspace_schema(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise ApplicationError("طلب اختيار التصميم غير صالح.")
    manager = require_workspace()
    try:
        context = manager.set_active(clean_text(payload.get("schema_id")))
    except WorkspaceError as exc:
        raise ApplicationError(str(exc)) from exc
    return {"ok": True, "active_schema_id": context.schema_id}


def manage_workspace_schema(payload: Any) -> dict[str, Any]:
    with WORKBOOK_LOCK:
        return _manage_workspace_schema_unlocked(payload)


def _manage_workspace_schema_unlocked(payload: Any) -> dict[str, Any]:
    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("طلب إدارة التصميم غير صالح.")
    manager = require_workspace()
    action = clean_text(payload.get("action")).casefold()
    try:
        if action in {"create", "duplicate"}:
            template_id = clean_text(payload.get("template_schema_id")) or None
            if action == "duplicate" and not template_id:
                raise ApplicationError("اختر التصميم المراد نسخ بنيته.")
            context = manager.create_schema(
                payload.get("name"),
                default_schema(),
                template_schema_id=template_id,
                profile_source_schema_id=clean_text(payload.get("profile_source_schema_id")) or None,
            )
            with use_context(context):
                ensure_storage()
            synchronize_workspace_profiles()
        elif action == "rename":
            context = manager.rename_schema(
                clean_text(payload.get("schema_id")), payload.get("name")
            )
        elif action in {"archive", "restore"}:
            create_backup(automatic=True)
            context = manager.archive_schema(
                clean_text(payload.get("schema_id")), action == "archive"
            )
        elif action == "delete":
            backup = create_backup(automatic=True)
            deleted = manager.delete_schema(
                clean_text(payload.get("schema_id")), payload.get("confirmation_name")
            )
            context = manager.context(deleted["active_schema_id"])
        else:
            raise ApplicationError("عملية إدارة التصميم غير مدعومة.")
    except WorkspaceError as exc:
        raise ApplicationError(str(exc)) from exc
    invalidate_all_dataset_caches()
    response = {"ok": True, "schema": manager.response(), "schema_id": context.schema_id}
    if action == "delete":
        response["deleted_schema_id"] = deleted["deleted_schema_id"]
        response["backup"] = backup
    return response


def _row_value(
    row: tuple[Any, ...],
    columns: dict[str, int],
    header: str,
) -> Any:
    """Read from an iter_rows tuple using a one-based header map."""
    column = columns.get(header)
    if column is None or column > len(row):
        return None
    return row[column - 1]


def read_excel_field_value(value: Any, field: dict[str, Any]) -> Any:
    raw = json_value(value)
    if field["type"] == "checkbox":
        return clean_text(raw).casefold() in {"نعم", "true", "1", "yes"}
    if field["type"] == "checkbox_group":
        return [clean_text(item) for item in parse_checkbox_group_value(raw) if clean_text(item)]
    if number_is_text(field):
        return clean_text(raw)
    return raw


def excel_boolean(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return clean_text(value).casefold() in {"نعم", "true", "1", "yes"}




def read_dataset_unlocked(schema: dict[str, Any]) -> list[dict[str, Any]]:
    workbook_path = _workbook_path()
    if not workbook_path.is_file():
        return []
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    except Exception as exc:
        raise ApplicationError(f"تعذّر قراءة ملف Excel: {exc}") from exc

    try:
        if MAIN_SHEET not in workbook.sheetnames:
            raise ApplicationError(
                "ملف Excel لا يتوافق مع إصدار التطبيق العام الحالي."
            )
        main_sheet = workbook[MAIN_SHEET]
        headers = technical_headers(main_sheet)
        missing = [header for header in MAIN_REQUIRED_INTERNAL_HEADERS if header not in headers]
        if missing:
            raise ApplicationError(
                "ملف Excel لا يحتوي على الأعمدة الداخلية المطلوبة."
            )
        columns = {header: index for index, header in enumerate(headers, start=1)}
        fields = main_fields(schema)
        records: list[dict[str, Any]] = []
        by_internal_id: dict[str, dict[str, Any]] = {}
        seen_record_codes: set[str] = set()

        for row in main_sheet.iter_rows(
            min_row=FIRST_DATA_ROW,
            values_only=True,
        ):
            record_id = clean_text(_row_value(row, columns, "_record_id"))
            record_code = clean_text(_row_value(row, columns, "record_code"))
            if not record_id and not record_code:
                continue
            if not INTERNAL_ID_PATTERN.fullmatch(record_id):
                raise ApplicationError("يوجد معرّف سجل داخلي تالف في Excel.")
            if not PERSON_CODE_PATTERN.fullmatch(record_code):
                raise ApplicationError("يوجد معرّف سجل ظاهر تالف في Excel.")
            if record_id in by_internal_id:
                raise ApplicationError(
                    "يوجد معرّف سجل داخلي مكرر في Excel. "
                    "استعد نسخة سليمة قبل متابعة التعديل."
                )
            if record_code in seen_record_codes:
                raise ApplicationError(
                    f'معرّف السجل الظاهر "{record_code}" مكرر في Excel.'
                )
            record = {
                "_record_id": record_id,
                "record_code": record_code,
                "created_at": clean_text(_row_value(row, columns, "created_at")),
                "updated_at": clean_text(_row_value(row, columns, "updated_at")),
                "archived": excel_boolean(_row_value(row, columns, "_archived")),
                "archived_at": clean_text(_row_value(row, columns, "_archived_at")),
                "values": {
                    field["id"]: read_excel_field_value(
                        _row_value(row, columns, field["id"]), field
                    )
                    for field in fields
                },
                "related": {},
            }
            records.append(record)
            by_internal_id[record_id] = record
            seen_record_codes.add(record_code)

        names = related_sheet_names(schema)
        seen_child_ids: set[str] = set()
        for category in schema["categories"]:
            if category["kind"] != "repeatable":
                continue
            for record in records:
                record["related"][category["id"]] = []
            sheet_name = names[category["id"]]
            if sheet_name not in workbook.sheetnames:
                continue
            worksheet = workbook[sheet_name]
            related_headers = technical_headers(worksheet)
            related_columns = {
                header: index
                for index, header in enumerate(related_headers, start=1)
            }
            missing = [
                header
                for header in RELATED_INTERNAL_HEADERS
                if header not in related_columns
            ]
            if missing:
                raise ApplicationError(
                    f'ورقة "{sheet_name}" لا تحتوي على الأعمدة الداخلية المطلوبة.'
                )
            for row in worksheet.iter_rows(
                min_row=FIRST_DATA_ROW,
                values_only=True,
            ):
                record_id = clean_text(
                    _row_value(row, related_columns, "_record_id")
                )
                child_id = clean_text(
                    _row_value(row, related_columns, "_child_id")
                )
                related_code = clean_text(
                    _row_value(row, related_columns, "record_code")
                )
                if not record_id and not child_id and not related_code:
                    continue
                if record_id not in by_internal_id:
                    raise ApplicationError(
                        f'يوجد صف في ورقة "{sheet_name}" لا يعود إلى سجل معروف.'
                    )
                if related_code != by_internal_id[record_id]["record_code"]:
                    raise ApplicationError(
                        f'يوجد صف في ورقة "{sheet_name}" يحمل معرّف سجل غير متطابق.'
                    )
                if not INTERNAL_ID_PATTERN.fullmatch(child_id):
                    raise ApplicationError(
                        f'يوجد معرّف صف داخلي تالف في ورقة "{sheet_name}".'
                    )
                if child_id in seen_child_ids:
                    raise ApplicationError(
                        f'يوجد معرّف صف داخلي مكرر في ورقة "{sheet_name}".'
                    )
                minor_id = _row_value(row, related_columns, "minor_id")
                try:
                    minor_number = int(minor_id)
                except (TypeError, ValueError) as exc:
                    raise ApplicationError(
                        f'يوجد minor_id غير صالح في ورقة "{sheet_name}".'
                    ) from exc
                if minor_number < 1:
                    raise ApplicationError(
                        f'يوجد minor_id غير صالح في ورقة "{sheet_name}".'
                    )
                child = {
                    "_child_id": child_id,
                    "minor_id": minor_number,
                    "created_at": clean_text(
                        _row_value(row, related_columns, "created_at")
                    ),
                    "updated_at": clean_text(
                        _row_value(row, related_columns, "updated_at")
                    ),
                    "linked_record_code": clean_text(
                        _row_value(row, related_columns, RELATED_LINK_HEADER)
                    ),
                    "parent_child_id": clean_text(
                        _row_value(row, related_columns, RELATED_PARENT_HEADER)
                    ),
                    "values": {
                        field["id"]: read_excel_field_value(
                            _row_value(row, related_columns, field["id"]), field
                        )
                        for field in data_fields(category)
                    },
                }
                by_internal_id[record_id]["related"][category["id"]].append(child)
                seen_child_ids.add(child_id)

            for record in records:
                record["related"][category["id"]].sort(
                    key=lambda child: (
                        int(child.get("minor_id") or 0),
                        child["_child_id"],
                    )
                )
        category_lookup = {category["id"]: category for category in schema["categories"]}
        for record in records:
            for category in schema["categories"]:
                if category["kind"] != "repeatable":
                    continue
                parent = category_lookup.get(category.get("parent_category_id"))
                if not parent or parent.get("kind") != "repeatable":
                    continue
                parent_rows = record["related"].get(parent["id"], [])
                fallback_parent_id = parent_rows[0]["_child_id"] if parent_rows else ""
                for child in record["related"].get(category["id"], []):
                    if not child.get("parent_child_id"):
                        child["parent_child_id"] = fallback_parent_id
        return records
    finally:
        workbook.close()


def _schema_cache_signature(schema: dict[str, Any]) -> str:
    payload = json.dumps(
        schema,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _workbook_cache_signature() -> tuple[int, int] | None:
    try:
        status = _workbook_path().stat()
    except FileNotFoundError:
        return None
    return status.st_mtime_ns, status.st_size


def invalidate_dataset_cache() -> None:
    global _DATASET_SNAPSHOT
    key = str(_workbook_path().resolve())
    _DATASET_SNAPSHOTS.pop(key, None)
    if active_context() is None:
        _DATASET_SNAPSHOT = None


def invalidate_all_dataset_caches() -> None:
    global _DATASET_SNAPSHOT
    _DATASET_SNAPSHOTS.clear()
    _DATASET_SNAPSHOT = None


def _search_index_value(field: dict[str, Any], value: Any) -> Any:
    field_type = field["type"]
    if field_type == "checkbox":
        return excel_boolean(value)
    if field_type == "checkbox_group":
        return frozenset(
            normalize_search_text(item)
            for item in parse_checkbox_group_value(value)
        )
    if field_type == "number" and number_is_text(field):
        return normalize_search_text(value)
    if field_type == "number":
        try:
            return float(value)
        except (TypeError, ValueError):
            return None
    return normalize_search_text(value)


def _system_field_value(record: dict[str, Any], field: dict[str, Any]) -> Any:
    return {
        "system_record_code": record.get("record_code", ""),
        "system_created_at": record.get("created_at", ""),
        "system_updated_at": record.get("updated_at", ""),
    }.get(field.get("type"), "")


def _snapshot_record_copy(
    schema: dict[str, Any], record: dict[str, Any]
) -> dict[str, Any]:
    snapshot_record = copy.deepcopy(record)
    for category in schema["categories"]:
        if category["kind"] == "main":
            values = snapshot_record.setdefault("values", {})
            for field in data_fields(category):
                if field["type"] in SYSTEM_FIELD_TYPES:
                    values.pop(field["id"], None)
                    continue
                values[field["id"]] = read_excel_field_value(
                    values.get(field["id"], ""), field
                )
            continue
        for row in snapshot_record.setdefault("related", {}).setdefault(
            category["id"], []
        ):
            values = row.setdefault("values", {})
            for field in data_fields(category):
                values[field["id"]] = read_excel_field_value(
                    values.get(field["id"], ""), field
                )
    return snapshot_record


def _build_dataset_snapshot(
    schema: dict[str, Any],
    records: Iterable[dict[str, Any]],
    *,
    workbook_signature: tuple[int, int] | None = None,
) -> DatasetSnapshot:
    record_tuple = tuple(
        _snapshot_record_copy(schema, record) for record in records
    )
    records_by_code: dict[str, dict[str, Any]] = {}
    records_by_id: dict[str, dict[str, Any]] = {}
    record_positions: dict[str, int] = {}
    unique_mutable: dict[str, dict[str, set[str]]] = {}
    search_main: dict[str, dict[str, Any]] = {}
    search_related: dict[str, dict[str, tuple[dict[str, Any], ...]]] = {}

    searchable_fields = {
        field["id"]: field
        for category in schema["categories"]
        for field in data_fields(category)
        if field["type"] != "file"
    }
    unique_fields = {
        field["id"]: (category, field)
        for category in schema["categories"]
        for field in data_fields(category)
        if field.get("unique")
    }

    for position, record in enumerate(record_tuple):
        record_id = record["_record_id"]
        record_code = record["record_code"]
        records_by_code[record_code] = record
        records_by_id[record_id] = record
        record_positions[record_id] = position

        main_values: dict[str, Any] = {}
        related_values: dict[str, tuple[dict[str, Any], ...]] = {}
        for category in schema["categories"]:
            if category["kind"] == "main":
                for field in data_fields(category):
                    if field["id"] in searchable_fields:
                        main_values[field["id"]] = _search_index_value(
                            field,
                            _system_field_value(record, field)
                            if field["type"] in SYSTEM_FIELD_TYPES
                            else record.get("values", {}).get(field["id"], ""),
                        )
            else:
                indexed_rows: list[dict[str, Any]] = []
                for row in record.get("related", {}).get(category["id"], []):
                    indexed_row = {
                            field["id"]: _search_index_value(
                                field, row.get("values", {}).get(field["id"], "")
                            )
                            for field in data_fields(category)
                            if field["id"] in searchable_fields
                    }
                    indexed_rows.append(indexed_row)
                related_values[category["id"]] = tuple(indexed_rows)
        search_main[record_id] = main_values
        search_related[record_id] = related_values

        for field_id, (category, field) in unique_fields.items():
            for value in unique_field_values_in_record(record, category, field):
                if value == "" or value is None or value == []:
                    continue
                token = canonical_unique_value(value)
                unique_mutable.setdefault(field_id, {}).setdefault(
                    token, set()
                ).add(record_id)

    unique_indexes = {
        field_id: {
            token: frozenset(owners)
            for token, owners in values.items()
        }
        for field_id, values in unique_mutable.items()
    }

    return DatasetSnapshot(
        workbook_path=str(_workbook_path().resolve()),
        workbook_signature=(
            _workbook_cache_signature()
            if workbook_signature is None
            else workbook_signature
        ),
        schema_signature=_schema_cache_signature(schema),
        records=record_tuple,
        records_by_code=records_by_code,
        records_by_id=records_by_id,
        record_positions=record_positions,
        unique_indexes=unique_indexes,
        search_main=search_main,
        search_related=search_related,
    )


def _dataset_snapshot_unlocked(
    schema: dict[str, Any],
    *,
    force_reload: bool = False,
) -> DatasetSnapshot:
    global _DATASET_SNAPSHOT
    workbook_path = str(_workbook_path().resolve())
    workbook_signature = _workbook_cache_signature()
    schema_signature = _schema_cache_signature(schema)
    cached = _DATASET_SNAPSHOTS.get(workbook_path)
    if active_context() is None and _DATASET_SNAPSHOT is not None:
        cached = _DATASET_SNAPSHOT

    if (
        not force_reload
        and cached is not None
        and cached.workbook_path == workbook_path
        and cached.workbook_signature == workbook_signature
        and cached.schema_signature == schema_signature
    ):
        return cached

    started = time.perf_counter()
    records = read_dataset_unlocked(schema)
    snapshot = _build_dataset_snapshot(
        schema, records, workbook_signature=workbook_signature
    )
    _DATASET_SNAPSHOTS[workbook_path] = snapshot
    _DATASET_SNAPSHOT = snapshot
    LOGGER.info(
        "Workbook loaded and indexed: %.3f seconds; records=%d",
        time.perf_counter() - started,
        len(records),
    )
    return snapshot


def _publish_dataset_snapshot(
    schema: dict[str, Any], records: Iterable[dict[str, Any]]
) -> DatasetSnapshot:
    global _DATASET_SNAPSHOT
    snapshot = _build_dataset_snapshot(schema, records)
    _DATASET_SNAPSHOTS[str(_workbook_path().resolve())] = snapshot
    _DATASET_SNAPSHOT = snapshot
    return snapshot


def _safe_publish_dataset_snapshot(
    schema: dict[str, Any], records: Iterable[dict[str, Any]]
) -> None:
    try:
        _publish_dataset_snapshot(schema, records)
    except Exception:
        # The workbook commit has already succeeded. A cache failure must never
        # turn a successful save into an apparent failure or remove attachments.
        invalidate_dataset_cache()
        LOGGER.exception("Failed to rebuild the in-memory dataset snapshot")


def read_dataset(schema: dict[str, Any]) -> list[dict[str, Any]]:
    with WORKBOOK_LOCK:
        # Public callers receive an isolated copy; only internal read paths use
        # the immutable published snapshot directly.
        return copy.deepcopy(list(_dataset_snapshot_unlocked(schema).records))


def record_count(schema: dict[str, Any] | None = None) -> int:
    schema = schema or read_schema_file()
    with WORKBOOK_LOCK:
        return len(_dataset_snapshot_unlocked(schema).records)


def _record_file_paths(
    schema: dict[str, Any], record: dict[str, Any]
) -> set[str]:
    paths: set[str] = set()
    for category, field in file_fields(schema):
        if category["kind"] == "main":
            raw_values = [record.get("values", {}).get(field["id"], "")]
        else:
            raw_values = [
                child.get("values", {}).get(field["id"], "")
                for child in record.get("related", {}).get(category["id"], [])
            ]
        for raw in raw_values:
            relative = attachment_relative_path(raw)
            if relative:
                paths.add(relative)
    return paths


def migrate_field_value(
    value: Any,
    old_field: dict[str, Any] | None,
    new_field: dict[str, Any],
) -> Any:
    if old_field is None:
        return value
    old_type = old_field["type"]
    new_type = new_field["type"]
    if (old_type == "file") != (new_type == "file"):
        return ""
    if old_type == new_type and new_type in {"select", "yes_no"}:
        old_option = option_by_value(old_field, value)
        if old_option:
            new_option = option_by_id(new_field, old_option["id"])
            return new_option["label"] if new_option else ""
    if old_type == new_type == "checkbox_group":
        migrated: list[str] = []
        for item in parse_checkbox_group_value(value):
            old_option = option_by_value(old_field, item)
            if not old_option:
                continue
            new_option = option_by_id(new_field, old_option["id"])
            if new_option and new_option["label"] not in migrated:
                migrated.append(new_option["label"])
        return migrated
    if new_type == "checkbox" and old_type != "checkbox":
        return False
    if old_type == "checkbox" and new_type != "checkbox":
        return ""
    if new_type == "checkbox_group" and old_type != "checkbox_group":
        return []
    if old_type == "checkbox_group" and new_type != "checkbox_group":
        return ""
    return value




def migrate_records_to_schema(
    old_schema: dict[str, Any],
    new_schema: dict[str, Any],
    records: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    old_indexes = schema_indexes(old_schema)
    new_indexes = schema_indexes(new_schema)

    for field_id, old_category_id in old_indexes["field_categories"].items():
        if field_id not in new_indexes["field_categories"]:
            continue
        new_category_id = new_indexes["field_categories"][field_id]
        old_kind = old_indexes["categories"][old_category_id]["kind"]
        new_kind = new_indexes["categories"][new_category_id]["kind"]
        if (old_category_id != new_category_id or old_kind != new_kind) and records and not (old_kind == new_kind == "main"):
            raise ApplicationError(
                "نقل حقل بين الفئات المتكررة أو بينها وبين الرئيسية بعد وجود سجلات "
                "يتطلب ترحيل قيم البطاقات. النقل بين الفئات الرئيسية يحافظ على القيم ومتاح مباشرة."
            )

    for category_id, old_category in old_indexes["categories"].items():
        if category_id in new_indexes["categories"]:
            new_category = new_indexes["categories"][category_id]
            if old_category["kind"] != new_category["kind"] and records:
                raise ApplicationError(
                    f'لا يمكن تغيير نوع الفئة "{old_category["label"]}" بعد وجود سجلات.'
                )

    old_fields = old_indexes["fields"]
    new_fields = new_indexes["fields"]
    new_main_ids = {
        field["id"]
        for category in new_schema["categories"]
        if category["kind"] == "main"
        for field in data_fields(category)
        if field["type"] not in SYSTEM_FIELD_TYPES
    }
    new_related_ids = {
        category["id"]: {field["id"] for field in data_fields(category)}
        for category in new_schema["categories"]
        if category["kind"] == "repeatable"
    }

    migrated = copy.deepcopy(records)
    for record in migrated:
        new_values: dict[str, Any] = {}
        for field_id in new_main_ids:
            value = record.get("values", {}).get(field_id, "")
            value = migrate_field_value(
                value,
                old_fields.get(field_id),
                new_fields[field_id],
            )
            new_values[field_id] = value
        record["values"] = new_values

        new_related: dict[str, list[dict[str, Any]]] = {}
        for category_id, field_id_set in new_related_ids.items():
            rows = record.get("related", {}).get(category_id, [])
            new_rows = []
            for row in rows:
                values = {}
                for field_id in field_id_set:
                    value = row.get("values", {}).get(field_id, "")
                    value = migrate_field_value(
                        value,
                        old_fields.get(field_id),
                        new_fields[field_id],
                    )
                    values[field_id] = value
                row["values"] = values
                row.pop("markers", None)
                if not new_indexes["categories"][category_id].get(
                    "related_person_enabled"
                ):
                    row["linked_record_code"] = ""
                new_rows.append(row)
            new_related[category_id] = new_rows
        record["related"] = new_related
    return migrated


def schema_response(schema: dict[str, Any]) -> dict[str, Any]:
    response = copy.deepcopy(schema)
    with WORKBOOK_LOCK:
        records = _dataset_snapshot_unlocked(schema).records
    access = builder_access_response()
    response["developer_mode"] = access["unlocked"]
    response["builder_access"] = access
    response["stats"] = {
        "record_count": len(records),
        "field_count": sum(
            len(category["fields"]) for category in schema["categories"]
        ),
        "category_count": len(schema["categories"]),
    }
    response["archive_stats"] = {
        "active_record_count": sum(not record.get("archived") for record in records),
        "archived_record_count": sum(bool(record.get("archived")) for record in records),
    }
    return response


_HOME_SCHEMA_CACHE: dict[str, tuple[str, dict[str, Any]]] = {}


def home_schema_response(known_signature: str = "") -> dict[str, Any]:
    """Revalidate Home from file metadata before opening/validating any workbook."""
    with WORKBOOK_LOCK:
        key = str(_schema_path().resolve())
        def signature() -> str:
            try:
                status = _schema_path().stat()
                schema_stamp = (status.st_mtime_ns, status.st_size)
            except FileNotFoundError:
                schema_stamp = None
            context = active_context()
            return hashlib.sha256(json.dumps([
                key, schema_stamp, _workbook_cache_signature(),
                context.name if context else "",
                builder_access_response(),
            ], sort_keys=True).encode()).hexdigest()
        current = signature()
        if known_signature and known_signature == current:
            return {"unchanged": True, "signature": current}
        cached = _HOME_SCHEMA_CACHE.get(key)
        if cached and cached[0] == current:
            return {"signature": current, "schema": copy.deepcopy(cached[1])}
        response = schema_response(read_schema_with_excel_labels())
        response["schema_id"] = current_schema_id()
        context = active_context()
        response["schema_name"] = context.name if context else ""
        # Excel label synchronization can update schema.json during the read.
        current = signature()
        if len(_HOME_SCHEMA_CACHE) >= 64:
            _HOME_SCHEMA_CACHE.pop(next(iter(_HOME_SCHEMA_CACHE)))
        _HOME_SCHEMA_CACHE[key] = (current, response)
        return {"signature": current, "schema": copy.deepcopy(response)}


def schema_change_is_destructive(
    old_schema: dict[str, Any], new_schema: dict[str, Any]
) -> bool:
    old_indexes = schema_indexes(old_schema)
    new_indexes = schema_indexes(new_schema)
    if set(old_indexes["categories"]) - set(new_indexes["categories"]):
        return True
    if set(old_indexes["fields"]) - set(new_indexes["fields"]):
        return True
    for category_id, old_category in old_indexes["categories"].items():
        new_category = new_indexes["categories"].get(category_id)
        if not new_category:
            continue
        if old_category["kind"] != new_category["kind"]:
            return True
    for field_id, old_field in old_indexes["fields"].items():
        new_field = new_indexes["fields"].get(field_id)
        if not new_field:
            continue
        if old_field["type"] != new_field["type"]:
            return True
        old_options = {option["id"] for option in old_field.get("options", [])}
        new_options = {option["id"] for option in new_field.get("options", [])}
        if old_options - new_options:
            return True
    return False




def save_schema(payload: Any) -> dict[str, Any]:
    operation_started = time.perf_counter()
    new_schema = validate_schema(payload)
    expected_revision = new_schema["revision"]
    with WORKBOOK_LOCK:
        current = read_schema_file()
        if expected_revision != current["revision"]:
            raise ApplicationError(
                "تغيّر الإعداد في نافذة أخرى. أعد تحميله قبل الحفظ."
            )
        # The migration routine owns the single defensive copy. Previously a
        # list-field save read the entire workbook once to sync labels, copied
        # every record here, then copied it all again during migration.
        records = list(_dataset_snapshot_unlocked(current).records)
        automatic_backup = None
        destructive_change = bool(records) and schema_change_is_destructive(current, new_schema)
        context = active_context()
        background_sync = WORKSPACE_MANAGER is not None and context is not None
        if destructive_change and not background_sync:
            automatic_backup = create_backup(automatic=True)
        old_paths = {
            path
            for record in records
            for path in _record_file_paths(current, record)
        }
        new_schema["revision"] = current["revision"] + 1
        migrated = migrate_records_to_schema(current, new_schema, records)
        new_paths = {
            path
            for record in migrated
            for path in _record_file_paths(new_schema, record)
        }

        if background_sync:
            # schema.json and the indexed snapshot are the transactional source
            # of truth. Excel is a projection and is rebuilt immediately after
            # the request, or recovered on the next cold launch after a crash.
            atomic_write_json(_schema_path(), new_schema)
        else:
            workbook_temp: Path | None = None
            schema_temp: Path | None = None
            try:
                with tempfile.NamedTemporaryFile(
                    prefix=".database-schema-",
                    suffix=".xlsx",
                    dir=DATA_DIR,
                    delete=False,
                ) as temporary:
                    workbook_temp = Path(temporary.name)
                write_dataset_workbook(new_schema, migrated, workbook_temp)
                schema_temp = _write_json_temp(_schema_path(), new_schema)
                _replace_storage_pair(workbook_temp, schema_temp)
                workbook_temp = None
                schema_temp = None
            except PermissionError as exc:
                raise ApplicationError(
                    "تعذّر تحديث الإعداد. أغلق database.xlsx في Excel ثم حاول مجددًا."
                ) from exc
            finally:
                if workbook_temp:
                    workbook_temp.unlink(missing_ok=True)
                if schema_temp:
                    schema_temp.unlink(missing_ok=True)
        removed_paths = old_paths - new_paths
        if not (background_sync and destructive_change):
            remove_attachment_files(removed_paths)
        _safe_publish_dataset_snapshot(new_schema, migrated)
    response = schema_response(new_schema)
    if automatic_backup:
        response["automatic_backup"] = automatic_backup
    elif background_sync and destructive_change:
        response["automatic_backup_pending"] = True
    if background_sync:
        response["workbook_sync_pending"] = True
        # Start only after the response snapshot is complete. Slow recovery ZIP
        # and Excel generation are coalesced by schema and run without holding
        # the request lock.
        schedule_workbook_schema_sync(
            context,
            new_schema["revision"],
            recovery_schema=current if destructive_change else None,
            recovery_records=records if destructive_change else None,
            recovery_paths=old_paths if destructive_change else None,
            cleanup_paths=removed_paths if destructive_change else None,
        )
    LOGGER.info(
        "Schema save completed in %.3f seconds; records=%d",
        time.perf_counter() - operation_started,
        len(records),
    )
    return response


def add_runtime_list_option(payload: Any) -> dict[str, Any]:
    """Append a user-entered list value without granting Builder access."""

    if not isinstance(payload, dict):
        raise ApplicationError("صيغة خيار القائمة غير صحيحة.")
    field_id = clean_text(payload.get("field_id"))
    label = clean_text(payload.get("label"))[:240]
    dependency_token = clean_text(payload.get("dependency_token"))
    if not field_id or not label:
        raise ApplicationError("اكتب قيمة القائمة المطلوب إضافتها.")
    with WORKBOOK_LOCK:
        schema = read_schema_file()
        field = schema_indexes(schema)["fields"].get(field_id)
        if field is None or field.get("type") not in {"select", "yes_no"}:
            raise ApplicationError("الحقل المطلوب ليس قائمة قابلة للتوسعة.")
        normalized = normalize_search_text(label)
        existing = next(
            (
                option
                for option in field.get("options", [])
                if normalize_search_text(option.get("label")) == normalized
            ),
            None,
        )
        option = existing or {
            "id": new_definition_id("opt"),
            "label": label,
            "active": True,
        }
        changed = existing is None
        if existing is None:
            field.setdefault("options", []).append(option)
        option_filter = field.get("option_filter")
        if isinstance(option_filter, dict):
            source_field = schema_indexes(schema)["fields"].get(
                clean_text(option_filter.get("source_field_id"))
            )
            valid_tokens: set[str] = set()
            if source_field and source_field.get("type") == "checkbox":
                valid_tokens = {"true", "false"}
            elif source_field and source_field.get("type") in {"select", "yes_no"}:
                valid_tokens = {
                    clean_text(candidate.get("id"))
                    for candidate in source_field.get("options", [])
                    if isinstance(candidate, dict) and candidate.get("active", True)
                }
            if dependency_token not in valid_tokens:
                raise ApplicationError(
                    "اختر قيمة الحقل المصدر قبل إضافة قيمة جديدة إلى القائمة التابعة."
                )
            mappings = option_filter.setdefault("mappings", {})
            mapped_options = mappings.setdefault(dependency_token, [])
            if option["id"] not in mapped_options:
                mapped_options.append(option["id"])
                changed = True
        if changed:
            schema["revision"] += 1
            atomic_write_json(_schema_path(), schema)
            invalidate_dataset_cache()
    return {
        "ok": True,
        "option": copy.deepcopy(option),
        "dependency_token": dependency_token,
        "revision": schema["revision"],
    }


def save_app_settings(payload: Any) -> dict[str, Any]:
    """Save application-wide settings without rewriting record rows."""

    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("صيغة إعدادات التطبيق غير صحيحة.")
    try:
        expected_revision = int(payload.get("revision"))
    except (TypeError, ValueError) as exc:
        raise ApplicationError("إصدار إعدادات التطبيق غير صالح.") from exc
    raw_app = payload.get("app")
    if not isinstance(raw_app, dict):
        raise ApplicationError("إعدادات التطبيق غير مكتملة.")
    with WORKBOOK_LOCK:
        current = _sync_excel_field_labels_unlocked(read_schema_file())
        if expected_revision != current["revision"]:
            raise ApplicationError(
                "تغيّرت الإعدادات في نافذة أخرى. أعد تحميلها قبل الحفظ."
            )
        candidate = copy.deepcopy(current)
        candidate["app"] = copy.deepcopy(raw_app)
        candidate["revision"] = current["revision"] + 1
        candidate = validate_schema(candidate)
        atomic_write_json(_schema_path(), candidate)
        invalidate_dataset_cache()
        preferences = read_workspace_settings()
        preferences["primary_color"] = candidate["app"]["primary_color"]
        atomic_write_json(WORKSPACE_SETTINGS_PATH, preferences)
    return schema_response(candidate)


def condition_value_is_empty(
    value: Any,
    source_field: dict[str, Any],
) -> bool:
    field_type = source_field["type"]

    if field_type == "checkbox_group":
        return not parse_checkbox_group_value(value)

    if field_type == "file":
        if isinstance(value, dict):
            return not (
                value.get("stored_path")
                or value.get("upload")
            )

        return not clean_text(value)

    if value is None:
        return True

    return clean_text(json_value(value)) == ""


def condition_matches(
    condition: dict[str, Any],
    main_values: dict[str, Any],
    row_values: dict[str, Any] | None,
    indexes: dict[str, Any],
    target_category_id: str,
) -> bool:
    source_id = condition["source_field_id"]
    source_category_id = indexes[
        "field_categories"
    ][source_id]
    source_category = indexes[
        "categories"
    ][source_category_id]

    if (
        source_category["kind"] == "repeatable"
        and source_category_id
        == target_category_id
    ):
        value = (row_values or {}).get(
            source_id,
            "",
        )
    else:
        value = main_values.get(
            source_id,
            "",
        )

    source_field = indexes["fields"][source_id]
    field_type = source_field["type"]
    operator = condition["operator"]
    expected = condition.get("value", "")

    empty = condition_value_is_empty(
        value,
        source_field,
    )

    if operator == "empty":
        result = empty

    elif operator == "not_empty":
        result = not empty

    elif empty:
        result = False

    elif field_type == "checkbox":
        actual = (
            "true"
            if excel_boolean(value)
            else "false"
        )

        result = (
            actual == expected
            if operator == "equals"
            else actual != expected
        )

    elif field_type in {"select", "yes_no"}:
        option = option_by_value(
            source_field,
            value,
        )
        actual = option["id"] if option else ""

        result = (
            actual == expected
            if operator == "equals"
            else actual != expected
        )

    elif field_type == "checkbox_group":
        actual_ids: set[str] = set()

        for item in parse_checkbox_group_value(value):
            option = option_by_value(
                source_field,
                item,
            )

            if option:
                actual_ids.add(option["id"])

        if operator == "contains":
            result = expected in actual_ids
        else:
            result = expected not in actual_ids

    elif field_type == "number" and number_is_text(source_field):
        actual_text = normalize_search_text(value)
        expected_text = normalize_search_text(expected)
        result = {
            "equals": actual_text == expected_text,
            "not_equals": actual_text != expected_text,
            "contains": expected_text in actual_text,
            "not_contains": expected_text not in actual_text,
        }.get(operator, False)

    elif field_type == "number":
        try:
            actual_number = float(value)
            expected_number = float(expected)
        except (TypeError, ValueError):
            result = False
        else:
            result = {
                "equals":
                    actual_number == expected_number,
                "not_equals":
                    actual_number != expected_number,
                "greater_than":
                    actual_number > expected_number,
                "greater_or_equal":
                    actual_number >= expected_number,
                "less_than":
                    actual_number < expected_number,
                "less_or_equal":
                    actual_number <= expected_number,
            }.get(operator, False)

    elif field_type.startswith("date_"):
        actual_date = clean_text(
            json_value(value)
        )

        result = {
            "equals":
                actual_date == expected,
            "not_equals":
                actual_date != expected,
            "before":
                actual_date < expected,
            "after":
                actual_date > expected,
            "on_or_before":
                actual_date <= expected,
            "on_or_after":
                actual_date >= expected,
        }.get(operator, False)

    else:
        actual_text = normalize_search_text(value)
        expected_text = normalize_search_text(
            expected
        )

        result = {
            "equals":
                actual_text == expected_text,
            "not_equals":
                actual_text != expected_text,
            "contains":
                expected_text in actual_text,
            "not_contains":
                expected_text not in actual_text,
        }.get(operator, False)

    if condition.get("negate", False):
        return not result

    return result


def target_visible(
    target_type: str,
    target_id: str,
    main_values: dict[str, Any],
    row_values: dict[str, Any] | None,
    indexes: dict[str, Any],
) -> bool:
    rules = indexes[
        "conditions_by_target"
    ].get(
        (target_type, target_id),
        [],
    )

    if not rules:
        return True

    target_category_id = (
        target_id
        if target_type == "category"
        else indexes["field_categories"][target_id]
    )

    groups: dict[
        str,
        list[dict[str, Any]],
    ] = {}

    for rule in rules:
        group_id = (
            rule.get("group_id")
            or stable_condition_group_id(
                target_type,
                target_id,
            )
        )

        groups.setdefault(
            group_id,
            [],
        ).append(rule)

    # OR between groups, AND inside each group.
    return any(
        all(
            condition_matches(
                rule,
                main_values,
                row_values,
                indexes,
                target_category_id,
            )
            for rule in group_rules
        )
        for group_rules in groups.values()
    )

def field_source_value(
    source_field_id: str,
    target_category_id: str,
    main_values: dict[str, Any],
    row_values: dict[str, Any] | None,
    indexes: dict[str, Any],
) -> Any:
    source_category_id = indexes["field_categories"][source_field_id]
    source_category = indexes["categories"][source_category_id]
    if source_category["kind"] == "repeatable" and source_category_id == target_category_id:
        return (row_values or {}).get(source_field_id, "")
    return main_values.get(source_field_id, "")


def allowed_option_ids(
    field: dict[str, Any],
    main_values: dict[str, Any],
    row_values: dict[str, Any] | None,
    indexes: dict[str, Any],
) -> set[str]:
    all_ids = {
        option["id"] for option in field.get("options", []) if option.get("active", True)
    }
    option_filter = field.get("option_filter")
    if not option_filter:
        return all_ids
    target_category_id = indexes["field_categories"][field["id"]]
    source_id = option_filter["source_field_id"]
    source = indexes["fields"][source_id]
    source_value = field_source_value(
        source_id,
        target_category_id,
        main_values,
        row_values,
        indexes,
    )
    token = option_token(source, source_value)
    mappings = option_filter.get("mappings", {})
    if token in mappings:
        return set(mappings[token]) & all_ids
    return set() if option_filter.get("unmatched") == "none" else all_ids


def validate_dependent_option_value(
    field: dict[str, Any],
    normalized: Any,
    main_values: dict[str, Any],
    row_values: dict[str, Any] | None,
    indexes: dict[str, Any],
) -> None:
    if field["type"] not in {"select", "checkbox_group"} or normalized is None or normalized == "":
        return
    allowed = allowed_option_ids(field, main_values, row_values, indexes)
    values = normalized if field["type"] == "checkbox_group" else [normalized]
    for value in values:
        option = option_by_value(field, value)
        if option is None or option["id"] not in allowed:
            raise ApplicationError(
                f'قيمة الحقل "{field["label"]}" غير متاحة وفق الحقل المتحكم.'
            )


def validate_unique_repeated_checkboxes(
    category: dict[str, Any], rows: list[dict[str, Any]]
) -> None:
    """Enforce card-selector checkboxes without relying on retired tags."""

    for field in category.get("fields", []):
        if field.get("type") != "checkbox" or not field.get(
            "unique_checked_across_cards"
        ):
            continue
        selected = sum(
            excel_boolean(row.get("values", {}).get(field["id"])) for row in rows
        )
        if selected > 1:
            raise ApplicationError(
                f'يمكن تحديد مربع "{field["label"]}" في بطاقة واحدة فقط من فئة "{category["label"]}".'
            )




def _row_has_data(
    values: dict[str, Any],
    fields: Iterable[dict[str, Any]],
) -> bool:
    for field in fields:
        value = values.get(field["id"], "")
        if field["type"] == "file" and isinstance(value, dict):
            if value.get("stored_path") or value.get("upload"):
                return True
            continue
        if field["type"] == "checkbox":
            if bool(value):
                return True
            continue
        if field["type"] == "checkbox_group":
            if parse_checkbox_group_value(value):
                return True
            continue
        if value is not None and clean_text(value):
            return True
    return False


def attachments_directory() -> Path:
    context = active_context()
    return context.attachments_path if context is not None else DATA_DIR / "attachments"


def attachment_relative_path(value: Any) -> str:
    if isinstance(value, dict):
        value = value.get("stored_path", "")
    text = clean_text(value).replace("\\", "/")
    if not text:
        return ""
    if not text.startswith("attachments/"):
        raise ApplicationError("مسار ملف المرفق غير صالح.")
    filename = text.removeprefix("attachments/")
    if (
        not filename
        or filename in {".", ".."}
        or "/" in filename
        or "\\" in filename
    ):
        raise ApplicationError("مسار ملف المرفق غير صالح.")
    return f"attachments/{filename}"


def attachment_absolute_path(value: Any) -> Path | None:
    relative = attachment_relative_path(value)
    if not relative:
        return None
    return attachments_directory() / relative.removeprefix("attachments/")


def sanitize_filename_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = INVALID_WINDOWS_FILENAME_PATTERN.sub(" ", text)
    return " ".join(text.split()).strip(" .")


def attachment_extension(original_name: Any) -> str:
    filename = str(original_name or "").replace("\\", "/").rsplit("/", 1)[-1]
    extension = Path(filename).suffix
    if (
        not extension
        or len(extension) > 20
        or INVALID_WINDOWS_FILENAME_PATTERN.search(extension)
        or any(character.isspace() for character in extension)
    ):
        return ""
    return extension


def decode_attachment_upload(upload: Any) -> tuple[str, bytes]:
    if not isinstance(upload, dict):
        raise ApplicationError("بيانات الملف غير صالحة.")
    original_name = clean_text(upload.get("name"))
    encoded = upload.get("data")
    if not original_name or not isinstance(encoded, str) or not encoded:
        raise ApplicationError("بيانات الملف غير مكتملة.")
    try:
        content = base64.b64decode(encoded, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise ApplicationError("تعذّر قراءة الملف.") from exc
    if len(content) > MAX_ATTACHMENT_BYTES:
        raise ApplicationError("حجم الملف يتجاوز 100 ميغابايت.")
    return original_name, content


def _field_display_value(
    field_id: str,
    main_values: dict[str, Any],
    row_values: dict[str, Any] | None,
    indexes: dict[str, Any],
) -> str:
    if row_values and field_id in row_values:
        value = row_values.get(field_id, "")
    else:
        value = main_values.get(field_id, "")
    if isinstance(value, dict):
        return ""
    field = indexes["fields"].get(field_id)
    if field and field["type"] in {"select", "yes_no", "checkbox_group"}:
        options = {
            str(option["id"]): str(option["label"])
            for option in field.get("options", [])
        }
        if field["type"] == "checkbox_group":
            if isinstance(value, str):
                try:
                    parsed = json.loads(value)
                except (TypeError, ValueError, json.JSONDecodeError):
                    parsed = [value]
            else:
                parsed = value
            if isinstance(parsed, list):
                return "، ".join(
                    options.get(str(item), str(item)) for item in parsed
                )
        return options.get(str(value), clean_text(json_value(value)))
    return clean_text(json_value(value))


def attachment_base_name(
    field: dict[str, Any],
    original_name: str,
    main_values: dict[str, Any],
    row_values: dict[str, Any] | None,
    indexes: dict[str, Any],
) -> str:
    naming = field["file_naming"]
    if naming["mode"] == "original":
        base = sanitize_filename_text(Path(original_name).stem)
        if not base:
            raise ApplicationError("اسم الملف الأصلي غير صالح.")
        return base[:180]

    pieces: list[str] = []
    missing: list[str] = []
    for part in naming["parts"]:
        source_id = part["field_id"]
        value = _field_display_value(
            source_id,
            main_values,
            row_values,
            indexes,
        )
        if not value:
            missing.append(indexes["fields"][source_id]["label"])
        prefix = str(part.get("prefix", ""))
        suffix = str(part.get("suffix", ""))

        pieces.append(
            f"{prefix}{value}{suffix}"
        )
    if missing:
        raise ApplicationError(
            f'لا يمكن رفع "{field["label"]}" قبل تعبئة: '
            + "، ".join(missing)
            + "."
        )
    base = sanitize_filename_text("".join(pieces))
    if not base:
        raise ApplicationError("تعذّر إنشاء اسم صالح للملف.")
    return base[:180].rstrip(" .")


def unique_attachment_destination(
    base: str,
    extension: str,
    reserved: set[Path],
) -> tuple[str, Path]:
    directory = attachments_directory()
    counter = 1
    while True:
        suffix = "" if counter == 1 else f" ({counter})"
        filename = f"{base}{suffix}{extension}"
        destination = directory / filename
        if destination not in reserved and not destination.exists():
            reserved.add(destination)
            return f"attachments/{filename}", destination
        counter += 1


def prepare_file_value(
    raw_value: Any,
    field: dict[str, Any],
    main_values: dict[str, Any],
    row_values: dict[str, Any] | None,
    indexes: dict[str, Any],
    mode: str,
    old_paths: set[str],
    kept_paths: set[str],
    staged: list[tuple[Path, Path]],
    reserved: set[Path],
) -> str:
    if not raw_value:
        return ""
    if not isinstance(raw_value, dict):
        raw_value = {"stored_path": raw_value}
    upload = raw_value.get("upload")
    stored_path = attachment_relative_path(raw_value.get("stored_path", ""))
    if upload:
        original_name, content = decode_attachment_upload(upload)
        base = attachment_base_name(
            field,
            original_name,
            main_values,
            row_values,
            indexes,
        )
        extension = attachment_extension(original_name)
        relative, destination = unique_attachment_destination(
            base, extension, reserved
        )
        attachments_directory().mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            prefix=".attachment-",
            suffix=".tmp",
            dir=attachments_directory(),
            delete=False,
        ) as temporary:
            temporary.write(content)
            temporary_path = Path(temporary.name)
        staged.append((temporary_path, destination))
        kept_paths.add(relative)
        return relative
    if stored_path:
        if mode != "update" or stored_path not in old_paths:
            raise ApplicationError("مرجع الملف لا يعود إلى هذا السجل.")
        kept_paths.add(stored_path)
        return stored_path
    return ""


def remove_attachment_files(paths: set[str]) -> int:
    deleted = 0
    for value in paths:
        path = attachment_absolute_path(value)
        if path is None:
            continue
        try:
            path.unlink()
            deleted += 1
        except FileNotFoundError:
            continue
        except OSError:
            continue
    return deleted


def create_backup(*, automatic: bool = False) -> dict[str, str]:
    """Create a consistent ZIP of the complete multi-schema workspace."""
    with WORKBOOK_LOCK:
        ensure_storage()
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        hide_packaged_support_paths()
        stamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S-%f")
        kind = "auto-backup" if automatic else "backup"
        filename = f"GenericSchemaCraft-{kind}-{stamp}.zip"
        destination = BACKUP_DIR / filename
        with tempfile.NamedTemporaryFile(
            prefix=".backup-",
            suffix=".zip",
            dir=BACKUP_DIR,
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        try:
            with zipfile.ZipFile(
                temporary_path,
                mode="w",
                compression=zipfile.ZIP_DEFLATED,
            ) as archive:
                if WORKSPACE_MANAGER is not None:
                    excluded_roots = {
                        BACKUP_DIR.resolve(),
                        (DATA_DIR / "release-2-original").resolve(),
                    }
                    for path in sorted(DATA_DIR.rglob("*")):
                        if not path.is_file():
                            continue
                        resolved = path.resolve()
                        if any(root == resolved or root in resolved.parents for root in excluded_roots):
                            continue
                        archive.write(path, path.relative_to(DATA_DIR).as_posix())
                else:
                    archive.write(_schema_path(), "schema.json")
                    archive.write(_workbook_path(), "database.xlsx")
                    attachment_files = sorted(
                        path
                        for path in attachments_directory().rglob("*")
                        if path.is_file()
                    )
                    if not attachment_files:
                        archive.writestr("attachments/", b"")
                    for path in attachment_files:
                        relative = path.relative_to(attachments_directory())
                        archive.write(
                            path,
                            (Path("attachments") / relative).as_posix(),
                        )
            os.replace(temporary_path, destination)
        finally:
            temporary_path.unlink(missing_ok=True)
        if automatic:
            automatic_backups = sorted(
                BACKUP_DIR.glob("GenericSchemaCraft-auto-backup-*.zip"),
                key=lambda path: path.stat().st_mtime,
                reverse=True,
            )
            for old in automatic_backups[10:]:
                old.unlink(missing_ok=True)
    return {
        "filename": filename,
        "download_url": f"/api/backups/{quote(filename)}",
        "automatic": automatic,
    }


def backup_file_path(filename: Any) -> Path | None:
    text = clean_text(filename)
    name = Path(text).name
    if name != text:
        return None
    if not re.fullmatch(
        r"GenericSchemaCraft-(?:auto-)?backup-\d{8}-\d{6}-\d{6}\.zip",
        name,
    ):
        return None
    return BACKUP_DIR / name


def required_value_missing(field: dict[str, Any], value: Any) -> bool:
    if field["type"] == "checkbox":
        return not bool(value)
    if field["type"] == "checkbox_group":
        return not bool(value)
    return value in {"", None}


def apply_auto_update_rules(
    schema: dict[str, Any],
    raw_main: dict[str, Any],
    raw_related: dict[str, Any],
    audit_name: str,
) -> None:
    """Apply saved value dependencies before validation and file naming."""
    indexes = schema_indexes(schema)
    main_values = dict(raw_main)

    def result_for(
        rule: dict[str, Any],
        target_field: dict[str, Any],
        source_row: dict[str, Any] | None,
    ) -> Any:
        action = rule.get("action")
        if action == "clear":
            return False if target_field["type"] == "checkbox" else ""
        if action == "current_user":
            if not audit_name:
                raise ApplicationError("اختر اسم المستخدم قبل تنفيذ قاعدة المستخدم الحالي.")
            return audit_name
        if action == "copy_source":
            source_id = rule["source_field_id"]
            source_category_id = indexes["field_categories"][source_id]
            source_category = indexes["categories"][source_category_id]
            if source_category["kind"] == "repeatable":
                return (source_row or {}).get(source_id, "")
            return main_values.get(source_id, "")
        return rule.get("result_value", "")

    def row_values(raw_row: Any) -> dict[str, Any] | None:
        if not isinstance(raw_row, dict):
            return None
        values = raw_row.get("values", raw_row)
        return values if isinstance(values, dict) else None

    def matching_source_row(
        rule: dict[str, Any],
        target_category_id: str,
        target_row: dict[str, Any] | None,
    ) -> tuple[bool, dict[str, Any] | None]:
        source_id = rule["source_field_id"]
        source_category_id = indexes["field_categories"][source_id]
        source_category = indexes["categories"][source_category_id]
        if source_category["kind"] == "main":
            return (
                condition_matches(
                    rule, main_values, None, indexes, target_category_id
                ),
                None,
            )
        if source_category_id == target_category_id:
            return (
                condition_matches(
                    rule, main_values, target_row, indexes, target_category_id
                ),
                target_row,
            )
        source_rows = raw_related.get(source_category_id, [])
        if not isinstance(source_rows, list):
            return False, None
        for raw_source_row in source_rows:
            source_values = row_values(raw_source_row)
            if source_values is None:
                continue
            if condition_matches(
                rule,
                main_values,
                source_values,
                indexes,
                source_category_id,
            ):
                return True, source_values
        return False, None

    # A few bounded passes allow simple dependency chains without permitting loops
    # to keep a save running forever.
    rule_count = sum(
        1
        for category in schema["categories"]
        for field in data_fields(category)
        if field.get("auto_update")
    )
    for _ in range(max(1, min(rule_count + 1, 20))):
        changed = False
        main_values.update(raw_main)
        for category in schema["categories"]:
            for field in data_fields(category):
                rule = field.get("auto_update")
                if not rule:
                    continue
                if category["kind"] == "main":
                    matched, source_row = matching_source_row(
                        rule, category["id"], None
                    )
                    if matched:
                        value = result_for(rule, field, source_row)
                        if raw_main.get(field["id"]) != value:
                            raw_main[field["id"]] = value
                            main_values[field["id"]] = value
                            changed = True
                    continue
                rows = raw_related.get(category["id"], [])
                if not isinstance(rows, list):
                    continue
                for raw_row in rows:
                    values = row_values(raw_row)
                    if values is None:
                        continue
                    matched, source_row = matching_source_row(
                        rule, category["id"], values
                    )
                    if matched:
                        value = result_for(rule, field, source_row)
                        if values.get(field["id"]) != value:
                            values[field["id"]] = value
                            changed = True
        if not changed:
            break


def _normalize_submission(
    schema: dict[str, Any],
    payload: dict[str, Any],
    mode: str,
    old_paths: set[str],
) -> tuple[
    dict[str, Any],
    dict[str, list[dict[str, Any]]],
    list[tuple[Path, Path]],
    set[str],
]:
    indexes = schema_indexes(schema)
    raw_main = payload.get("main", {})
    raw_related = payload.get("related", {})
    if not isinstance(raw_main, dict) or not isinstance(raw_related, dict):
        raise ApplicationError("صيغة بيانات السجل غير صحيحة.")

    main_values = {
        field["id"]: raw_main.get(field["id"], "")
        for category in schema["categories"]
        if category["kind"] == "main"
        for field in data_fields(category)
    }
    staged: list[tuple[Path, Path]] = []
    kept_paths: set[str] = set()
    reserved: set[Path] = set()
    normalized_main: dict[str, Any] = {}
    normalized_related: dict[str, list[dict[str, Any]]] = {}
    visible_main_fields: list[dict[str, Any]] = []
    referenced_parent_ids = {
        clean_text(row.get("parent_child_id") or row.get(RELATED_PARENT_HEADER))
        for rows in raw_related.values()
        if isinstance(rows, list)
        for row in rows
        if isinstance(row, dict)
        and clean_text(row.get("parent_child_id") or row.get(RELATED_PARENT_HEADER))
    }

    try:
        for category in schema["categories"]:
            if category["kind"] != "main":
                continue
            category_visible = target_visible(
                "category", category["id"], main_values, None, indexes
            )
            for field in data_fields(category):
                if field["type"] in SYSTEM_FIELD_TYPES:
                    continue
                visible = category_visible and target_visible(
                    "field",
                    field["id"],
                    main_values,
                    None,
                    indexes,
                )

                # Visibility must never erase an entered or previously saved value.
                raw = main_values.get(field["id"], "")

                if visible:
                    visible_main_fields.append(field)
                if field["type"] == "file":
                    normalized = prepare_file_value(
                        raw,
                        field,
                        main_values,
                        None,
                        indexes,
                        mode,
                        old_paths,
                        kept_paths,
                        staged,
                        reserved,
                    )
                else:
                    normalized = normalize_field_value(raw, field)

                    # A hidden dependent field may retain a value that is not
                    # currently allowed by its controlling field.
                    if visible:
                        validate_dependent_option_value(
                            field,
                            normalized,
                            main_values,
                            None,
                            indexes,
                        )
                if field["required"] and visible and required_value_missing(field, normalized):
                    raise ApplicationError(f'الحقل "{field["label"]}" مطلوب.')
                normalized_main[field["id"]] = normalized

        validate_cross_field_constraints(
            visible_main_fields,
            normalized_main,
        )

        for category in schema["categories"]:
            if category["kind"] != "repeatable":
                continue
            category_visible = target_visible(
                "category", category["id"], main_values, None, indexes
            )
            raw_rows = raw_related.get(category["id"], [])
            if not isinstance(raw_rows, list):
                rows: list[dict[str, Any]] = []
                validate_unique_repeated_checkboxes(category, rows)
                normalized_related[category["id"]] = rows
                continue

            rows = []
            for raw_row in raw_rows:
                if not isinstance(raw_row, dict):
                    continue
                raw_values = raw_row.get("values", raw_row)
                if not isinstance(raw_values, dict):
                    continue
                linked_record_code = clean_text(
                    raw_row.get("linked_record_code")
                )
                related_person_mode = clean_text(
                    raw_row.get("related_person_mode")
                )
                if linked_record_code:
                    related_person_mode = "existing"
                elif related_person_mode == "existing":
                    raise ApplicationError(
                        f'أدخل معرّف الشخص المرتبط في فئة "{category["label"]}".'
                    )
                else:
                    related_person_mode = "manual"
                raw_child_id = clean_text(raw_row.get("_child_id"))
                if (
                    not _row_has_data(raw_values, data_fields(category))
                    and not linked_record_code
                    and raw_child_id not in referenced_parent_ids
                ):
                    continue
                condition_row_values = dict(raw_values)
                mode_source = related_person_mode_source_field(category)
                if mode_source:
                    condition_row_values[mode_source["id"]] = related_person_mode
                values: dict[str, Any] = {}
                visible_row_fields: list[dict[str, Any]] = []
                for field in data_fields(category):
                    visible = category_visible and target_visible(
                        "field",
                        field["id"],
                        main_values,
                        condition_row_values,
                        indexes,
                    )

                    raw = raw_values.get(field["id"], "")

                    if visible:
                        visible_row_fields.append(field)
                    if field["type"] == "file":
                        normalized = prepare_file_value(
                            raw,
                            field,
                            main_values,
                            raw_values,
                            indexes,
                            mode,
                            old_paths,
                            kept_paths,
                            staged,
                            reserved,
                        )
                    else:
                        normalized = normalize_field_value(raw, field)

                        if visible:
                            validate_dependent_option_value(
                                field,
                                normalized,
                                main_values,
                                raw_values,
                                indexes,
                            )
                    if field["required"] and visible and required_value_missing(field, normalized):
                        raise ApplicationError(
                            f'الحقل "{field["label"]}" مطلوب في فئة '
                            f'"{category["label"]}".'
                        )
                    values[field["id"]] = normalized
                validate_cross_field_constraints(
                    visible_row_fields,
                    values,
                )
                child_id = raw_child_id
                if child_id:
                    child_id = validate_internal_id(child_id, "معرّف الصف")
                rows.append(
                    {
                        "_child_id": child_id,
                        "_client_generated": bool(raw_row.get("_client_generated")),
                        "parent_child_id": clean_text(
                            raw_row.get("parent_child_id")
                            or raw_row.get(RELATED_PARENT_HEADER)
                        ),
                        "linked_record_code": linked_record_code,
                        "values": values,
                    }
                )
            validate_unique_repeated_checkboxes(category, rows)
            normalized_related[category["id"]] = rows

        for category in schema["categories"]:
            if category["kind"] != "repeatable":
                continue
            parent = indexes["categories"].get(category.get("parent_category_id"))
            rows = normalized_related.get(category["id"], [])
            if not parent or parent.get("kind") != "repeatable":
                for row in rows:
                    row["parent_child_id"] = ""
                continue
            parent_ids = {
                row.get("_child_id")
                for row in normalized_related.get(parent["id"], [])
                if row.get("_child_id")
            }
            for row in rows:
                parent_child_id = row.get("parent_child_id", "")
                if not parent_child_id or parent_child_id not in parent_ids:
                    raise ApplicationError(
                        f'بطاقة في الفئة "{category["label"]}" لا ترتبط ببطاقة أم صالحة.'
                    )
    except Exception:
        for temporary_path, _ in staged:
            temporary_path.unlink(missing_ok=True)
        raise

    return normalized_main, normalized_related, staged, kept_paths


def validate_related_person_links(
    schema: dict[str, Any],
    records: list[dict[str, Any]],
    current_record_code: str,
    related_values: dict[str, list[dict[str, Any]]],
) -> None:
    known_codes = {record["record_code"] for record in records}
    for category in schema["categories"]:
        if category["kind"] != "repeatable":
            continue
        enabled = bool(category.get("related_person_enabled"))
        for row in related_values.get(category["id"], []):
            linked_code = clean_text(row.get("linked_record_code"))
            if not linked_code:
                continue
            if not enabled:
                raise ApplicationError(
                    f'الفئة "{category["label"]}" غير مهيأة لربط الأشخاص.'
                )
            linked_code = validate_person_code(linked_code)
            if linked_code == current_record_code:
                raise ApplicationError("لا يمكن ربط السجل بنفسه كشخص مرتبط.")
            if linked_code not in known_codes:
                raise ApplicationError(
                    f'لم يُعثر على الشخص المرتبط ذي المعرّف "{linked_code}".'
                )
            row["linked_record_code"] = linked_code


def _find_record(
    records: list[dict[str, Any]], record_code: str
) -> dict[str, Any] | None:
    return next(
        (record for record in records if record["record_code"] == record_code),
        None,
    )


def canonical_unique_value(value: Any) -> str:
    if isinstance(value, list):
        return json.dumps(
            sorted(normalize_search_text(item) for item in value),
            ensure_ascii=False,
        )
    return normalize_search_text(value)


def unique_field_values_in_record(
    record: dict[str, Any], category: dict[str, Any], field: dict[str, Any]
) -> list[Any]:
    if category["kind"] == "main":
        return [record.get("values", {}).get(field["id"], "")]
    return [
        row.get("values", {}).get(field["id"], "")
        for row in record.get("related", {}).get(category["id"], [])
    ]


def validate_unique_fields(
    schema: dict[str, Any],
    records: list[dict[str, Any]],
    current_record_id: str | None,
    main_values: dict[str, Any],
    related_values: dict[str, list[dict[str, Any]]],
    unique_indexes: dict[str, dict[str, frozenset[str]]] | None = None,
) -> None:
    for category in schema["categories"]:
        for field in data_fields(category):
            if not field.get("unique"):
                continue
            if category["kind"] == "main":
                submitted = [main_values.get(field["id"], "")]
            else:
                submitted = [
                    row.get("values", {}).get(field["id"], "")
                    for row in related_values.get(category["id"], [])
                ]
            submitted_tokens = [
                canonical_unique_value(value)
                for value in submitted
                if value != "" and value is not None and value != []
            ]
            if len(submitted_tokens) != len(set(submitted_tokens)):
                raise ApplicationError(
                    f'قيمة الحقل الفريد "{field["label"]}" مكررة داخل السجل.'
                )
            if not submitted_tokens:
                continue

            if unique_indexes is not None:
                field_index = unique_indexes.get(field["id"], {})
                for token in submitted_tokens:
                    owners = field_index.get(token, frozenset())
                    if any(owner != current_record_id for owner in owners):
                        raise ApplicationError(
                            f'قيمة الحقل "{field["label"]}" مستخدمة في سجل آخر.'
                        )
                continue

            existing_tokens: set[str] = set()
            for record in records:
                if current_record_id and record.get("_record_id") == current_record_id:
                    continue
                for value in unique_field_values_in_record(record, category, field):
                    if value == "" or value is None or value == []:
                        continue
                    existing_tokens.add(canonical_unique_value(value))
            if any(token in existing_tokens for token in submitted_tokens):
                raise ApplicationError(
                    f'قيمة الحقل "{field["label"]}" مستخدمة في سجل آخر.'
                )




def save_record(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise ApplicationError("صيغة البيانات المرسلة غير صحيحة.")
    mode = payload.get("mode", "create")
    if mode not in {"create", "update"}:
        raise ApplicationError("نوع عملية الحفظ غير صالح.")
    requested_code = clean_text(payload.get("record_code"))
    record_code = validate_person_code(requested_code) if requested_code else ""

    operation_started = time.perf_counter()
    staged: list[tuple[Path, Path]] = []
    promoted: list[Path] = []
    obsolete_paths: set[str] = set()
    with WORKBOOK_LOCK:
        schema = read_schema_file()
        snapshot = _dataset_snapshot_unlocked(schema)
        records = list(snapshot.records)
        existing = snapshot.records_by_code.get(record_code) if record_code else None

        if mode == "create":
            if record_code and record_code in snapshot.records_by_code:
                if payload.get("client_generated_code"):
                    record_code = ""
                else:
                    raise ApplicationError(
                        "معرّف السجل مستخدم مسبقًا. افتح نموذجًا جديدًا للحصول على معرّف آخر."
                    )
            if (
                record_code
                and WORKSPACE_MANAGER is not None
                and WORKSPACE_MANAGER.person_id_in_use(record_code)
                and payload.get("client_generated_code")
                and not payload.get("link_existing")
            ):
                record_code = ""
            if not record_code:
                record_code = generate_person_code()
                while (
                    record_code in snapshot.records_by_code
                    or (
                        WORKSPACE_MANAGER is not None
                        and WORKSPACE_MANAGER.person_id_in_use(record_code)
                    )
                ):
                    record_code = generate_person_code()
            if WORKSPACE_MANAGER is not None:
                try:
                    WORKSPACE_MANAGER.assert_profile_creation(
                        record_code,
                        current_schema_id(),
                        link_existing=bool(payload.get("link_existing")),
                    )
                except WorkspaceError as exc:
                    raise ApplicationError(str(exc)) from exc
            existing = None
        elif not record_code:
            raise ApplicationError("اختر سجلًا قبل حفظ التعديلات.")
        elif existing is None:
            raise ApplicationError("لم يُعثر على السجل المطلوب تعديله.")

        expected_updated_at = clean_text(payload.get("expected_updated_at"))
        if (
            mode == "update"
            and expected_updated_at
            and clean_text(existing.get("updated_at")) != expected_updated_at
        ):
            raise ApplicationError(
                "تم تعديل السجل في نافذة أخرى بعد فتحه. أعد تحميله قبل الحفظ."
            )

        old_paths = _record_file_paths(schema, existing) if existing else set()
        try:
            payload = copy.deepcopy(payload)
            raw_main = payload.get("main") if isinstance(payload.get("main"), dict) else {}
            payload["main"] = raw_main
            audit_name = current_audit_user()
            has_last_update_date = any(
                field.get("type") == "system_updated_at"
                for category in schema.get("categories", [])
                if category.get("kind") == "main"
                for field in category.get("fields", [])
            )
            for category in schema.get("categories", []):
                if category.get("kind") != "main":
                    continue
                for field in category.get("fields", []):
                    previous_value = (
                        existing.get("values", {}).get(field["id"], "")
                        if existing
                        else ""
                    )
                    if (
                        field.get("type") == "date_gregorian"
                        and field.get("date_value_mode") == "on_checkbox"
                    ):
                        trigger_id = field.get("date_trigger_field_id")
                        checked_now = excel_boolean(raw_main.get(trigger_id))
                        checked_before = excel_boolean(
                            existing.get("values", {}).get(trigger_id, "")
                            if existing
                            else False
                        )
                        raw_main[field["id"]] = (
                            date.today().isoformat()
                            if checked_now and not checked_before
                            else previous_value
                        )
                        continue
                    if field.get("type") != "user_name":
                        continue
                    value_mode = field.get("user_value_mode")
                    if value_mode == "current_on_save":
                        if (
                            has_last_update_date
                            and mode == "update"
                            and clean_text(existing.get("updated_at"))
                        ):
                            if not audit_name:
                                raise ApplicationError(
                                    "اختر اسم المستخدم قبل حفظ تعديل سجل يحتوي على حقل آخر محرر."
                                )
                            raw_main[field["id"]] = audit_name
                        else:
                            raw_main[field["id"]] = previous_value
                    elif value_mode == "current_on_checkbox":
                        trigger_id = field.get("user_trigger_field_id")
                        checked_now = excel_boolean(raw_main.get(trigger_id))
                        checked_before = excel_boolean(
                            existing.get("values", {}).get(trigger_id, "")
                            if existing
                            else False
                        )
                        if checked_now and not checked_before:
                            if not audit_name:
                                raise ApplicationError(
                                    "اختر اسم المستخدم قبل تأكيد مربع اختيار يحدّث اسم المستخدم."
                                )
                            raw_main[field["id"]] = audit_name
                        else:
                            raw_main[field["id"]] = previous_value
                    elif existing and not field.get("user_editable", True) and previous_value:
                        raw_main[field["id"]] = previous_value
                    elif mode == "create" or not clean_text(raw_main.get(field["id"])):
                        if not audit_name:
                            raise ApplicationError(
                                "اختر اسم المستخدم قبل حفظ سجل يحتوي على حقل المستخدم."
                            )
                        raw_main[field["id"]] = audit_name
            apply_auto_update_rules(schema, raw_main, payload.get("related", {}), audit_name)
            normalize_started = time.perf_counter()
            (
                main_values,
                related_values,
                staged,
                kept_paths,
            ) = _normalize_submission(schema, payload, mode, old_paths)
            normalize_seconds = time.perf_counter() - normalize_started

            validate_started = time.perf_counter()
            validate_unique_fields(
                schema,
                records,
                existing.get("_record_id") if existing else None,
                main_values,
                related_values,
                snapshot.unique_indexes,
            )
            validate_related_person_links(
                schema,
                records,
                record_code,
                related_values,
            )
            validate_seconds = time.perf_counter() - validate_started
            timestamp = now_iso()

            if existing is None:
                record = {
                    "_record_id": new_internal_id(),
                    "record_code": record_code,
                    "created_at": timestamp,
                    "updated_at": timestamp,
                    "archived": False,
                    "archived_at": "",
                    "values": main_values,
                    "related": {},
                }
                records.append(record)
            else:
                # Never mutate an object owned by the published cache before
                # the Excel replacement succeeds.
                record = copy.deepcopy(existing)
                record["updated_at"] = timestamp
                record["values"] = main_values
                records[snapshot.record_positions[record["_record_id"]]] = record

            occupied_child_ids = {
                child.get("_child_id")
                for candidate in records
                if candidate.get("_record_id") != record.get("_record_id")
                for rows in candidate.get("related", {}).values()
                for child in rows
                if child.get("_child_id")
            }
            submitted_child_ids: set[str] = set()
            for category in schema["categories"]:
                if category["kind"] != "repeatable":
                    continue
                previous_rows = {
                    child["_child_id"]: child
                    for child in record.get("related", {}).get(category["id"], [])
                }
                new_rows = []
                for row in related_values.get(category["id"], []):
                    child_id = row["_child_id"]
                    if child_id in submitted_child_ids or child_id in occupied_child_ids:
                        raise ApplicationError(
                            f'معرّف صف مكرر في الفئة "{category["label"]}".'
                        )
                    if (
                        child_id
                        and child_id not in previous_rows
                        and not row.get("_client_generated")
                    ):
                        raise ApplicationError(
                            f'معرّف صف في الفئة "{category["label"]}" لا يعود إلى هذا السجل.'
                        )
                    if child_id:
                        submitted_child_ids.add(child_id)
                    previous = previous_rows.get(child_id)
                    new_rows.append(
                        {
                            "_child_id": child_id or new_internal_id(),
                            "minor_id": len(new_rows) + 1,
                            "created_at": (
                                previous["created_at"] if previous else timestamp
                            ),
                            "updated_at": timestamp,
                            "linked_record_code": row.get(
                                "linked_record_code", ""
                            ),
                            "parent_child_id": row.get("parent_child_id", ""),
                            "values": row["values"],
                        }
                    )
                record.setdefault("related", {})[category["id"]] = new_rows

            attachment_started = time.perf_counter()
            for temporary_path, destination in staged:
                os.replace(temporary_path, destination)
                promoted.append(destination)
            attachment_seconds = time.perf_counter() - attachment_started

            workbook_started = time.perf_counter()
            atomic_write_workbook(schema, records)
            workbook_seconds = time.perf_counter() - workbook_started
            obsolete_paths = old_paths - kept_paths
            deleted_files = remove_attachment_files(obsolete_paths)
            _safe_publish_dataset_snapshot(schema, records)
            if mode == "create" and WORKSPACE_MANAGER is not None:
                WORKSPACE_MANAGER.register_profile(
                    record_code, current_schema_id()
                )
            LOGGER.info(
                "Record save timings: mode=%s normalize=%.3f validate=%.3f "
                "attachments=%.3f workbook=%.3f total=%.3f",
                mode,
                normalize_seconds,
                validate_seconds,
                attachment_seconds,
                workbook_seconds,
                time.perf_counter() - operation_started,
            )
        except PermissionError as exc:
            for temporary_path, _ in staged:
                temporary_path.unlink(missing_ok=True)
            for destination in promoted:
                destination.unlink(missing_ok=True)
            invalidate_dataset_cache()
            raise ApplicationError(
                "تعذّر الحفظ. أغلق database.xlsx في Excel ثم حاول مرة أخرى."
            ) from exc
        except Exception:
            for temporary_path, _ in staged:
                temporary_path.unlink(missing_ok=True)
            for destination in promoted:
                destination.unlink(missing_ok=True)
            invalidate_dataset_cache()
            raise

    attachment_paths = sorted(_record_file_paths(schema, record))
    related_count = sum(
        len(record.get("related", {}).get(category["id"], []))
        for category in schema["categories"]
        if category["kind"] == "repeatable"
    )
    response = {
        "ok": True,
        "action": "updated" if mode == "update" else "created",
        "record_code": record_code,
        "_record_id": record["_record_id"],
        "related_rows": related_count,
        "attachment_files": attachment_paths,
        "deleted_attachment_files": deleted_files,
        "updated_at": record.get("updated_at", ""),
    }
    if mode == "update" and payload.get("propagate_global_values"):
        response["propagation"] = propagate_global_values(
            record_code,
            schema,
            record,
            payload.get("propagate_global_refs"),
        )
    return response


def propagate_global_values(
    record_code: str,
    source_schema: dict[str, Any],
    source_record: dict[str, Any],
    requested_refs: Any = None,
) -> dict[str, Any]:
    """Copy chosen global-field values to this person's other schema profiles."""

    manager = WORKSPACE_MANAGER
    if manager is None:
        return {"updated_profiles": 0, "schemas": []}
    selected = {
        clean_text(item)
        for item in (requested_refs if isinstance(requested_refs, list) else [])
        if clean_text(item).startswith("gfld_")
    }
    source_values: dict[str, Any] = {}
    for category in source_schema.get("categories", []):
        if category.get("kind") != "main":
            continue
        for field in category.get("fields", []):
            global_ref = clean_text(field.get("global_ref"))
            if global_ref and (not selected or global_ref in selected):
                source_values[global_ref] = copy.deepcopy(
                    source_record.get("values", {}).get(field["id"], "")
                )
    if not source_values:
        return {"updated_profiles": 0, "schemas": []}
    updated_schemas: list[str] = []
    source_id = current_schema_id()
    for membership in manager.identity_search(record_code):
        target_id = membership["schema_id"]
        if target_id == source_id:
            continue
        context = manager.context(target_id)
        with use_context(context):
            target_schema = read_schema_file()
            try:
                target_record = load_record(record_code)
            except ApplicationError:
                continue
            main = copy.deepcopy(target_record["main"])
            changed = False
            for category in target_schema.get("categories", []):
                if category.get("kind") != "main":
                    continue
                for field in category.get("fields", []):
                    global_ref = clean_text(field.get("global_ref"))
                    if global_ref in source_values and main.get(field["id"]) != source_values[global_ref]:
                        main[field["id"]] = copy.deepcopy(source_values[global_ref])
                        changed = True
            if not changed:
                continue
            save_record(
                {
                    "mode": "update",
                    "record_code": record_code,
                    "expected_updated_at": target_record.get("updated_at", ""),
                    "main": main,
                    "related": target_record.get("related", {}),
                    "propagate_global_values": False,
                }
            )
            updated_schemas.append(target_id)
    return {"updated_profiles": len(updated_schemas), "schemas": updated_schemas}


def load_record(record_code: Any) -> dict[str, Any]:
    code = validate_person_code(record_code)
    schema = read_schema_file()
    with WORKBOOK_LOCK:
        record = _dataset_snapshot_unlocked(schema).records_by_code.get(code)
    if record is None:
        raise ApplicationError("لم يُعثر على السجل المطلوب.")
    return {
        "record_code": code,
        "_record_id": record["_record_id"],
        "created_at": record.get("created_at", ""),
        "updated_at": record.get("updated_at", ""),
        "archived": bool(record.get("archived")),
        "main": copy.deepcopy(record["values"]),
        "related": {
            category["id"]: [
                {
                    "_child_id": child["_child_id"],
                    "parent_child_id": child.get("parent_child_id", ""),
                    "minor_id": index + 1,
                    "linked_record_code": child.get(
                        "linked_record_code", ""
                    ),
                    "values": copy.deepcopy(child["values"]),
                }
                for index, child in enumerate(
                    record.get("related", {}).get(category["id"], [])
                )
            ]
            for category in schema["categories"]
            if category["kind"] == "repeatable"
        },
    }


def _field_matches(field: dict[str, Any], query: Any, value: Any) -> bool:
    if isinstance(query, dict):
        operator = clean_text(query.get("operator"))
        if query.get("empty") is True or operator == "empty":
            return condition_value_is_empty(value, field)
        if operator == "not_empty":
            return not condition_value_is_empty(value, field)
        if operator:
            return _field_operator_matches(field, query, value)
    if isinstance(query, dict) and isinstance(query.get("values"), list):
        query = query["values"]
    if isinstance(query, list) and field["type"] != "checkbox_group":
        return any(_field_matches(field, item, value) for item in query)
    if field["type"] == "checkbox":
        return excel_boolean(query) == excel_boolean(value)
    if field["type"] == "checkbox_group":
        try:
            expected = {
                normalize_search_text(item)
                for item in normalize_field_value(query, field)
            }
        except ApplicationError:
            return False
        actual = {
            normalize_search_text(item)
            for item in parse_checkbox_group_value(value)
        }
        # Multiple values chosen for one list field are alternatives (OR).
        # AND remains the relationship between different filter fields.
        return bool(expected & actual)
    if field["type"] in {"select", "yes_no"}:
        try:
            query = normalize_choice_value(query, field)
        except ApplicationError:
            return False
    if field["type"] in {"system_created_at", "system_updated_at"}:
        return normalize_search_text(query) in normalize_search_text(value)
    if field["search_match"] == "contains" and field["type"] != "number":
        return normalize_search_text(query) in normalize_search_text(value)
    if field["type"] == "number" and number_is_text(field):
        if field.get("search_match") == "contains":
            return normalize_search_text(query) in normalize_search_text(value)
        return normalize_search_text(query) == normalize_search_text(value)
    if field["type"] == "number":
        try:
            return float(query) == float(value)
        except (TypeError, ValueError):
            return False
    return normalize_search_text(query) == normalize_search_text(value)


def _field_operator_matches(
    field: dict[str, Any], query: dict[str, Any], value: Any
) -> bool:
    """Evaluate an explicit search/export comparison against a raw value."""

    operator = clean_text(query.get("operator"))
    expected = query.get("value", "")
    upper = query.get("to", "")
    if condition_value_is_empty(value, field):
        return False
    if field["type"] == "number" and not number_is_text(field):
        try:
            actual_number = float(value)
            expected_number = float(expected)
        except (TypeError, ValueError):
            return False
        if operator == "between":
            try:
                upper_number = float(upper)
            except (TypeError, ValueError):
                return False
            lower_bound, upper_bound = sorted((expected_number, upper_number))
            return lower_bound <= actual_number <= upper_bound
        return {
            "equals": actual_number == expected_number,
            "greater_than": actual_number > expected_number,
            "greater_or_equal": actual_number >= expected_number,
            "less_than": actual_number < expected_number,
            "less_or_equal": actual_number <= expected_number,
        }.get(operator, False)
    if field["type"].startswith("date_") or field["type"] in {
        "system_created_at",
        "system_updated_at",
    }:
        actual_date = clean_text(json_value(value))
        expected_date = clean_text(expected)
        if operator == "between":
            upper_date = clean_text(upper)
            if not expected_date or not upper_date:
                return False
            lower_bound, upper_bound = sorted((expected_date, upper_date))
            return lower_bound <= actual_date <= upper_bound
        return {
            "equals": actual_date == expected_date,
            "before": actual_date < expected_date,
            "after": actual_date > expected_date,
            "on_or_before": actual_date <= expected_date,
            "on_or_after": actual_date >= expected_date,
        }.get(operator, False)
    return _field_matches(field, expected, value) if operator == "equals" else False


def _indexed_field_matches(
    field: dict[str, Any], query: Any, indexed_value: Any
) -> bool:
    if isinstance(query, dict):
        operator = clean_text(query.get("operator"))
        empty = not indexed_value if field["type"] == "checkbox_group" else indexed_value in {"", None}
        if query.get("empty") is True or operator == "empty":
            return empty
        if operator == "not_empty":
            return not empty
        if operator:
            return _indexed_field_operator_matches(field, query, indexed_value)
    if isinstance(query, dict) and isinstance(query.get("values"), list):
        query = query["values"]
    if isinstance(query, list) and field["type"] != "checkbox_group":
        return any(_indexed_field_matches(field, item, indexed_value) for item in query)
    if field["type"] == "checkbox":
        return excel_boolean(query) == indexed_value
    if field["type"] == "checkbox_group":
        try:
            expected = {
                normalize_search_text(item)
                for item in normalize_field_value(query, field)
            }
        except ApplicationError:
            return False
        return bool(expected & (indexed_value or frozenset()))
    if field["type"] in {"select", "yes_no"}:
        try:
            query = normalize_choice_value(query, field)
        except ApplicationError:
            return False
    if field["type"] in {"system_created_at", "system_updated_at"}:
        return normalize_search_text(query) in (indexed_value or "")
    if field["search_match"] == "contains" and field["type"] != "number":
        return normalize_search_text(query) in (indexed_value or "")
    if field["type"] == "number" and number_is_text(field):
        query_text = normalize_search_text(query)
        if field.get("search_match") == "contains":
            return query_text in (indexed_value or "")
        return query_text == (indexed_value or "")
    if field["type"] == "number":
        try:
            return float(query) == indexed_value
        except (TypeError, ValueError):
            return False
    return normalize_search_text(query) == (indexed_value or "")


def _indexed_field_operator_matches(
    field: dict[str, Any], query: dict[str, Any], indexed_value: Any
) -> bool:
    """Evaluate an explicit comparison against a normalized snapshot value."""

    operator = clean_text(query.get("operator"))
    expected = query.get("value", "")
    upper = query.get("to", "")
    if indexed_value in {"", None}:
        return False
    if field["type"] == "number" and not number_is_text(field):
        try:
            expected_number = float(expected)
            actual_number = float(indexed_value)
        except (TypeError, ValueError):
            return False
        if operator == "between":
            try:
                upper_number = float(upper)
            except (TypeError, ValueError):
                return False
            lower_bound, upper_bound = sorted((expected_number, upper_number))
            return lower_bound <= actual_number <= upper_bound
        return {
            "equals": actual_number == expected_number,
            "greater_than": actual_number > expected_number,
            "greater_or_equal": actual_number >= expected_number,
            "less_than": actual_number < expected_number,
            "less_or_equal": actual_number <= expected_number,
        }.get(operator, False)
    if field["type"].startswith("date_") or field["type"] in {
        "system_created_at",
        "system_updated_at",
    }:
        actual_date = clean_text(indexed_value)
        expected_date = clean_text(expected)
        if operator == "between":
            upper_date = clean_text(upper)
            if not expected_date or not upper_date:
                return False
            lower_bound, upper_bound = sorted((expected_date, upper_date))
            return lower_bound <= actual_date <= upper_bound
        return {
            "equals": actual_date == expected_date,
            "before": actual_date < expected_date,
            "after": actual_date > expected_date,
            "on_or_before": actual_date <= expected_date,
            "on_or_after": actual_date >= expected_date,
        }.get(operator, False)
    return _indexed_field_matches(field, expected, indexed_value) if operator == "equals" else False


def _search_criterion_is_active(field: dict[str, Any], value: Any) -> bool:
    if isinstance(value, dict):
        operator = clean_text(value.get("operator"))
        if value.get("empty") is True or operator in {"empty", "not_empty"}:
            return True
        if isinstance(value.get("values"), list):
            return bool(value["values"])
        if operator == "between":
            return bool(clean_text(value.get("value"))) and bool(clean_text(value.get("to")))
        if operator:
            return bool(clean_text(value.get("value")))
    if isinstance(value, list):
        return bool(value)
    if field["type"] in {"checkbox", "checkbox_group"}:
        return bool(value)
    return bool(clean_text(value))


def field_display_value(field: dict[str, Any], value: Any) -> Any:
    if field["type"] == "checkbox":
        return (field.get("checkbox_true_label") or "نعم") if excel_boolean(value) else (field.get("checkbox_false_label") or "لا")
    option_labels = {
        option["id"]: option["label"] for option in field.get("options", [])
    }
    if field["type"] == "checkbox_group":
        return "، ".join(
            option_labels.get(clean_text(item), clean_text(item))
            for item in parse_checkbox_group_value(value)
        )
    if field["type"] in {"select", "yes_no"}:
        return option_labels.get(clean_text(value), clean_text(value))
    if (
        field["type"] == "number"
        and field.get("number_behavior", {}).get("format_thousands")
        and clean_text(value)
    ):
        text = clean_text(value).replace(",", "").replace("٬", "")
        if re.fullmatch(r"[+-]?\d+(?:\.\d+)?", text):
            integer, separator, fraction = text.partition(".")
            sign = ""
            if integer.startswith(("+", "-")):
                sign, integer = integer[0], integer[1:]
            grouped = f"{int(integer):,}" if integer else "0"
            return f"{sign}{grouped}{separator}{fraction}"
    return json_value(value)


def _field_values_for_record(
    record: dict[str, Any],
    category: dict[str, Any],
    field: dict[str, Any],
) -> list[Any]:
    if category["kind"] == "main":
        if field["type"] in SYSTEM_FIELD_TYPES:
            return [_system_field_value(record, field)]
        return [record.get("values", {}).get(field["id"], "")]
    return [
        row.get("values", {}).get(field["id"], "")
        for row in record.get("related", {}).get(category["id"], [])
    ]


def field_value_suggestions(
    field_id: Any,
    query: Any = "",
    limit: Any = 30,
    selected_values: Any = None,
) -> dict[str, Any]:
    """Return distinct, live field values for shared search/export filters."""

    schema = read_schema_file()
    indexes = schema_indexes(schema)
    clean_id = clean_text(field_id)
    field = indexes["fields"].get(clean_id)
    if not field or field["type"] == "file":
        raise ApplicationError("الحقل المطلوب غير متاح للاقتراحات.")
    category = indexes["categories"][indexes["field_categories"][clean_id]]
    try:
        maximum = max(1, min(100, int(limit)))
    except (TypeError, ValueError):
        maximum = 30
    needle = normalize_search_text(query)
    selected_tokens = {
        normalize_search_text(json_value(item))
        for item in selected_values
        if clean_text(item)
    } if isinstance(selected_values, (list, tuple, set)) else set()
    counts: dict[str, dict[str, Any]] = {}
    records_with_value = 0
    checked_record_count = 0
    matching_record_count = 0
    matching_value_count = 0
    with WORKBOOK_LOCK:
        records = _dataset_snapshot_unlocked(schema).records
    for record in records:
        record_values = _field_values_for_record(record, category, field)
        record_matches_selected = False
        if any(not condition_value_is_empty(value, field) for value in record_values):
            records_with_value += 1
        if field["type"] == "checkbox" and any(
            value is True
            or clean_text(value).casefold() in {"1", "true", "yes", "نعم"}
            for value in record_values
        ):
            checked_record_count += 1
        for raw_value in record_values:
            raw_items = (
                parse_checkbox_group_value(raw_value)
                if field["type"] == "checkbox_group"
                else [raw_value]
            )
            for item in raw_items:
                if condition_value_is_empty(item, field):
                    continue
                raw = json_value(item)
                display = clean_text(field_display_value(field, item))
                if not display:
                    continue
                if needle and needle not in normalize_search_text(display) and needle not in normalize_search_text(raw):
                    continue
                token = normalize_search_text(raw) or normalize_search_text(display)
                if selected_tokens and token in selected_tokens:
                    record_matches_selected = True
                    matching_value_count += 1
                current = counts.setdefault(
                    token,
                    {"value": raw, "label": display, "count": 0},
                )
                current["count"] += 1
        if record_matches_selected:
            matching_record_count += 1
    values = sorted(
        counts.values(),
        key=lambda item: (-int(item["count"]), normalize_search_text(item["label"])),
    )[:maximum]
    return {
        "ok": True,
        "field_id": clean_id,
        "field_label": field["label"],
        "values": values,
        "total_distinct": len(counts),
        "record_count": len(records),
        "records_with_value": records_with_value,
        "checked_record_count": checked_record_count,
        "matching_record_count": matching_record_count,
        "matching_value_count": matching_value_count,
    }


def search_records(criteria: Any) -> dict[str, Any]:
    if not isinstance(criteria, dict):
        raise ApplicationError("صيغة البحث غير صحيحة.")
    schema = read_schema_file()
    indexes = schema_indexes(schema)
    requested_field_ids = criteria.get("_search_field_ids")
    if requested_field_ids is None:
        selected_field_ids = [
            field_id
            for field_id, field in indexes["fields"].items()
            if field["searchable"]
        ]
    else:
        if not isinstance(requested_field_ids, list):
            raise ApplicationError("حقول البحث المؤقتة غير صحيحة.")
        selected_field_ids = []
        for raw_field_id in requested_field_ids:
            field_id = clean_text(raw_field_id)
            field = indexes["fields"].get(field_id)
            if not field or field["type"] == "file":
                raise ApplicationError("أحد حقول البحث المؤقتة غير متاح.")
            if field_id not in selected_field_ids:
                selected_field_ids.append(field_id)
    searchable = {
        field_id: indexes["fields"][field_id]
        for field_id in selected_field_ids
    }
    include_archived = bool(criteria.get("_include_archived"))
    record_code_query = clean_text(criteria.get("_record_code")).upper()

    active = {
        field_id: value
        for field_id, value in criteria.items()
        if field_id in searchable
        and _search_criterion_is_active(searchable[field_id], value)
    }
    allow_empty = bool(criteria.get("_allow_empty"))
    if (
        not active
        and not record_code_query
        and not allow_empty
    ):
        raise ApplicationError("أدخل معيار بحث واحدًا على الأقل.")

    try:
        offset = max(0, int(criteria.get("_offset", 0)))
    except (TypeError, ValueError):
        offset = 0
    try:
        limit = int(criteria.get("_limit", MAX_SEARCH_RESULTS))
    except (TypeError, ValueError):
        limit = MAX_SEARCH_RESULTS
    limit = min(MAX_SEARCH_PAGE_SIZE, max(1, limit))

    main_criteria: list[tuple[dict[str, Any], Any]] = []
    related_criteria: dict[str, list[tuple[dict[str, Any], Any]]] = {}
    for field_id, query in active.items():
        category_id = indexes["field_categories"][field_id]
        category = indexes["categories"][category_id]
        pair = (searchable[field_id], query)
        if category["kind"] == "main":
            main_criteria.append(pair)
        else:
            related_criteria.setdefault(category_id, []).append(pair)

    requested_result_field_ids = criteria.get("_result_field_ids")
    custom_result_fields = requested_result_field_ids is not None
    if custom_result_fields:
        if not isinstance(requested_result_field_ids, list):
            raise ApplicationError("أعمدة نتائج البحث غير صحيحة.")
        result_fields = []
        for raw_field_id in requested_result_field_ids:
            field_id = clean_text(raw_field_id)
            field = indexes["fields"].get(field_id)
            if (
                not field
                or field["type"] == "file"
            ):
                raise ApplicationError("أحد أعمدة نتائج البحث غير متاح.")
            if field not in result_fields:
                result_fields.append(field)
    else:
        result_fields = [
            field
            for category in schema["categories"]
            for field in data_fields(category)
            if field["show_in_results"]
        ]
        if not result_fields:
            result_fields = [
                field
                for category in schema["categories"]
                if category["kind"] == "main"
                for field in data_fields(category)
                if field["type"] != "file"
                and field["type"] not in SYSTEM_FIELD_TYPES
            ][:4]
    title_fields = [
        field
        for category in schema["categories"]
        for field in data_fields(category)
        if field["result_title"]
    ]
    presentation_fields = list(result_fields)
    for field in title_fields:
        if field not in presentation_fields:
            presentation_fields.append(field)

    matches: list[dict[str, Any]] = []
    total = 0
    search_started = time.perf_counter()
    with WORKBOOK_LOCK:
        snapshot = _dataset_snapshot_unlocked(schema)
        records = snapshot.records

    for record in records:
        if record.get("archived") and not include_archived:
            continue
        if record_code_query and record_code_query not in record["record_code"].upper():
            continue
        record_id = record["_record_id"]
        main_index = snapshot.search_main.get(record_id, {})
        if any(
            not _indexed_field_matches(
                field, query, main_index.get(field["id"])
            )
            for field, query in main_criteria
        ):
            continue

        related_match = True
        related_index = snapshot.search_related.get(record_id, {})
        for category_id, pairs in related_criteria.items():
            indexed_rows = related_index.get(category_id, ())
            if not any(
                all(
                    _indexed_field_matches(
                        field, query, row.get(field["id"])
                    )
                    for field, query in pairs
                )
                for row in indexed_rows
            ):
                related_match = False
                break
        if not related_match:
            continue

        total += 1
        if total <= offset or len(matches) >= limit:
            continue

        details = []
        title_parts = []
        title_field_ids: list[str] = []
        for field in presentation_fields:
            category = indexes["categories"][
                indexes["field_categories"][field["id"]]
            ]
            values = _field_values_for_record(record, category, field)
            value = next(
                (
                    field_display_value(field, item)
                    for item in values
                    if item != "" and item is not None and item != []
                ),
                "",
            )
            if field["result_title"] and value:
                title_parts.append(str(value))
                title_field_ids.append(field["id"])
            details.append(
                {
                    "field_id": field["id"],
                    "label": field["label"],
                    "value": value,
                }
            )
        if not title_parts:
            fallback_title_items = [
                item for item in details if clean_text(item["value"])
            ][:2]
            title_parts = [str(item["value"]) for item in fallback_title_items]
            title_field_ids = [item["field_id"] for item in fallback_title_items]
        if custom_result_fields:
            requested_result_ids = {field["id"] for field in result_fields}
            detail_items = [
                item for item in details if item["field_id"] in requested_result_ids
            ]
        else:
            detail_items = [
                item for item in details
                if item["field_id"] not in title_field_ids
            ]
        matches.append(
            {
                "record_code": record["record_code"],
                "title": " ".join(title_parts) or record["record_code"],
                "archived": bool(record.get("archived")),
                "details": detail_items,
            }
        )
    LOGGER.info(
        "Search completed: %.3f seconds; scanned=%d; matches=%d; criteria=%d",
        time.perf_counter() - search_started,
        len(records),
        total,
        len(active) + bool(record_code_query),
    )
    return {
        "matches": matches,
        "total": total,
        "offset": offset,
        "limit": limit,
        "truncated": offset + len(matches) < total,
    }


def archive_record(
    record_code: Any,
    archived: bool = True,
    expected_updated_at: Any = None,
) -> dict[str, Any]:
    code = validate_person_code(record_code)
    with WORKBOOK_LOCK:
        schema = read_schema_file()
        snapshot = _dataset_snapshot_unlocked(schema)
        cached_record = snapshot.records_by_code.get(code)
        if cached_record is None:
            raise ApplicationError("لم يُعثر على السجل المطلوب.")
        expected = clean_text(expected_updated_at)
        if expected and clean_text(cached_record.get("updated_at")) != expected:
            raise ApplicationError(
                "تم تعديل السجل في نافذة أخرى. أعد تحميله قبل تغيير حالة الأرشفة."
            )
        records = list(snapshot.records)
        record = copy.deepcopy(cached_record)
        records[snapshot.record_positions[record["_record_id"]]] = record
        timestamp = now_iso()
        record["archived"] = bool(archived)
        record["archived_at"] = timestamp if archived else ""
        record["updated_at"] = timestamp
        try:
            atomic_write_workbook(schema, records)
            _safe_publish_dataset_snapshot(schema, records)
        except PermissionError as exc:
            invalidate_dataset_cache()
            raise ApplicationError(
                "تعذّر تحديث حالة الأرشفة. أغلق database.xlsx في Excel ثم حاول مرة أخرى."
            ) from exc
    return {
        "ok": True,
        "record_code": code,
        "archived": bool(archived),
        "archived_at": record["archived_at"],
        "updated_at": record["updated_at"],
    }


def delete_record(record_code: Any, *, synchronized: bool = False) -> dict[str, Any]:
    if not synchronized:
        require_builder_access()
    code = validate_person_code(record_code)
    with WORKBOOK_LOCK:
        schema = read_schema_file()
        snapshot = _dataset_snapshot_unlocked(schema)
        record = snapshot.records_by_code.get(code)
        if record is None:
            raise ApplicationError("لم يُعثر على السجل المطلوب حذفه.")
        paths = _record_file_paths(schema, record)
        related_rows = sum(
            len(record.get("related", {}).get(category["id"], []))
            for category in schema["categories"]
            if category["kind"] == "repeatable"
        )
        records = [
            candidate
            for candidate in snapshot.records
            if candidate["_record_id"] != record["_record_id"]
        ]
        try:
            atomic_write_workbook(schema, records)
            _safe_publish_dataset_snapshot(schema, records)
        except PermissionError as exc:
            invalidate_dataset_cache()
            raise ApplicationError(
                "تعذّر الحذف. أغلق database.xlsx في Excel ثم حاول مرة أخرى."
            ) from exc
        deleted_files = remove_attachment_files(paths)
        if WORKSPACE_MANAGER is not None:
            WORKSPACE_MANAGER.unregister_profile(code, current_schema_id())
        purged_search_results = 0
        if SEARCH_HISTORY is not None:
            try:
                purged_search_results = SEARCH_HISTORY.purge_record(
                    code, current_schema_id()
                )
            except AdvancedFeatureError:
                LOGGER.exception(
                    "تعذّر تنظيف سجل البحث بعد حذف الملف الشخصي %s.", code
                )
    return {
        "ok": True,
        "record_code": code,
        "deleted_main_rows": 1,
        "deleted_related_rows": related_rows,
        "deleted_attachment_files": deleted_files,
        "purged_search_results": purged_search_results,
    }


def create_filtered_export(payload: Any) -> tuple[str, bytes, int]:
    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("صيغة طلب التصدير غير صحيحة.")
    criteria = payload.get("criteria", {})
    if not isinstance(criteria, dict):
        raise ApplicationError("مرشحات التصدير غير صالحة.")
    selected_field_ids = payload.get("field_ids")
    if selected_field_ids is not None and not isinstance(selected_field_ids, list):
        raise ApplicationError("حقول التصدير غير صالحة.")

    with WORKBOOK_LOCK:
        schema = read_schema_file()
        snapshot = _dataset_snapshot_unlocked(schema)
        export_criteria = copy.deepcopy(criteria)
        export_criteria["_allow_empty"] = True
        export_criteria["_limit"] = MAX_SEARCH_PAGE_SIZE
        export_criteria["_offset"] = 0
        record_codes: list[str] = []
        while True:
            page = search_records(export_criteria)
            record_codes.extend(match["record_code"] for match in page["matches"])
            if len(record_codes) >= page["total"]:
                break
            export_criteria["_offset"] = len(record_codes)
        records = [
            copy.deepcopy(snapshot.records_by_code[code])
            for code in record_codes
            if code in snapshot.records_by_code
        ]
        selected_related_categories = {
            clean_text(item)
            for item in payload.get("related_category_ids", [])
            if clean_text(item)
        } if isinstance(payload.get("related_category_ids", []), list) else set()
        if selected_related_categories:
            for record in records:
                for category in schema.get("categories", []):
                    if category.get("kind") != "repeatable":
                        continue
                    if selected_related_categories and category["id"] not in selected_related_categories:
                        record.get("related", {})[category["id"]] = []
        content = export_workbook_bytes(
            schema,
            records,
            selected_field_ids,
            include_related=bool(payload.get("include_related", True)),
        )
    stamp = datetime.now().astimezone().strftime('%Y%m%d-%H%M%S')
    filename = f"SchemaCraft-export-{stamp}.xlsx"
    if bool(payload.get("include_attachments")):
        workbook_name = filename
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(workbook_name, content)
            added_paths: set[str] = set()
            for record in records:
                for relative in sorted(_record_file_paths(schema, record)):
                    source = attachment_absolute_path(relative)
                    if source is None or not source.is_file():
                        continue
                    archive_name = f'attachments/{record["record_code"]}/{source.name}'
                    if archive_name in added_paths:
                        continue
                    archive.write(source, archive_name)
                    added_paths.add(archive_name)
        content = output.getvalue()
        filename = f"SchemaCraft-export-{stamp}-with-attachments.zip"
    return filename, content, len(records)


def _record_report_fields(
    schema: dict[str, Any], record: dict[str, Any], selected_ids: set[str], global_refs: set[str]
) -> list[dict[str, str]]:
    fields = [{"label": "ID", "value": record["record_code"]}]
    for category in schema.get("categories", []):
        for field in category.get("fields", []):
            if field["type"] == "file" or field["type"] in SYSTEM_FIELD_TYPES:
                continue
            if (
                (selected_ids or global_refs)
                and field["id"] not in selected_ids
                and field.get("global_ref") not in global_refs
            ):
                continue
            values = _field_values_for_record(record, category, field)
            display = [
                clean_text(field_display_value(field, value))
                for value in values
                if value != "" and value is not None and value != []
            ]
            if display:
                fields.append(
                    {
                        "label": (
                            field["label"]
                            if category["kind"] == "main"
                            else f'{category["label"]} — {field["label"]}'
                        ),
                        "value": " | ".join(display),
                    }
                )
    return fields


def _record_report_sections(
    schema: dict[str, Any],
    record: dict[str, Any],
    selected_ids: set[str],
    global_refs: set[str],
) -> list[dict[str, Any]]:
    """Read-only sections: visible data fields, Builder widths and card order."""
    indexes = schema_indexes(schema)
    sections: list[dict[str, Any]] = []
    for category in schema.get("categories", []):
        chosen = [field for field in category.get("fields", [])
            if field["type"] not in SYSTEM_FIELD_TYPES | {"spacer"}
            and (not selected_ids and not global_refs
                 or field["id"] in selected_ids
                 or field.get("global_ref") in global_refs)]
        source_rows = [{"values":record.get("values", {})}] if category["kind"] == "main" else record.get("related", {}).get(category["id"], [])
        cards = []
        for row_index, row in enumerate(source_rows):
            fields = []
            for field in chosen:
                value = row.get("values", {}).get(field["id"])
                if value is None or value == [] or isinstance(value, str) and not value.strip():
                    continue
                if isinstance(value, list) and not any(str(item or "").strip() for item in value):
                    continue
                display = str(value).replace("\\", "/").rsplit("/", 1)[-1] if field["type"] == "file" else str(field_display_value(field, value))
                fields.append({"label":field["label"], "value":display,
                    "width":field.get("width", "1"), "start_new_line":field.get("start_new_line", False), "type":field["type"], "field_id":field["id"]})
            if not fields:
                continue
            title = ""
            if category["kind"] == "repeatable":
                title_field = indexes["fields"].get(category.get("card_title_field_id"))
                raw_title = row.get("values", {}).get(title_field["id"]) if title_field else None
                title = str(field_display_value(title_field, raw_title)).strip() if raw_title is not None else ""
                title = title or f'{category.get("card_name_prefix") or category["label"]} {len(cards) + 1}'
            cards.append({"title":title, "fields":fields, "row_index":row_index})
        if cards:
            sections.append({"label":category["label"], "kind":category["kind"], "category_id":category["id"], "cards":cards})
    return sections


def _profile_pdf_media(schema, record, sections, *, show_profile_image=False, show_attachments=False):
    """Resolve only this record's attachments in its own schema context."""
    attachments = []
    seen = set()
    section_map = {section["category_id"]:section for section in sections}
    for category in schema.get("categories", []):
        rows = [{"values":record.get("values", {})}] if category["kind"] == "main" else record.get("related", {}).get(category["id"], [])
        for row_index, row in enumerate(rows):
            for field in category.get("fields", []):
                if field["type"] != "file":
                    continue
                value = row.get("values", {}).get(field["id"])
                if not value:
                    continue
                portrait = show_profile_image and field.get("image_display") in {"profile", "card"}
                if not portrait and not show_attachments:
                    continue
                source = attachment_absolute_path(value)
                if source is None or not source.is_file():
                    raise ApplicationError(f"تعذّر العثور على المرفق: {Path(str(value)).name}")
                # attachment_relative_path validates the stored path; resolve
                # additionally prevents a symlink escaping the attachments folder.
                if not source.resolve().is_relative_to(attachments_directory().resolve()):
                    raise ApplicationError("مسار المرفق غير صالح.")
                if portrait:
                    section = section_map.get(category["id"])
                    if section is None:
                        section = {"category_id":category["id"], "label":category["label"], "kind":category["kind"], "cards":[]}
                        section_map[category["id"]] = section
                    if category["kind"] == "main":
                        section["profile_image"] = source
                    else:
                        card = next((item for item in section["cards"] if item.get("row_index") == row_index), None)
                        if card is None:
                            title_field = next((item for item in category["fields"] if item["id"] == category.get("card_title_field_id")), None)
                            raw_title = row.get("values", {}).get(title_field["id"]) if title_field else None
                            title = str(field_display_value(title_field, raw_title)).strip() if raw_title is not None else ""
                            card = {"row_index":row_index,"title":title or f'{category.get("card_name_prefix") or category["label"]} {row_index+1}',"fields":[]}
                            section["cards"].append(card)
                            section["cards"].sort(key=lambda item:item.get("row_index",0))
                        card["profile_image"] = source
                    for card in section["cards"]:
                        if card.get("row_index") == row_index:
                            card["fields"] = [item for item in card["fields"] if item.get("field_id") != field["id"]]
                if show_attachments and str(source.resolve()) not in seen:
                    attachments.append({"path":source, "name":source.name, "category":category["label"], "field":field["label"]})
                    seen.add(str(source.resolve()))
    sections[:] = [section_map[category["id"]] for category in schema.get("categories", []) if category["id"] in section_map]
    return attachments



def report_api(payload):
    """Advanced reports share the existing authenticated export permission."""
    require_builder_access()
    from schemacraft_reports import TemplateStore, DraftStore, ReportError, decode_template, resolve_report, validate_template, pdf_bytes
    if not isinstance(payload, dict): raise ApplicationError("طلب التقرير غير صالح.")
    manager = require_workspace()
    store = TemplateStore(manager.data_dir / "report-templates")
    try:
        action = payload.get("action")
        if isinstance(action,str) and (action.startswith("document_") or action.startswith("recipe_")):
            import schemacraft_report_api
            import sys
            return schemacraft_report_api.handle(sys.modules[__name__],manager,payload)
        drafts = DraftStore(manager.data_dir / "report-drafts")
        if action == "draft_list": return {"drafts": drafts.list()}
        if action == "draft_read": return {"saved": drafts.read(payload.get("id"))}
        if action == "draft_save": return {"saved": drafts.save(payload.get("draft"), payload.get("id"), payload.get("revision"))}
        if action == "draft_delete":
            drafts.delete(payload.get("id"), payload.get("revision")); return {"ok": True}
        if action == "pdf_preview":
            return {"pdf": base64.b64encode(pdf_bytes(payload.get("draft"), PDF_FONT_PATH)).decode("ascii")}
        if action == "validate":
            template = validate_template(payload.get("template"))
            if "document" in template:
                import schemacraft_report_api
                import sys
                return schemacraft_report_api.handle(sys.modules[__name__], manager, {"action":"document_validate","template":template})
            # Verify referenced fields against the current schemas before saving a template.
            active = set(re.findall(r"\{\{([A-Za-z][A-Za-z0-9_]*)\}\}", template["body"]))
            for name, item in template["placeholders"].items():
                if name not in active: continue
                with use_context(manager.context(item["schema_id"])):
                    field_ids = {f["id"] for c in read_schema_file().get("categories", []) for f in c.get("fields", [])}
                    needed = item["fields"] + [c["field"] for c in item.get("criteria", [])] + ([item["group_field"]] if item.get("mode") == "grouped" else [])
                    if any(f not in field_ids for f in needed): raise ReportError("حقل غير موجود في العنصر: " + name)
            return {"ok": True, "unused": sorted(set(template["placeholders"]) - active)}
        if action == "profiles":
            schema_ids = payload.get("schema_ids")
            query = payload.get("query", "")
            offset = payload.get("offset", 0)
            if not isinstance(schema_ids, list) or not 1 <= len(schema_ids) <= 20 or any(not isinstance(v,str) for v in schema_ids): raise ReportError("اختر التصاميم المطلوبة للمجموعة.")
            if not isinstance(query,str) or len(query)>120 or isinstance(offset,bool) or not isinstance(offset,int) or offset<0: raise ReportError("طلب بحث الملفات غير صالح.")
            snapshots=[]
            with WORKBOOK_LOCK:
                for schema_id in dict.fromkeys(schema_ids):
                    with use_context(manager.context(schema_id)):
                        schema=read_schema_file()
                        snapshots.append((schema, _dataset_snapshot_unlocked(schema).records_by_code))
                codes=set(snapshots[0][1])
                for _,records in snapshots[1:]: codes.intersection_update(records)
                fields=[f for c in snapshots[0][0].get("categories",[]) if c["kind"]=="main" for f in c.get("fields",[]) if f["type"] not in SYSTEM_FIELD_TYPES and f["type"]!="file"][:3]
                matches=[]
                for code in sorted(codes):
                    if not payload.get("include_archived") and any(records[code].get("archived") for _,records in snapshots): continue
                    record=snapshots[0][1][code]
                    description=" · ".join(clean_text(field_display_value(f,record.get("values",{}).get(f["id"],""))) for f in fields).strip(" ·")
                    if query.casefold() not in (code+" "+description).casefold(): continue
                    matches.append({"id":code,"description":description[:500]})
            return {"profiles":matches[offset:offset+40],"total":len(matches),"offset":offset,"has_more":offset+40<len(matches)}
        if action == "list": return {"templates": store.list()}
        if action == "save": return {"template": store.save(payload.get("template"), payload.get("id"), payload.get("revision"))}
        if action == "import": return {"template": store.save(decode_template(payload.get("markdown")))}
        if action == "delete":
            store.delete(payload.get("id"), payload.get("revision")); return {"ok": True}
        if action == "catalog":
            result = []
            for entry in manager.response()["schemas"]:
                if entry.get("archived"): continue
                with use_context(manager.context(entry["id"])):
                    schema = read_schema_file()
                    fields = [{"id": f["id"], "label": c["label"] + " / " + f["label"], "type": f["type"], "repeated": c["kind"] != "main", "category_id": c["id"]}
                              for c in schema.get("categories", []) for f in c.get("fields", [])
                              if f["type"] != "file" and f["type"] not in SYSTEM_FIELD_TYPES]
                    result.append({"id": entry["id"], "name": entry["name"], "fields": fields, "categories": [{"id":c["id"],"label":c["label"],"kind":c["kind"]} for c in schema.get("categories",[])]})
            return {"schemas": result}
        if action == "generate":
            def load(schema_id, ids):
                with use_context(manager.context(schema_id)):
                    schema = read_schema_file()
                    snapshot = _dataset_snapshot_unlocked(schema)
                    fields = [(c, f) for c in schema.get("categories", []) for f in c.get("fields", []) if f["type"] != "file" and f["type"] not in SYSTEM_FIELD_TYPES]
                    labels = {f["id"]: c["label"] + " / " + f["label"] for c, f in fields}
                    rows = []
                    for code in ids:
                        validate_person_code(code)
                        record = snapshot.records_by_code.get(code)
                        if record is None: raise ReportError("لم يُعثر على ID " + code + " في " + manager.context(schema_id).name)
                        values = {}
                        for category, field in fields:
                            raw = _field_values_for_record(record, category, field)
                            values[field["id"]] = [field_display_value(field, v) for v in raw if v is not None and v != "" and v != []]
                        rows.append({"id": code, "values": values})
                    return labels, rows
            with WORKBOOK_LOCK:
                draft = resolve_report(store.read(payload.get("id")), payload.get("bindings"), load)
            return {"draft": draft}
        raise ReportError("إجراء التقرير غير صالح.")
    except (ReportError, AdvancedFeatureError, WorkspaceError) as exc:
        raise ApplicationError(str(exc)) from exc


def create_advanced_report_export(payload):
    require_builder_access()
    from schemacraft_reports import pdf_bytes, ReportError
    try:
        draft = payload.get("draft")
        content = pdf_bytes(draft, PDF_FONT_PATH)
    except ReportError as exc: raise ApplicationError(str(exc)) from exc
    # The draft is intentionally detached: manual edits never write back to records.
    raw_count = draft.get("record_count", 0)
    count = raw_count if isinstance(raw_count, int) and 0 <= raw_count <= 50000 else 0
    raw_schemas = draft.get("schema_ids", [])
    known = {s["id"] for s in require_workspace().response()["schemas"]}
    schemas = [s for s in raw_schemas if isinstance(s, str) and s in known] if isinstance(raw_schemas, list) else []
    return "SchemaCraft-report-" + datetime.now().strftime("%Y%m%d-%H%M%S") + ".pdf", content, count, schemas


def create_advanced_report_batch(payload):
    require_builder_access()
    import io
    import zipfile
    drafts=payload.get("drafts")
    if not isinstance(drafts,list) or not 1<=len(drafts)<=20:
        raise ApplicationError("حدد من تقرير واحد إلى 20 تقريرًا.")
    buffer=io.BytesIO(); count=0; schemas=set()
    with zipfile.ZipFile(buffer,"w",zipfile.ZIP_DEFLATED) as archive:
        for index,draft in enumerate(drafts,1):
            _,content,n,ids=create_advanced_report_export({"draft":draft})
            title=re.sub(r'[\\/:*?"<>|\x00-\x1f]',"_",draft.get("title","report"))[:80].strip('. ') or "report"
            archive.writestr(f"{index:02d}-{title}.pdf",content)
            count+=n;schemas.update(ids)
    return "SchemaCraft-reports-"+datetime.now().strftime("%Y%m%d-%H%M%S")+".zip",buffer.getvalue(),count,sorted(schemas)


def create_profile_pdf_export(payload: Any) -> tuple[str, bytes, int, list[str]]:
    require_builder_access()
    manager = require_workspace()
    if not isinstance(payload, dict):
        raise ApplicationError("طلب تقرير الملف غير صالح.")
    code = validate_person_code(clean_text(payload.get("record_code")).upper())
    raw_schema_ids = payload.get("schema_ids", [])
    schema_ids = [clean_text(item) for item in raw_schema_ids] if isinstance(raw_schema_ids, list) else []
    if not schema_ids:
        schema_ids = [item["schema_id"] for item in manager.identity_search(code)]
    fields_by_schema = payload.get("field_ids_by_schema", {})
    fields_by_schema = fields_by_schema if isinstance(fields_by_schema, dict) else {}
    raw_global_refs = payload.get("global_refs", [])
    global_refs = {
        clean_text(item)
        for item in (raw_global_refs if isinstance(raw_global_refs, list) else [])
    }
    profiles = []
    used_schemas: list[str] = []
    for schema_id in schema_ids:
        context = manager.context(schema_id)
        with use_context(context):
            schema = read_schema_file()
            try:
                record = _dataset_snapshot_unlocked(schema).records_by_code[code]
            except KeyError:
                continue
            raw_selected = fields_by_schema.get(schema_id, [])
            selected = {
                clean_text(item)
                for item in (raw_selected if isinstance(raw_selected, list) else [])
            }
            sections = _record_report_sections(schema, record, selected, global_refs)
            attachments = _profile_pdf_media(schema, record, sections,
                show_profile_image=payload.get("show_profile_image") is True,
                show_attachments=payload.get("show_attachments") is True)
            profiles.append({
                "schema_name":context.name,
                "title":f'{schema["app"].get("entity_singular", "سجل")} {code}',
                "sections":sections, "attachments":attachments,
            })
            used_schemas.append(schema_id)
    if not profiles:
        raise ApplicationError("لم يُعثر على ملف لهذا الشخص في التصاميم المختارة.")
    title = f"SchemaCraft — {code}"
    try:
        content = profile_pdf_bytes(title, profiles, PDF_FONT_PATH)
    except AdvancedFeatureError as exc:
        raise ApplicationError(str(exc)) from exc
    filename = f"SchemaCraft-profile-{code}-{datetime.now().astimezone().strftime('%Y%m%d-%H%M%S')}.pdf"
    return filename, content, 1, used_schemas


def profile_pdf_batch_codes(payload: Any) -> list[str]:
    raw_codes = payload.get("record_codes") if isinstance(payload, dict) else None
    if not isinstance(raw_codes, list) or not raw_codes:
        raise ApplicationError("أدخل قائمة IDs صالحة لتصدير تقارير PDF.")
    return list(dict.fromkeys(validate_person_code(clean_text(code).upper()) for code in raw_codes))


def create_profile_pdf_batch_export(payload: Any) -> tuple[str, bytes, int, list[str]]:
    """Create one independent report per identity, saving the entire batch together."""
    require_builder_access()
    codes = profile_pdf_batch_codes(payload)
    by_record = payload.get("schema_ids_by_record")
    if by_record is not None and not isinstance(by_record, dict):
        raise ApplicationError("اختيار التصاميم لكل ID غير صالح.")
    output = io.BytesIO()
    used_schemas = []
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for code in codes:
            single = {**payload, "type":"profile_pdf", "record_code":code}
            if by_record is not None:
                selected = by_record.get(code)
                if not isinstance(selected, list) or not selected:
                    raise ApplicationError(f"اختر تصميمًا واحدًا على الأقل للمعرّف {code}.")
                single["schema_ids"] = selected
            try:
                _, content, _, schemas = create_profile_pdf_export(single)
            except ApplicationError as exc:
                raise ApplicationError(f"{code}: {exc}") from exc
            archive.writestr(f"SchemaCraft-profile-{code}.pdf", content)
            used_schemas.extend(item for item in schemas if item not in used_schemas)
    stamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S")
    return f"SchemaCraft-profiles-{stamp}.zip", output.getvalue(), len(codes), used_schemas


def create_portable_export(payload: Any) -> tuple[str, bytes, int, list[str]]:
    require_builder_access()
    manager = require_workspace()
    if not isinstance(payload, dict):
        raise ApplicationError("طلب الحزمة المحمولة غير صالح.")
    schema_id = clean_text(payload.get("schema_id")) or current_schema_id()
    context = manager.context(schema_id)
    with use_context(context):
        schema = read_schema_file()
        count = record_count(schema)
    definitions = GLOBAL_DEFINITIONS.response() if GLOBAL_DEFINITIONS is not None else {}
    category_refs = {
        category.get("global_ref")
        for category in schema.get("categories", [])
        if category.get("global_ref")
    }
    field_refs = {
        field.get("global_ref")
        for category in schema.get("categories", [])
        for field in category.get("fields", [])
        if field.get("global_ref")
    }
    if definitions:
        definitions = {
            **{
                key: copy.deepcopy(value)
                for key, value in definitions.items()
                if key not in {"categories", "fields"}
            },
            "categories": {
                ref: copy.deepcopy(item)
                for ref, item in definitions.get("categories", {}).items()
                if ref in category_refs
            },
            "fields": {
                ref: copy.deepcopy(item)
                for ref, item in definitions.get("fields", {}).items()
                if ref in field_refs
            },
        }
    try:
        content = portable_package_bytes(
            schema_id=schema_id,
            schema_name=context.name,
            schema_path=context.schema_path,
            workbook_path=context.workbook_path,
            attachments_path=context.attachments_path,
            global_definitions=definitions,
        )
    except (OSError, AdvancedFeatureError) as exc:
        raise ApplicationError(f"تعذّر إنشاء الحزمة المحمولة: {exc}") from exc
    filename = f"SchemaCraft-{context.name}-{datetime.now().astimezone().strftime('%Y%m%d-%H%M%S')}.zip"
    return filename, content, count, [schema_id]


def choose_export_destination_response(payload: Any) -> dict[str, Any]:
    """Choose a destination before generation so the page workflow stays explicit."""

    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("طلب مكان التصدير غير صالح.")
    export_type = clean_text(payload.get("type"))
    stamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S")
    if export_type == "advanced_report":
        filename = f"SchemaCraft-report-{stamp}.pdf"
        file_types = [("PDF report", "*.pdf")]
    elif export_type == "advanced_report_batch":
        filename = f"SchemaCraft-reports-{stamp}.zip"
        file_types = [("PDF reports ZIP", "*.zip")]
    elif export_type == "profile_pdf":
        code = validate_person_code(clean_text(payload.get("record_code")).upper())
        filename = f"SchemaCraft-profile-{code}-{stamp}.pdf"
        file_types = [("PDF report", "*.pdf")]
    elif export_type == "profile_pdf_batch":
        profile_pdf_batch_codes(payload)
        filename = f"SchemaCraft-profiles-{stamp}.zip"
        file_types = [("PDF reports ZIP", "*.zip")]
    elif export_type == "table":
        include_attachments = bool(payload.get("include_attachments"))
        filename = f"SchemaCraft-export-{stamp}{'-with-attachments.zip' if include_attachments else '.xlsx'}"
        file_types = [("ZIP archive", "*.zip")] if include_attachments else [("Excel workbook", "*.xlsx")]
    elif export_type == "portable_zip":
        manager = require_workspace()
        schema_id = clean_text(payload.get("schema_id")) or current_schema_id()
        context = manager.context(schema_id)
        filename = f"SchemaCraft-{context.name}-{stamp}.zip"
        file_types = [("SchemaCraft package", "*.zip")]
    else:
        raise ApplicationError("نوع التصدير غير صالح.")
    try:
        destination = choose_export_destination(filename, file_types)
    except (OSError, AdvancedFeatureError) as exc:
        raise ApplicationError(f"تعذّر اختيار مكان التصدير: {exc}") from exc
    if destination is None:
        return {"ok": False, "cancelled": True}
    return {"ok": True, "cancelled": False, "destination": str(destination)}


def save_export(payload: Any) -> dict[str, Any]:
    """Generate an export, ask for its native destination, and audit it."""

    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("طلب التصدير غير صالح.")
    export_type = clean_text(payload.get("type"))
    if export_type == "table":
        filename, content, count = create_filtered_export(payload)
        schemas = [clean_text(payload.get("schema_id")) or current_schema_id()]
        file_types = (
            [("ZIP archive", "*.zip")]
            if bool(payload.get("include_attachments"))
            else [("Excel workbook", "*.xlsx")]
        )
    elif export_type == "advanced_report":
        filename, content, count, schemas = create_advanced_report_export(payload)
        file_types = [("PDF report", "*.pdf")]
    elif export_type == "advanced_report_batch":
        filename, content, count, schemas = create_advanced_report_batch(payload)
        file_types = [("PDF reports ZIP", "*.zip")]
    elif export_type == "profile_pdf":
        filename, content, count, schemas = create_profile_pdf_export(payload)
        file_types = [("PDF report", "*.pdf")]
    elif export_type == "profile_pdf_batch":
        filename, content, count, schemas = create_profile_pdf_batch_export(payload)
        file_types = [("PDF reports ZIP", "*.zip")]
    elif export_type == "portable_zip":
        filename, content, count, schemas = create_portable_export(payload)
        file_types = [("SchemaCraft package", "*.zip")]
    else:
        raise ApplicationError("نوع التصدير غير صالح.")
    raw_destination = clean_text(payload.get("destination"))
    selected_destination = Path(raw_destination) if raw_destination else None
    if selected_destination is not None:
        if not selected_destination.is_absolute():
            raise ApplicationError("مكان حفظ التصدير يجب أن يكون مسارًا كاملًا.")
        try:
            selected_destination = normalize_export_destination(
                selected_destination, file_types
            )
        except AdvancedFeatureError as exc:
            raise ApplicationError(str(exc)) from exc
    try:
        destination = save_export_bytes(content, filename, file_types, destination=selected_destination)
    except (OSError, AdvancedFeatureError) as exc:
        raise ApplicationError(f"تعذّر حفظ ملف التصدير: {exc}") from exc
    if destination is None:
        return {"ok": False, "cancelled": True}
    checksum = hashlib.sha256(content).hexdigest()
    history_id = hashlib.sha256(os.urandom(32)).hexdigest()[:16]
    backup_archive = ""
    try:
        backup_archive = _archive_export_content(history_id, destination.name, content)
    except OSError as exc:
        LOGGER.warning("Export saved but data-folder backup failed: %s", exc)
    history = None
    if EXPORT_HISTORY is not None:
        try:
            schema_names = []
            for schema_id in schemas:
                try:
                    context = WORKSPACE_MANAGER.context(schema_id) if WORKSPACE_MANAGER else None
                except AdvancedFeatureError:
                    context = None
                schema_names.append(context.name if context else schema_id)
            history = EXPORT_HISTORY.append(
                {
                    "id": history_id,
                    "type": export_type,
                    "schemas": schemas,
                    "schema_names": schema_names,
                    "person_id": clean_text(payload.get("record_code")),
                    "person_ids": copy.deepcopy(payload.get("record_codes", [])),
                    "row_count": count,
                    "filename": destination.name,
                    "destination": str(destination),
                    "status": "success",
                    "checksum": checksum,
                    "backup_archive": backup_archive,
                    "notes": payload.get("notes", ""),
                    "user_name": current_audit_user(),
                    "configuration": {
                        key: copy.deepcopy(payload.get(key))
                        for key in (
                            "type", "schema_id", "criteria", "field_ids",
                            "include_related", "related_category_ids",
                            "include_attachments", "record_code", "schema_ids",
                            "field_ids_by_schema", "global_refs", "show_profile_image", "show_attachments",
                            "record_codes", "schema_ids_by_record",
                        )
                        if key in payload
                    },
                }
            )
        except AdvancedFeatureError as exc:
            LOGGER.warning("Export saved but history failed: %s", exc)
    return {
        "ok": True,
        "cancelled": False,
        "filename": destination.name,
        "destination": str(destination),
        "row_count": count,
        "checksum": checksum,
        "backup_archive": backup_archive,
        "history": history,
    }


def _archive_export_content(entry_id: str, filename: str, content: bytes) -> str:
    """Retain the exact generated export beneath the application data folder."""

    suffix = Path(filename).suffix.casefold()
    if suffix not in {".xlsx", ".zip", ".pdf"}:
        raise OSError("Unsupported export backup extension.")
    safe_filename = sanitize_filename_text(filename) or f"export{suffix}"
    archive_root = DATA_DIR / "exported-files"
    folder = archive_root / entry_id
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / safe_filename
    with tempfile.NamedTemporaryFile(
        mode="wb", prefix=".export-", suffix=suffix, dir=folder, delete=False
    ) as temporary:
        temporary.write(content)
        temporary_path = Path(temporary.name)
    try:
        os.replace(temporary_path, target)
    finally:
        temporary_path.unlink(missing_ok=True)
    return f"{entry_id}/{safe_filename}"


def export_history_response(payload: Any = None) -> dict[str, Any]:
    require_builder_access()
    if EXPORT_HISTORY is None:
        return {"entries": [], "total": 0}
    raw_limit = payload.get("limit", 20) if isinstance(payload, dict) else 20
    limit = None if clean_text(raw_limit).casefold() == "all" else int(raw_limit)
    try:
        result = EXPORT_HISTORY.response(limit)
        for entry in result.get("entries", []):
            if entry.get("schema_names"):
                continue
            names = []
            for schema_id in entry.get("schemas", []):
                try:
                    context = WORKSPACE_MANAGER.context(schema_id) if WORKSPACE_MANAGER else None
                except AdvancedFeatureError:
                    context = None
                names.append(context.name if context else schema_id)
            entry["schema_names"] = names
        return result
    except (AdvancedFeatureError, TypeError, ValueError) as exc:
        raise ApplicationError(str(exc)) from exc


def import_history_response(payload: Any = None) -> dict[str, Any]:
    require_builder_access()
    if IMPORT_HISTORY is None:
        return {"entries": [], "total": 0}
    raw_limit = payload.get("limit", 20) if isinstance(payload, dict) else 20
    limit = None if clean_text(raw_limit).casefold() == "all" else int(raw_limit)
    try:
        return IMPORT_HISTORY.response(limit)
    except (AdvancedFeatureError, TypeError, ValueError) as exc:
        raise ApplicationError(str(exc)) from exc


def update_operation_history_notes(payload: Any) -> dict[str, Any]:
    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("طلب تحديث ملاحظات السجل غير صالح.")
    kind = clean_text(payload.get("kind"))
    entry_id = clean_text(payload.get("id"))
    store = IMPORT_HISTORY if kind == "import" else EXPORT_HISTORY if kind == "export" else None
    if store is None or not entry_id:
        raise ApplicationError("عنصر سجل العمليات المطلوب غير صالح.")
    try:
        return {"ok": True, "entry": store.update_notes(entry_id, payload.get("notes"))}
    except AdvancedFeatureError as exc:
        raise ApplicationError(str(exc)) from exc


def delete_operation_history(kind: str, entry_id: Any) -> dict[str, Any]:
    require_builder_access()
    clean_id = clean_text(entry_id)
    store = IMPORT_HISTORY if kind == "import" else EXPORT_HISTORY if kind == "export" else None
    if store is None or not clean_id:
        raise ApplicationError("عنصر سجل العمليات المطلوب غير صالح.")
    try:
        removed = store.delete(clean_id)
    except AdvancedFeatureError as exc:
        raise ApplicationError(str(exc)) from exc
    if kind == "import":
        archive = clean_text(removed.get("source_archive"))
        if archive:
            archive_target = (IMPORT_ARCHIVE_DIR / archive).resolve()
            try:
                archive_target.relative_to(IMPORT_ARCHIVE_DIR.resolve())
            except ValueError:
                archive_target = None
            if archive_target is not None:
                if archive_target.is_file():
                    archive_target.unlink(missing_ok=True)
                parent = archive_target.parent
                if parent != IMPORT_ARCHIVE_DIR.resolve() and parent.is_dir():
                    try:
                        parent.rmdir()
                    except OSError:
                        pass
    elif kind == "export":
        archive = clean_text(removed.get("backup_archive"))
        archive_root = (DATA_DIR / "exported-files").resolve()
        if archive:
            archive_target = (archive_root / archive).resolve()
            try:
                archive_target.relative_to(archive_root)
            except ValueError:
                archive_target = None
            if archive_target is not None:
                if archive_target.is_file():
                    archive_target.unlink(missing_ok=True)
                parent = archive_target.parent
                if parent != archive_root and parent.is_dir():
                    try:
                        parent.rmdir()
                    except OSError:
                        pass
    return {"ok": True, "entry": removed}


def _remove_history_archive_root(target: Path) -> None:
    """Remove one known history archive directory without touching its parent."""

    data_root = DATA_DIR.resolve()
    resolved = target.resolve()
    if resolved == data_root or data_root not in resolved.parents:
        raise ApplicationError("مسار أرشيف السجل غير آمن.")
    if resolved.is_dir():
        shutil.rmtree(resolved)


def clear_history(kind: str) -> dict[str, Any]:
    """Permanently clear one authoritative history and its retained files."""

    normalized = clean_text(kind).casefold()
    if normalized == "search":
        if SEARCH_HISTORY is None:
            return {"ok": True, "kind": normalized, "removed": 0}
        try:
            removed = SEARCH_HISTORY.clear()
        except AdvancedFeatureError as exc:
            raise ApplicationError(str(exc)) from exc
        return {"ok": True, "kind": normalized, "removed": len(removed)}
    require_builder_access()
    if normalized == "import":
        store = IMPORT_HISTORY
        archive_root = IMPORT_ARCHIVE_DIR
    elif normalized == "export":
        store = EXPORT_HISTORY
        archive_root = DATA_DIR / "exported-files"
    else:
        raise ApplicationError("نوع سجل العمليات المطلوب غير صالح.")
    if store is None:
        return {"ok": True, "kind": normalized, "removed": 0}
    try:
        removed = store.clear()
        _remove_history_archive_root(archive_root)
    except (AdvancedFeatureError, OSError) as exc:
        raise ApplicationError(f"تعذّر مسح السجل نهائيًا: {exc}") from exc
    return {"ok": True, "kind": normalized, "removed": len(removed)}


def open_export_history_file(entry_id: Any) -> dict[str, Any]:
    require_builder_access()
    clean_id = clean_text(entry_id)
    if EXPORT_HISTORY is None or not clean_id:
        raise ApplicationError("عنصر سجل التصدير المطلوب غير صالح.")
    entry = next((item for item in EXPORT_HISTORY.read() if item.get("id") == clean_id), None)
    if entry is None:
        raise ApplicationError("عنصر سجل التصدير المطلوب غير موجود.")
    target = Path(clean_text(entry.get("destination")))
    if not target.is_file():
        relative = clean_text(entry.get("backup_archive"))
        archive_root = (DATA_DIR / "exported-files").resolve()
        backup_target = (archive_root / relative).resolve() if relative else None
        if backup_target is not None:
            try:
                backup_target.relative_to(archive_root)
            except ValueError:
                backup_target = None
        if backup_target is None or not backup_target.is_file():
            raise ApplicationError("ملف التصدير لم يعد موجودًا في مكانه أو في النسخة الاحتياطية.")
        target = backup_target
    try:
        if os.name == "nt":
            os.startfile(str(target))  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(target)])
        else:
            subprocess.Popen(["xdg-open", str(target)])
    except OSError as exc:
        raise ApplicationError(f"تعذّر فتح ملف التصدير: {exc}") from exc
    return {
        "ok": True,
        "filename": target.name,
        "from_backup": target != Path(clean_text(entry.get("destination"))),
    }


def _record_import_history(payload: dict[str, Any], result: dict[str, Any], import_type: str) -> dict[str, Any] | None:
    if IMPORT_HISTORY is None:
        return None
    context = active_context()
    try:
        entry = IMPORT_HISTORY.append(
            {
                "type": import_type,
                "schema_id": current_schema_id(),
                "schema_name": context.name if context else "",
                "filename": clean_text(payload.get("filename")),
                "status": "success",
                "added": result.get("imported", result.get("created", 0)),
                "updated": result.get("updated", 0),
                "skipped": result.get("skipped", 0),
                "rejected": result.get("rejected", 0),
                "details": result.get("changes", []) + result.get("errors", []),
                "notes": payload.get("notes", ""),
                "user_name": current_audit_user(),
            }
        )
        source = _archive_import_source(entry["id"], payload)
        if source:
            entry = IMPORT_HISTORY.update_source_archive(entry["id"], source)
        return entry
    except AdvancedFeatureError as exc:
        LOGGER.warning("Import completed but history failed: %s", exc)
        return None


def log_import_history(payload: Any) -> dict[str, Any]:
    require_builder_access()
    if IMPORT_HISTORY is None or not isinstance(payload, dict):
        raise ApplicationError("طلب تسجيل الاستيراد غير صالح.")
    context = active_context()
    status = clean_text(payload.get("status")) or "aborted"
    if status not in {"aborted", "failed"}:
        raise ApplicationError("حالة سجل الاستيراد غير صالحة.")
    try:
        entry = IMPORT_HISTORY.append(
            {
                "type": clean_text(payload.get("type")) or "excel",
                "schema_id": current_schema_id(),
                "schema_name": context.name if context else "",
                "filename": clean_text(payload.get("filename")),
                "status": status,
                "details": payload.get("details") or [],
                "notes": payload.get("notes", ""),
                "user_name": current_audit_user(),
            }
        )
    except AdvancedFeatureError as exc:
        raise ApplicationError(str(exc)) from exc
    return {"ok": True, "entry": entry}


def _archive_import_source(entry_id: str, payload: dict[str, Any]) -> str:
    """Keep the exact imported source under the application data directory."""
    encoded = payload.get("file_data")
    filename = sanitize_filename_text(payload.get("filename"))
    if not isinstance(encoded, str) or not encoded or not filename:
        return ""
    try:
        content = base64.b64decode(encoded, validate=True)
    except (ValueError, binascii.Error):
        return ""
    if not content or len(content) > MAX_REQUEST_BYTES:
        return ""
    suffix = Path(filename).suffix.casefold()
    if suffix not in {".xlsx", ".zip"}:
        return ""
    safe_name = f"source{suffix}"
    folder = IMPORT_ARCHIVE_DIR / entry_id
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / safe_name
    with tempfile.NamedTemporaryFile(
        mode="wb", prefix=".source-", suffix=suffix, dir=folder, delete=False
    ) as temporary:
        temporary.write(content)
        temporary_path = Path(temporary.name)
    try:
        os.replace(temporary_path, target)
    finally:
        temporary_path.unlink(missing_ok=True)
    return f"{entry_id}/{safe_name}"


def import_source_path(entry_id: str) -> tuple[Path, str] | None:
    if IMPORT_HISTORY is None:
        return None
    entry = next((item for item in IMPORT_HISTORY.read() if item.get("id") == entry_id), None)
    relative = clean_text(entry.get("source_archive")) if entry else ""
    if not relative:
        return None
    candidate = (IMPORT_ARCHIVE_DIR / relative).resolve()
    root = IMPORT_ARCHIVE_DIR.resolve()
    if root not in candidate.parents or not candidate.is_file():
        return None
    original = clean_text(entry.get("filename")) or candidate.name
    return candidate, original


def open_import_history_file(entry_id: Any) -> dict[str, Any]:
    """Open the retained import source locally without sending its bytes."""

    require_builder_access()
    clean_id = clean_text(entry_id)
    if not clean_id:
        raise ApplicationError("عنصر سجل الاستيراد المطلوب غير صالح.")
    resolved = import_source_path(clean_id)
    if resolved is None:
        raise ApplicationError("ملف الاستيراد المحفوظ لم يعد موجودًا.")
    target, original_name = resolved
    try:
        if os.name == "nt":
            os.startfile(str(target))  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(target)])
        else:
            subprocess.Popen(["xdg-open", str(target)])
    except OSError as exc:
        raise ApplicationError(f"تعذّر فتح ملف الاستيراد: {exc}") from exc
    return {"ok": True, "filename": original_name}


def search_history_response(payload: Any = None) -> dict[str, Any]:
    if SEARCH_HISTORY is None:
        return {"entries": [], "total": 0}
    raw_limit = payload.get("limit", 50) if isinstance(payload, dict) else 50
    limit = None if clean_text(raw_limit).casefold() == "all" else int(raw_limit)
    query = clean_text(payload.get("query")) if isinstance(payload, dict) else ""
    try:
        return SEARCH_HISTORY.response(limit, query)
    except (AdvancedFeatureError, TypeError, ValueError) as exc:
        raise ApplicationError(str(exc)) from exc


def save_search_history(payload: Any) -> dict[str, Any]:
    if SEARCH_HISTORY is None or not isinstance(payload, dict):
        raise ApplicationError("طلب حفظ البحث غير صالح.")
    mode = clean_text(payload.get("mode")) or "schema"
    if mode not in {"schema", "global"}:
        raise ApplicationError("نوع البحث المحفوظ غير صالح.")
    try:
        entry = SEARCH_HISTORY.append(
            {
                **payload,
                "mode": mode,
                "user_name": current_audit_user(),
            }
        )
    except AdvancedFeatureError as exc:
        raise ApplicationError(str(exc)) from exc
    return {"ok": True, "entry": entry}


def delete_search_history(entry_id: Any) -> dict[str, Any]:
    if SEARCH_HISTORY is None:
        raise ApplicationError("سجل البحث غير مفعّل.")
    try:
        return SEARCH_HISTORY.delete(clean_text(entry_id))
    except AdvancedFeatureError as exc:
        raise ApplicationError(str(exc)) from exc


def inspect_import(payload: Any) -> dict[str, Any]:
    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("صيغة طلب الاستيراد غير صحيحة.")
    schema = read_schema_file()
    try:
        result = inspect_import_workbook(
            payload.get("file_data"),
            schema,
            clean_text(payload.get("sheet_name")) or None,
        )
    except WorkbookExchangeError as exc:
        raise ApplicationError(str(exc)) from exc
    result.update(inspect_import_sheets(payload.get("file_data"), schema))
    result["schema_revision"] = schema["revision"]
    return result


def _decode_portable_data(value: Any) -> bytes:
    if not isinstance(value, str) or not value:
        raise ApplicationError("بيانات حزمة الاستيراد غير مكتملة.")
    try:
        content = base64.b64decode(value, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise ApplicationError("تعذّر قراءة حزمة الاستيراد.") from exc
    if not content.startswith(b"PK"):
        raise ApplicationError("اختر حزمة SchemaCraft بصيغة ZIP.")
    if len(content) > MAX_REQUEST_BYTES:
        raise ApplicationError("حجم حزمة الاستيراد يتجاوز الحد المسموح.")
    return content


def _portable_package_records(package: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    schema = validate_schema(package["schema"])
    with tempfile.TemporaryDirectory(prefix="schemacraft-package-") as temporary:
        root = Path(temporary)
        schema_path = root / "schema.json"
        workbook_path = root / "database.xlsx"
        schema_path.write_text(
            json.dumps(schema, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        workbook_path.write_bytes(package["record_workbook_bytes"])
        context = SchemaContext(
            schema_id="0" * 32,
            name="portable",
            folder=root,
            schema_path=schema_path,
            workbook_path=workbook_path,
            attachments_path=root / "attachments",
        )
        with use_context(context):
            records = read_dataset_unlocked(schema)
    return schema, records


def inspect_portable_import(payload: Any) -> dict[str, Any]:
    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("طلب فحص الحزمة غير صالح.")
    try:
        package = inspect_portable_package(_decode_portable_data(payload.get("file_data")))
        package_schema, package_records = _portable_package_records(package)
    except AdvancedFeatureError as exc:
        raise ApplicationError(str(exc)) from exc
    target = read_schema_file()
    target_fields = schema_indexes(target)["fields"]
    incoming_fields = schema_indexes(package_schema)["fields"]
    incompatible = [
        field_id
        for field_id, field in incoming_fields.items()
        if field_id in target_fields and target_fields[field_id]["type"] != field["type"]
    ]
    with WORKBOOK_LOCK:
        local_codes = set(_dataset_snapshot_unlocked(target).records_by_code)
    incoming_codes = {record["record_code"] for record in package_records}
    return {
        "ok": True,
        "manifest": package["manifest"],
        "record_count": len(package_records),
        "new_count": len(incoming_codes - local_codes),
        "matching_count": len(incoming_codes & local_codes),
        "attachment_count": len(package["attachments"]),
        "schema_revision": target["revision"],
        "package_schema_revision": package_schema["revision"],
        "incompatible_field_ids": incompatible,
        "can_import": not incompatible,
    }


def commit_portable_import(payload: Any) -> dict[str, Any]:
    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("طلب استيراد الحزمة غير صالح.")
    policy = clean_text(payload.get("conflict_policy")) or "newer"
    if policy not in {"newer", "overwrite"}:
        raise ApplicationError("سياسة تعارض الحزمة غير صالحة.")
    try:
        package = inspect_portable_package(_decode_portable_data(payload.get("file_data")))
        package_schema, package_records = _portable_package_records(package)
    except AdvancedFeatureError as exc:
        raise ApplicationError(str(exc)) from exc
    current = read_schema_file()
    current_index = schema_indexes(current)
    incoming_index = schema_indexes(package_schema)
    incompatible = [
        field_id
        for field_id, field in incoming_index["fields"].items()
        if field_id in current_index["fields"]
        and current_index["fields"][field_id]["type"] != field["type"]
    ]
    if incompatible:
        raise ApplicationError("تحتوي الحزمة على أنواع حقول غير متوافقة مع التصميم الهدف.")
    backup = create_backup(automatic=True)
    if policy == "overwrite" or package_schema["revision"] > current["revision"]:
        candidate = copy.deepcopy(package_schema)
        candidate["revision"] = current["revision"]
        save_schema(candidate)
        current = read_schema_file()
    with WORKBOOK_LOCK:
        snapshot = _dataset_snapshot_unlocked(current)
        records = [copy.deepcopy(record) for record in snapshot.records]
        positions = {record["record_code"]: index for index, record in enumerate(records)}
        created: list[str] = []
        updated: list[str] = []
        skipped: list[str] = []
        for incoming in package_records:
            code = incoming["record_code"]
            position = positions.get(code)
            if position is None:
                candidate = copy.deepcopy(incoming)
                candidate["_record_id"] = new_internal_id()
                for rows in candidate.get("related", {}).values():
                    for row in rows:
                        row["_child_id"] = new_internal_id()
                records.append(candidate)
                positions[code] = len(records) - 1
                created.append(code)
                continue
            local = records[position]
            if policy == "newer" and clean_text(incoming.get("updated_at")) <= clean_text(local.get("updated_at")):
                skipped.append(code)
                continue
            candidate = copy.deepcopy(incoming)
            candidate["_record_id"] = local["_record_id"]
            existing_children = {
                (category_id, row.get("minor_id")): row.get("_child_id")
                for category_id, rows in local.get("related", {}).items()
                for row in rows
            }
            for category_id, rows in candidate.get("related", {}).items():
                for row in rows:
                    row["_child_id"] = (
                        existing_children.get((category_id, row.get("minor_id")))
                        or new_internal_id()
                    )
            records[position] = candidate
            updated.append(code)
        atomic_write_workbook(current, records)
        _safe_publish_dataset_snapshot(current, records)
        if WORKSPACE_MANAGER is not None:
            WORKSPACE_MANAGER.register_profiles(created, current_schema_id())
        for name, content in package["attachments"].items():
            relative = Path(*PurePosixPath(name).parts[1:])
            destination = attachments_directory() / relative
            destination.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(
                prefix=f".{destination.stem}-",
                suffix=destination.suffix,
                dir=destination.parent,
                delete=False,
            ) as temporary:
                temporary.write(content)
                temporary_path = Path(temporary.name)
            os.replace(temporary_path, destination)
    if GLOBAL_DEFINITIONS is not None:
        incoming_globals = package.get("global_definitions", {})
        current_globals = GLOBAL_DEFINITIONS.response()
        for kind, collection_name in (("category", "categories"), ("field", "fields")):
            incoming_collection = incoming_globals.get(collection_name, {})
            if not isinstance(incoming_collection, dict):
                continue
            for global_ref, item in incoming_collection.items():
                if not isinstance(item, dict) or not isinstance(item.get("definition"), dict):
                    continue
                local_item = current_globals.get(collection_name, {}).get(global_ref)
                if (
                    policy == "newer"
                    and local_item
                    and clean_text(item.get("updated_at")) <= clean_text(local_item.get("updated_at"))
                ):
                    continue
                try:
                    GLOBAL_DEFINITIONS.save_definition(kind, global_ref, item["definition"])
                except AdvancedFeatureError as exc:
                    LOGGER.warning("Portable global definition skipped: %s", exc)
    result = {
        "ok": True,
        "created": len(created),
        "updated": len(updated),
        "skipped": len(skipped),
        "created_ids": created,
        "updated_ids": updated,
        "backup": backup,
    }
    result["history"] = _record_import_history(payload, result, "portable_zip")
    return result


# Import previews are short-lived, server-owned proposals. No workbook is changed here.
_IMPORT_REVIEWS: dict[str, dict[str, Any]] = {}


def _merge_import_cards(previous, incoming, mode, clear_blanks):
    if mode == 'replace': return copy.deepcopy(incoming)
    rows=copy.deepcopy(previous)
    by_minor={row.get('minor_id',index):row for index,row in enumerate(rows,1)}
    for patch in incoming:
        minor=patch.get('minor_id')
        if minor is not None and minor not in by_minor:
            raise ApplicationError(f'رقم البطاقة {minor} غير موجود؛ اترك رقم البطاقة فارغًا لإضافة بطاقة جديدة.')
        if minor is None:
            rows.append(copy.deepcopy(patch)); continue
        row=by_minor[minor]
        for field,value in patch.get('values',{}).items():
            if clear_blanks or clean_text(value): row.setdefault('values',{})[field]=value
        for key in ['linked_record_code','parent_minor_id','parent_child_id']:
            if key in patch: row[key]=patch[key]
    return rows


def _prepare_import_cards(schema, related):
    for rows in related.values():
        for row in rows: row['_child_id']=row.get('_child_id') or new_internal_id()
    for category in schema['categories']:
        parent_id=category.get('parent_category_id')
        parents=related.get(parent_id,[])
        for row in related.get(category['id'],[]):
            if row.get('parent_minor_id'):
                try: index=int(row['parent_minor_id'])-1
                except (TypeError,ValueError): raise ApplicationError('رقم بطاقة الفئة الأم غير صالح.')
                if index<0 or index>=len(parents): raise ApplicationError('بطاقة الفئة الأم غير موجودة.')
                row['parent_child_id']=parents[index]['_child_id']
    return related


def _import_review_value(value):
    if isinstance(value, (datetime,date)): return value.isoformat()
    if isinstance(value, dict): return {key:_import_review_value(item) for key,item in value.items()}
    if isinstance(value, (list,tuple)): return [_import_review_value(item) for item in value]
    return value


def _import_changes(schema, before, after):
    before=before or {'values':{},'related':{},'archived':False}
    changes=[]
    def add(kind, label, old, new, **extra):
        old,new=_import_review_value(old),_import_review_value(new)
        if old!=new: changes.append({'id':str(len(changes)), 'kind':kind,'label':label,'before':copy.deepcopy(old),'after':copy.deepcopy(new),**extra})
    for category in schema['categories']:
        fields={f['id']:f['label'] for f in category.get('fields',[])}
        if category['kind']=='main':
            for fid,label in fields.items(): add('field',f'{category["label"]} — {label}',before.get('values',{}).get(fid,''),after.get('values',{}).get(fid,''),field_id=fid)
        else:
            old={r['_child_id']:r for r in before.get('related',{}).get(category['id'],[])}
            new={r['_child_id']:r for r in after.get('related',{}).get(category['id'],[])}
            def display(row):
                return {fields.get(fid,fid):value for fid,value in row.get('values',{}).items()} if row else None
            for child in dict.fromkeys([*old,*new]):
                a,b=old.get(child),new.get(child)
                label=f'{category["label"]} — البطاقة {(b or a).get("minor_id", "")}'
                if a is None or b is None:
                    add('card',label,display(a),display(b),category_id=category['id'],child_id=child)
                else:
                    for fid,field_label in fields.items(): add('card_field',f'{label} — {field_label}',a['values'].get(fid,''),b['values'].get(fid,''),category_id=category['id'],child_id=child,field_id=fid)
                    for key,title in [('linked_record_code','الشخص المرتبط'),('parent_child_id','البطاقة الأم')]: add('card_property',f'{label} — {title}',a.get(key,''),b.get(key,''),category_id=category['id'],child_id=child,field_id=key)
    add('archived','حالة الأرشفة',bool(before.get('archived')),bool(after.get('archived')))
    return changes


def _store_import_review(schema, records, accepted, errors, skipped, payload):
    timestamp=time.monotonic()
    for token,plan in list(_IMPORT_REVIEWS.items()):
        if timestamp-plan['created']>1800: _IMPORT_REVIEWS.pop(token,None)
    while len(_IMPORT_REVIEWS)>=4: _IMPORT_REVIEWS.pop(next(iter(_IMPORT_REVIEWS)))
    entries=[]
    for item in accepted:
        changes=_import_changes(schema,item.get('previous'),item['record'])
        entries.append({**item,'changes':changes})
    token=secrets.token_urlsafe(32)
    _IMPORT_REVIEWS[token]={'created':timestamp,'schema_id':current_schema_id(),'user':current_audit_user(),
        'schema_signature':_schema_cache_signature(schema),'workbook_signature':_workbook_cache_signature(), 'workbook_path':str(_workbook_path().resolve()),
        'payload':copy.deepcopy(payload),'entries':entries,'errors':errors,'skipped':skipped}
    return {'ok':True,'review_token':token,'entries':[{'record_code':item['record']['record_code'],
        'action':'updated' if item.get('updated') and item['changes'] else 'unchanged' if item.get('updated') else 'added',
        'sources':item.get('sources',[]),'changes':item['changes']} for item in entries], 'errors':errors,
        'skipped':skipped,'rejected':len(errors)-skipped}


def apply_import_review(payload):
    require_builder_access()
    token=clean_text(payload.get('review_token'))
    with WORKBOOK_LOCK:
        plan=_IMPORT_REVIEWS.get(token)
        if not plan or time.monotonic()-plan['created']>1800: raise ApplicationError('انتهت صلاحية المراجعة. أعد المعاينة.')
        if plan['schema_id']!=current_schema_id() or plan['user']!=current_audit_user() or plan['workbook_path']!=str(_workbook_path().resolve()): raise ApplicationError('هذه المراجعة لا تخص المستخدم أو التصميم الحالي.')
        schema=read_schema_file()
        if plan['schema_signature']!=_schema_cache_signature(schema) or plan['workbook_signature']!=_workbook_cache_signature(): raise ApplicationError('تغيّرت البيانات أو الفئات بعد المعاينة. أعد المعاينة قبل الحفظ.')
        selections=payload.get('selections')
        if not isinstance(selections,list): raise ApplicationError('اختيارات المراجعة غير صالحة.')
        selected={}
        for item in selections:
            if not isinstance(item,dict) or not isinstance(item.get('change_ids'),list): raise ApplicationError('اختيارات المراجعة غير صالحة.')
            code=clean_text(item.get('record_code'))
            if code in selected: raise ApplicationError('ID مكرر في اختيارات المراجعة.')
            selected[code]=set(map(str,item['change_ids']))
        entries={item['record']['record_code']:item for item in plan['entries']}
        if not selected or not set(selected)<=set(entries): raise ApplicationError('اختر ملفًا صالحًا واحدًا على الأقل.')
        records=[copy.deepcopy(record) for record in _dataset_snapshot_unlocked(schema).records]
        applied=[]; audit=[]; added=updated=unchanged=0
        for code,item in entries.items():
            if code not in selected: audit.append({"record_code":code,"action":"skipped","sources":item.get("sources",[]),"changes":[],"message":"استُبعد من المراجعة"})
        for code,choice in selected.items():
            item=entries[code]; candidate=item['record']; previous=item.get('previous')
            available={change['id'] for change in item['changes']}
            if not choice<=available: raise ApplicationError('اختيار تغيير غير صالح.')
            if previous is None:
                if choice!=available: raise ApplicationError('يجب اعتماد جميع بيانات الملف الجديد أو استبعاده.')
                record=copy.deepcopy(candidate)
                if WORKSPACE_MANAGER:
                    if item.get('generated') and WORKSPACE_MANAGER.person_id_in_use(code): raise ApplicationError('استُخدم معرّف مولّد بعد المراجعة. أعد المعاينة.')
                    WORKSPACE_MANAGER.assert_profile_creation(code,current_schema_id(),link_existing=True)
            else:
                record=copy.deepcopy(previous)
                for change in item['changes']:
                    if change['id'] not in choice: continue
                    kind=change['kind']
                    if kind=='field': record['values'][change['field_id']]=copy.deepcopy(change['after'])
                    elif kind=='archived': record['archived']=change['after'];record['archived_at']=candidate['archived_at']
                    else:
                        category=change['category_id']; child=change['child_id']; rows=record['related'].setdefault(category,[])
                        row=next((r for r in rows if r['_child_id']==child),None)
                        source=next((r for r in candidate['related'].get(category,[]) if r['_child_id']==child),None)
                        if kind=='card':
                            if row: rows.remove(row)
                            if source: rows.append(copy.deepcopy(source))
                        elif kind=='card_field': row['values'][change['field_id']]=copy.deepcopy(change['after'])
                        else: row[change['field_id']]=copy.deepcopy(change['after'])
                for rows in record['related'].values():
                    for index,row in enumerate(rows,1): row['minor_id']=index
            try:
                main,related,staged,_kept=_normalize_submission(schema,{'main':record['values'],'related':record['related']},'update' if previous else 'create',_record_file_paths(schema,previous) if previous else set())
                if staged:
                    for temporary,_destination in staged: temporary.unlink(missing_ok=True)
                    raise ApplicationError('لا يمكن استيراد ملفات مرفقة من Excel.')
                record['values']=main
                for category,rows in related.items():
                    existing={r['_child_id']:r for r in record['related'].get(category,[])}
                    record['related'][category]=[{**existing.get(r['_child_id'],{}),**r,'minor_id':index} for index,r in enumerate(rows,1)]
            except ApplicationError as exc: raise ApplicationError(f'{code}: {exc}') from exc
            differences=_import_changes(schema,previous,record)
            if previous and not differences:
                unchanged+=1
                audit.append({"record_code":code,"action":"unchanged","sources":item.get("sources",[]),"changes":[],"message":"تم الاحتفاظ بالقيم الحالية"})
                continue
            # Revalidation must not introduce changes the user did not approve.
            expected=_import_changes(schema,previous,candidate)
            expected_keys={(c['kind'],c.get('category_id'),c.get('child_id'),c.get('field_id')) for c in expected if c['id'] in choice}
            if previous and any((c['kind'],c.get('category_id'),c.get('child_id'),c.get('field_id')) not in expected_keys for c in differences): raise ApplicationError(f'{code}: الاختيارات تؤثر على حقول أخرى؛ راجع الحقول التابعة معًا.')
            record['updated_at']=now_iso()
            if previous:
                index=next(i for i,r in enumerate(records) if r['record_code']==code);records[index]=record;updated+=1
            else: records.append(record);added+=1
            applied.append({'record_code':code,'action':'updated' if previous else 'added','sources':item.get('sources',[]),'changes':differences,'message':'تم تحديث الملف' if previous else 'تمت إضافة الملف'})
        # Validate the final combined proposal, including links between selected profiles.
        for entry in applied:
            record=next(r for r in records if r['record_code']==entry['record_code'])
            try:
                validate_unique_fields(schema,records,record['_record_id'],record['values'],record['related'],None)
                validate_related_person_links(schema,records,record['record_code'],record['related'])
            except ApplicationError as exc: raise ApplicationError(f'{record["record_code"]}: {exc}') from exc
        backup=None
        if applied:
            backup=create_backup(automatic=True)
            try:
                atomic_write_workbook(schema,records)
                _safe_publish_dataset_snapshot(schema,records)
                if WORKSPACE_MANAGER: WORKSPACE_MANAGER.register_profiles([e['record_code'] for e in applied if e['action']=='added'],current_schema_id())
            except PermissionError as exc:
                invalidate_dataset_cache();raise ApplicationError('أغلق database.xlsx في Excel ثم حاول مجددًا.') from exc
        _IMPORT_REVIEWS.pop(token,None)
        result={'ok':True,'imported':added,'updated':updated,'unchanged':unchanged,'skipped':plan['skipped']+len(entries)-len(selected),
            'rejected':len(plan['errors'])-plan['skipped'],'errors':plan['errors'],'changes':applied+audit,'backup':backup,'record_count':len(records)}
        result['history']=_record_import_history(plan['payload'],result,'excel')
        return result


def commit_import(payload: Any, *, _preview=False) -> dict[str, Any]:
    require_builder_access()
    if not isinstance(payload, dict):
        raise ApplicationError("صيغة طلب الاستيراد غير صحيحة.")
    if "sheet_mappings" in payload and not _preview:
        raise ApplicationError("راجع التغييرات أولًا ثم اعتمدها من حوار المراجعة.")
    mapping = payload.get("mapping")
    if not isinstance(mapping, dict) and not isinstance(payload.get("sheet_mappings"), list):
        raise ApplicationError("خريطة أعمدة الاستيراد غير صالحة.")
    duplicate_policy = clean_text(payload.get("duplicate_policy")) or "update"
    if duplicate_policy not in {"update", "skip", "reject"}:
        raise ApplicationError("سياسة السجلات المكررة غير صالحة.")
    clear_blank_values = bool(payload.get("clear_blank_values"))
    generate_missing_ids = bool(payload.get("generate_missing_ids", True))
    try:
        expected_revision = int(payload.get("schema_revision"))
    except (TypeError, ValueError) as exc:
        raise ApplicationError("إصدار تصميم الاستيراد غير صالح.") from exc

    with WORKBOOK_LOCK:
        schema = read_schema_file()
        if expected_revision != schema["revision"]:
            raise ApplicationError(
                "تغيّر تصميم التطبيق بعد فحص الملف. افحص ملف الاستيراد مجددًا."
            )
        try:
            imported_rows = parse_import_sheets(payload.get("file_data"), schema, payload["sheet_mappings"]) if "sheet_mappings" in payload else parse_import_rows(
                payload.get("file_data"),
                schema,
                sheet_name=clean_text(payload.get("sheet_name")),
                mapping=mapping,
            )
        except WorkbookExchangeError as exc:
            raise ApplicationError(str(exc)) from exc

        snapshot = _dataset_snapshot_unlocked(schema)
        records = [copy.deepcopy(record) for record in snapshot.records]
        existing_count = len(records)
        codes_in_use = {record["record_code"] for record in records}
        accepted: list[dict[str, Any]] = []
        updated = 0
        errors: list[dict[str, Any]] = []
        skipped = 0
        timestamp = now_iso()

        for imported in imported_rows:
            source_row = int(imported.get("source_row") or 0)
            requested_code = clean_text(imported.get("record_code")).upper()
            try:
                if imported.get("_import_error"):
                    raise ApplicationError(imported["_import_error"])
                if requested_code:
                    record_code = validate_person_code(requested_code)
                else:
                    if not generate_missing_ids:
                        raise ApplicationError(
                            "هذا الصف لا يحتوي على ID. أكّد توليد معرّف جديد أولًا."
                        )
                    record_code = generate_person_code()
                    while (
                        record_code in codes_in_use
                        or (
                            WORKSPACE_MANAGER is not None
                            and WORKSPACE_MANAGER.person_id_in_use(record_code)
                        )
                    ):
                        record_code = generate_person_code()
                existing_record = _find_record(records, record_code)
                if existing_record is not None:
                    if duplicate_policy != "update":
                        if duplicate_policy == "skip": skipped += 1
                        reason = (
                            "تم تخطي ID موجود مسبقًا."
                            if duplicate_policy == "skip"
                            else "ID مستخدم مسبقًا."
                        )
                        errors.append(
                            {"row": source_row, "record_code": record_code, "message": reason}
                        )
                        continue
                    merged_main = copy.deepcopy(existing_record.get("values", {}))
                    for field_id, value in imported.get("main", {}).items():
                        if clear_blank_values or clean_text(value):
                            merged_main[field_id] = value
                    merged_related = {
                        category_id: [
                            {
                                **copy.deepcopy(row),
                                "values": copy.deepcopy(row.get("values", {})),
                                "linked_record_code": row.get("linked_record_code", ""),
                            }
                            for row in existing_record.get("related", {}).get(category_id, [])
                        ]
                        for category_id in {
                            category["id"]
                            for category in schema["categories"]
                            if category["kind"] == "repeatable"
                        }
                    }
                    for category_id, rows in imported.get("related", {}).items():
                        merged_related[category_id] = _merge_import_cards(existing_record.get("related", {}).get(category_id, []), rows, imported.get("related_modes", {}).get(category_id, "replace"), clear_blank_values)
                    main_values, related_values, staged, _kept = _normalize_submission(
                        schema,
                        {"main": merged_main, "related": _prepare_import_cards(schema, merged_related)},
                        "update",
                        _record_file_paths(schema, existing_record),
                    )
                    if staged:
                        for temporary_path, _destination in staged:
                            temporary_path.unlink(missing_ok=True)
                        raise ApplicationError("لا يمكن استيراد ملفات مرفقة من جدول Excel.")
                    validate_unique_fields(
                        schema,
                        records,
                        existing_record["_record_id"],
                        main_values,
                        related_values,
                        None,
                    )
                    replacement = copy.deepcopy(existing_record)
                    replacement["values"] = main_values
                    replacement["updated_at"] = timestamp
                    replacement["archived"] = bool(imported.get("archived", existing_record.get("archived")))
                    replacement["archived_at"] = timestamp if replacement["archived"] else ""
                    replacement["related"] = {}
                    for category in schema["categories"]:
                        if category["kind"] != "repeatable":
                            continue
                        previous_rows = {
                            row.get("minor_id"): row
                            for row in existing_record.get("related", {}).get(category["id"], [])
                        }
                        new_rows = []
                        for position, row in enumerate(related_values.get(category["id"], []), start=1):
                            previous = previous_rows.get(position, {})
                            new_rows.append(
                                {
                                    "_child_id": row.get("_child_id") or previous.get("_child_id") or new_internal_id(),
                                    "parent_child_id": row.get("parent_child_id", ""),
                                    "minor_id": position,
                                    "created_at": previous.get("created_at") or timestamp,
                                    "updated_at": timestamp,
                                    "linked_record_code": row.get("linked_record_code", ""),
                                    "values": row["values"],
                                }
                            )
                        replacement["related"][category["id"]] = new_rows
                    records[records.index(existing_record)] = replacement
                    accepted.append(
                        {
                            "record": replacement,
                            "related_values": related_values,
                            "source_row": source_row,
                            "sources": imported.get("source_sheets", []),
                            "updated": True,
                            "previous": existing_record,
                        }
                    )
                    updated += 1
                    continue
                if WORKSPACE_MANAGER is not None:
                    try:
                        WORKSPACE_MANAGER.assert_profile_creation(
                            record_code,
                            current_schema_id(),
                            link_existing=True,
                        )
                    except WorkspaceError as exc:
                        raise ApplicationError(str(exc)) from exc

                main_values, related_values, staged, _kept = _normalize_submission(
                    schema,
                    {
                        "main": imported.get("main", {}),
                        "related": _prepare_import_cards(schema, copy.deepcopy(imported.get("related", {}))),
                    },
                    "create",
                    set(),
                )
                if staged:
                    for temporary_path, _destination in staged:
                        temporary_path.unlink(missing_ok=True)
                    raise ApplicationError("لا يمكن استيراد ملفات مرفقة من جدول Excel.")
                validate_unique_fields(
                    schema,
                    records,
                    None,
                    main_values,
                    related_values,
                    None,
                )

                record = {
                    "_record_id": new_internal_id(),
                    "record_code": record_code,
                    "created_at": timestamp,
                    "updated_at": timestamp,
                    "archived": bool(imported.get("archived")),
                    "archived_at": timestamp if imported.get("archived") else "",
                    "values": main_values,
                    "related": {},
                }
                for category in schema["categories"]:
                    if category["kind"] != "repeatable":
                        continue
                    new_rows = []
                    for row in related_values.get(category["id"], []):
                        new_rows.append(
                            {
                                "_child_id": row.get("_child_id") or new_internal_id(),
                                "parent_child_id": row.get("parent_child_id", ""),
                                "minor_id": len(new_rows) + 1,
                                "created_at": timestamp,
                                "updated_at": timestamp,
                                "linked_record_code": row.get("linked_record_code", ""),
                                "values": row["values"],
                            }
                        )
                    record["related"][category["id"]] = new_rows
                records.append(record)
                codes_in_use.add(record_code)
                accepted.append(
                    {
                        "record": record,
                        "related_values": related_values,
                        "source_row": source_row,
                        "sources": imported.get("source_sheets", []),
                        "updated": False,
                        "generated": not requested_code,
                    }
                )
            except ApplicationError as exc:
                errors.append(
                    {
                        "row": source_row,
                        "record_code": requested_code,
                        "message": str(exc),
                        "sources": imported.get("source_sheets", []),
                    }
                )

        # Linked-person validation is performed after every candidate has an
        # ID, so records in the same import may reference one another.
        changed = True
        while changed:
            changed = False
            for item in list(accepted):
                record = item["record"]
                try:
                    validate_related_person_links(
                        schema,
                        records,
                        record["record_code"],
                        item["related_values"],
                    )
                except ApplicationError as exc:
                    records.remove(record)
                    if item.get("updated") and item.get("previous"):
                        records.append(item["previous"])
                        updated = max(0, updated - 1)
                    accepted.remove(item)
                    codes_in_use.discard(record["record_code"])
                    errors.append(
                        {
                            "row": item["source_row"],
                            "record_code": record["record_code"],
                            "message": str(exc),
                        }
                    )
                    changed = True

        if _preview:
            return _store_import_review(schema, snapshot.records, accepted, errors, skipped, payload)

        backup = None
        if len(records) > existing_count or updated:
            backup = create_backup(automatic=True)
            try:
                atomic_write_workbook(schema, records)
                _safe_publish_dataset_snapshot(schema, records)
                if WORKSPACE_MANAGER is not None:
                    WORKSPACE_MANAGER.register_profiles(
                        [
                            item["record"]["record_code"]
                            for item in accepted
                            if not item.get("updated")
                        ],
                        current_schema_id(),
                    )
            except PermissionError as exc:
                invalidate_dataset_cache()
                raise ApplicationError(
                    "تعذّر الاستيراد. أغلق database.xlsx في Excel ثم حاول مرة أخرى."
                ) from exc

    result = {
        "ok": True,
        "imported": len(records) - existing_count,
        "updated": updated,
        "skipped": skipped,
        "rejected": len(errors) - skipped,
        "errors": errors,
        "backup": backup,
        "record_count": len(records),
    }
    result["history"] = _record_import_history(payload, result, "excel")
    return result


def multi_schema_search(payload: Any) -> dict[str, Any]:
    """Run independent indexed searches and return one result table per schema."""
    manager = require_workspace()
    if not isinstance(payload, dict) or not isinstance(payload.get("queries"), list):
        raise ApplicationError("طلب البحث متعدد التصاميم غير صالح.")
    queries = payload["queries"]
    if not queries or len(queries) > 50:
        raise ApplicationError("اختر تصميمًا واحدًا على الأقل للبحث.")
    results: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in queries:
        if not isinstance(item, dict):
            raise ApplicationError("إعداد أحد التصاميم في البحث غير صالح.")
        schema_id = clean_text(item.get("schema_id"))
        if schema_id in seen:
            raise ApplicationError("لا يمكن تكرار التصميم نفسه في البحث.")
        seen.add(schema_id)
        try:
            context = manager.context(schema_id)
        except WorkspaceError as exc:
            raise ApplicationError(str(exc)) from exc
        criteria = item.get("criteria", {})
        if not isinstance(criteria, dict):
            raise ApplicationError("مرشحات أحد التصاميم غير صالحة.")
        with use_context(context):
            result = search_records(criteria)
        results.append(
            {
                "schema_id": schema_id,
                "schema_name": context.name,
                **result,
            }
        )
    return {"ok": True, "results": results}


def global_field_search(payload: Any) -> dict[str, Any]:
    """Search global field criteria once across every applicable schema."""

    manager = require_workspace()
    if not isinstance(payload, dict):
        raise ApplicationError("طلب البحث العام غير صالح.")
    raw_criteria = payload.get("criteria", {})
    if not isinstance(raw_criteria, dict):
        raise ApplicationError("مرشحات البحث العام غير صالحة.")
    global_criteria = {
        clean_text(key): value
        for key, value in raw_criteria.items()
        if clean_text(key).startswith("gfld_")
    }
    record_code = clean_text(payload.get("record_code")).upper()
    requested_columns = [
        clean_text(item)
        for item in payload.get("column_refs", [])
        if clean_text(item).startswith("gfld_")
    ] if isinstance(payload.get("column_refs", []), list) else []
    results: list[dict[str, Any]] = []
    for context in manager.contexts(include_archived=False):
        with use_context(context):
            schema = read_schema_file()
            field_by_ref = {
                field["global_ref"]: field
                for category in schema.get("categories", [])
                for field in category.get("fields", [])
                if field.get("global_ref")
            }
            required_refs = set(global_criteria) | set(requested_columns)
            if required_refs and not required_refs.issubset(field_by_ref):
                continue
            criteria = {
                field_by_ref[global_ref]["id"]: value
                for global_ref, value in global_criteria.items()
            }
            criteria["_record_code"] = record_code
            criteria["_allow_empty"] = bool(payload.get("allow_empty"))
            criteria["_include_archived"] = bool(payload.get("include_archived"))
            criteria["_search_field_ids"] = [
                field_by_ref[global_ref]["id"] for global_ref in global_criteria
            ]
            column_ids = [
                field_by_ref[global_ref]["id"]
                for global_ref in requested_columns
                if global_ref in field_by_ref
            ]
            if column_ids:
                criteria["_result_field_ids"] = column_ids
            result = search_records(criteria)
            results.append(
                {
                    "schema_id": context.schema_id,
                    "schema_name": context.name,
                    **result,
                }
            )
    return {"ok": True, "mode": "global", "results": results}


def _query_value_matches(field: dict[str, Any], query: str, value: Any) -> bool:
    if value in (None, "", []):
        return False
    query_text = normalize_search_text(query)
    displayed = normalize_search_text(field_display_value(field, value))
    if query_text and query_text in displayed:
        return True
    if field.get("type") == "number":
        try:
            return float(query.replace(",", "").replace("٬", "")) == float(value)
        except (TypeError, ValueError):
            return False
    return False


def query_across_schemas(payload: Any) -> dict[str, Any]:
    """Free query search across selected schemas and field values.

    This mode is intentionally unrelated to reusable global field definitions.
    It reports the exact category/field/value locations that matched.
    """
    manager = require_workspace()
    if not isinstance(payload, dict):
        raise ApplicationError("طلب البحث العام غير صالح.")
    query = clean_text(payload.get("query"))
    if not query:
        raise ApplicationError("اكتب عبارة البحث.")
    raw_configs = payload.get("schemas")
    if not isinstance(raw_configs, list) or not raw_configs:
        raise ApplicationError("اختر تصميمًا واحدًا على الأقل.")
    results: list[dict[str, Any]] = []
    grand_total = 0
    for raw_config in raw_configs[:50]:
        if not isinstance(raw_config, dict):
            continue
        schema_id = clean_text(raw_config.get("schema_id"))
        try:
            context = manager.context(schema_id)
        except WorkspaceError as exc:
            raise ApplicationError(str(exc)) from exc
        with use_context(context):
            schema = read_schema_file()
            indexes = schema_indexes(schema)
            requested_search = raw_config.get("search_field_ids")
            requested_display = raw_config.get("display_field_ids")
            available = [
                field
                for category in schema.get("categories", [])
                for field in category.get("fields", [])
                if field.get("type") != "file"
            ]
            if isinstance(requested_search, list) and requested_search:
                allowed_search = {
                    clean_text(item) for item in requested_search
                    if clean_text(item) in indexes["fields"]
                }
                search_fields = [field for field in available if field["id"] in allowed_search]
            else:
                search_fields = available
            if isinstance(requested_display, list) and requested_display:
                allowed_display = {
                    clean_text(item) for item in requested_display
                    if clean_text(item) in indexes["fields"]
                }
                display_fields = [field for field in available if field["id"] in allowed_display]
            else:
                display_fields = [field for field in available if field.get("show_in_results")]
                if not display_fields:
                    display_fields = [
                        field for field in available
                        if field.get("type") not in SYSTEM_FIELD_TYPES
                    ][:4]
            with WORKBOOK_LOCK:
                records = _dataset_snapshot_unlocked(schema).records
            cards: list[dict[str, Any]] = []
            for record in records:
                if record.get("archived") and not bool(raw_config.get("include_archived")):
                    continue
                locations: list[dict[str, str]] = []
                for field in search_fields:
                    category = indexes["categories"][indexes["field_categories"][field["id"]]]
                    for value in _field_values_for_record(record, category, field):
                        if _query_value_matches(field, query, value):
                            locations.append(
                                {
                                    "category": category["label"],
                                    "field": field["label"],
                                    "value": clean_text(field_display_value(field, value)),
                                }
                            )
                if not locations:
                    continue
                details: list[dict[str, str]] = []
                for field in display_fields:
                    category = indexes["categories"][indexes["field_categories"][field["id"]]]
                    values = [
                        clean_text(field_display_value(field, value))
                        for value in _field_values_for_record(record, category, field)
                        if value not in (None, "", [])
                    ]
                    if values:
                        details.append(
                            {
                                "field_id": field["id"],
                                "label": field["label"],
                                "value": " | ".join(values),
                            }
                        )
                cards.append(
                    {
                        "record_code": record["record_code"],
                        "archived": bool(record.get("archived")),
                        "details": details,
                        "matched_at": locations,
                    }
                )
                if len(cards) >= MAX_SEARCH_PAGE_SIZE:
                    break
            grand_total += len(cards)
            results.append(
                {
                    "schema_id": schema_id,
                    "schema_name": context.name,
                    "matches": cards,
                    "total": len(cards),
                }
            )
    return {"ok": True, "mode": "query", "query": query, "results": results, "total": grand_total}


def _merge_global_configuration(
    local: dict[str, Any], definition: dict[str, Any], kind: str, global_ref: str
) -> dict[str, Any]:
    preserved = {"id": local["id"], "global_ref": global_ref}
    if kind == "category":
        preserved.update(
            {
                "fields": copy.deepcopy(local.get("fields", [])),
                "parent_category_id": local.get("parent_category_id"),
                "parent_field_id": local.get("parent_field_id"),
                "anchor_field_id": local.get("anchor_field_id"),
            }
        )
    merged = copy.deepcopy(definition)
    # A category definition can carry an entire reusable tree. The tree is
    # library metadata, not a property of one local category row.
    merged.pop("category_tree", None)
    merged.update(preserved)
    return merged


def _merge_global_category_tree(
    schema: dict[str, Any], definition: dict[str, Any], global_ref: str
) -> bool:
    """Update linked instances of a reusable nested category tree in place.

    Local category/field IDs are deliberately preserved so record values and
    relationships remain attached while configuration changes propagate.
    """

    raw_tree = definition.get("category_tree")
    if not isinstance(raw_tree, list) or not raw_tree:
        return False
    nodes = {
        clean_text(node.get("key")): node
        for node in raw_tree
        if isinstance(node, dict) and clean_text(node.get("key"))
    }
    changed = False
    category_ids_by_key: dict[str, str] = {}
    field_ids_by_key: dict[str, str] = {}
    for category in schema.get("categories", []):
        if category.get("global_tree_ref") != global_ref:
            continue
        node = nodes.get(clean_text(category.get("global_tree_key")))
        if not node:
            continue
        category_ids_by_key[clean_text(category.get("global_tree_key"))] = category["id"]
        category_definition = copy.deepcopy(node.get("definition") or {})
        category_definition.pop("fields", None)
        category_definition.pop("category_tree", None)
        preserved_category = {
            "id": category["id"],
            "global_ref": category.get("global_ref"),
            "global_tree_ref": global_ref,
            "global_tree_key": category.get("global_tree_key"),
            "parent_category_id": category.get("parent_category_id"),
            "parent_field_id": category.get("parent_field_id"),
            "anchor_field_id": category.get("anchor_field_id"),
        }
        field_nodes = {
            clean_text(item.get("key")): item
            for item in node.get("fields", [])
            if isinstance(item, dict) and clean_text(item.get("key"))
        }
        updated_fields = []
        for field in category.get("fields", []):
            field_node = field_nodes.get(clean_text(field.get("global_tree_key")))
            if not field_node or field.get("global_tree_ref") != global_ref:
                updated_fields.append(field)
                continue
            field_ids_by_key[clean_text(field.get("global_tree_key"))] = field["id"]
            field_definition = copy.deepcopy(field_node.get("definition") or {})
            field_definition.update(
                {
                    "id": field["id"],
                    "global_ref": field.get("global_ref"),
                    "global_tree_ref": global_ref,
                    "global_tree_key": field.get("global_tree_key"),
                }
            )
            updated_fields.append(field_definition)
        category_definition.update(preserved_category)
        category_definition["fields"] = updated_fields
        category.clear()
        category.update(category_definition)
        changed = True
    if changed:
        retained_conditions = [
            condition
            for condition in schema.get("conditions", [])
            if condition.get("global_tree_ref") != global_ref
        ]
        group_ids: dict[str, str] = {}
        imported_conditions: list[dict[str, Any]] = []
        for index, condition in enumerate(definition.get("conditions", [])):
            if not isinstance(condition, dict):
                continue
            source_id = field_ids_by_key.get(clean_text(condition.get("source_field_key")))
            target_type = clean_text(condition.get("target_type"))
            target_key = clean_text(condition.get("target_key"))
            target_id = (
                category_ids_by_key.get(target_key)
                if target_type == "category"
                else field_ids_by_key.get(target_key)
            )
            if not source_id or not target_id or target_type not in {"category", "field"}:
                continue
            group_key = clean_text(condition.get("group_key")) or f"group-{index + 1}"
            group_ids.setdefault(group_key, new_definition_id("grp"))
            imported_conditions.append(
                {
                    "id": new_definition_id("cond"),
                    "group_id": group_ids[group_key],
                    "negate": bool(condition.get("negate", False)),
                    "target_type": target_type,
                    "target_id": target_id,
                    "source_field_id": source_id,
                    "operator": clean_text(condition.get("operator")) or "equals",
                    "value": copy.deepcopy(condition.get("value", "")),
                    "global_tree_ref": global_ref,
                    "global_condition_key": clean_text(condition.get("key")) or f"condition-{index + 1}",
                }
            )
        schema["conditions"] = retained_conditions + imported_conditions
    return changed


def _backup_global_metadata(definitions: dict[str, Any], schemas: list[tuple[Any, dict[str, Any]]]) -> dict[str, Any]:
    """Recover a definition deletion without copying unchanged record files."""
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S-%f")
    filename = f"GenericSchemaCraft-auto-backup-{stamp}.zip"
    destination = BACKUP_DIR / filename
    with tempfile.NamedTemporaryFile(dir=BACKUP_DIR, suffix=".zip", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(temporary_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("global-definitions.json", json.dumps(definitions, ensure_ascii=False))
            for context, schema in schemas:
                archive.writestr(context.schema_path.relative_to(DATA_DIR).as_posix(), json.dumps(schema, ensure_ascii=False))
            archive.writestr("RECOVERY.txt", "Metadata-only recovery for a general-definition deletion. Records and attachments were not changed. With the application closed, restore the included JSON files to their matching workspace paths to restore the definitions and links. This is not a full workspace backup.")
        os.replace(temporary_path, destination)
    finally:
        temporary_path.unlink(missing_ok=True)
    return {"filename": filename, "download_url": f"/api/backups/{quote(filename)}", "automatic": True, "metadata_only": True}


def _save_detached_schema_metadata(previous: dict[str, Any], schema: dict[str, Any]) -> None:
    """Only global-link metadata changed; keep values and search indexes intact."""
    global _DATASET_SNAPSHOT
    schema["revision"] = previous["revision"] + 1
    atomic_write_json(_schema_path(), schema)
    key = str(_workbook_path().resolve())
    cached = _DATASET_SNAPSHOTS.get(key)
    if cached and cached.schema_signature == _schema_cache_signature(previous):
        updated = copy.copy(cached)
        updated.schema_signature = _schema_cache_signature(schema)
        _DATASET_SNAPSHOTS[key] = updated
        _DATASET_SNAPSHOT = updated
    context = active_context()
    if context is not None:
        schedule_workbook_schema_sync(context, schema["revision"])


def manage_global_definition(payload: Any) -> dict[str, Any]:
    """Save a global configuration and migrate every linked schema safely."""

    require_builder_access()
    if GLOBAL_DEFINITIONS is None:
        raise ApplicationError("مكتبة التعريفات العامة غير مفعّلة.")
    if not isinstance(payload, dict):
        raise ApplicationError("طلب التعريف العام غير صالح.")
    action = clean_text(payload.get("action")) or "save"
    kind = clean_text(payload.get("kind"))
    if kind not in {"category", "field"}:
        raise ApplicationError("نوع التعريف العام غير صالح.")
    global_ref = clean_text(payload.get("global_ref"))
    try:
        expected = payload.get("expected_revision")
        expected_revision = int(expected) if expected is not None else None
    except (TypeError, ValueError) as exc:
        raise ApplicationError("إصدار مكتبة التعريفات العامة غير صالح.") from exc
    manager = require_workspace()
    if action == "reorder_all":
        try:
            return GLOBAL_DEFINITIONS.reorder_definitions(kind, payload.get("order"), expected_revision=expected_revision)
        except AdvancedFeatureError as exc:
            raise ApplicationError(str(exc)) from exc
    if action == "reorder":
        direction = clean_text(payload.get("direction"))
        try:
            return GLOBAL_DEFINITIONS.reorder_definition(
                kind,
                global_ref,
                direction,
                expected_revision=expected_revision,
            )
        except AdvancedFeatureError as exc:
            raise ApplicationError(str(exc)) from exc
    if action == "delete":
        # Hold the same lock as record/schema writes across recovery and detach.
        with WORKBOOK_LOCK:
            definitions = GLOBAL_DEFINITIONS.read()
            if expected_revision is not None and definitions["revision"] != expected_revision:
                raise ApplicationError("تغيّرت مكتبة التعريفات العامة. أعد تحميلها ثم حاول مجددًا.")
            collection = definitions["categories" if kind == "category" else "fields"]
            if global_ref not in collection:
                raise ApplicationError("التعريف العام المطلوب غير موجود.")
            originals = []
            for context in manager.contexts(include_archived=True):
                with use_context(context):
                    candidate = read_schema_file()
                linked = any(
                    (kind == "category" and (category.get("global_ref") == global_ref or category.get("global_tree_ref") == global_ref))
                    or (kind == "field" and any(field.get("global_ref") == global_ref for field in category.get("fields", [])))
                    for category in candidate.get("categories", [])
                ) or (kind == "category" and any(condition.get("global_tree_ref") == global_ref for condition in candidate.get("conditions", [])))
                if linked:
                    originals.append((context, candidate))
            backup = _backup_global_metadata(definitions, originals)
            try:
                result = GLOBAL_DEFINITIONS.remove_definition(kind, global_ref, expected_revision=expected_revision)
            except AdvancedFeatureError as exc:
                raise ApplicationError(str(exc)) from exc
            detached: list[str] = []
            for context, previous in originals:
                with use_context(context):
                    schema = copy.deepcopy(previous)
                    changed = False
                    for category in schema.get("categories", []):
                        if kind == "category" and category.get("global_ref") == global_ref:
                            category["global_ref"] = None
                            changed = True
                        if kind == "category" and category.get("global_tree_ref") == global_ref:
                            category["global_tree_ref"] = None
                            category["global_tree_key"] = None
                            for field in category.get("fields", []):
                                if field.get("global_tree_ref") == global_ref:
                                    field["global_tree_ref"] = None
                                    field["global_tree_key"] = None
                            changed = True
                        if kind == "field":
                            for field in category.get("fields", []):
                                if field.get("global_ref") == global_ref:
                                    field["global_ref"] = None
                                    changed = True
                    if kind == "category":
                        for condition in schema.get("conditions", []):
                            if condition.get("global_tree_ref") != global_ref:
                                continue
                            condition.pop("global_tree_ref", None)
                            condition.pop("global_condition_key", None)
                            changed = True
                    if changed:
                        _save_detached_schema_metadata(previous, schema)
                        detached.append(context.schema_id)
        return {
            **result,
            "backup": backup,
            "detached_schema_ids": detached,
            "global_definitions": GLOBAL_DEFINITIONS.response(),
        }
    definition = payload.get("definition")
    # Each destructive linked-schema migration is backed up inside save_schema.
    # Avoid a second whole-workspace archive here: it was the dominant delay
    # when editing list definitions shared by several schemas.
    backup = None
    try:
        saved = GLOBAL_DEFINITIONS.save_definition(
            kind,
            global_ref,
            definition,
            expected_revision=expected_revision,
        )
    except AdvancedFeatureError as exc:
        raise ApplicationError(str(exc)) from exc
    updated: list[str] = []
    normalized_definition = saved["definition"]
    for context in manager.contexts(include_archived=True):
        with use_context(context):
            schema = read_schema_file()
            changed = False
            for category in schema.get("categories", []):
                if kind == "category" and category.get("global_ref") == global_ref:
                    merged_category = _merge_global_configuration(
                        category,
                        normalized_definition,
                        kind,
                        global_ref,
                    )
                    category.clear()
                    category.update(merged_category)
                    changed = True
                if kind == "field":
                    for index, field in enumerate(category.get("fields", [])):
                        if field.get("global_ref") != global_ref:
                            continue
                        category["fields"][index] = _merge_global_configuration(
                            field, normalized_definition, kind, global_ref
                        )
                        changed = True
            if kind == "category":
                changed = _merge_global_category_tree(
                    schema, normalized_definition, global_ref
                ) or changed
            if changed:
                # The current revision is the optimistic concurrency token.
                schema["revision"] = read_schema_file()["revision"]
                save_schema(schema)
                updated.append(context.schema_id)
    return {
        "ok": True,
        "definition": saved,
        "updated_schema_ids": updated,
        "backup": backup,
        "global_definitions": GLOBAL_DEFINITIONS.response(),
    }


def inspect_identity_profiles(payload: Any) -> dict[str, Any]:
    manager = require_workspace()
    if not isinstance(payload, dict):
        raise ApplicationError("طلب فحص الشخص غير صالح.")
    code = validate_person_code(clean_text(payload.get("record_code")).upper())
    target_id = clean_text(payload.get("target_schema_id")) or current_schema_id()
    profiles = []
    for membership in manager.identity_search(code):
        schema_id = membership["schema_id"]
        if schema_id == target_id:
            continue
        context = manager.context(schema_id)
        with use_context(context):
            schema = read_schema_file()
            record = load_record(code)
            shape = {"values": record["main"], "related": record["related"]}
            profiles.append(
                {
                    "schema_id": schema_id,
                    "schema_name": context.name,
                    "fields": _transferable_fields(schema, shape, include_files=bool(payload.get("include_files"))),
                }
            )
    values_by_global: dict[str, list[dict[str, Any]]] = {}
    for profile in profiles:
        for field in profile["fields"]:
            global_ref = clean_text(field.get("global_ref"))
            values = field.get("values") or []
            if not global_ref or not any(value not in (None, "", []) for value in values):
                continue
            values_by_global.setdefault(global_ref, []).append(
                {
                    "schema_id": profile["schema_id"],
                    "schema_name": profile["schema_name"],
                    "field_id": field["field_id"],
                    "label": field["label"],
                    "values": copy.deepcopy(values),
                    "display_values": copy.deepcopy(field.get("display_values") or values),
                }
            )
    conflicts = []
    for global_ref, sources in values_by_global.items():
        distinct = {
            json.dumps(source["values"], ensure_ascii=False, sort_keys=True, default=str)
            for source in sources
        }
        if len(distinct) > 1:
            conflicts.append(
                {
                    "global_ref": global_ref,
                    "label": sources[0]["label"],
                    "sources": sources,
                }
            )
    return {
        "ok": True,
        "record_code": code,
        "target_schema_id": target_id,
        "profiles": profiles,
        "global_conflicts": conflicts,
    }


def _transferable_fields(
    schema: dict[str, Any], record: dict[str, Any] | None = None, *, include_files: bool = False
) -> list[dict[str, Any]]:
    fields: list[dict[str, Any]] = []
    for category in schema["categories"]:
        for field in data_fields(category):
            if (field["type"] == "file" and not include_files) or field["type"] in SYSTEM_FIELD_TYPES:
                continue
            values: list[Any] = []
            if record is not None:
                values = _field_values_for_record(record, category, field)
            fields.append(
                {
                    "category_id": category["id"],
                    "category": category["label"],
                    "category_kind": category["kind"],
                    "field_id": field["id"],
                    "label": field["label"],
                    "type": field["type"],
                    "global_ref": field.get("global_ref"),
                    "values": copy.deepcopy(values),
                    "display_values": [
                        _transfer_value(field, value) for value in values
                    ],
                }
            )
    return fields


def inspect_profile_transfer(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise ApplicationError("طلب ربط الملف غير صالح.")
    manager = require_workspace()
    source_id = clean_text(payload.get("source_schema_id"))
    target_id = clean_text(payload.get("target_schema_id")) or current_schema_id()
    code = validate_person_code(clean_text(payload.get("record_code")).upper())
    if source_id == target_id:
        raise ApplicationError("اختر تصميم مصدر مختلفًا عن التصميم الحالي.")
    try:
        source_context = manager.context(source_id)
        target_context = manager.context(target_id)
    except WorkspaceError as exc:
        raise ApplicationError(str(exc)) from exc
    with use_context(source_context):
        source_schema = read_schema_file()
        source_record = load_record(code)
        # Convert the public shape back to the record shape expected by the
        # field extractor without exposing internal IDs to the browser.
        transfer_record = {
            "values": source_record["main"],
            "related": source_record["related"],
        }
        sources = _transferable_fields(source_schema, transfer_record)
    with use_context(target_context):
        target_schema = read_schema_file()
        try:
            load_record(code)
        except ApplicationError:
            pass
        else:
            raise ApplicationError("يوجد ملف لهذا الشخص داخل التصميم الهدف بالفعل.")
        targets = _transferable_fields(target_schema)

    target_by_id = {item["field_id"]: item for item in targets}
    target_by_global = {
        item["global_ref"]: item
        for item in targets
        if item.get("global_ref")
    }
    suggestions: dict[str, str] = {}
    for source in sources:
        global_match = target_by_global.get(source.get("global_ref"))
        if global_match:
            suggestions[source["field_id"]] = global_match["field_id"]
            continue
        direct = target_by_id.get(source["field_id"])
        if direct:
            suggestions[source["field_id"]] = direct["field_id"]
            continue
        label = normalize_search_text(source["label"])
        match = next(
            (
                target
                for target in targets
                if normalize_search_text(target["label"]) == label
                and target["type"] == source["type"]
            ),
            None,
        )
        if match:
            suggestions[source["field_id"]] = match["field_id"]
    for mapping in manager.mapping_profile(source_id, target_id):
        source_field_id = clean_text(mapping.get("source_field_id"))
        target_field_id = clean_text(mapping.get("target_field_id"))
        if (
            any(item["field_id"] == source_field_id for item in sources)
            and target_field_id in target_by_id
        ):
            suggestions[source_field_id] = target_field_id
    return {
        "ok": True,
        "record_code": code,
        "source_schema_id": source_id,
        "source_schema_name": source_context.name,
        "target_schema_id": target_id,
        "target_schema_name": target_context.name,
        "source_fields": sources,
        "target_fields": targets,
        "suggestions": suggestions,
    }


def _transfer_value(field: dict[str, Any], value: Any) -> Any:
    if field["type"] == "checkbox_group":
        labels = {
            option["id"]: option["label"] for option in field.get("options", [])
        }
        return [
            labels.get(clean_text(item), clean_text(item))
            for item in parse_checkbox_group_value(value)
        ]
    if field["type"] in {"select", "yes_no"}:
        return field_display_value(field, value)
    if field["type"] == "checkbox":
        return excel_boolean(value)
    return copy.deepcopy(json_value(value))


def create_linked_profile(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict) or not isinstance(payload.get("mappings"), list):
        raise ApplicationError("خريطة ربط الملف غير صالحة.")
    manager = require_workspace()
    source_id = clean_text(payload.get("source_schema_id"))
    target_id = clean_text(payload.get("target_schema_id")) or current_schema_id()
    code = validate_person_code(clean_text(payload.get("record_code")).upper())
    if source_id == target_id:
        raise ApplicationError("لا يمكن إنشاء ملف مرتبط داخل التصميم نفسه.")
    try:
        source_context = manager.context(source_id)
        target_context = manager.context(target_id)
    except WorkspaceError as exc:
        raise ApplicationError(str(exc)) from exc

    with use_context(source_context):
        source_schema = read_schema_file()
        source_record = load_record(code)
        source_shape = {
            "values": source_record["main"],
            "related": source_record["related"],
        }
        source_index = schema_indexes(source_schema)

    with use_context(target_context):
        target_schema = read_schema_file()
        target_index = schema_indexes(target_schema)
        try:
            load_record(code)
        except ApplicationError:
            pass
        else:
            raise ApplicationError("يوجد ملف لهذا الشخص داخل التصميم الهدف بالفعل.")

        main: dict[str, Any] = {}
        related: dict[str, list[dict[str, Any]]] = {}
        used_targets: set[str] = set()
        applied: list[dict[str, str]] = []
        for mapping in payload["mappings"]:
            if not isinstance(mapping, dict):
                continue
            source_field_id = clean_text(mapping.get("source_field_id"))
            target_field_id = clean_text(mapping.get("target_field_id"))
            if not source_field_id or not target_field_id:
                continue
            if target_field_id in used_targets:
                raise ApplicationError("لا يمكن توجيه حقلين مصدر إلى الحقل الهدف نفسه.")
            source_field = source_index["fields"].get(source_field_id)
            target_field = target_index["fields"].get(target_field_id)
            if (
                not source_field
                or not target_field
                or source_field["type"] == "file"
                or target_field["type"] == "file"
                or source_field["type"] in SYSTEM_FIELD_TYPES
                or target_field["type"] in SYSTEM_FIELD_TYPES
            ):
                raise ApplicationError("أحد الحقول المختارة للربط غير متاح.")
            source_category = source_index["categories"][
                source_index["field_categories"][source_field_id]
            ]
            target_category = target_index["categories"][
                target_index["field_categories"][target_field_id]
            ]
            source_values = _field_values_for_record(
                source_shape, source_category, source_field
            )
            converted = [
                _transfer_value(source_field, value)
                for value in source_values
                if value is not None and value != "" and value != []
            ]
            if not converted:
                continue
            if target_category["kind"] == "main":
                main[target_field_id] = converted[0]
            else:
                rows = related.setdefault(target_category["id"], [])
                while len(rows) < len(converted):
                    rows.append({"values": {}})
                for index, value in enumerate(converted):
                    rows[index]["values"][target_field_id] = value
            used_targets.add(target_field_id)
            applied.append(
                {
                    "source_field_id": source_field_id,
                    "target_field_id": target_field_id,
                }
            )

        if not applied and not bool(payload.get("allow_empty")):
            raise ApplicationError("اختر مطابقة حقل واحدة على الأقل.")
        create_backup(automatic=True)
        result = save_record(
            {
                "mode": "create",
                "record_code": code,
                "link_existing": True,
                "main": main,
                "related": related,
            }
        )
    result["source_schema_id"] = source_id
    result["target_schema_id"] = target_id
    result["mappings_applied"] = applied
    try:
        manager.save_mapping_profile(source_id, target_id, applied)
    except OSError:
        LOGGER.exception("Linked profile saved but mapping profile could not be persisted")
    return result



class DataEntryHTTPServer(ThreadingHTTPServer):
    daemon_threads = True
    allow_reuse_address = True

    def __init__(
        self,
        server_address,
        handler_class,
        *,
        require_session: bool = False,
        startup_ready: bool = True,
    ):
        super().__init__(server_address, handler_class)

        self._browser_session_lock = threading.RLock()
        self._last_browser_heartbeat = 0.0
        self._request_condition = threading.Condition(threading.RLock())
        self._active_request_count = 0
        self._shutdown_pending = False
        self.require_session = bool(require_session)
        self.sessions = BrowserSessionManager()
        self._startup_state_lock = threading.RLock()
        self._startup_ready = bool(startup_ready)
        self._startup_error = ""

    def startup_status(self) -> dict[str, Any]:
        """Return the non-sensitive initialization state used by the splash."""

        with self._startup_state_lock:
            return {
                "ready": self._startup_ready,
                "error": self._startup_error,
            }

    def mark_startup_ready(self) -> None:
        """Publish successful workspace initialization to the splash window."""

        with self._startup_state_lock:
            self._startup_ready = True
            self._startup_error = ""

    def mark_startup_error(self, message: Any) -> None:
        """Publish a display-safe initialization error to the splash window."""

        with self._startup_state_lock:
            self._startup_ready = False
            self._startup_error = clean_text(message)[:500]

    def process_request_thread(self, request, client_address) -> None:
        with self._request_condition:
            self._active_request_count += 1
        try:
            super().process_request_thread(request, client_address)
        finally:
            with self._request_condition:
                self._active_request_count = max(0, self._active_request_count - 1)
                self._request_condition.notify_all()

    def schedule_shutdown_when_idle(self, idle_seconds: float = 2.0) -> bool:
        """Stop the server after every request has been idle for a grace period."""
        with self._request_condition:
            if self._shutdown_pending:
                return False
            self._shutdown_pending = True

        def shutdown_when_idle() -> None:
            idle_started: float | None = None
            while True:
                with self._request_condition:
                    active = self._active_request_count
                    if active == 0:
                        if idle_started is None:
                            idle_started = time.monotonic()
                        remaining = idle_seconds - (time.monotonic() - idle_started)
                        if remaining <= 0:
                            break
                        self._request_condition.wait(timeout=remaining)
                    else:
                        idle_started = None
                        self._request_condition.wait(timeout=0.1)
            LOGGER.info(
                "Application shutdown started after %.1f idle seconds.",
                idle_seconds,
            )
            self.shutdown()

        threading.Thread(
            target=shutdown_when_idle,
            daemon=True,
            name="graceful-application-shutdown",
        ).start()
        return True

    def note_browser_heartbeat(self) -> None:
        with self._browser_session_lock:
            self._last_browser_heartbeat = time.monotonic()

    def browser_is_active(self) -> bool:
        with self._browser_session_lock:
            last_seen = self._last_browser_heartbeat

        if last_seen <= 0:
            return False

        return (
            time.monotonic() - last_seen
            <= BROWSER_ACTIVE_SECONDS
        )

    def touch_client(self) -> None:
        # Kept for compatibility with existing request code.
        # Normal API requests do not control application shutdown.
        return

    def note_disconnect(self) -> None:
        # Closing a browser page marks it inactive, but does not stop
        # the Python server.
        with self._browser_session_lock:
            self._last_browser_heartbeat = 0.0

    def handle_error(self, request, client_address) -> None:
        LOGGER.exception(
            "Unhandled request error from %s:%s",
            client_address[0],
            client_address[1],
        )

class DataEntryRequestHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(APP_DIR), **kwargs)

    def log_message(self, _format: str, *args) -> None:
        return

    def guess_type(self, path: str) -> str:
        content_type = super().guess_type(path)
        suffix = Path(urlparse(path).path).suffix.casefold()

        if suffix == ".html":
            return "text/html; charset=utf-8"
        if suffix == ".css":
            return "text/css; charset=utf-8"
        if suffix == ".js":
            return "text/javascript; charset=utf-8"

        return content_type

    def send_head(self):
        """Serve editable application wording without rebuilding frontend files."""

        requested = Path(self.translate_path(self.path))
        if requested.is_dir() and urlparse(self.path).path.endswith("/"):
            requested = requested / "index.html"
        try:
            resolved = requested.resolve()
            inside_app = resolved == APP_DIR.resolve() or APP_DIR.resolve() in resolved.parents
        except OSError:
            inside_app = False
        if (
            inside_app
            and resolved.is_file()
            and resolved.name != UI_TEXT_PATH.name
            and resolved.suffix.casefold() in {".html", ".js", ".css"}
        ):
            try:
                content = localized_ui_asset(resolved)
            except (OSError, UnicodeError):
                return super().send_head()
            self.send_response(200)
            self.send_header("Content-Type", self.guess_type(str(resolved)))
            self.send_header("Content-Length", str(len(content)))
            self.send_header("Last-Modified", self.date_time_string(resolved.stat().st_mtime))
            self.end_headers()
            return io.BytesIO(content)
        return super().send_head()

    def request_origin_is_allowed(self) -> bool:
        origin = self.headers.get("Origin")
        if not origin:
            return True

        expected_port = self.server.server_address[1]
        allowed_origins = {
            f"http://127.0.0.1:{expected_port}",
            f"http://localhost:{expected_port}",
        }
        return origin.rstrip("/") in allowed_origins

    def request_host_is_allowed(self) -> bool:
        """Reject DNS-rebinding requests before routing any local resource."""

        expected_port = self.server.server_address[1]
        host = clean_text(self.headers.get("Host")).casefold()
        return host in {
            f"127.0.0.1:{expected_port}",
            f"localhost:{expected_port}",
        }

    def reject_invalid_host(self) -> bool:
        if self.request_host_is_allowed():
            return False
        LOGGER.warning("Rejected invalid Host header.")
        self.send_json(403, {"error": "عنوان الطلب غير مسموح."})
        return True

    def reject_foreign_origin(self) -> bool:
        if self.request_origin_is_allowed():
            return False

        LOGGER.warning(
            "Rejected foreign request origin: %s",
            self.headers.get("Origin"),
        )
        self.send_json(403, {"error": "مصدر الطلب غير مسموح."})
        return True

    def startup_token_is_valid(self) -> bool:
        return isinstance(self.server, DataEntryHTTPServer) and self.server.sessions.startup_token_matches(
            self.headers.get("X-SchemaCraft-Startup")
        )

    def browser_session_user(self) -> str:
        if not isinstance(self.server, DataEntryHTTPServer):
            return ""
        if not self.server.require_session:
            return current_audit_user()
        return self.server.sessions.authenticated_user(self.headers.get("Cookie"))

    def reject_unready_application(self) -> bool:
        if not isinstance(self.server, DataEntryHTTPServer):
            return False
        status = self.server.startup_status()
        if status["ready"]:
            return False
        self.send_json(
            503,
            {
                "error": status["error"] or "لا يزال التطبيق قيد التجهيز.",
                "ready": False,
            },
        )
        return True

    def reject_unauthenticated_session(self) -> bool:
        if not isinstance(self.server, DataEntryHTTPServer) or not self.server.require_session:
            return False
        if self.browser_session_user():
            return False
        self.send_json(401, {"error": "انتهت جلسة التطبيق. أعد تشغيل SchemaCraft."})
        return True

    def end_headers(self) -> None:
        path = unquote(urlparse(self.path).path)

        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("X-Frame-Options", "DENY")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header("Cross-Origin-Resource-Policy", "same-origin")
        self.send_header("Permissions-Policy", "camera=(), microphone=(), geolocation=()")
        self.send_header(
            "Content-Security-Policy",
            "default-src 'self'; connect-src 'self'; script-src 'self'; "
            "style-src 'self'; img-src 'self' data:; object-src 'none'; "
            "base-uri 'none'; form-action 'self'; frame-ancestors 'none'",
        )

        if not path.startswith("/api/"):
            self.send_header(
                "Cache-Control",
                "no-store, no-cache, must-revalidate, max-age=0",
            )
            self.send_header("Pragma", "no-cache")
            self.send_header("Expires", "0")

        super().end_headers()

    def send_json(
        self,
        status: int,
        body: dict[str, Any],
        *,
        headers: dict[str, str] | None = None,
    ) -> None:
        payload = json.dumps(body, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.send_header("Cache-Control", "no-store")
        for name, value in (headers or {}).items():
            self.send_header(name, value)
        self.end_headers()
        self.wfile.write(payload)

    def send_bytes(
        self,
        status: int,
        content: bytes,
        content_type: str,
        filename: str,
    ) -> None:
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(content)))
        self.send_header(
            "Content-Disposition",
            f"attachment; filename*=UTF-8''{quote(filename)}",
        )
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(content)

    def touch_client(self) -> None:
        if isinstance(self.server, DataEntryHTTPServer):
            self.server.touch_client()

    def schema_context(self, payload: Any = None):
        if WORKSPACE_MANAGER is None:
            return nullcontext()
        parsed = urlparse(self.path)
        query = parse_qs(parsed.query)
        requested = clean_text(self.headers.get("X-Schema-ID"))
        if not requested:
            requested = clean_text((query.get("schema_id") or [""])[0])
        if not requested and isinstance(payload, dict):
            requested = clean_text(payload.get("schema_id"))
        try:
            context = WORKSPACE_MANAGER.context(requested or None)
        except WorkspaceError as exc:
            raise ApplicationError(str(exc)) from exc
        return use_context(context)

    def _read_json(self) -> Any:
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError as exc:
            raise ApplicationError("حجم الطلب غير صالح.") from exc
        if length > MAX_REQUEST_BYTES:
            raise ApplicationError("حجم الطلب يتجاوز الحد المسموح.")
        raw = self.rfile.read(length)
        if not raw:
            return {}
        try:
            return json.loads(raw.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ApplicationError("صيغة JSON غير صحيحة.") from exc

    def do_GET(self) -> None:
        if self.reject_invalid_host():
            return

        path = unquote(urlparse(self.path).path)
        if path == "/api/health":
            browser_active = (
                self.server.browser_is_active()
                if isinstance(self.server, DataEntryHTTPServer)
                else False
            )

            self.send_json(
                200,
                {
                    "ok": True,
                    "browser_active": browser_active,
                },
            )
            return
        if path == "/api/startup/status":
            if not self.startup_token_is_valid():
                self.send_json(403, {"error": "رمز تشغيل التطبيق غير صالح."})
                return
            if isinstance(self.server, DataEntryHTTPServer):
                self.server.note_browser_heartbeat()
                self.send_json(200, self.server.startup_status())
            else:
                self.send_json(503, {"ready": False, "error": "الخادم غير جاهز."})
            return
        if path == "/api/session/bootstrap":
            if not self.startup_token_is_valid():
                self.send_json(403, {"error": "رمز تشغيل التطبيق غير صالح."})
                return
            if self.reject_unready_application():
                return
            audit_users = AUDIT_USERS.read() if AUDIT_USERS is not None else {
                "version": 1,
                "current_user": "",
                "users": [],
            }
            self.send_json(200, {"ok": True, "users": audit_users.get("users", [])})
            return
        if path == "/api/session/status":
            if self.reject_unready_application():
                return
            if self.reject_unauthenticated_session():
                return
            self.send_json(
                200,
                {
                    "ok": True,
                    "authenticated": True,
                    "user_name": current_audit_user(),
                },
            )
            return
        if path.startswith("/api/") and (
            self.reject_unready_application() or self.reject_unauthenticated_session()
        ):
            return
        if path == "/api/workspace":
            self.touch_client()
            try:
                self.send_json(200, workspace_response())
            except ApplicationError as exc:
                self.send_json(400, {"error": str(exc)})
            return
        if path == "/api/settings/background-image":
            self.touch_client()
            query = parse_qs(urlparse(self.path).query)
            settings = read_workspace_settings()
            requested_id = clean_text((query.get("id") or [settings.get("background_image_id")])[0])
            background_path = workspace_background_image_path() if requested_id == "legacy" else workspace_background_image_file(requested_id)
            if not background_path.is_file():
                self.send_error(404)
                return
            try:
                content = background_path.read_bytes()
                content_type = detect_background_image_mime(content[:16])
                if not content_type:
                    self.send_error(404)
                    return
                self.send_response(200)
                self.send_header("Content-Type", content_type)
                self.send_header("Content-Length", str(len(content)))
                self.send_header("Content-Disposition", "inline")
                self.send_header("Cache-Control", "no-store")
                self.send_header("X-Content-Type-Options", "nosniff")
                self.end_headers()
                self.wfile.write(content)
            except OSError:
                self.send_error(404)
            return
        if path == "/api/global-definitions":
            self.touch_client()
            try:
                if GLOBAL_DEFINITIONS is None:
                    raise ApplicationError("مكتبة التعريفات العامة غير مفعّلة.")
                self.send_json(200, GLOBAL_DEFINITIONS.response())
            except ApplicationError as exc:
                self.send_json(400, {"error": str(exc)})
            return
        if path == "/api/export/history":
            self.touch_client()
            try:
                query = parse_qs(urlparse(self.path).query)
                self.send_json(
                    200,
                    export_history_response({"limit": (query.get("limit") or [20])[0]}),
                )
            except ApplicationError as exc:
                self.send_json(400, {"error": str(exc)})
            return
        if path == "/api/import/history":
            self.touch_client()
            try:
                query = parse_qs(urlparse(self.path).query)
                self.send_json(
                    200,
                    import_history_response({"limit": (query.get("limit") or [20])[0]}),
                )
            except ApplicationError as exc:
                self.send_json(400, {"error": str(exc)})
            return
        if path == "/api/search/history":
            self.touch_client()
            try:
                query = parse_qs(urlparse(self.path).query)
                self.send_json(
                    200,
                    search_history_response(
                        {
                            "limit": (query.get("limit") or [50])[0],
                            "query": (query.get("query") or [""])[0],
                        }
                    ),
                )
            except ApplicationError as exc:
                self.send_json(400, {"error": str(exc)})
            return
        if path == "/api/search/field-values":
            self.touch_client()
            try:
                query = parse_qs(urlparse(self.path).query)
                with self.schema_context():
                    result = field_value_suggestions(
                        (query.get("field_id") or [""])[0],
                        (query.get("query") or [""])[0],
                        (query.get("limit") or [30])[0],
                        query.get("selected_value") or [],
                    )
                self.send_json(200, result)
            except ApplicationError as exc:
                self.send_json(400, {"error": str(exc)})
            return
        if path == "/api/home/schema":
            self.touch_client()
            try:
                query = parse_qs(urlparse(self.path).query)
                with self.schema_context():
                    response = home_schema_response((query.get("signature") or [""])[0])
                self.send_json(200, response)
            except ApplicationError as exc:
                self.send_json(400, {"error": str(exc)})
            return
        if path == "/api/schema":
            self.touch_client()
            try:
                with self.schema_context():
                    response = schema_response(read_schema_with_excel_labels())
                    response["schema_id"] = current_schema_id()
                    context = active_context()
                    response["schema_name"] = context.name if context else ""
                self.send_json(200, response)
            except ApplicationError as exc:
                self.send_json(400, {"error": str(exc)})
            return
        if path.startswith("/api/records/"):
            self.touch_client()
            try:
                code = path.removeprefix("/api/records/")
                with self.schema_context():
                    self.send_json(200, load_record(code))
            except ApplicationError as exc:
                self.send_json(400, {"error": str(exc)})
            return
        if path.startswith("/api/attachments/"):
            self.touch_client()
            try:
                filename = path.removeprefix("/api/attachments/")
                with self.schema_context():
                    attachment = attachment_absolute_path(f"attachments/{filename}")
                if attachment is None or not attachment.is_file():
                    self.send_error(404)
                    return
                content = attachment.read_bytes()
                inline = (
                    attachment.suffix.casefold()
                    in INLINE_ATTACHMENT_EXTENSIONS
                )
                content_type = (
                    mimetypes.guess_type(attachment.name)[0]
                    if inline
                    else "application/octet-stream"
                ) or "application/octet-stream"
                self.send_response(200)
                self.send_header("Content-Type", content_type)
                self.send_header("Content-Length", str(len(content)))
                self.send_header(
                    "Content-Disposition",
                    f"{'inline' if inline else 'attachment'}; "
                    f"filename*=UTF-8''{quote(attachment.name)}",
                )
                self.send_header("X-Content-Type-Options", "nosniff")
                if not inline:
                    self.send_header("Content-Security-Policy", "sandbox")
                    self.send_header("X-Download-Options", "noopen")
                self.end_headers()
                self.wfile.write(content)
            except (ApplicationError, OSError):
                self.send_error(404)
            return
        if path.startswith("/api/backups/"):
            self.touch_client()
            if not builder_is_unlocked():
                self.send_error(404)
                return
            backup = backup_file_path(path.removeprefix("/api/backups/"))
            if backup is None or not backup.is_file():
                self.send_error(404)
                return
            try:
                content = backup.read_bytes()
                self.send_response(200)
                self.send_header("Content-Type", "application/zip")
                self.send_header("Content-Length", str(len(content)))
                self.send_header(
                    "Content-Disposition",
                    f"attachment; filename*=UTF-8''{quote(backup.name)}",
                )
                self.send_header("X-Content-Type-Options", "nosniff")
                self.end_headers()
                self.wfile.write(content)
            except OSError:
                self.send_error(404)
            return
        super().do_GET()

    def do_POST(self) -> None:
        if self.reject_invalid_host() or self.reject_foreign_origin():
            return

        path = unquote(urlparse(self.path).path)
        if path == "/api/session/login":
            if self.reject_unready_application():
                return
            startup_token = self.headers.get("X-SchemaCraft-Startup")
            if not self.startup_token_is_valid():
                self.send_json(403, {"error": "رمز تشغيل التطبيق غير صالح."})
                return
            payload = self._read_json()
            audit_users = select_audit_user(payload)
            if not isinstance(self.server, DataEntryHTTPServer):
                self.send_json(503, {"error": "الخادم غير جاهز."})
                return
            try:
                session_token = self.server.sessions.create_session(
                    audit_users["audit_users"]["current_user"],
                    startup_token,
                )
            except (PermissionError, ValueError):
                self.send_json(403, {"error": "تعذّر إنشاء جلسة التطبيق."})
                return
            self.send_json(
                200,
                {"ok": True, "audit_users": audit_users["audit_users"]},
                headers={"Set-Cookie": self.server.sessions.set_cookie_header(session_token)},
            )
            return
        if self.reject_unready_application() or self.reject_unauthenticated_session():
            return
        if path == "/api/heartbeat":
            if isinstance(self.server, DataEntryHTTPServer):
                self.server.note_browser_heartbeat()

            self.send_json(
                200,
                {
                    "ok": True,
                    "browser_active": True,
                },
            )
            return
        if path == "/api/disconnect":
            if isinstance(self.server, DataEntryHTTPServer):
                self.server.note_disconnect()

            try:
                self.send_json(200, {"ok": True})
            except (
                BrokenPipeError,
                ConnectionResetError,
            ):
                pass

            return
        if path == "/api/shutdown":
            payload = self._read_json()
            requested_delay = payload.get("idle_seconds", 2)
            try:
                idle_seconds = float(requested_delay)
            except (TypeError, ValueError):
                idle_seconds = 2.0
            idle_seconds = min(5.0, max(2.0, idle_seconds))
            if isinstance(self.server, DataEntryHTTPServer):
                self.server.sessions.invalidate_all()
            scheduled = (
                self.server.schedule_shutdown_when_idle(idle_seconds)
                if isinstance(self.server, DataEntryHTTPServer)
                else False
            )
            self.send_json(
                200,
                {
                    "ok": True,
                    "scheduled": scheduled,
                    "idle_seconds": idle_seconds,
                },
                headers={
                    "Set-Cookie": self.server.sessions.clear_cookie_header()
                } if isinstance(self.server, DataEntryHTTPServer) else None,
            )
            return
        if path == "/api/backup":
            self.touch_client()
            try:
                require_builder_access()
                with self.schema_context():
                    result = create_backup()
                self.send_json(200, result)
            except ApplicationError as exc:
                LOGGER.warning(
                    "Backup request rejected: %s",
                    exc,
                )

                self.send_json(
                    400,
                    {"error": str(exc)},
                )

            except Exception:
                LOGGER.exception(
                    "Unexpected backup creation failure"
                )

                self.send_json(
                    500,
                    {
                        "error":
                            "تعذّر إنشاء النسخة الاحتياطية بسبب خطأ داخلي."
                    },
                )
            return
        try:
            self.touch_client()
            payload = self._read_json()
            if path == "/api/client-log":
                level = clean_text(
                    payload.get("level", "error")
                ).casefold()

                category = clean_text(
                    payload.get("category", "browser")
                )[:80]

                message = clean_text(
                    payload.get("message", "Unknown browser error")
                )[:4000]

                location = clean_text(
                    payload.get("location", "")
                )[:4000]

                occurred_at = clean_text(
                    payload.get("occurred_at", "")
                )[:80]

                log_text = (
                    f"Browser [{category}] {message}"
                )

                if occurred_at:
                    log_text += f" | occurred_at={occurred_at}"

                if location:
                    log_text += f"\n{location}"

                if level == "warning":
                    LOGGER.warning(log_text)
                else:
                    LOGGER.error(log_text)

                self.send_json(
                    200,
                    {"ok": True},
                )
                return
            if path == "/api/workspace/select":
                self.send_json(200, select_workspace_schema(payload))
                return
            if path == "/api/workspace/schemas":
                with self.schema_context(payload):
                    result = manage_workspace_schema(payload)
                self.send_json(200, result)
                return
            if path == "/api/users/select":
                self.send_json(200, select_audit_user(payload))
                return
            if path == "/api/builder/unlock":
                access = unlock_builder(
                    payload.get("password"),
                    initialize=bool(payload.get("initialize")),
                )
                self.send_json(200, {"ok": True, "builder_access": access})
                return
            if path == "/api/builder/lock":
                self.send_json(200, {"ok": True, "builder_access": lock_builder()})
                return
            if path == "/api/builder/password":
                self.send_json(
                    200,
                    {
                        "ok": True,
                        "builder_access": change_builder_password(
                            payload.get("current_password"),
                            payload.get("new_password"),
                        ),
                    },
                )
                return
            if path == "/api/settings":
                with self.schema_context(payload):
                    result = save_app_settings(payload)
                self.send_json(200, result)
                return
            if path == "/api/settings/shortcuts":
                self.send_json(200, save_workspace_settings(payload))
                return
            if path == "/api/home/custom-stats":
                self.send_json(200, save_home_custom_stats(payload))
                return
            if path == "/api/settings/default-app/destination":
                self.send_json(200, choose_default_app_parent())
                return
            if path == "/api/settings/default-app/create":
                self.send_json(200, create_default_app_copy(payload))
                return
            if path == "/api/history/notes":
                self.send_json(200, update_operation_history_notes(payload))
                return
            if path == "/api/export/history/open":
                self.send_json(200, open_export_history_file(payload.get("id")))
                return
            if path == "/api/import/history/log":
                with self.schema_context(payload):
                    self.send_json(200, log_import_history(payload))
                return
            if path == "/api/records":
                with self.schema_context(payload):
                    result = save_record(payload)
                self.send_json(200, result)
                return
            if path == "/api/schema/options":
                with self.schema_context(payload):
                    result = add_runtime_list_option(payload)
                self.send_json(200, result)
                return
            if path == "/api/archive":
                with self.schema_context(payload):
                    result = archive_record(
                        payload.get("record_code"),
                        bool(payload.get("archived", True)),
                        payload.get("expected_updated_at"),
                    )
                self.send_json(200, result)
                return
            if path == "/api/search":
                with self.schema_context(payload):
                    result = search_records(payload)
                self.send_json(200, result)
                return
            if path == "/api/search/query":
                self.send_json(200, query_across_schemas(payload))
                return
            if path == "/api/search/history":
                self.send_json(200, save_search_history(payload))
                return
            if path == "/api/search/multi":
                self.send_json(200, multi_schema_search(payload))
                return
            if path == "/api/search/global":
                self.send_json(200, global_field_search(payload))
                return
            if path == "/api/global-definitions":
                self.send_json(200, manage_global_definition(payload))
                return
            if path == "/api/identities/inspect":
                with self.schema_context(payload):
                    self.send_json(200, inspect_identity_profiles(payload))
                return
            if path == "/api/profiles/inspect":
                with self.schema_context(payload):
                    result = inspect_profile_transfer(payload)
                self.send_json(200, result)
                return
            if path == "/api/profiles/link":
                with self.schema_context(payload):
                    result = create_linked_profile(payload)
                self.send_json(200, result)
                return
            if path == "/api/export":
                with self.schema_context(payload):
                    filename, content, _count = create_filtered_export(payload)
                self.send_bytes(
                    200,
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    filename,
                )
                return
            if path == "/api/reports":
                self.send_json(200, report_api(payload))
                return
            if path == "/api/export/save":
                with self.schema_context(payload):
                    self.send_json(200, save_export(payload))
                return
            if path == "/api/export/destination":
                with self.schema_context(payload):
                    self.send_json(200, choose_export_destination_response(payload))
                return
            if path == "/api/import/inspect":
                with self.schema_context(payload):
                    result = inspect_import(payload)
                self.send_json(200, result)
                return
            if path == "/api/import/preview":
                with self.schema_context(payload):
                    self.send_json(200, commit_import(payload, _preview=True))
                return
            if path == "/api/import/apply-review":
                with self.schema_context(payload):
                    self.send_json(200, apply_import_review(payload))
                return
            if path == "/api/import/commit":
                with self.schema_context(payload):
                    result = commit_import(payload)
                self.send_json(200, result)
                return
            if path == "/api/import/history/open":
                self.send_json(200, open_import_history_file(payload.get("id")))
                return
            if path == "/api/import/portable/inspect":
                with self.schema_context(payload):
                    self.send_json(200, inspect_portable_import(payload))
                return
            if path == "/api/import/portable/commit":
                with self.schema_context(payload):
                    self.send_json(200, commit_portable_import(payload))
                return
            self.send_json(404, {"error": "المسار المطلوب غير موجود."})
        except ApplicationError as exc:
            LOGGER.warning(
                "Request rejected: %s %s | %s",
                self.command,
                path,
                exc,
            )

            self.send_json(
                400,
                {"error": str(exc)},
            )

        except Exception:
            LOGGER.exception(
                "Unexpected API error: %s %s",
                self.command,
                path,
            )

            self.send_json(
                500,
                {"error": "حدث خطأ داخلي غير متوقع في التطبيق."},
            )

    def do_PUT(self) -> None:
        if self.reject_invalid_host() or self.reject_foreign_origin():
            return

        path = unquote(urlparse(self.path).path)
        if self.reject_unready_application() or self.reject_unauthenticated_session():
            return
        try:
            self.touch_client()
            if path != "/api/schema":
                self.send_json(404, {"error": "المسار المطلوب غير موجود."})
                return
            require_builder_access()
            payload = self._read_json()
            with self.schema_context(payload):
                result = save_schema(payload)
            self.send_json(200, result)
        except ApplicationError as exc:
            LOGGER.warning(
                "Request rejected: %s %s | %s",
                self.command,
                path,
                exc,
            )
            self.send_json(400, {"error": str(exc)})
        except Exception:
            LOGGER.exception(
                "Unexpected API error: %s %s",
                self.command,
                path,
            )
            self.send_json(
                500,
                {"error": "حدث خطأ داخلي غير متوقع في التطبيق."},
            )

    def do_DELETE(self) -> None:
        if self.reject_invalid_host() or self.reject_foreign_origin():
            return

        path = unquote(urlparse(self.path).path)
        if self.reject_unready_application() or self.reject_unauthenticated_session():
            return
        self.touch_client()
        try:
            if path.startswith("/api/history/"):
                kind = path.removeprefix("/api/history/")
                self.send_json(200, clear_history(kind))
                return
            if path.startswith("/api/search/history/"):
                entry_id = path.removeprefix("/api/search/history/")
                self.send_json(200, delete_search_history(entry_id))
                return
            if path.startswith("/api/import/history/"):
                entry_id = path.removeprefix("/api/import/history/")
                self.send_json(200, delete_operation_history("import", entry_id))
                return
            if path.startswith("/api/export/history/"):
                entry_id = path.removeprefix("/api/export/history/")
                self.send_json(200, delete_operation_history("export", entry_id))
                return
            if not path.startswith("/api/records/"):
                self.send_json(404, {"error": "المسار المطلوب غير موجود."})
                return
            code = path.removeprefix("/api/records/")
            with self.schema_context():
                result = delete_record(code)
            self.send_json(200, result)
        except ApplicationError as exc:
            LOGGER.warning(
                "Request rejected: %s %s | %s",
                self.command,
                path,
                exc,
            )

            self.send_json(
                400,
                {"error": str(exc)},
            )

        except Exception:
            LOGGER.exception(
                "Unexpected API error: %s %s",
                self.command,
                path,
            )

            self.send_json(
                500,
                {"error": "حدث خطأ داخلي غير متوقع في التطبيق."},
            )


def acquire_single_instance() -> bool:
    """Return False when this application folder is already running."""
    global _WINDOWS_MUTEX_HANDLE

    if os.name != "nt":
        return True

    folder_key = hashlib.sha256(
        str(BASE_DIR).casefold().encode("utf-8")
    ).hexdigest()[:20]

    mutex_name = f"Local\\GenericSchemaCraft-{folder_key}"

    kernel32 = ctypes.windll.kernel32

    kernel32.CreateMutexW.argtypes = [
        ctypes.c_void_p,
        ctypes.c_bool,
        ctypes.c_wchar_p,
    ]
    kernel32.CreateMutexW.restype = ctypes.c_void_p

    kernel32.CloseHandle.argtypes = [ctypes.c_void_p]
    kernel32.CloseHandle.restype = ctypes.c_bool

    handle = kernel32.CreateMutexW(None, False, mutex_name)

    if not handle:
        raise OSError("تعذّر إنشاء قفل تشغيل البرنامج.")

    if kernel32.GetLastError() == WINDOWS_ALREADY_EXISTS:
        kernel32.CloseHandle(ctypes.c_void_p(handle))
        return False

    _WINDOWS_MUTEX_HANDLE = int(handle)
    return True

def _application_browser_candidates() -> list[Path]:
    """Return installed Chromium-family browsers that support --app mode."""
    candidates: list[Path] = []
    commands = (
        "msedge",
        "microsoft-edge",
        "google-chrome",
        "chrome",
        "chromium",
        "chromium-browser",
    )
    for command in commands:
        executable = shutil.which(command)
        if executable:
            candidates.append(Path(executable))

    if os.name == "nt":
        roots = [
            os.environ.get("PROGRAMFILES(X86)"),
            os.environ.get("PROGRAMFILES"),
            os.environ.get("LOCALAPPDATA"),
        ]
        relative_paths = (
            Path("Microsoft/Edge/Application/msedge.exe"),
            Path("Google/Chrome/Application/chrome.exe"),
        )
        for root in roots:
            if not root:
                continue
            for relative_path in relative_paths:
                candidate = Path(root) / relative_path
                if candidate.is_file():
                    candidates.append(candidate)

    unique: list[Path] = []
    seen: set[str] = set()
    for candidate in candidates:
        key = str(candidate).casefold()
        if key not in seen:
            seen.add(key)
            unique.append(candidate)
    return unique


def open_application_window(
    url: str,
    *,
    window_size: tuple[int, int] | None = None,
    full_screen: bool = True,
) -> subprocess.Popen[Any]:
    """Open one isolated, borderless Chromium kiosk for the entire lifecycle."""
    global _APPLICATION_BROWSER_PROCESS
    size_arguments = (
        [f"--window-size={window_size[0]},{window_size[1]}"]
        if window_size and not full_screen
        else []
    )
    display_arguments = (
        ["--kiosk", "--start-fullscreen"]
        if full_screen
        else []
    )
    for executable in _application_browser_candidates():
        profile_directory = Path(tempfile.mkdtemp(prefix="SchemaCraft-window-"))
        try:
            process = subprocess.Popen(
                [
                    str(executable),
                    f"--app={url}",
                    "--new-window",
                    f"--user-data-dir={profile_directory}",
                    "--no-first-run",
                    "--no-default-browser-check",
                    "--no-service-autorun",
                    "--disable-background-mode",
                    "--disable-session-crashed-bubble",
                    "--disable-pinch",
                    "--overscroll-history-navigation=0",
                    *display_arguments,
                    *size_arguments,
                ],
                close_fds=True,
            )
            _APPLICATION_BROWSER_PROCESS = process
            threading.Thread(
                target=_remove_application_profile_when_closed,
                args=(process, profile_directory),
                daemon=False,
                name="application-window-profile-cleanup",
            ).start()
            LOGGER.info("Opened application-mode browser window with %s", executable)
            return process
        except OSError:
            shutil.rmtree(profile_directory, ignore_errors=True)
            LOGGER.exception("Could not open application-mode browser %s", executable)

    raise ApplicationError(
        "تعذّر فتح نافذة SchemaCraft بلا شريط عنوان. "
        "ثبّت Microsoft Edge أو Google Chrome أو Chromium ثم أعد المحاولة."
    )


def _remove_application_profile_when_closed(
    process: subprocess.Popen[Any],
    profile_directory: Path,
) -> None:
    """Delete the isolated kiosk profile after its browser process exits."""
    try:
        process.wait()
    except (OSError, subprocess.SubprocessError):
        LOGGER.exception("Could not wait for the application browser process.")
    finally:
        shutil.rmtree(profile_directory, ignore_errors=True)


def close_application_window_process() -> None:
    """Ensure the dedicated kiosk process and all of its children have exited."""
    global _APPLICATION_BROWSER_PROCESS
    process = _APPLICATION_BROWSER_PROCESS
    _APPLICATION_BROWSER_PROCESS = None
    if process is None:
        return
    try:
        process.wait(timeout=1.5)
        return
    except subprocess.TimeoutExpired:
        pass
    except OSError:
        return

    if os.name == "nt":
        subprocess.run(
            ["taskkill", "/PID", str(process.pid), "/T", "/F"],
            check=False,
            capture_output=True,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    else:
        process.terminate()
        try:
            process.wait(timeout=3)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=3)


def open_running_instance() -> bool:
    """
    Connect to the already-running local server.

    Open a browser page only if the server has no active page.
    """
    port = application_port()
    url = f"http://{HOST}:{port}/"
    health: dict[str, Any] | None = None

    def read_health() -> dict[str, Any] | None:
        connection = HTTPConnection(HOST, port, timeout=0.5)
        try:
            connection.request("GET", "/api/health")
            response = connection.getresponse()
            if response.status != 200:
                response.read()
                return None
            return json.loads(response.read().decode("utf-8") or "{}")
        except (OSError, json.JSONDecodeError):
            return None
        finally:
            connection.close()

    # The first instance may still be starting.
    for _attempt in range(20):
        health = read_health()
        if health is not None:
            break
        time.sleep(0.15)

    if health is None:
        return False

    if health.get("browser_active", False):
        LOGGER.info(
            "Application and browser session are already open."
        )
        return True

    # Give a recently opened or restored page time to send a heartbeat.
    time.sleep(2.5)

    refreshed_health = read_health()
    if refreshed_health and refreshed_health.get(
        "browser_active",
        False,
    ):
        LOGGER.info(
            "Browser session became active; "
            "no duplicate page was opened."
        )
        return True

    LOGGER.info(
        "Server is running without a browser session; "
        "opening %s",
        url,
    )

    open_application_window(url, full_screen=True)
    return True
def release_single_instance() -> None:
    global _WINDOWS_MUTEX_HANDLE
    if os.name == "nt" and _WINDOWS_MUTEX_HANDLE is not None:
        ctypes.windll.kernel32.CloseHandle(
            ctypes.c_void_p(_WINDOWS_MUTEX_HANDLE)
        )
    _WINDOWS_MUTEX_HANDLE = None


def report_startup_failure(exc: BaseException) -> None:
    details = (
        f"[{now_iso()}]\n"
        f"{''.join(traceback.format_exception(type(exc), exc, exc.__traceback__))}\n"
    )
    log_written = False
    try:
        STARTUP_ERROR_LOG.write_text(details, encoding="utf-8")
        hide_packaged_support_paths()
        log_written = True
    except OSError:
        pass
    message = f"تعذّر تشغيل البرنامج:\n{exc}"
    if log_written:
        message += f"\n\nحُفظت التفاصيل في:\n{STARTUP_ERROR_LOG}"
    if os.name != "nt":
        print(message, file=sys.stderr)
        return

    try:
        import tkinter as tk

        dialog = tk.Tk()
        dialog.overrideredirect(True)
        dialog.attributes("-topmost", True)
        dialog.configure(background="#ffffff")
        dialog.resizable(False, False)

        body = tk.Frame(dialog, background="#ffffff", padx=28, pady=24)
        body.pack(fill="both", expand=True)
        tk.Label(
            body,
            text="تعذّر تشغيل SchemaCraft",
            background="#ffffff",
            foreground="#b42318",
            font=("Segoe UI Semibold", 15),
            justify="right",
            anchor="e",
        ).pack(fill="x")
        tk.Label(
            body,
            text=message,
            background="#ffffff",
            foreground="#26334b",
            font=("Segoe UI", 10),
            justify="right",
            anchor="e",
            wraplength=520,
        ).pack(fill="x", pady=(12, 20))
        tk.Button(
            body,
            text="إغلاق",
            command=dialog.destroy,
            padx=18,
            pady=7,
            relief="flat",
            background="#2f78c4",
            foreground="#ffffff",
            activebackground="#245f9d",
            activeforeground="#ffffff",
            font=("Segoe UI Semibold", 10),
        ).pack(anchor="w")

        dialog.update_idletasks()
        width = dialog.winfo_reqwidth()
        height = dialog.winfo_reqheight()
        x = max(0, (dialog.winfo_screenwidth() - width) // 2)
        y = max(0, (dialog.winfo_screenheight() - height) // 2)
        dialog.geometry(f"{width}x{height}+{x}+{y}")
        dialog.grab_set()
        dialog.focus_force()
        dialog.mainloop()
    except Exception:
        print(message, file=sys.stderr)

def run() -> None:
    if not APP_DIR.is_dir():
        raise ApplicationError(
            "مجلد واجهة التطبيق app غير موجود."
        )

    port = application_port()

    server = DataEntryHTTPServer(
        (HOST, port),
        DataEntryRequestHandler,
        require_session=True,
        startup_ready=False,
    )

    url = f"http://{HOST}:{port}/"
    server_thread = threading.Thread(
        target=server.serve_forever,
        daemon=True,
        name="application-http-server",
    )
    server_thread.start()

    splash_url = (
        f"{url}startup.html"
        f"#token={server.sessions.startup_token}"
    )
    open_application_window(
        splash_url,
        full_screen=True,
    )

    try:
        initialize_workspace()
    except Exception as exc:
        server.mark_startup_error(exc)
        # Give the already-visible splash a brief opportunity to present the
        # safe error message before the process reports the fatal startup error.
        time.sleep(2.0)
        server.shutdown()
        server_thread.join(timeout=5)
        server.server_close()
        close_application_window_process()
        raise

    server.mark_startup_ready()

    LOGGER.info(
        "Application server started at %s",
        url,
    )

    print(
        f"Data Entry Builder is running at {url}"
    )

    try:
        server_thread.join()

    except KeyboardInterrupt:
        LOGGER.info(
            "Application interrupted from the console."
        )
        server.shutdown()
        server_thread.join(timeout=5)

    finally:
        if server_thread.is_alive():
            server.shutdown()
            server_thread.join(timeout=5)
        server.server_close()
        close_application_window_process()

        LOGGER.info(
            "Application server stopped."
        )
if __name__ == "__main__":
    try:
        configure_logging()
        install_exception_logging()

        LOGGER.info(
            "Application process starting. "
            "Frozen=%s Base=%s",
            getattr(sys, "frozen", False),
            BASE_DIR,
        )

        owns_instance = acquire_single_instance()

        if not owns_instance:
            if open_running_instance():
                raise SystemExit(0)

            raise ApplicationError(
                "البرنامج مفتوح بالفعل، "
                "لكن تعذّر الاتصال بنافذته الحالية."
            )

        run()

    except SystemExit:
        raise

    except Exception as error:
        LOGGER.exception(
            "Application startup or runtime failure"
        )

        report_startup_failure(error)
        raise SystemExit(1) from error

    finally:
        release_single_instance()

        if LOGGER.handlers:
            LOGGER.info(
                "Application process finished."
            )
