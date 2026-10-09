"""Fail-closed, byte-backed migration of an existing Excel workspace.

Preflight never invokes workspace initialization, schema rewrite, registry repair,
or workbook writers. Originals remain in place. PostgreSQL is enabled only after
an atomic logical import, exact readback, and a second full source hash check.
"""
from __future__ import annotations

import copy
import hashlib
import json
import math
import os
import re
import shutil
import sys
import tempfile
import uuid
import zipfile
from contextlib import nullcontext
from datetime import date, datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import load_workbook

from schemacraft_workspace import REGISTRY_HEADERS, REGISTRY_SHEET, WorkspaceManager


class MigrationBlocked(ValueError):
    def __init__(self, code: str, location: str):
        self.code, self.location = code, location
        super().__init__(f"{code}: {location}")


def _block(code: str, location: Any) -> None:
    raise MigrationBlocked(code, str(location))


def _exception_diagnostic(exc: BaseException) -> dict:
    """Describe code provenance without formatting any exception or its data.

    Exception messages, arguments, driver diagnostics, source lines and frame
    locals can all contain workbook values or credentials. Only code metadata
    and explicitly approved machine codes are copied into migration reports.
    """
    chain, seen = [], set()
    storage = sys.modules.get("schemacraft_storage")
    storage_error = getattr(storage, "StorageError", None)
    current, relation = exc, "exception"
    while current is not None and id(current) not in seen and len(chain) < 16:
        seen.add(id(current))
        name = type(current).__name__
        item = {"exception_type": name if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]{0,127}", name) else "Exception",
                "relation": relation, "traceback": []}
        sqlstate = _diagnostic_attribute(current, "sqlstate")
        if type(sqlstate) is str and re.fullmatch(r"[0-9A-Z]{5}", sqlstate):
            item["sqlstate"] = sqlstate
        for key in ("errno", "winerror"):
            value = _diagnostic_attribute(current, key)
            if type(value) is int:
                item[key] = value
        # A validation code is trusted only when storage explicitly publishes
        # it as safe; arbitrary code/path/detail attributes are never reported.
        if isinstance(storage_error, type) and isinstance(current, storage_error):
            safe_codes = getattr(storage, "SAFE_VALIDATION_CODES", frozenset())
            validation_code = _diagnostic_attribute(current, "validation_code")
            if (type(validation_code) is str and isinstance(safe_codes, (set, frozenset))
                    and validation_code in safe_codes and re.fullmatch(r"[a-z][a-z0-9_]{0,95}", validation_code)):
                item["validation_code"] = validation_code
        tb = current.__traceback__
        while tb is not None and len(item["traceback"]) < 32:
            code = tb.tb_frame.f_code
            # Basenames describe the executing code without disclosing the
            # developer's home directory or a selected company directory.
            filename = code.co_filename.replace("\\", "/").rsplit("/", 1)[-1]
            if not re.fullmatch(r"[A-Za-z0-9_.-]{1,160}\.py[cw]?", filename):
                filename = "<dynamic>"
            function = code.co_name
            if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]{0,127}|<module>|<lambda>|<genexpr>|<listcomp>|<dictcomp>|<setcomp>", function):
                function = "<dynamic>"
            item["traceback"].append({"filename": filename, "function": function, "line": tb.tb_lineno})
            tb = tb.tb_next
        chain.append(item)
        if current.__cause__ is not None:
            current, relation = current.__cause__, "cause"
        elif not current.__suppress_context__:
            current, relation = current.__context__, "context"
        else:
            current = None
    diagnostic = {"exception_type": chain[0]["exception_type"],
                  "root_cause_type": chain[-1]["exception_type"], "cause_chain": chain}
    for key in ("sqlstate", "validation_code", "errno", "winerror"):
        for item in reversed(chain):
            if key in item:
                diagnostic[key] = item[key]
                break
    return diagnostic


def _diagnostic_attribute(exc: BaseException, key: str) -> Any:
    try:
        return getattr(exc, key, None)
    except Exception:
        return None


def _failure_issue(exc: BaseException, code: str, location: str, *, stage: str, operation: str) -> dict:
    return {"code": exc.code if isinstance(exc, MigrationBlocked) else code,
            "location": exc.location if isinstance(exc, MigrationBlocked) else location,
            "stage": stage, "operation": operation, "diagnostic": _exception_diagnostic(exc)}


def _record_failure(report: dict, exc: BaseException, code: str, *, stage: str, operation: str,
                    location: str | None = None) -> None:
    report.update(status="blocked", activated=False)
    report.setdefault("stage", stage)
    report.setdefault("operation", operation)
    report.setdefault("issues", []).append(_failure_issue(exc, code, location or type(exc).__name__,
                                                         stage=stage, operation=operation))


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def _equal(a: Any, b: Any) -> bool:
    # JSON spelling distinguishes false/0, integer/float, and negative zero.
    return _json(a) == _json(b)


def _digest(value: Any) -> str:
    return hashlib.sha256(_json(value).encode("utf-8")).hexdigest()


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _inventory(root: Path) -> dict[str, dict[str, Any]]:
    result = {}
    for folder, dirs, files in os.walk(root, followlinks=False):
        current = Path(folder)
        if current == root:
            dirs[:] = [name for name in dirs if name != ".postgresql"]
        for name in dirs:
            if (current / name).is_symlink():
                _block("symlink_directory", (current / name).relative_to(root))
        for name in files:
            path = current / name
            relative = path.relative_to(root).as_posix()
            if path.is_symlink() or not path.is_file():
                _block("unsafe_source_file", relative)
            result[relative] = {"size": path.stat().st_size, "sha256": _file_hash(path)}
    return dict(sorted(result.items()))


def readonly_workspace(data_dir: Path, *, require_workbooks: bool = True) -> WorkspaceManager:
    """Load only catalog metadata; do not initialize or reconcile registry."""
    root = Path(data_dir).resolve(strict=True)
    manager = WorkspaceManager(root)
    manager._catalog = _read_json(manager.catalog_path, "workspace.json")
    catalog = manager._catalog
    if catalog.get("workspace_version") != 1 or not isinstance(catalog.get("schemas"), list) or not catalog["schemas"]:
        _block("invalid_workspace_catalog", "workspace.json")
    ids, names = set(), set()
    for entry in catalog["schemas"]:
        if not isinstance(entry, dict):
            _block("invalid_schema_catalog_entry", "workspace.json")
        sid, name = entry.get("id"), entry.get("name")
        if not isinstance(sid, str) or not re.fullmatch(r"[a-f0-9]{32}", sid) or sid in ids:
            _block("invalid_or_duplicate_schema_id", "workspace.json")
        if not isinstance(name, str) or not name.strip() or name.casefold() in names:
            _block("invalid_or_duplicate_schema_name", "workspace.json")
        ids.add(sid); names.add(name.casefold())
        for key in ("folder", "workbook"):
            value = entry.get(key)
            if not isinstance(value, str) or value in {"", ".", ".."} or "/" in value or "\\" in value:
                _block("unsafe_schema_path", f"workspace.json:{sid}:{key}")
        context = manager._context_from_entry(entry)
        if not context.schema_path.is_file() or (require_workbooks and not context.workbook_path.is_file()):
            _block("missing_schema_or_workbook", f"schemas/{entry['folder']}")
    if catalog.get("active_schema_id") not in ids:
        _block("invalid_active_schema", "workspace.json")
    parents = {entry["id"]: entry.get("profile_source_schema_id", "") for entry in catalog["schemas"]}
    for sid in parents:
        seen, parent = {sid}, parents[sid]
        while parent:
            if parent not in parents or parent in seen:
                _block("invalid_profile_source_chain", f"workspace.json:{sid}")
            seen.add(parent); parent = parents[parent]
    return manager


def _read_json(path: Path, location: str) -> dict:
    def pairs(items):
        result = {}
        for key, value in items:
            if key in result:
                _block("duplicate_json_key", location)
            result[key] = value
        return result
    try:
        value = json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=pairs,
                           parse_constant=lambda _: _block("nonfinite_json_number", location))
    except MigrationBlocked:
        raise
    except (OSError, ValueError):
        _block("invalid_or_missing_json", location)
    if not isinstance(value, dict):
        _block("json_object_required", location)
    return value


def _cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return ""
    if type(value) not in (str, bool, int, float) or (type(value) is float and not math.isfinite(value)):
        _block("unsupported_excel_value", type(value).__name__)
    return value


def _text(value: Any, location: str) -> str:
    if value is None:
        return ""
    if not isinstance(value, str) or value != value.strip():
        _block("metadata_normalization_loss", location)
    return value


def _boolean(value: Any, location: str) -> bool:
    if value is None or value == "":
        return False
    if type(value) is bool:
        return value
    if type(value) in (int, float) and value in (0, 1):
        return bool(value)
    if isinstance(value, str) and value.casefold() in {"نعم", "لا", "true", "false", "1", "0", "yes", "no"}:
        return value.casefold() in {"نعم", "true", "1", "yes"}
    _block("unrecognized_boolean", location)


def _field_value(app, value: Any, field: dict, location: str) -> Any:
    raw = _cell_value(value)
    kind = field["type"]
    if kind == "checkbox":
        return _boolean(value, location)
    if kind == "checkbox_group":
        parsed = app.read_excel_field_value(value, field)
        if value not in (None, ""):
            # Only the documented separator encoding is accepted. Never drop
            # empty choices, surrounding whitespace, or duplicate choices.
            if not isinstance(raw, str) or " | ".join(parsed) != raw or len(set(parsed)) != len(parsed):
                _block("checkbox_group_normalization_loss", location)
        return parsed
    if app.number_is_text(field):
        if raw != "" and (not isinstance(raw, str) or raw != raw.strip()):
            _block("number_text_normalization_loss", location)
        return raw
    return raw


def _headers(sheet, allowed: set[str], required: set[str], first_data: int, location: str) -> dict[str, int]:
    positions = {}
    for column in range(1, sheet.max_column + 1):
        header = sheet.cell(1, column).value
        populated = any(sheet.cell(row, column).value not in (None, "") for row in range(first_data, sheet.max_row + 1))
        if header in allowed:
            if header in positions:
                _block("duplicate_excel_header", f"{location}:{sheet.cell(1,column).coordinate}")
            positions[header] = column
        elif populated:
            _block("unknown_populated_column", f"{location}:{sheet.cell(1,column).coordinate}")
    if required - positions.keys():
        _block("missing_required_columns", location)
    return positions


def _raw_at(sheet, row: int, columns: dict, header: str) -> Any:
    return sheet.cell(row, columns[header]).value if header in columns else None


def _visible_labels(app, sheet, fields: list, columns: dict, location: str) -> None:
    for field in fields:
        if field["id"] not in columns:
            continue
        value = sheet.cell(app.VISIBLE_HEADER_ROW, columns[field["id"]]).value
        if value not in (None, "") and str(value).strip() != field["label"]:
            _block("visible_field_label_requires_review", f"{location}:{sheet.title}:{field['id']}")


def _workbook_evidence(workbook, relative: str, evidence: list) -> None:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                location = f"{relative}:{sheet.title}:{cell.coordinate}"
                if cell.data_type in {"f", "e"}:
                    _block("formula_or_excel_error", location)
                value = _cell_value(cell.value)
                evidence.append({"location": location, "type": cell.data_type,
                                 "python_type": type(cell.value).__name__, "value_sha256": _digest(value)})


def _strict_dataset(app, context, raw_schema: dict, root: Path, evidence: list) -> list:
    location = context.workbook_path.relative_to(root).as_posix()
    schema = app.validate_schema(copy.deepcopy(raw_schema))
    raw_mapping = [(c.get("id"), c.get("kind"), [(f.get("id"), f.get("type")) for f in c.get("fields", [])]) for c in raw_schema.get("categories", [])]
    mapped = [(c.get("id"), c.get("kind"), [(f.get("id"), f.get("type")) for f in c.get("fields", [])]) for c in schema["categories"]]
    if raw_mapping != mapped:
        _block("schema_mapping_normalization_loss", context.schema_path.relative_to(root))
    names = app.related_sheet_names(schema)
    allowed_sheets = {app.MAIN_SHEET, app.META_SHEET, *names.values()}
    workbook = load_workbook(context.workbook_path, data_only=False, read_only=False)
    expected, by_id, children = [], {}, set()
    try:
        _workbook_evidence(workbook, location, evidence)
        for sheet in workbook.worksheets:
            if sheet.title not in allowed_sheets and any(cell.value not in (None, "") for row in sheet.iter_rows() for cell in row):
                _block("unknown_populated_sheet", f"{location}:{sheet.title}")
        if app.MAIN_SHEET not in workbook.sheetnames:
            _block("missing_main_sheet", location)
        main = workbook[app.MAIN_SHEET]
        fields = app.main_fields(schema)
        columns = _headers(main, set(app.MAIN_INTERNAL_HEADERS) | {f["id"] for f in fields}, set(app.MAIN_REQUIRED_INTERNAL_HEADERS), app.FIRST_DATA_ROW, f"{location}:{main.title}")
        _visible_labels(app, main, fields, columns, location)
        for index in range(app.FIRST_DATA_ROW, main.max_row + 1):
            if not any(cell.value not in (None, "") for cell in main[index]):
                continue
            where = f"{location}:{main.title}:{index}"
            get = lambda header: _raw_at(main, index, columns, header)
            rid, code = _text(get("_record_id"), where), _text(get("record_code"), where)
            if not re.fullmatch(r"[a-fA-F0-9]{32}", rid) or not re.fullmatch(r"[A-Z][A-Z0-9]{7}", code):
                _block("invalid_or_ignored_main_row", where)
            if rid in by_id or any(r["record_code"] == code for r in expected):
                _block("duplicate_record_identity", where)
            record = {"_record_id": rid, "record_code": code,
                      "created_at": _text(get("created_at"), where), "updated_at": _text(get("updated_at"), where),
                      "archived": _boolean(get("_archived"), where), "archived_at": _text(get("_archived_at"), where),
                      "values": {f["id"]: _field_value(app, get(f["id"]), f, f"{where}:{f['id']}") for f in fields},
                      "related": {c["id"]: [] for c in schema["categories"] if c["kind"] == "repeatable"}}
            expected.append(record); by_id[rid] = record
        for category in schema["categories"]:
            if category["kind"] != "repeatable" or names[category["id"]] not in workbook.sheetnames:
                continue
            sheet = workbook[names[category["id"]]]
            fields = app.data_fields(category)
            columns = _headers(sheet, set(app.RELATED_INTERNAL_HEADERS) | {app.RELATED_LINK_HEADER, app.RELATED_PARENT_HEADER, "_profile_link", "_transaction_origin"} | {f["id"] for f in fields}, set(app.RELATED_INTERNAL_HEADERS), app.FIRST_DATA_ROW, f"{location}:{sheet.title}")
            _visible_labels(app, sheet, fields, columns, location)
            for index in range(app.FIRST_DATA_ROW, sheet.max_row + 1):
                if not any(cell.value not in (None, "") for cell in sheet[index]):
                    continue
                where = f"{location}:{sheet.title}:{index}"
                get = lambda header: _raw_at(sheet, index, columns, header)
                rid, cid, code = (_text(get(key), where) for key in ("_record_id", "_child_id", "record_code"))
                if rid not in by_id or code != by_id[rid]["record_code"] or not re.fullmatch(r"[a-fA-F0-9]{32}", cid) or cid in children:
                    _block("invalid_or_orphan_child_row", where)
                minor = get("minor_id")
                if type(minor) is bool or not isinstance(minor, (int, float, str)) or not re.fullmatch(r"[1-9][0-9]*", str(minor)):
                    _block("minor_id_normalization_loss", where)
                child = {"_child_id": cid, "minor_id": int(minor), "created_at": _text(get("created_at"), where),
                         "updated_at": _text(get("updated_at"), where), "linked_record_code": _text(get(app.RELATED_LINK_HEADER), where),
                         "parent_child_id": _text(get(app.RELATED_PARENT_HEADER), where),
                         "values": {f["id"]: _field_value(app, get(f["id"]), f, f"{where}:{f['id']}") for f in fields}}
                for header, key in (("_profile_link", "profile_link"), ("_transaction_origin", "transaction_origin")):
                    value = get(header)
                    if value not in (None, ""):
                        try:
                            parsed = json.loads(value)
                        except (ValueError, TypeError):
                            _block("invalid_child_metadata_json", f"{where}:{header}")
                        if not isinstance(parsed, dict):
                            _block("invalid_child_metadata_json", f"{where}:{header}")
                        child[key] = parsed
                by_id[rid]["related"][category["id"]].append(child); children.add(cid)
            for record in expected:
                rows = record["related"][category["id"]]
                if len({r["minor_id"] for r in rows}) != len(rows):
                    _block("duplicate_minor_id", f"{location}:{sheet.title}")
                rows.sort(key=lambda row: (row["minor_id"], row["_child_id"]))
        if app.META_SHEET in workbook.sheetnames:
            meta = workbook[app.META_SHEET]
            allowed = {"schema_version", "schema_revision", "category_id", *names.keys()}
            seen = set()
            for row in meta.iter_rows(values_only=True):
                if not any(v not in (None, "") for v in row):
                    continue
                key = row[0]
                if key not in allowed or key in seen or any(v not in (None, "") for v in row[2:]):
                    _block("unaccounted_workbook_metadata", f"{location}:{app.META_SHEET}")
                seen.add(key)
                value = row[1] if len(row) > 1 else None
                if key in names and value != names[key]:
                    _block("sheet_mapping_mismatch", f"{location}:{app.META_SHEET}")
                if key == "schema_revision" and value != raw_schema.get("revision", 0):
                    _block("schema_revision_mismatch", f"{location}:{app.META_SHEET}")
                if key == "schema_version" and value not in app.SUPPORTED_SCHEMA_VERSIONS:
                    _block("unsupported_workbook_version", f"{location}:{app.META_SHEET}")
                if key == "category_id" and value != "sheet_name":
                    _block("unaccounted_workbook_metadata", f"{location}:{app.META_SHEET}")
        reader = getattr(app, "read_excel_dataset_unlocked", None) or app.read_dataset_unlocked
        with app.use_context(context):
            actual = reader(schema)
        if not _equal(expected, actual):
            _block("excel_reader_normalization_loss", location)
        # Storage validation checks nested row ownership and all JSON values
        # before any PostgreSQL connection or original backup is modified.
        from schemacraft_storage import _prepare_dataset
        _prepare_dataset(context.schema_id, raw_schema, expected)
        return expected
    finally:
        workbook.close()


def _strict_registry(manager, datasets: dict, root: Path, evidence: list) -> dict:
    location = manager.registry_path.relative_to(root).as_posix()
    if not manager.registry_path.is_file():
        _block("missing_identity_registry", location)
    workbook = load_workbook(manager.registry_path, data_only=False)
    registry, identities = {}, set()
    try:
        _workbook_evidence(workbook, location, evidence)
        if REGISTRY_SHEET not in workbook.sheetnames:
            _block("missing_registry_sheet", location)
        for sheet in workbook.worksheets:
            if sheet.title != REGISTRY_SHEET and any(c.value not in (None, "") for row in sheet.iter_rows() for c in row):
                _block("unknown_registry_sheet", f"{location}:{sheet.title}")
        sheet = workbook[REGISTRY_SHEET]
        columns = _headers(sheet, set(REGISTRY_HEADERS), set(REGISTRY_HEADERS), 2, location)
        for index in range(2, sheet.max_row + 1):
            if not any(cell.value not in (None, "") for cell in sheet[index]):
                continue
            where = f"{location}:{index}"
            get = lambda header: _raw_at(sheet, index, columns, header)
            code, identity = _text(get("person_id"), where), _text(get("person_uuid"), where)
            if not re.fullmatch(r"[A-Z][A-Z0-9]{7}", code) or not re.fullmatch(r"[a-fA-F0-9]{32}", identity) or code in registry or identity in identities:
                _block("invalid_or_duplicate_registry_identity", where)
            membership_text = _text(get("schema_ids"), where)
            memberships = membership_text.split(",") if membership_text else []
            if len(set(memberships)) != len(memberships) or any(sid not in datasets for sid in memberships):
                _block("invalid_registry_membership", where)
            retired = get("retired")
            if type(retired) is not bool:
                _block("registry_boolean_normalization_loss", where)
            registry[code] = {"person_id": code, "person_uuid": identity, "schema_ids": sorted(memberships),
                              "created_at": _text(get("created_at"), where), "updated_at": _text(get("updated_at"), where), "retired": retired}
            identities.add(identity)
        memberships = {}
        for sid, dataset in datasets.items():
            for record in dataset["records"]:
                memberships.setdefault(record["record_code"], []).append(sid)
        for code, expected in memberships.items():
            entry = registry.get(code)
            if not entry or entry["schema_ids"] != sorted(expected) or entry["retired"]:
                _block("registry_workbook_membership_mismatch", location)
        for code, entry in registry.items():
            if code not in memberships and (entry["schema_ids"] or not entry["retired"]):
                _block("registry_orphan_membership", location)
        return registry
    finally:
        workbook.close()


def _attachments(app, manager, datasets: dict, root: Path, inventory: dict) -> dict:
    hashed, references = {}, 0
    for context in manager.contexts(include_archived=True):
        definition = datasets[context.schema_id]["schema"]
        field_types = {f["id"]: f.get("type") for c in definition["categories"] for f in c.get("fields", [])}
        for record in datasets[context.schema_id]["records"]:
            values = [record["values"], *(child["values"] for rows in record["related"].values() for child in rows)]
            for group in values:
                for fid, value in group.items():
                    if field_types.get(fid) != "file" or value in (None, ""):
                        continue
                    with app.use_context(context):
                        relative = app.attachment_relative_path(value)
                    if not isinstance(value, str) or relative != value:
                        _block("attachment_path_normalization_loss", context.schema_id + ":" + fid)
                    path = context.folder / relative
                    key = path.relative_to(root).as_posix()
                    if key not in inventory:
                        _block("missing_referenced_attachment", key)
                    hashed[key] = inventory[key]; references += 1
        for key, item in inventory.items():
            prefix = context.attachments_path.relative_to(root).as_posix() + "/"
            if key.startswith(prefix):
                hashed[key] = item
    return {"references": references, "files": dict(sorted(hashed.items()))}


def _atomic_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, prefix=".migration-", delete=False) as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n"); stream.flush(); os.fsync(stream.fileno())
        temporary = Path(stream.name)
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _backup(root: Path, migration_id: str, inventory: dict, evidence: list) -> Path:
    parent = root / ".postgresql" / "migration-backups"
    parent.mkdir(parents=True, exist_ok=True)
    temporary = parent / (migration_id + ".incomplete")
    destination = parent / migration_id
    temporary.mkdir(exist_ok=False)
    originals = temporary / "originals"
    originals.mkdir()
    for relative, details in inventory.items():
        target = originals / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(root / relative, target)
        if target.stat().st_size != details["size"] or _file_hash(target) != details["sha256"]:
            _block("backup_hash_mismatch", relative)
    _atomic_json(temporary / "manifest.json", {"version": 1, "migration_id": migration_id, "files": inventory,
                                               "cell_inventory_sha256": _digest(evidence)})
    _atomic_json(temporary / "cell-inventory.json", evidence)
    for folder, _, files in os.walk(temporary):
        for name in files:
            os.chmod(Path(folder) / name, 0o400)
    os.replace(temporary, destination)
    for folder, _, _ in os.walk(destination, topdown=False):
        os.chmod(folder, 0o500)
    return destination


def migrate_workspace(app, store, *, activate: bool = False, diagnostic_root: Path | None = None) -> dict:
    """Return a private-value-free report; blocked sources are never repaired."""
    report = {"version": 1, "status": "blocked", "activated": False, "readback_verified": False,
              "issues": [], "migration_id": uuid.uuid4().hex, "checked_at": datetime.now(timezone.utc).isoformat()}
    original_snapshot, staged = None, False
    stage, operation = "source_preflight", "resolve_workspace"
    try:
        root = Path(app.WORKSPACE_MANAGER.data_dir if app.WORKSPACE_MANAGER else app.DATA_DIR).resolve()
        # The graphical updater verifies a sibling copy before moving it into place.
        # Backup/report references use its final location, while the server log
        # points to the actual workspace used by this migration attempt.
        final_root = Path(diagnostic_root).resolve() if diagnostic_root is not None else root
        if activate:
            report["postgres_log_path"] = str(root / ".postgresql" / "postgres.log")
        lock = getattr(app, "WORKBOOK_LOCK", None)
        operation = "acquire_workbook_lock"
        with lock if lock is not None else nullcontext():
            operation = "inventory_source"
            source_inventory = _inventory(root)
            if "workspace.json" not in source_inventory:
                _block("workspace_catalog_required", "workspace.json")
            operation = "read_workspace_catalog"
            manager = readonly_workspace(root)
            evidence, datasets = [], {}
            for context in manager.contexts(include_archived=True):
                operation = "read_schema_definition"
                relative = context.schema_path.relative_to(root).as_posix()
                schema = _read_json(context.schema_path, relative)
                operation = "validate_excel_dataset"
                records = _strict_dataset(app, context, schema, root, evidence)
                datasets[context.schema_id] = {"schema": schema, "records": records}
            operation = "validate_excel_registry"
            registry = _strict_registry(manager, datasets, root, evidence)
            operation = "validate_attachments"
            attachments = _attachments(app, manager, datasets, root, source_inventory)
            operation = "summarize_source"
            source_signature = _digest(source_inventory)
            logical_signature = _digest({"datasets": datasets, "registry": registry})
            report.update({"status": "verified_source", "source_sha256": source_signature,
                           "logical_sha256": logical_signature, "schemas": len(datasets),
                           "archived_schemas": sum(c.archived for c in manager.contexts()),
                           "records": sum(len(d["records"]) for d in datasets.values()),
                           "children": sum(len(rows) for d in datasets.values() for r in d["records"] for rows in r["related"].values()),
                           "registry_entries": len(registry), "attachment_references": attachments["references"],
                           "attachment_files": len(attachments["files"]), "attachment_sha256": _digest(attachments["files"]),
                           "source_files": len(source_inventory), "excel_nonempty_cells": len(evidence),
                           "cell_inventory_sha256": _digest(evidence)})
            operation = "recheck_source_inventory"
            if _inventory(root) != source_inventory:
                _block("source_changed_during_preflight", "workspace")
            if not activate:
                return report
            if store is None:
                _block("postgres_store_required", "target")
            stage, operation = "backup", "backup_originals"
            backup = _backup(root, report["migration_id"], source_inventory, evidence)
            backup_reference = final_root / backup.relative_to(root)
            report["backup_path"] = str(backup_reference)
            stage, operation = "target_initialization", "store.initialize"
            store.initialize()
            stage, operation = "target_inspection", "store.export_snapshot"
            original_snapshot = store.export_snapshot()
            operation = "verify_target_empty_or_identical"
            nonempty = bool(original_snapshot.get("datasets") or original_snapshot.get("registry")
                            or set(original_snapshot.get("metadata", {})) - {"storage_format_version"})
            existing = original_snapshot.get("metadata", {}).get("excel_migration")
            if nonempty and not (isinstance(existing, dict) and existing.get("logical_sha256") == logical_signature
                                 and _equal(original_snapshot["datasets"], datasets) and _equal(original_snapshot["registry"], registry)):
                _block("target_contains_unrelated_or_changed_data", "target")
            seed = {"format_version": 1, "datasets": datasets, "registry": registry,
                    "metadata": {**original_snapshot.get("metadata", {}), "excel_migration":
                                 {"migration_id": report["migration_id"], "source_sha256": source_signature,
                                  "logical_sha256": logical_signature, "backup_path": str(backup_reference)}}}
            if not nonempty:
                staged = True
                stage, operation = "target_import", "store.restore_snapshot"
                store.restore_snapshot(seed)
            stage, operation = "target_readback", "store.export_snapshot"
            readback = store.export_snapshot()
            operation = "compare_snapshot"
            if not _equal(readback["datasets"], datasets) or not _equal(readback["registry"], registry):
                _block("postgres_readback_mismatch", "target")
            for sid, dataset in datasets.items():
                operation = "store.read_schema"
                read_schema = store.read_schema(sid)
                operation = "compare_schema"
                if not _equal(read_schema, dataset["schema"]):
                    _block("postgres_dataset_readback_mismatch", sid)
                operation = "store.read_dataset"
                read_records = store.read_dataset(sid)
                operation = "compare_dataset"
                if not _equal(read_records, dataset["records"]):
                    _block("postgres_dataset_readback_mismatch", sid)
            operation = "store.read_registry"
            read_registry = store.read_registry()
            operation = "compare_registry"
            if not _equal(read_registry, registry):
                _block("postgres_registry_readback_mismatch", "target")
            report["readback_verified"] = True
            stage, operation = "source_recheck", "inventory_before_activation"
            if _inventory(root) != source_inventory:
                _block("source_changed_before_activation", "workspace")
            report.update({"status": "activated", "activated": True, "verified_at": datetime.now(timezone.utc).isoformat()})
            report_path = root / ".postgresql" / "migration-reports" / (report["migration_id"] + ".json")
            stage, operation = "activation", "write_migration_report"
            _atomic_json(report_path, report)
            operation = "write_storage_marker"
            _atomic_json(root / "storage.json", {"version": 1, "backend": "managed-postgresql",
                                                "migration_id": report["migration_id"], "verified_at": report["verified_at"],
                                                "report_path": str(final_root / report_path.relative_to(root))})
            return report
    except Exception as exc:
        report.update({"status": "blocked", "activated": False, "stage": stage, "operation": operation})
        if staged and original_snapshot is not None:
            try:
                store.restore_snapshot(original_snapshot)
                report["target_rolled_back"] = True
            except Exception as rollback_exc:
                report["target_rolled_back"] = False
                report["issues"].append(_failure_issue(rollback_exc, "target_rollback_failed", "target",
                                                        stage="target_rollback", operation="store.restore_snapshot"))
        report["issues"].append(_failure_issue(exc, "preflight_or_storage_error", type(exc).__name__,
                                               stage=stage, operation=operation))
        return report


def validate_postgres_seed(app, snapshot: dict, data_dir: Path) -> dict:
    """Validate a portable logical backup without starting PostgreSQL.

    The snapshot is authoritative: PostgreSQL workspaces may have stale Excel
    projections or no workbook at all. Values are never normalized or repaired.
    """
    from schemacraft_storage import PostgresStore, _prepare_dataset
    root = Path(data_dir).resolve(strict=True)
    manager = readonly_workspace(root, require_workbooks=False)
    if not isinstance(snapshot, dict) or snapshot.get("format_version") != 1:
        _block("invalid_snapshot_format", "postgresql/snapshot.json")
    datasets, registry, metadata = (snapshot.get(key) for key in ("datasets", "registry", "metadata"))
    if not isinstance(datasets, dict) or not isinstance(registry, dict) or not isinstance(metadata, dict):
        _block("invalid_snapshot_shape", "postgresql/snapshot.json")
    contexts = manager.contexts(include_archived=True)
    if set(datasets) != {context.schema_id for context in contexts}:
        _block("snapshot_catalog_schema_mismatch", "workspace.json")
    for context in contexts:
        dataset = datasets[context.schema_id]
        if not isinstance(dataset, dict) or set(dataset) != {"schema", "records"}:
            _block("invalid_snapshot_dataset", context.schema_id)
        schema, records = dataset["schema"], dataset["records"]
        app.validate_schema(copy.deepcopy(schema))
        _prepare_dataset(context.schema_id, schema, records)
        for record in records:
            if not re.fullmatch(r"[a-fA-F0-9]{32}", str(record.get("_record_id", ""))) or not re.fullmatch(r"[A-Z][A-Z0-9]{7}", str(record.get("record_code", ""))):
                _block("invalid_snapshot_record_identity", context.schema_id)
            for rows in record.get("related", {}).values():
                for child in rows:
                    if not re.fullmatch(r"[a-fA-F0-9]{32}", str(child.get("_child_id", ""))):
                        _block("invalid_snapshot_child_identity", context.schema_id)
    PostgresStore._prepare_registry(registry)
    memberships, identities = {}, set()
    for sid, dataset in datasets.items():
        for record in dataset["records"]:
            memberships.setdefault(record["record_code"], []).append(sid)
    for code, entry in registry.items():
        identity = entry.get("person_uuid", "")
        if not re.fullmatch(r"[A-Z][A-Z0-9]{7}", code) or not re.fullmatch(r"[a-fA-F0-9]{32}", identity) or identity in identities:
            _block("invalid_snapshot_registry_identity", "postgresql/snapshot.json")
        if entry.get("schema_ids") != sorted(memberships.get(code, [])) or type(entry.get("retired")) is not bool or entry["retired"] != (code not in memberships):
            _block("snapshot_registry_membership_mismatch", "postgresql/snapshot.json")
        identities.add(identity)
    if set(memberships) - set(registry):
        _block("snapshot_registry_missing_identity", "postgresql/snapshot.json")
    from schemacraft_storage import _validate_json
    _validate_json(snapshot)
    inventory = _inventory(root)
    attachments = _attachments(app, manager, datasets, root, inventory)
    return {"schemas": len(datasets), "records": sum(len(d["records"]) for d in datasets.values()),
            "attachment_files": len(attachments["files"]), "attachment_references": attachments["references"],
            "logical_sha256": _digest({"datasets": datasets, "registry": registry}),
            "attachment_sha256": _digest(attachments["files"])}


def _extract_backup(archive_path: Path, destination: Path) -> dict:
    """Bounded ZIP extraction, with Windows and POSIX path checks."""
    total, seen = 0, set()
    with zipfile.ZipFile(archive_path) as archive:
        entries = archive.infolist()
        if len(entries) > 100000:
            _block("backup_file_count_limit", "archive")
        for info in entries:
            name = info.filename
            relative = PurePosixPath(name)
            if (not name or len(name) > 1024 or name.startswith("/") or "\\" in name or ":" in name or "\x00" in name
                    or any(part in {"", ".", "..", ".postgresql"} for part in name.rstrip("/").split("/"))
                    or relative.is_absolute() or (info.external_attr >> 16) & 0o170000 == 0o120000):
                _block("unsafe_backup_path", "archive")
            key = name.rstrip("/").casefold()
            if key in seen:
                _block("duplicate_backup_path", "archive")
            seen.add(key)
            total += info.file_size
            if info.file_size > 4 * 1024**3 or total > 64 * 1024**3:
                _block("backup_size_limit", "archive")
            if name.startswith("postgresql/") and name not in {"postgresql/snapshot.json", "postgresql/README.txt", "postgresql/"}:
                _block("unknown_postgresql_backup_entry", "archive")
        for info in entries:
            target = destination / info.filename
            if info.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with archive.open(info) as source, target.open("xb") as output:
                shutil.copyfileobj(source, output, 1024 * 1024)
    return _read_json(destination / "postgresql" / "snapshot.json", "postgresql/snapshot.json")


def restore_backup(app, store, archive_path: Path, *, data_dir: Path | None = None) -> dict:
    """Restore a portable PostgreSQL ZIP only into an empty deployment."""
    report = {"version": 1, "status": "blocked", "activated": False, "readback_verified": False, "issues": []}
    original_snapshot, staged, installed = None, False, []
    stage, operation = "restore_preflight", "resolve_workspace"
    try:
        root = Path(data_dir or app.DATA_DIR).resolve()
        report["postgres_log_path"] = str(root / ".postgresql" / "postgres.log")
        operation = "prepare_empty_workspace"
        root.mkdir(parents=True, exist_ok=True)
        operation = "inventory_restore_target"
        if _inventory(root):
            _block("restore_requires_empty_workspace", "target")
        if store is None:
            _block("postgres_store_required", "target")
        operation = "create_restore_staging_directory"
        with tempfile.TemporaryDirectory(prefix="schemacraft-restore-", dir=root.parent) as temporary:
            extracted = Path(temporary)
            operation = "extract_backup"
            snapshot = _extract_backup(Path(archive_path), extracted)
            operation = "validate_postgres_seed"
            report.update(validate_postgres_seed(app, snapshot, extracted))
            # Every original backup byte is kept except portable logical
            # snapshot documentation; schema projections come from authority.
            operation = "prepare_schema_projections"
            for context in readonly_workspace(extracted, require_workbooks=False).contexts():
                _atomic_json(context.schema_path, snapshot["datasets"][context.schema_id]["schema"])
            (extracted / "storage.json").unlink(missing_ok=True)
            stage, operation = "target_initialization", "store.initialize"
            store.initialize()
            stage, operation = "target_inspection", "store.export_snapshot"
            original_snapshot = store.export_snapshot()
            operation = "verify_restore_target_empty"
            if original_snapshot.get("datasets") or original_snapshot.get("registry") or set(original_snapshot.get("metadata", {})) - {"storage_format_version"}:
                _block("restore_requires_empty_database", "target")
            staged = True
            stage, operation = "target_import", "store.restore_snapshot"
            store.restore_snapshot(snapshot)
            stage, operation = "target_readback", "store.export_snapshot"
            readback = store.export_snapshot()
            operation = "compare_restore_snapshot"
            if not _equal(readback, snapshot):
                _block("restore_postgres_readback_mismatch", "target")
            report["readback_verified"] = True
            stage, operation = "restore_install", "inventory_before_install"
            if _inventory(root):
                _block("restore_target_changed", "target")
            for path in list(extracted.iterdir()):
                if path.name == "postgresql":
                    continue
                target = root / path.name
                operation = "install_restored_files"
                os.replace(path, target)
                installed.append(target)
            stage, operation = "target_readback", "store.export_snapshot"
            readback = store.export_snapshot()
            operation = "validate_installed_postgres_seed"
            validate_postgres_seed(app, readback, root)
            restore_id = uuid.uuid4().hex
            report.update({"status": "activated", "activated": True, "restore_id": restore_id,
                           "verified_at": datetime.now(timezone.utc).isoformat()})
            report_path = root / ".postgresql" / "migration-reports" / ("restore-" + restore_id + ".json")
            stage, operation = "activation", "write_restore_report"
            _atomic_json(report_path, report)
            operation = "write_storage_marker"
            _atomic_json(root / "storage.json", {"version": 1, "backend": "managed-postgresql", "restore_id": restore_id,
                                                "verified_at": report["verified_at"], "report_path": str(report_path)})
            return report
    except Exception as exc:
        report.update(status="blocked", activated=False, stage=stage, operation=operation)
        if staged and original_snapshot is not None:
            try:
                store.restore_snapshot(original_snapshot)
                report["target_rolled_back"] = True
            except Exception as rollback_exc:
                report["target_rolled_back"] = False
                report["issues"].append(_failure_issue(rollback_exc, "target_rollback_failed", "target",
                                                        stage="target_rollback", operation="store.restore_snapshot"))
        for path in installed:
            try:
                if path.is_dir():
                    shutil.rmtree(path)
                else:
                    path.unlink(missing_ok=True)
            except Exception as cleanup_exc:
                report["issues"].append(_failure_issue(cleanup_exc, "restore_cleanup_failed", "target",
                                                        stage="restore_cleanup", operation="remove_installed_files"))
        report["issues"].append(_failure_issue(exc, "backup_restore_error", type(exc).__name__,
                                               stage=stage, operation=operation))
        return report


def cli_main(argv=None, *, app=None) -> int:
    """Shared source/frozen command line; no implicit data directory."""
    import argparse
    import sys
    parser = argparse.ArgumentParser(description="Verify an Excel workspace and optionally migrate it to managed PostgreSQL.")
    parser.add_argument("--data-dir", required=True, type=Path, help="Existing workspace directory; close SchemaCraft before applying.")
    action = parser.add_mutually_exclusive_group(required=True)
    action.add_argument("--verify-only", action="store_true", help="Read-only preflight; no database or workspace writes.")
    action.add_argument("--apply", action="store_true", help="Backup originals, import, verify, then activate PostgreSQL.")
    action.add_argument("--restore-backup", type=Path, help="Restore a portable PostgreSQL ZIP into an empty --data-dir.")
    parser.add_argument("--report", type=Path, help="Optional JSON report outside the selected workspace.")
    parser.add_argument("--final-data-dir", type=Path, help="Final location for diagnostic references when an updater verifies a staged workspace.")
    args = parser.parse_args(argv)
    root = args.data_dir.expanduser().resolve()
    if not args.restore_backup and not root.is_dir():
        parser.error("--data-dir must be an existing directory")
    if args.restore_backup and root.exists() and (not root.is_dir() or any(root.iterdir())):
        parser.error("backup restore requires a new or completely empty --data-dir")
    if args.report and args.report.expanduser().resolve().is_relative_to(root):
        parser.error("--report must be outside the source workspace")
    imported_app = app is None
    report = {"version": 1, "status": "blocked", "activated": False, "issues": []}
    runtime, store = None, None
    stage, operation = "migration_setup", "import_application"
    try:
        if imported_app:
            import SchemaCraft as app
        operation = "acquire_application_instance"
        if imported_app and (args.apply or args.restore_backup) and not app.acquire_single_instance():
            _block("close_running_application_first", "application")
        app.DATA_DIR = root
        if args.apply or args.restore_backup:
            report["postgres_log_path"] = str(root / ".postgresql" / "postgres.log")
        if args.restore_backup:
            # Reject damaged/unsafe archives before provisioning a server.
            stage, operation = "restore_preflight", "create_restore_staging_directory"
            with tempfile.TemporaryDirectory(prefix="schemacraft-restore-preflight-") as temporary:
                operation = "extract_backup"
                snapshot = _extract_backup(args.restore_backup.expanduser().resolve(), Path(temporary))
                operation = "validate_postgres_seed"
                validate_postgres_seed(app, snapshot, Path(temporary))
            report["status"] = "verified_backup"
        else:
            stage, operation = "source_preflight", "read_workspace_catalog"
            app.WORKSPACE_MANAGER = readonly_workspace(root)
            operation = "migrate_workspace"
            report = migrate_workspace(app, None, activate=False)
        if args.restore_backup or (args.apply and report["status"] == "verified_source"):
            stage, operation = "migration_setup", "import_postgres_runtime"
            from schemacraft_postgres_runtime import ManagedPostgres
            operation = "import_postgres_store"
            from schemacraft_storage import PostgresStore
            operation = "create_managed_runtime"
            report["postgres_log_path"] = str(root / ".postgresql" / "postgres.log")
            runtime = ManagedPostgres(app.BASE_DIR, root)
            operation = "runtime.start"
            dsn = runtime.start()
            operation = "create_postgres_store"
            store = PostgresStore(dsn)
            stage = "restore" if args.restore_backup else "migration"
            operation = "restore_backup" if args.restore_backup else "migrate_workspace"
            report = (restore_backup(app, store, args.restore_backup.expanduser().resolve(), data_dir=root)
                      if args.restore_backup else migrate_workspace(app, store, activate=True, diagnostic_root=args.final_data_dir))
    except Exception as exc:
        _record_failure(report, exc, "migration_setup_failed", stage=stage, operation=operation)
    finally:
        # Cleanup failures must not replace the primary migration failure. Both
        # cleanup steps run independently, including when connection setup fails.
        if store is not None:
            try:
                store.close()
            except Exception as exc:
                _record_failure(report, exc, "migration_cleanup_failed", stage="migration_cleanup", operation="store.close")
        if runtime is not None:
            try:
                runtime.stop()
            except Exception as exc:
                _record_failure(report, exc, "migration_cleanup_failed", stage="migration_cleanup", operation="runtime.stop")
    output = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    if args.report:
        try:
            args.report.expanduser().write_text(output, encoding="utf-8")
        except OSError as exc:
            report.setdefault("issues", []).append(_failure_issue(exc, "report_write_failed", "report",
                                                                   stage="report_write", operation="write_requested_report"))
            report["report_write_failed"] = True
            output = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
            sys.stdout.write(output)
            sys.stderr.write("Unable to write the requested report file.\n")
            return 2
    sys.stdout.write(output)
    return 0 if report["status"] in {"verified_source", "activated"} else 2
