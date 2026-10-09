"""Read-only exports of schema view categories through the existing Export page."""
from __future__ import annotations
import io
from datetime import datetime
import schemacraft_finance as F
import schemacraft_finance_api as API


def export_view(app, payload):
    app.require_builder_access()
    schema=app.read_schema_file()
    category=next((c for c in schema['categories'] if c['id']==payload.get('category_id') and c.get('view_table')),None)
    if category is None:raise app.ApplicationError('اختر جدول عرض من التصميم الحالي.')
    loaded=app.load_record(payload.get('record_code'))
    if loaded.get('archived'):raise app.ApplicationError('لا تصدّر جدول عرض من سجل مؤرشف.')
    record={**loaded,'values':loaded['main']}
    try:
        errors,engine=API.recompute(app,schema,record,strict=True)
        data=engine.view(category['id'])
    except F.FinanceError as exc:raise app.ApplicationError(str(exc)) from exc
    fmt=payload.get('format','xlsx')
    if fmt not in {'xlsx','pdf'}:raise app.ApplicationError('صيغة تصدير جدول العرض غير صالحة.')
    stamp=datetime.now().astimezone();title=category['label'];code=record['record_code']
    rows=[[row['display'].get(c['id'],'') for c in data['columns']] for row in data['rows']]
    if fmt=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Alignment,Font,PatternFill,Border,Side
        from openpyxl.utils import get_column_letter
        wb=Workbook();ws=wb.active;ws.title='جدول العرض';ws.sheet_view.rightToLeft=True;ws.sheet_view.showGridLines=False
        headers=[c['label'] for c in data['columns']];ws.append(headers)
        for row in rows:ws.append(row)
        # All values are a fixed, faithful snapshot; never execute a source string as a formula.
        for row in ws:
            for cell in row:
                cell.data_type='s';cell.alignment=Alignment(horizontal='right',vertical='top',wrap_text=True)
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='194F79');c.font=Font(color='FFFFFF',bold=True)
        ws.row_dimensions[1].height=30;ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{max(1,ws.max_row)}'
        for i in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(i)].width=24
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border=Border(bottom=Side(style='hair',color='DCE4EC'))
                if c.row%2==0:c.fill=PatternFill('solid',fgColor='F4F8FC')
        summary=wb.create_sheet('المجاميع');summary.sheet_view.rightToLeft=True;summary.append(['حقل الجمع','القيمة','العملة'])
        for t in data['totals']:summary.append([t['label'],t.get('display',t['value']),t['currency']])
        info=wb.create_sheet('تعريف التصدير');info.sheet_view.rightToLeft=True
        info.append(['جدول العرض',title]);info.append(['السجل',code]);info.append(['تاريخ الإنشاء',stamp.isoformat()]);info.append(['عدد الصفوف',str(data['count'])]);info.append(['المحتوى','نسخة من البيانات المحفوظة مع شروط التصميم؛ المرشحات المؤقتة لا تُستخدم.'])
        for sheet in [summary,info]:
            for row in sheet:
                for c in row:c.data_type='s';c.alignment=Alignment(horizontal='right',wrap_text=True)
            for col in ['A','B','C']:sheet.column_dimensions[col].width=28 if col!='B' else 48
            for c in sheet[1]:c.font=Font(bold=True)
        output=io.BytesIO();wb.save(output);content=output.getvalue()
    else:
        if len(data['columns'])>12:raise app.ApplicationError('جدول PDF واسع جدًا؛ اختر حتى 12 عمودًا أو استخدم Excel.')
        if len(rows)>10000:raise app.ApplicationError('جدول PDF كبير جدًا؛ ضيّق شروط العرض أو استخدم Excel.')
        from schemacraft_document_v3 import pdf_bytes
        doc={'title':title,'page':{'size':'A4','orientation':'landscape' if len(data['columns'])>4 else 'portrait','margin':16,'font_size':10},'body':[
            {'type':'heading','text':title},
            {'type':'paragraph','text':f'{code} · {stamp.date().isoformat()} · {data["count"]}'},
            {'type':'table','columns':[{'label':c['label']} for c in data['columns']],'rows':rows},
            *[{'type':'paragraph','text':f'{t["label"]}: {t.get("display",t["value"])} {t["currency"]}'} for t in data['totals']]
        ]}
        # The document renderer consumes the "blocks" array after structured evaluation.
        doc['blocks']=doc.pop('body')
        content=pdf_bytes(doc,app.PDF_FONT_PATH)
    name=app.sanitize_filename_text(title) or 'view-table'
    return f'{name}-{code}-{stamp.strftime("%Y%m%d-%H%M%S")}.{fmt}',content,data['count'],[app.current_schema_id()]
