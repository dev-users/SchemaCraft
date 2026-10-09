"""Excel import/export helpers for SchemaCraft.

This module intentionally has no dependency on the HTTP application.  It
receives an already validated schema and plain record dictionaries, which
keeps workbook interchange separate from record validation and persistence.
"""

from __future__ import annotations
import copy
import schemacraft_composite_text

import base64
import binascii
import io
import json
import re
import unicodedata
from collections.abc import Iterable
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXPORT_META_SHEET = "_schemacraft_export"
EXPORT_MAIN_SHEET = "السجلات المصدرة"
MAX_IMPORT_BYTES = 50 * 1024 * 1024
SPECIAL_TARGETS = {
    "record_code": "__record_code__",
    "archived": "__archived__",
}


class WorkbookExchangeError(ValueError):
    """An import/export error that is safe to display to the user."""


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKC", _text(value)).casefold()
    text = re.sub(r"[\u0610-\u061A\u064B-\u065F\u0670\u06D6-\u06ED]", "", text)
    text = text.translate(
        {
            ord("أ"): "ا",
            ord("إ"): "ا",
            ord("آ"): "ا",
            ord("ٱ"): "ا",
            ord("ى"): "ي",
            ord("ـ"): "",
        }
    )
    return " ".join(text.split())


def _excel_value(value: Any, field: dict[str, Any]) -> Any:
    if value is None:
        return ""
    if field.get("type") == "checkbox":
        return "نعم" if bool(value) else "لا"
    if field.get("type") == "checkbox_group" and isinstance(value, (list, tuple)):
        return " | ".join(_text(item) for item in value)
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _style_sheet(worksheet, column_count: int) -> None:
    worksheet.row_dimensions[1].hidden = True
    worksheet.row_dimensions[2].height = 28
    worksheet.freeze_panes = "A3"
    worksheet.sheet_view.showGridLines = False
    for index in range(1, column_count + 1):
        cell = worksheet.cell(2, index)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.column_dimensions[get_column_letter(index)].width = max(
            13, min(40, len(_text(cell.value)) + 5)
        )


def _safe_sheet_name(label: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", label).strip(" '") or "فئة"
    base = base[:31]
    candidate = base
    sequence = 2
    while candidate.casefold() in used:
        suffix = f"-{sequence}"
        candidate = f"{base[:31-len(suffix)]}{suffix}"
        sequence += 1
    used.add(candidate.casefold())
    return candidate


def export_workbook_bytes(
    schema: dict[str, Any],
    records: Iterable[dict[str, Any]],
    selected_field_ids: Iterable[str] | None = None,
    *,
    include_related: bool = True,
) -> bytes:
    """Build a readable, round-trippable filtered workbook.

    Modern schemas expose profile metadata (record ID / created / updated) as
    normal system fields.  When those fields exist, their presence in the
    workbook follows the user's field selection just like any other field.
    Legacy schemas without explicit system fields keep the historical behavior
    and always receive those technical metadata columns so old round-trips stay
    compatible.
    """

    records = [schemacraft_composite_text.recompute_record(schema, copy.deepcopy(record)) for record in records]
    normal_selectable = {
        field["id"]: (category, field)
        for category in schema.get("categories", [])
        for field in category.get("fields", [])
        if field.get("type")
        not in {
            "file",
            "spacer", "field_group",
            "system_record_code",
            "system_created_at",
            "system_updated_at",
        }
    }
    system_fields_by_type: dict[str, list[dict[str, Any]]] = {
        "system_record_code": [],
        "system_created_at": [],
        "system_updated_at": [],
    }
    for category in schema.get("categories", []):
        for field in category.get("fields", []):
            if field.get("type") in system_fields_by_type:
                system_fields_by_type[field["type"]].append(field)

    if selected_field_ids is None:
        requested = list(normal_selectable) + [
            field["id"]
            for fields in system_fields_by_type.values()
            for field in fields
        ]
    else:
        requested = list(selected_field_ids)
    requested_set = set(requested)
    selected = {field_id for field_id in requested if field_id in normal_selectable}

    has_explicit_system_fields = any(system_fields_by_type.values())

    def include_system_field(field_type: str) -> bool:
        # Old schemas have no selectable metadata fields, so retain the former
        # always-on metadata columns.  New schemas honor the user's selection.
        if not has_explicit_system_fields or selected_field_ids is None:
            return True
        return any(field["id"] in requested_set for field in system_fields_by_type[field_type])

    def system_label(field_type: str, fallback: str) -> str:
        fields = system_fields_by_type[field_type]
        return fields[0].get("label") or fallback if fields else fallback

    metadata_columns: list[tuple[str, str, Any]] = []
    if include_system_field("system_record_code"):
        metadata_columns.append(("record_code", system_label("system_record_code", "ID"), lambda record: record.get("record_code", "")))
    # Archive state is internal export metadata rather than a configurable
    # schema field, so it remains present for round-trip compatibility.
    metadata_columns.append(("_archived", "مؤرشف", lambda record: "نعم" if record.get("archived") else "لا"))
    if include_system_field("system_created_at"):
        metadata_columns.append(("created_at", system_label("system_created_at", "تاريخ الإنشاء"), lambda record: record.get("created_at", "")))
    if include_system_field("system_updated_at"):
        metadata_columns.append(("updated_at", system_label("system_updated_at", "تاريخ التعديل"), lambda record: record.get("updated_at", "")))

    workbook = Workbook()
    main = workbook.active
    main.title = EXPORT_MAIN_SHEET
    main_fields = [
        field
        for category in schema.get("categories", [])
        if category.get("kind") == "main"
        for field in category.get("fields", [])
        if field["id"] in selected
    ]
    technical = [
        *(column[0] for column in metadata_columns),
        *(field["id"] for field in main_fields),
    ]
    labels = [
        *(column[1] for column in metadata_columns),
        *(field["label"] for field in main_fields),
    ]
    main.append(technical)
    main.append(labels)
    for record in records:
        main.append(
            [
                *(getter(record) for _technical, _label, getter in metadata_columns),
                *(
                    _excel_value(record.get("values", {}).get(field["id"], ""), field)
                    for field in main_fields
                ),
            ]
        )
    for row in range(3, main.max_row + 1):
        schemacraft_composite_text.literal_cells(main, main_fields, row, {key:i+1 for i,key in enumerate(technical)})
    _style_sheet(main, len(technical))
    main.auto_filter.ref = f"A2:{get_column_letter(len(technical))}{max(2, main.max_row)}"

    used_names = {EXPORT_MAIN_SHEET.casefold(), EXPORT_META_SHEET.casefold()}
    category_sheets: list[tuple[str, str]] = []
    if include_related:
        for category in schema.get("categories", []):
            if category.get("kind") != "repeatable":
                continue
            fields = [
                field
                for field in category.get("fields", [])
                if field["id"] in selected and field.get("type") not in {"file", "spacer", "field_group"}
            ]
            if not fields:
                continue
            sheet_name = _safe_sheet_name(category["label"], used_names)
            category_sheets.append((category["id"], sheet_name))
            worksheet = workbook.create_sheet(sheet_name)
            headers = [
                "record_code",
                "minor_id",
                "_linked_record_code",
                *(field["id"] for field in fields),
            ]
            visible = [
                "ID",
                "minor_id",
                "ID الشخص المرتبط",
                *(field["label"] for field in fields),
            ]
            worksheet.append(headers)
            worksheet.append(visible)
            for record in records:
                for position, row in enumerate(
                    record.get("related", {}).get(category["id"], []), start=1
                ):
                    worksheet.append(
                        [
                            record.get("record_code", ""),
                            row.get("minor_id") or position,
                            row.get("linked_record_code", ""),
                            *(
                                _excel_value(row.get("values", {}).get(field["id"], ""), field)
                                for field in fields
                            ),
                        ]
                    )
            for row in range(3, worksheet.max_row + 1):
                schemacraft_composite_text.literal_cells(worksheet, fields, row, {key:i+1 for i,key in enumerate(headers)})
            _style_sheet(worksheet, len(headers))
            worksheet.auto_filter.ref = (
                f"A2:{get_column_letter(len(headers))}{max(2, worksheet.max_row)}"
            )

    meta = workbook.create_sheet(EXPORT_META_SHEET)
    meta.sheet_state = "hidden"
    meta.append(["format", "SchemaCraft Export"])
    meta.append(["version", 1])
    meta.append(["schema_revision", schema.get("revision", 0)])
    meta.append(["main_sheet", EXPORT_MAIN_SHEET])
    meta.append(["category_id", "sheet_name"])
    for category_id, sheet_name in category_sheets:
        meta.append([category_id, sheet_name])

    output = io.BytesIO()
    try:
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


def decode_workbook_data(encoded: Any) -> bytes:
    if not isinstance(encoded, str) or not encoded:
        raise WorkbookExchangeError("بيانات ملف الاستيراد غير مكتملة.")
    try:
        content = base64.b64decode(encoded, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise WorkbookExchangeError("تعذّر قراءة ملف الاستيراد.") from exc
    if not content.startswith(b"PK"):
        raise WorkbookExchangeError("اختر ملف Excel بصيغة .xlsx.")
    if len(content) > MAX_IMPORT_BYTES:
        raise WorkbookExchangeError("حجم ملف الاستيراد يتجاوز 50 ميغابايت.")
    return content


def _open_import_workbook(encoded: Any):
    content = decode_workbook_data(encoded)
    try:
        return load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise WorkbookExchangeError(f"تعذّر فتح ملف Excel: {exc}") from exc


def _sheet_headers(worksheet, field_ids: set[str]) -> dict[str, Any]:
    first = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    first_values = [_text(value) for value in first]
    technical_tokens = field_ids | {
        "record_code",
        "_archived",
        "created_at",
        "updated_at",
    }
    technical = any(value in field_ids for value in first_values) or ("record_code" in first_values and "_archived" in first_values)
    if technical:
        second = next(
            worksheet.iter_rows(min_row=2, max_row=2, values_only=True), ()
        )
        labels = [_text(value) for value in second]
        return {
            "technical": first_values,
            "labels": labels,
            "header_row": 2,
            "data_row": 3,
        }
    return {
        "technical": ["" for _ in first_values],
        "labels": first_values,
        "header_row": 1,
        "data_row": 2,
    }


def _main_fields(schema: dict[str, Any]) -> list[tuple[dict[str, Any], dict[str, Any]]]:
    return [
        (category, field)
        for category in schema.get("categories", [])
        if category.get("kind") == "main"
        for field in category.get("fields", [])
        if field.get("type")
        not in {
            "file",
            "spacer", "field_group",
            "system_record_code",
            "system_created_at",
            "system_updated_at",
        }
    ]


def inspect_import_workbook(
    encoded: Any,
    schema: dict[str, Any],
    requested_sheet: str | None = None,
    *, _workbook=None,
) -> dict[str, Any]:
    workbook = _workbook or _open_import_workbook(encoded)
    try:
        visible_sheets = [
            name for name in workbook.sheetnames if name != EXPORT_META_SHEET
        ]
        if not visible_sheets:
            raise WorkbookExchangeError("لا توجد ورقة قابلة للاستيراد في الملف.")
        sheet_name = requested_sheet if requested_sheet in visible_sheets else None
        if sheet_name is None and EXPORT_MAIN_SHEET in visible_sheets:
            sheet_name = EXPORT_MAIN_SHEET
        if sheet_name is None:
            sheet_name = visible_sheets[0]
        worksheet = workbook[sheet_name]
        fields = _main_fields(schema)
        field_ids = {field["id"] for _, field in fields}
        headers = _sheet_headers(worksheet, field_ids)

        by_label: dict[str, list[str]] = {}
        for category, field in fields:
            for label in {
                field["label"],
                f'{category["label"]} — {field["label"]}',
                f'{category["label"]} - {field["label"]}',
            }:
                by_label.setdefault(_normalize(label), []).append(field["id"])

        columns = []
        problems: list[dict[str, Any]] = []
        field_names = {
            field["id"]: f'{category["label"]} — {field["label"]}'
            for category, field in fields
        }
        used_suggestions: dict[str, int] = {}
        maximum = max(len(headers["technical"]), len(headers["labels"]))
        for offset in range(maximum):
            technical = (
                headers["technical"][offset]
                if offset < len(headers["technical"])
                else ""
            )
            label = headers["labels"][offset] if offset < len(headers["labels"]) else ""
            if not technical and not label:
                continue
            suggestion = ""
            if technical in field_ids:
                suggestion = technical
            elif technical == "record_code" or _normalize(label) in {
                "id",
                "record_code",
                "معرف السجل",
                "معرّف السجل",
            }:
                suggestion = SPECIAL_TARGETS["record_code"]
            else:
                candidates = list(dict.fromkeys(by_label.get(_normalize(label), [])))
                if len(candidates) == 1:
                    suggestion = candidates[0]
                elif len(candidates) > 1:
                    choices = "، ".join(field_names.get(item, item) for item in candidates)
                    problems.append(
                        {
                            "column": offset + 1,
                            "code": "ambiguous_header",
                            "message": f'العمود «{label or technical}» يطابق أكثر من حقل: {choices}. اختر المطابقة يدويًا.',
                        }
                    )
            if suggestion and suggestion in used_suggestions:
                first_column = used_suggestions[suggestion]
                problems.append(
                    {
                        "column": offset + 1,
                        "code": "duplicate_target",
                        "message": f'العمود «{label or technical}» يكرر مطابقة العمود {first_column}. اختر حقلًا مختلفًا أو تجاهله.',
                    }
                )
                suggestion = ""
            elif suggestion:
                used_suggestions[suggestion] = offset + 1
            if not suggestion and not any(
                item.get("column") == offset + 1 for item in problems
            ):
                problems.append(
                    {
                        "column": offset + 1,
                        "code": "unmatched_header",
                        "message": f'لم تُحدّد مطابقة تلقائية للعمود «{label or technical}». اختر الحقل المقابل أو تجاهل العمود.',
                    }
                )
            columns.append(
                {
                    "column": offset + 1,
                    "header": label or technical or f"عمود {offset + 1}",
                    "technical_header": technical,
                    "suggested_target": suggestion,
                }
            )

        preview = []
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=headers["data_row"], values_only=True),
            start=headers["data_row"],
        ):
            values = [_text(value) for value in row]
            if not any(values):
                continue
            preview.append({"row": row_index, "values": values})
            if len(preview) >= 20:
                break

        data_rows = 0
        for row in worksheet.iter_rows(min_row=headers["data_row"], values_only=True):
            if any(_text(value) for value in row):
                data_rows += 1

        if not columns:
            problems.append(
                {
                    "column": 0,
                    "code": "no_columns",
                    "message": "لم يُعثر على أعمدة معنونة قابلة للمطابقة في الورقة.",
                }
            )
        if not data_rows:
            problems.append(
                {
                    "column": 0,
                    "code": "no_data_rows",
                    "message": "لا تحتوي الورقة المختارة على صفوف بيانات.",
                }
            )

        return {
            "sheet_name": sheet_name,
            "sheets": visible_sheets,
            "columns": columns,
            "fields": [
                {
                    "id": field["id"],
                    "label": field["label"],
                    "category": category["label"],
                    "type": field["type"],
                    "required": bool(field.get("required")),
                }
                for category, field in fields
            ],
            "preview": preview,
            "data_rows": data_rows,
            "problems": problems,
            "native_export": EXPORT_META_SHEET in workbook.sheetnames,
        }
    finally:
        if _workbook is None:
            workbook.close()


def _truthy(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return _normalize(value) in {"1", "true", "yes", "نعم", "مكتمل"}


def parse_import_rows(
    encoded: Any,
    schema: dict[str, Any],
    *,
    sheet_name: str,
    mapping: dict[str, str],
) -> list[dict[str, Any]]:
    """Return raw record submissions; application validation runs later."""

    workbook = _open_import_workbook(encoded)
    try:
        if sheet_name not in workbook.sheetnames or sheet_name == EXPORT_META_SHEET:
            raise WorkbookExchangeError("ورقة الاستيراد المختارة غير موجودة.")
        worksheet = workbook[sheet_name]
        fields = _main_fields(schema)
        field_ids = {field["id"] for _, field in fields}
        headers = _sheet_headers(worksheet, field_ids)
        valid_targets = field_ids | set(SPECIAL_TARGETS.values())
        normalized_mapping: dict[int, str] = {}
        used_targets: set[str] = set()
        for raw_column, raw_target in mapping.items():
            try:
                column = int(raw_column)
            except (TypeError, ValueError):
                continue
            target = _text(raw_target)
            if column < 1 or target not in valid_targets or target in used_targets:
                continue
            normalized_mapping[column] = target
            used_targets.add(target)
        if not any(target in field_ids for target in normalized_mapping.values()):
            raise WorkbookExchangeError("اربط عمود بيانات واحدًا على الأقل بحقل في التطبيق.")

        imported: list[dict[str, Any]] = []
        by_code: dict[str, dict[str, Any]] = {}
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=headers["data_row"], values_only=True),
            start=headers["data_row"],
        ):
            if not any(_text(value) for value in row):
                continue
            item = {
                "source_row": row_index,
                "record_code": "",
                "archived": False,
                "main": {},
                "related": {},
            }
            for column, target in normalized_mapping.items():
                value = row[column - 1] if column <= len(row) else None
                if target == SPECIAL_TARGETS["record_code"]:
                    item["record_code"] = _text(value).upper()
                elif target == SPECIAL_TARGETS["archived"]:
                    item["archived"] = _truthy(value)
                else:
                    item["main"][target] = value
            imported.append(item)
            if item["record_code"]:
                by_code[item["record_code"]] = item

        if EXPORT_META_SHEET in workbook.sheetnames and by_code:
            meta = workbook[EXPORT_META_SHEET]
            category_sheets = {
                _text(row[0]): _text(row[1])
                for row in meta.iter_rows(min_row=6, values_only=True)
                if len(row) >= 2 and _text(row[0]) and _text(row[1])
            }
            categories = {
                category["id"]: category
                for category in schema.get("categories", [])
                if category.get("kind") == "repeatable"
            }
            for category_id, related_sheet_name in category_sheets.items():
                category = categories.get(category_id)
                if category is None or related_sheet_name not in workbook.sheetnames:
                    continue
                related_sheet = workbook[related_sheet_name]
                related_field_ids = {
                    field["id"]
                    for field in category.get("fields", [])
                    if field.get("type") not in {"file", "spacer", "field_group"}
                }
                related_headers = _sheet_headers(
                    related_sheet, related_field_ids
                )
                header_map = {
                    value: index
                    for index, value in enumerate(related_headers["technical"], start=1)
                    if value
                }
                code_column = header_map.get("record_code")
                if not code_column:
                    continue
                for related_row in related_sheet.iter_rows(
                    min_row=related_headers["data_row"], values_only=True
                ):
                    code = _text(
                        related_row[code_column - 1]
                        if code_column <= len(related_row)
                        else ""
                    ).upper()
                    owner = by_code.get(code)
                    if owner is None:
                        continue
                    values = {
                        field_id: (
                            related_row[column - 1]
                            if column <= len(related_row)
                            else ""
                        )
                        for field_id in related_field_ids
                        if (column := header_map.get(field_id))
                    }
                    linked_column = header_map.get("_linked_record_code")
                    linked = (
                        _text(related_row[linked_column - 1]).upper()
                        if linked_column and linked_column <= len(related_row)
                        else ""
                    )
                    owner["related"].setdefault(category_id, []).append(
                        {
                            "values": values,
                            "linked_record_code": linked,
                        }
                    )
        return imported
    finally:
        workbook.close()


def inspect_import_sheets(encoded, schema):
    """Inspect every worksheet once and expose category-scoped mapping metadata."""
    workbook = _open_import_workbook(encoded)
    try:
        categories = [c for c in schema.get('categories', []) if c.get('kind') in {'main','repeatable'}]
        meta = {}
        if EXPORT_META_SHEET in workbook.sheetnames:
            meta = {_text(row[1]): _text(row[0]) for row in workbook[EXPORT_META_SHEET].iter_rows(min_row=6, values_only=True) if len(row)>1 and row[0] and row[1]}
        names = [name for name in workbook.sheetnames if name != EXPORT_META_SHEET]
        sheets = []
        fields = [{**f, 'category_id':c['id'], 'category':c['label'], 'kind':c['kind']} for c in categories for f in c.get('fields',[]) if f.get('type') not in {'file','spacer','field_group','system_record_code','system_created_at','system_updated_at'}]
        for name in names:
            matches = [c['id'] for c in categories if _normalize(c['label']) == _normalize(name)]
            category_id = meta.get(name) or (matches[0] if len(matches)==1 else '__main__' if name==EXPORT_MAIN_SHEET or len(names)==1 else '')
            selected = [c for c in categories if c['id']==category_id or category_id=='__main__' and c['kind']=='main']
            fake = {**schema, 'categories':[{**c,'kind':'main'} for c in (selected or categories)]}
            item = inspect_import_workbook(encoded, fake, name, _workbook=workbook)
            item['category_id'] = category_id
            for column in item['columns']:
                header = column['technical_header'] or column['header']
                if header in {'minor_id','_linked_record_code','parent_minor_id','_parent_child_id'}:
                    column['suggested_target'] = {'minor_id':'__minor_id__','_linked_record_code':'__linked_record_code__','parent_minor_id':'__parent_minor_id__','_parent_child_id':'__parent_child_id__'}[header]
            sheets.append(item)
        return {'sheet_inspections':sheets, 'categories':[{'id':c['id'],'label':c['label'],'kind':c['kind']} for c in categories], 'all_fields':fields}
    finally:
        workbook.close()


def parse_import_sheets(encoded, schema, mappings):
    """Join main-sheet rows by ID and retain repeated rows as category patches."""
    if not isinstance(mappings,list) or not mappings:
        raise WorkbookExchangeError('اختر ورقة واحدة على الأقل للاستيراد.')
    workbook = _open_import_workbook(encoded)
    try:
        categories={c['id']:c for c in schema['categories']}
        owners={}; imported=[]; used_sheets=set(); repeated_keys=set()
        for config in mappings:
            name=config.get('sheet_name'); category_id=config.get('category_id')
            if name not in workbook.sheetnames or name==EXPORT_META_SHEET or name in used_sheets:
                raise WorkbookExchangeError('ورقة استيراد مفقودة أو مختارة أكثر من مرة.')
            used_sheets.add(name)
            category=categories.get(category_id)
            if category_id!='__main__' and category is None:
                raise WorkbookExchangeError(f'{name}: اختر الفئة المقابلة للورقة.')
            repeated=category is not None and category['kind']=='repeatable'
            mode=config.get('repeated_mode','merge')
            if mode not in {'merge','replace'}: raise WorkbookExchangeError('طريقة دمج البطاقات غير صالحة.')
            selected=[category] if category else [c for c in categories.values() if c['kind']=='main']
            fields={f['id'] for c in selected for f in c.get('fields',[]) if f.get('type') not in {'file','spacer','field_group','system_record_code','system_created_at','system_updated_at'}}
            allowed=fields | {'__record_code__','__archived__'} | ({'__minor_id__','__linked_record_code__','__parent_minor_id__','__parent_child_id__'} if repeated else set())
            mapping={}; targets=set()
            for col,target in config.get('mapping',{}).items():
                if target in {'','__ignore__'}: continue
                if target not in allowed or target in targets or not str(col).isdigit() or int(col)<1:
                    raise WorkbookExchangeError(f'{name}: مطابقة مكررة أو حقل لا ينتمي إلى الفئة المختارة.')
                mapping[int(col)]=target; targets.add(target)
            if not fields.intersection(targets): raise WorkbookExchangeError(f'{name}: اربط حقل بيانات واحدًا على الأقل.')
            headers=_sheet_headers(workbook[name],fields)
            for row_number,row in enumerate(workbook[name].iter_rows(min_row=headers['data_row'],values_only=True),start=headers['data_row']):
                if not any(_text(v) for v in row): continue
                values={target:row[col-1] if col<=len(row) else None for col,target in mapping.items()}
                code=_text(values.get('__record_code__')).upper()
                if code and code in owners: item=owners[code]
                else:
                    item={'source_row':row_number,'source_sheets':[],'record_code':code,'main':{},'related':{},'related_modes':{}}
                    imported.append(item)
                    if code: owners[code]=item
                source=f'{name}: {row_number}'; item['source_sheets'].append(source)
                if not code and (len(mappings)>1 or repeated):
                    item['_import_error']='يلزم ID لربط الصفوف بين الأوراق والبطاقات المتكررة.'; continue
                if '__archived__' in values and _text(values['__archived__']):
                    archived=_truthy(values['__archived__'])
                    if 'archived' in item and item['archived']!=archived: item['_import_error']='قيم أرشفة متعارضة بين الأوراق.'
                    item['archived']=archived
                data={key:value for key,value in values.items() if key in fields}
                if repeated:
                    raw_minor=_text(values.get('__minor_id__')); minor=None
                    if raw_minor:
                        try:
                            minor=int(raw_minor)
                            if minor<1: raise ValueError()
                        except ValueError:
                            item['_import_error']=f'{source}: رقم البطاقة يجب أن يكون عددًا صحيحًا موجبًا.'; continue
                        key=(code,category_id,minor)
                        if key in repeated_keys: item['_import_error']=f'{source}: رقم بطاقة مكرر في الفئة نفسها.'
                        repeated_keys.add(key)
                    previous_mode=item['related_modes'].get(category_id)
                    if previous_mode and previous_mode!=mode: item['_import_error']='طرق دمج متعارضة للفئة نفسها.'
                    item['related_modes'][category_id]=mode
                    child={'values':data,'minor_id':minor}
                    if '__parent_minor_id__' in values: child['parent_minor_id']=_text(values['__parent_minor_id__'])
                    if '__parent_child_id__' in values: child['parent_child_id']=_text(values['__parent_child_id__'])
                    if '__linked_record_code__' in values: child['linked_record_code']=_text(values['__linked_record_code__']).upper()
                    item['related'].setdefault(category_id,[]).append(child)
                else:
                    for key,value in data.items():
                        if key in item['main'] and _text(item['main'][key])!=_text(value): item['_import_error']=f'{source}: قيم متعارضة للحقل نفسه في أكثر من صف أو ورقة.'
                        item['main'][key]=value
        return imported
    finally:
        workbook.close()
